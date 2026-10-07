#!/usr/bin/env python3
"""
Run every script in the kit on the samples and check the results. Takes about a minute.

Run it after installing the requirements, and again after you change any script:
    python scripts/selftest.py

It writes only to a temporary folder (the samples are not changed) and prints PASS or
FAIL for each check. Exit code 0 means everything passed.
"""

from __future__ import annotations

import csv
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

KIT = Path(__file__).resolve().parent.parent
SCRIPTS = KIT / "scripts"
SAMPLES = KIT / "samples"
DICT = SAMPLES / "02_redcap" / "lakeside_data_dictionary.csv"
EVENTS = SAMPLES / "02_redcap" / "instrument_event_mapping.csv"
REPEATING = SAMPLES / "02_redcap" / "repeating_instruments.csv"
EVENT_DAYS = SAMPLES / "02_redcap" / "events.csv"
XLSFORM = SAMPLES / "03_xlsform" / "lakeside_health_check.xlsx"
XLS_SOURCE = SAMPLES / "03_xlsform" / "source"
QUESTIONNAIRE = SAMPLES / "01_client_input" / "lakeside_questionnaire_v1.2.docx"
FAKE = SAMPLES / "05_test_data" / "fake_export.csv"
ANSWER_KEY = SAMPLES / "05_test_data" / "fake_export_answer_key.csv"
RULES = SAMPLES / "05_test_data" / "cleaning_rules.csv"
BROKEN_DICT = SAMPLES / "07_checker_demo" / "broken_dictionary.csv"
BROKEN_EVENTS = SAMPLES / "07_checker_demo" / "broken_event_mapping.csv"
BROKEN_XLS = SAMPLES / "07_checker_demo" / "broken_xlsform.xlsx"

results = []


def run(*args):
    command = [sys.executable, "-B", *[str(a) for a in args]]
    return subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace",
                          cwd=KIT, timeout=300)


def check(name, ok, detail=""):
    results.append((name, bool(ok), detail))
    print(f"{'PASS' if ok else 'FAIL'}  {name}" + (f"  ({detail})" if detail and not ok else ""))


def logic_tests():
    sys.path.insert(0, str(SCRIPTS))
    import redcap_logic as rl
    values = {"a": "5", "b": "", "chk___2": "1"}

    def get(ref):
        key = ref.field + (f"___{ref.code}" if ref.code else "")
        return values.get(key, "")

    cases = [("[a] > 3 and [a] < 10", True), ("[b] < 3", False), ("[b] = ''", True), ("[chk(2)] = '1'", True),
             ("round(10 / 3, 2)", 3.33), ("sum([a], [b], 2)", 7.0), ("if([a] >= 5, 1, 0)", 1.0),
             ("datediff('2026-01-01', '2026-01-31', 'd')", 30.0), ("rounddown(datediff('1960-05-20', '2026-02-10', 'y'), 0)", 65.0)]
    for text, expected in cases:
        tree, _ = rl.parse(text)
        got = rl.evaluate(tree, get)
        if got != expected:
            return False, f"{text} gave {got!r}, expected {expected!r}"
    for bad in ["[a] = '1' &&", "[a = 1", "[a] = \u20181\u2019", "1 < [a] < 5"]:
        try:
            rl.parse(bad)
            return False, f"no error for {bad!r}"
        except rl.LogicError:
            pass
    return True, ""


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    tmp = Path(tempfile.mkdtemp(prefix="redcap_kit_selftest_"))
    try:
        ok, detail = logic_tests()
        check("logic parser and evaluator", ok, detail)

        r = run(SCRIPTS / "validate_redcap.py", DICT, "--events", EVENTS, "--repeating", REPEATING, "--strict")
        check("sample dictionary passes validate_redcap.py (strict)", r.returncode == 0, r.stdout[-400:])

        r = run(SCRIPTS / "validate_redcap.py", BROKEN_DICT, "--events", BROKEN_EVENTS)
        expected = ["does not exist", "curly quote", "used 2 times", "circle", "identical choices",
                    "not an event in the event mapping", "looks like it collects identifying"]
        missing = [e for e in expected if e not in r.stdout]
        check("broken dictionary fails with the expected errors", r.returncode == 1 and not missing,
              f"exit {r.returncode}, missing: {missing}")

        rebuilt = tmp / "rebuilt.xlsx"
        r = run(SCRIPTS / "build_xlsform.py", XLS_SOURCE, "-o", rebuilt)
        same = False
        if r.returncode == 0:
            import openpyxl
            a = openpyxl.load_workbook(rebuilt, read_only=True)
            b = openpyxl.load_workbook(XLSFORM, read_only=True)
            same = all([list(x.iter_rows(values_only=True)) == list(y.iter_rows(values_only=True))
                        for x, y in zip(a.worksheets, b.worksheets)]) and a.sheetnames == b.sheetnames
        check("build_xlsform.py rebuilds the sample XLSForm from its CSVs", same,
              "the .xlsx differs from samples/03_xlsform/source; rebuild it")

        r = run(SCRIPTS / "build_xlsform.py", "--split", XLSFORM, "-o", tmp / "split")
        check("build_xlsform.py --split writes one CSV per sheet",
              r.returncode == 0 and (tmp / "split" / "survey.csv").exists(), r.stdout[-300:])

        r = run(SCRIPTS / "check_xlsform.py", XLSFORM, "--xml", tmp / "form.xml")
        odk = "ODK Validate run" in r.stdout
        check("sample XLSForm passes check_xlsform.py" + ("" if odk else " (ODK Validate skipped: no Java)"),
              r.returncode == 0 and (tmp / "form.xml").exists(), r.stdout[-400:])

        r = run(SCRIPTS / "check_xlsform.py", BROKEN_XLS)
        check("broken XLSForm fails with plain-English errors",
              r.returncode == 1 and "REDCap-style [field]" in r.stdout and "In plain English" in r.stdout,
              f"exit {r.returncode}")

        r = run(SCRIPTS / "make_codebook.py", "--dictionary", DICT, "--events", EVENTS, "--repeating", REPEATING,
                "--xlsform", XLSFORM, "--out", tmp / "codebook", "--logic-table", tmp / "logic.md")
        files = [tmp / f"codebook_{k}.{e}" for k in ("redcap", "xlsform") for e in ("docx", "html")]
        good = r.returncode == 0 and all(f.exists() for f in files)
        if good:
            from docx import Document
            text = "\n".join(p.text for p in Document(files[0]).paragraphs)
            html = files[1].read_text(encoding="utf-8")
            good = "Codebook" in text and "conditions___1" in html and "fu_bp_change" in html
        check("make_codebook.py writes .docx and .html for REDCap and XLSForm", good, r.stdout[-300:] + r.stderr[-300:])
        check("make_codebook.py --logic-table lists skip logic",
              (tmp / "logic.md").exists() and "cig_per_day" in (tmp / "logic.md").read_text(encoding="utf-8"))

        fake = tmp / "fake.csv"
        r = run(SCRIPTS / "make_fake_export.py", "--dictionary", DICT, "--events", EVENTS, "--repeating", REPEATING,
                "--event-days", EVENT_DAYS, "--event-logic", "visit_*=[screening_arm_1][fu_agree] = '1'",
                "--repeat-logic", "health_event=[fu_health_event] = '1'", "--records", "30", "--errors", "16",
                "--seed", "99", "--today", "2026-10-07", "--start", "2026-02-01", "--out", fake)
        check("make_fake_export.py writes fake data and an answer key",
              r.returncode == 0 and fake.exists() and (tmp / "fake_answer_key.csv").exists(), r.stdout + r.stderr[-300:])

        r = run(SCRIPTS / "clean_export.py", "--dictionary", DICT, "--export", fake, "--events", EVENTS,
                "--rules", RULES, "--today", "2026-10-07", "--outdir", tmp / "clean",
                "--answer-key", tmp / "fake_answer_key.csv")
        check("clean_export.py finds every deliberate mistake in new fake data",
              r.returncode == 0 and "deliberate mistakes found" in r.stdout, r.stdout[-500:])

        r = run(SCRIPTS / "clean_export.py", "--dictionary", DICT, "--export", FAKE, "--events", EVENTS,
                "--rules", RULES, "--today", "2026-10-07", "--outdir", tmp / "clean_sample", "--answer-key", ANSWER_KEY)
        labelled = tmp / "clean_sample" / "fake_export_labelled.csv"
        good = r.returncode == 0 and labelled.exists()
        if good:
            with open(labelled, newline="", encoding="utf-8") as fh:
                first = next(csv.DictReader(fh))
            good = first.get("consent") == "Yes" and first.get("screening_complete") in ("Complete", "Incomplete")
        check("clean_export.py on the sample export: all mistakes found, labels applied", good, r.stdout[-500:])

        r = run(SCRIPTS / "check_wording.py", "--source", QUESTIONNAIRE, "--dictionary", DICT, "--xlsform", XLSFORM)
        check("check_wording.py compares labels with the questionnaire",
              r.returncode == 0 and "match the source word for word" in r.stdout, r.stdout[-300:])

        r = run(SCRIPTS / "extract_text.py", QUESTIONNAIRE, "--python-docx")
        check("extract_text.py reads the questionnaire", "What is your date of birth?" in r.stdout, r.stderr[-300:])

        r = run(SCRIPTS / "md_to_docx.py", KIT / "templates" / "delivery_note.md", "--outdir", tmp)
        check("md_to_docx.py converts a template", r.returncode == 0 and (tmp / "delivery_note.docx").exists(),
              r.stderr[-300:])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    failed = [name for name, ok, _ in results if not ok]
    print()
    print(f"{len(results) - len(failed)} of {len(results)} checks passed.")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
