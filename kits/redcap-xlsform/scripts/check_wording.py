#!/usr/bin/env python3
"""
Check that question wording matches the client's approved questionnaire word for word.

Ethics committees approve exact wording. This script takes every label, section header,
note and answer choice from a REDCap data dictionary and/or an XLSForm and looks for it
in the client's document (.docx or .txt). It lists:
  * texts that match word for word
  * texts that differ only in capitals or punctuation
  * texts that are not in the document at all (added or changed - each needs a reason,
    and anything participants see may need the client's or the committee's approval)

Usage (from the kit folder):
    python scripts/check_wording.py --source samples/01_client_input/lakeside_questionnaire_v1.2.docx \
        --dictionary samples/02_redcap/lakeside_data_dictionary.csv \
        --xlsform samples/03_xlsform/lakeside_health_check.xlsx --report out/wording_check.txt

Exit code: 0 (use --strict to get 1 when any text is not an exact match).
Needs python-docx for .docx sources and openpyxl for XLSForms.
"""

from __future__ import annotations

import argparse
import difflib
import re
import sys
import unicodedata
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import redcap_dictionary as rd  # noqa: E402

NUMBERING = re.compile(r"^\s*[A-Z]?\d+[a-z]?[.)]\s+")
REPLACEMENTS = {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-",
                "\u00a0": " ", "\u2026": "..."}


BOXES = re.compile("[\u2610\u2611\u2612\u25a1\u25a2\uf0a8\uf06f]")


def normalise(text: str) -> str:
    text = unicodedata.normalize("NFKC", text or "")
    for old, new in REPLACEMENTS.items():
        text = text.replace(old, new)
    text = BOXES.sub(" ", text)                    # tick boxes
    text = re.sub(r"_{3,}|\|-{3,}\|", " ", text)     # answer lines and drawn scales
    return re.sub(r"\s+", " ", text).strip()


def loose(text: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", normalise(text).lower())).strip()


def read_source(path: Path) -> list[str]:
    if path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(path)
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    lines += [p.text for p in cell.paragraphs]
        for section in doc.sections:
            lines += [p.text for p in section.header.paragraphs]
        return [normalise(line) for line in lines if line.strip()]
    if path.suffix.lower() in (".txt", ".md", ".csv"):
        return [normalise(line) for line in path.read_text(encoding="utf-8", errors="replace").splitlines() if line.strip()]
    raise SystemExit(f"ERROR: cannot read {path.suffix} files. Save the questionnaire as .docx or .txt first "
                     "(LibreOffice: soffice --headless --convert-to docx file.doc).")


def items_from_dictionary(path) -> list[tuple[str, str, str]]:
    dd, issues = rd.load_dictionary(path)
    if dd is None or not dd.header_ok:
        raise SystemExit("ERROR: cannot read the data dictionary; run validate_redcap.py first.")
    items = []
    for f in dd.fields:
        items.append((f.name, "label", rd.strip_html(f.label)))
        if f.section:
            items.append((f.name, "section header", rd.strip_html(f.section)))
        if f.note:
            items.append((f.name, "field note", rd.strip_html(f.note)))
        if f.ftype in rd.CHOICE_TYPES:
            items += [(f.name, f"choice {code}", rd.strip_html(label)) for code, label in f.choices]
        if f.ftype == "slider" and f.choices_raw:
            items += [(f.name, "slider label", p.strip()) for p in f.choices_raw.split("|") if p.strip()]
    return items


def items_from_xlsform(path) -> list[tuple[str, str, str]]:
    import xlsform_reader as xr
    form = xr.read_xlsform(path)
    language = form.settings.get("default_language") or (form.languages[0] if form.languages else None)
    items = []
    for row in form.survey:
        name = row.get("name", "")
        for column in ("label", "hint"):
            text = re.sub(r"\*\*|__", "", form.text(row, column, language))
            text = re.sub(r"\$\{\w+\}", "", text).strip(" :")
            if text:
                items.append((name, column, NUMBERING.sub("", text)))
    for list_name, rows in form.choice_lists().items():
        for row in rows:
            text = form.text(row, "label", language)
            if text:
                items.append((f"{list_name} list", f"choice {row.get('name', '')}", text))
    return items


def check(items, source_lines):
    exact_source = " \n ".join(source_lines)
    loose_source = " | ".join(loose(line) for line in source_lines)
    loose_lines = [loose(line) for line in source_lines]
    seen = {}
    for name, where, text in items:
        text = normalise(text)
        if not text:
            continue
        place = f"{name} ({where})"
        if place not in seen.setdefault(text, []):
            seen[text].append(place)
    results = {"exact": [], "near": [], "missing": []}
    for text, places in seen.items():
        if text in exact_source:
            results["exact"].append((text, places, None, 1.0))
        elif loose(text) and loose(text) in loose_source and _covers(text, loose_lines):
            line = next((source_lines[i] for i, ln in enumerate(loose_lines) if loose(text) in ln), "")
            results["near"].append((text, places, line, 1.0))
        else:
            best, score = "", 0.0
            for line in source_lines:
                ratio = difflib.SequenceMatcher(None, loose(text), loose(line)).ratio()
                if ratio > score:
                    best, score = line, ratio
            results["missing"].append((text, places, best, score))
    return results


def _covers(text: str, loose_lines: list) -> bool:
    """A near match must be a real part of a source line, not two words inside a long sentence."""
    small = loose(text)
    return any(small in line and (len(small) >= 0.4 * len(line) or len(small.split()) >= 4)
               for line in loose_lines)


def word_diff(new: str, old: str) -> str:
    """Show the words that differ: [-removed-] {+added+}."""
    a, b = old.split(), new.split()
    out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b).get_opcodes():
        if op == "equal":
            out.append(" ".join(a[i1:i2]))
        if op in ("delete", "replace"):
            out.append("[-" + " ".join(a[i1:i2]) + "-]")
        if op in ("insert", "replace"):
            out.append("{+" + " ".join(b[j1:j2]) + "+}")
    return " ".join(out)


def report(results, source, checked, strict) -> tuple[str, int]:
    total = sum(len(v) for v in results.values())
    lines = ["Wording check", "=============", "", f"Source:   {source}", f"Checked:  {checked}", ""]
    lines.append(f"RESULT: {len(results['exact'])} of {total} distinct texts match the source word for word; "
                 f"{len(results['near'])} differ only in capitals or punctuation; "
                 f"{len(results['missing'])} are not in the source.")
    lines.append("")
    lines += ["Differs only in capitals or punctuation (check these)",
              "-----------------------------------------------------"]
    if not results["near"]:
        lines.append("  None.")
    for text, places, line, _ in results["near"]:
        lines.append(f'* "{text}"')
        lines.append(f"    used by: {', '.join(places)}")
        lines.append(f'    source:  "{line}"')
    lines.append("")
    lines += ["Not found in the source (added or changed: give a reason, and get approval for anything",
              "participants see)", "-" * 86]
    if not results["missing"]:
        lines.append("  None.")
    for text, places, best, score in sorted(results["missing"], key=lambda r: -r[3]):
        lines.append(f'* "{text}"')
        lines.append(f"    used by: {', '.join(places)}")
        if best and score >= 0.6:
            lines.append(f"    closest source text ({score:.0%} similar): \"{best}\"")
            lines.append(f"    difference: {word_diff(text, best)}")
        else:
            lines.append("    nothing similar in the source (probably text added by the form builder)")
    lines.append("")
    lines += ["Exact matches", "-------------"]
    lines.append("  " + "; ".join(f'"{t[:60]}{"..." if len(t) > 60 else ""}"' for t, *_ in results["exact"])
                 if results["exact"] else "  None.")
    lines += ["", "Normal additions: labels of calculated fields, staff instructions, and adapted instructions",
              "such as 'move the slider' for 'mark the line'. Record them in the delivery note."]
    failed = strict and (results["near"] or results["missing"])
    return "\n".join(lines) + "\n", 1 if failed else 0


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Check form wording against the client's approved questionnaire.")
    parser.add_argument("--source", required=True, help="approved questionnaire (.docx or .txt)")
    parser.add_argument("--dictionary", help="REDCap data dictionary CSV")
    parser.add_argument("--xlsform", help="XLSForm .xlsx")
    parser.add_argument("--report", help="also save the report to this text file")
    parser.add_argument("--strict", action="store_true", help="exit with code 1 unless everything matches exactly")
    args = parser.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    if not args.dictionary and not args.xlsform:
        parser.error("give --dictionary and/or --xlsform")
    source = Path(args.source)
    if not source.exists():
        print(f"ERROR: file not found: {source}")
        return 2
    source_lines = read_source(source)
    items, checked = [], []
    if args.dictionary:
        items += items_from_dictionary(args.dictionary)
        checked.append(f"{Path(args.dictionary).name} (REDCap)")
    if args.xlsform:
        items += items_from_xlsform(args.xlsform)
        checked.append(f"{Path(args.xlsform).name} (XLSForm)")
    text, code = report(check(items, source_lines), source.name, ", ".join(checked), args.strict)
    print(text, end="")
    if args.report:
        Path(args.report).parent.mkdir(parents=True, exist_ok=True)
        Path(args.report).write_text(text, encoding="utf-8")
        print(f"Report saved to {args.report}")
    return code


if __name__ == "__main__":
    sys.exit(main())
