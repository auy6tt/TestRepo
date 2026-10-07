#!/usr/bin/env python3
"""
Turn a simple Markdown file (like the templates in this kit) into a Word document.

Fill in a template for a client in Markdown, then make a .docx to send them.
Supports: # headings, paragraphs, - bullets, 1. numbered lists, - [ ] tick boxes,
| tables |, **bold**, *italic*, `code`, > notes and --- lines.

Usage (from the kit folder):
    python scripts/md_to_docx.py templates/client_test_checklist.md -o out/client_test_checklist.docx
    python scripts/md_to_docx.py templates/*.md --outdir out/

Needs python-docx (pip install python-docx).
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

BOX = "\u2610"
TICKED = "\u2611"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_inline(paragraph, text, size=None):
    from docx.shared import Pt
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)


def convert(source: Path, target: Path, page: str = "a4") -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = (Cm(21.0), Cm(29.7)) if page == "a4" else (Cm(21.59), Cm(27.94))
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.0))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    def shade(cell, color):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tc_pr.append(shd)

    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or stripped.startswith("<!--"):
            i += 1
            continue
        if stripped.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(lines[i].rstrip())
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                i += 1
            i += 1
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)", stripped)
        if heading:
            level = len(heading.group(1))
            if level == 1 and first_heading:
                doc.add_paragraph(heading.group(2), style="Title")
            else:
                doc.add_heading(heading.group(2), level=min(level, 3))
            first_heading = False
            i += 1
            continue
        if re.fullmatch(r"-{3,}|\*{3,}", stripped):
            p = doc.add_paragraph()
            p_pr = p._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "B0B8C1")):
                bottom.set(qn(key), value)
            border.append(bottom)
            p_pr.append(border)
            i += 1
            continue
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [[c.strip() for c in row.strip("|").split("|")] for row in block
                    if not re.fullmatch(r"\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?", row)]
            header = True
            if rows and not any(cell for cell in rows[0]):
                rows, header = rows[1:], False  # key-value table without a header row
            if not rows:
                continue
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for r, row in enumerate(rows):
                for c in range(width):
                    cell = table.rows[r].cells[c]
                    cell.text = ""
                    text = row[c] if c < len(row) else ""
                    text = text.replace("[ ]", BOX).replace("[x]", TICKED)
                    add_inline(cell.paragraphs[0], text, size=9.5)
                    if (r == 0 and header) or (not header and c == 0):
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                        shade(cell, "E8EEF4")
            doc.add_paragraph()
            continue
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, stripped.lstrip("> ").strip())
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x44, 0x4F, 0x5A)
            i += 1
            continue
        box = re.match(r"^(\s*)[-*]\s+\[( |x|X)\]\s+(.*)", line)
        bullet = re.match(r"^(\s*)[-*]\s+(.*)", line)
        number = re.match(r"^(\s*)(\d+)[.)]\s+(.*)", line)
        if box:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4 + len(box.group(1)) * 0.3)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, f"{TICKED if box.group(2).lower() == 'x' else BOX}  {box.group(3)}")
        elif bullet:
            style = "List Bullet 2" if len(bullet.group(1)) >= 2 else "List Bullet"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, bullet.group(2))
        elif number:  # numbers are written out so every list starts at its own first number
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7 + len(number.group(1)) * 0.3)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, f"{number.group(2)}.\t{number.group(3)}")
        else:
            text = stripped
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(
                    r"^\s*(#|[-*]\s|\d+[.)]\s|\||>|-{3,})", lines[i + 1]):
                i += 1
                text += " " + lines[i].strip()
            add_inline(doc.add_paragraph(), text)
        i += 1
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Convert simple Markdown files to Word (.docx).")
    parser.add_argument("files", nargs="+", help="Markdown files")
    parser.add_argument("-o", "--output", help="output .docx (only with one input file)")
    parser.add_argument("--outdir", help="folder for the .docx files (default: next to each input)")
    parser.add_argument("--page", choices=["a4", "letter"], default="a4")
    args = parser.parse_args(argv)
    if args.output and len(args.files) > 1:
        parser.error("-o works with one input file; use --outdir for several")
    for name in args.files:
        source = Path(name)
        if not source.exists():
            print(f"ERROR: not found: {source}")
            return 2
        if args.output:
            target = Path(args.output)
        else:
            folder = Path(args.outdir) if args.outdir else source.parent
            target = folder / (source.stem + ".docx")
        convert(source, target, args.page)
        print(f"Wrote {target}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
