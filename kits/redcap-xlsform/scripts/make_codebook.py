#!/usr/bin/env python3
"""
Make a readable codebook (.docx and .html) from a REDCap data dictionary and/or an XLSForm.

For every variable the codebook shows the question text, the type, the stored codes
and their labels, limits, whether it is required or an identifier, and when it is
shown (branching logic written out in plain English, with the original logic too).

Usage (from the kit folder):
    python scripts/make_codebook.py --dictionary samples/02_redcap/lakeside_data_dictionary.csv \
        --events samples/02_redcap/instrument_event_mapping.csv \
        --repeating samples/02_redcap/repeating_instruments.csv \
        --title "Lakeside Community Health Check" --out samples/04_codebook/lakeside_codebook

    python scripts/make_codebook.py --xlsform samples/03_xlsform/lakeside_health_check.xlsx \
        --out samples/04_codebook/lakeside_kobo_codebook

With one input you get OUT.docx and OUT.html. With both inputs you get OUT_redcap.* and OUT_xlsform.*.
Needs python-docx for .docx (pip install python-docx) and openpyxl for XLSForms.
"""

from __future__ import annotations

import argparse
import datetime as dt
import html
import re
import sys
from dataclasses import dataclass, field as dc_field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import redcap_dictionary as rd  # noqa: E402
import redcap_logic as rl  # noqa: E402


# ---------------------------------------------------------------------------
# A neutral codebook model, filled from REDCap or XLSForm, drawn as docx or html
# ---------------------------------------------------------------------------

@dataclass
class Row:
    kind: str                      # "field" or "header"
    variable: str = ""
    flags: list = dc_field(default_factory=list)   # e.g. ["Required", "Identifier", "Q 5"]
    label: str = ""
    note: str = ""
    type: str = ""
    values: list = dc_field(default_factory=list)  # lines
    shown: str = ""                # plain English
    shown_raw: str = ""            # original logic
    muted: bool = False            # descriptive text / notes that store no data


@dataclass
class Section:
    title: str
    intro: str
    rows: list


@dataclass
class Codebook:
    title: str
    source: str
    generated: str
    facts: list                    # [(label, value)]
    reading_notes: list            # bullet points
    overview_head: list
    overview_rows: list
    sections: list
    identifiers: list


def humanise(name: str) -> str:
    text = name.replace("_", " ").strip()
    return text[:1].upper() + text[1:]


# ---------------------------------------------------------------------------
# REDCap
# ---------------------------------------------------------------------------

def redcap_type(f: rd.Field) -> str:
    base = rd.FIELD_TYPES.get(f.ftype, f.ftype)
    if f.ftype == "text" and f.validation:
        info = rd.VALIDATION_TYPES.get(rd.LEGACY_VALIDATION.get(f.validation, f.validation))
        return f"Text: {info[1] if info else f.validation}"
    if f.ftype == "file" and f.validation == "signature":
        return "Signature"
    return base


def redcap_values(f: rd.Field) -> list:
    lines = []
    if f.ftype == "checkbox":
        lines.append("One column per option (1 = ticked, 0 = not ticked):")
        lines += [f"{rd.checkbox_column(f.name, code)} = {label}" for code, label in f.choices]
    elif f.choices:
        lines += [f"{code} = {rd.strip_html(label)}" for code, label in f.choices]
    if f.ftype == "calc":
        lines.append(f"Formula: {f.choices_raw}")
    if f.ftype == "slider":
        labels = [p.strip() for p in f.choices_raw.split("|")] if f.choices_raw else []
        low, high = f.vmin or "0", f.vmax or "100"
        lines.append(f"{low} to {high}" + (f" ({labels[0]} ... {labels[-1]})" if labels else ""))
    if f.ftype == "text" and (f.vmin or f.vmax):
        if f.vmin and f.vmax:
            lines.append(f"Range: {f.vmin} to {f.vmax}")
        elif f.vmin:
            lines.append(f"Minimum: {f.vmin}")
        else:
            lines.append(f"Maximum: {f.vmax}")
    if f.is_date:
        lines.append("Stored as YYYY-MM-DD" + (" HH:MM" if "datetime" in f.validation else ""))
    for tag, arg in rd.action_tags(f.annotation):
        if tag == "@NONEOFTHEABOVE":
            lines.append(f"'{f.choice_map.get(arg, arg)}' cannot be ticked with other options")
    return lines


def redcap_shown(f: rd.Field, dd: rd.Dictionary) -> tuple[str, str]:
    if not f.branching:
        return "Always", ""
    try:
        tree, _ = rl.parse(f.branching)
        return rl.describe(tree, dd.by_name), f.branching
    except rl.LogicError:
        return "See logic", f.branching


def codebook_from_redcap(dd, title, event_map=None, repeating=None, when="") -> Codebook:
    sections = []
    for form in dd.forms:
        fields = dd.form_fields(form)
        rows = []
        for f in fields:
            if f.section and not (f.matrix and f is not fields[0] and _matrix_continues(f, dd)):
                rows.append(Row("header", label=rd.strip_html(f.section)))
            flags = []
            if f.qnum:
                flags.append(f"Q {f.qnum}")
            if f.is_required:
                flags.append("Required")
            if f.is_identifier:
                flags.append("Identifier")
            if f.matrix:
                flags.append(f"Matrix: {f.matrix}")
            if f is dd.fields[0]:
                flags.append("Record ID")
            shown, raw = redcap_shown(f, dd)
            notes = [rd.strip_html(f.note)] if f.note else []
            tags = [t for t, _ in rd.action_tags(f.annotation)]
            if tags:
                notes.append("Action tags: " + ", ".join(tags))
            rows.append(Row("field", f.name, flags, rd.strip_html(f.label), " | ".join(n for n in notes if n),
                            redcap_type(f), redcap_values(f), shown, raw, muted=f.ftype == "descriptive"))
        rows.append(Row("field", f"{form}_complete", ["Added by REDCap"], "Form status", "", "Form status",
                        [f"{k} = {v}" for k, v in rd.FORM_STATUS.items()], "Always", ""))
        intro = []
        if event_map:
            events = event_map.events_for_form(form)
            if events:
                intro.append("Events: " + ", ".join(events))
        if repeating:
            reps = sorted(e for e, frm in repeating if frm == form)
            if reps:
                intro.append("Repeating instrument (several instances per event) in: " + ", ".join(reps))
        data_fields = [f for f in fields if f.ftype != "descriptive"]
        intro.append(f"{len(data_fields)} variables")
        sections.append(Section(f"{humanise(form)} ({form})", ". ".join(intro) + ".", rows))

    overview_head = ["Instrument", "Variables"]
    if event_map:
        overview_head += list(event_map.events)
    overview_rows = []
    for form in dd.forms:
        row = [f"{humanise(form)} ({form})", str(sum(1 for f in dd.form_fields(form) if f.ftype != "descriptive"))]
        if event_map:
            for event in event_map.events:
                cell = ""
                if form in event_map.forms_by_event.get(event, []):
                    cell = "Repeating" if repeating and (event, form) in repeating else "Yes"
                row.append(cell)
        overview_rows.append(row)
    identifiers = [f"{f.name}: {rd.strip_html(f.label)}" for f in dd.fields if f.is_identifier]
    facts = [("Instruments", str(len(dd.forms))),
             ("Variables", str(sum(1 for f in dd.fields if f.ftype != "descriptive"))),
             ("Identifiers flagged", str(len(identifiers)))]
    if event_map:
        facts.append(("Events", str(len(event_map.events))))
    reading = [
        "Variable is the column name in a REDCap export. Codes are what is stored; labels are what people see.",
        "Checkbox questions export one column per option, named variable___code (three underscores): "
        "1 = ticked, 0 = not ticked.",
        "Shown when: the question only appears when this is true. A blank value can mean 'not shown' "
        "as well as 'not answered'.",
        "Dates are stored as YYYY-MM-DD whatever the display format.",
        "Each instrument has a status column (instrument_complete): 0 = Incomplete, 1 = Unverified, 2 = Complete.",
        "Identifiers are flagged so they can be removed from de-identified exports.",
    ]
    if event_map:
        reading.append("Longitudinal project: each row of an export is one record at one event "
                       "(column redcap_event_name).")
    if repeating:
        reading.append("Repeating instruments get extra rows with redcap_repeat_instrument and "
                       "redcap_repeat_instance filled in.")
    return Codebook(title, Path(dd.path).name, when, facts, reading, overview_head, overview_rows,
                    sections, identifiers)


def _matrix_continues(f, dd) -> bool:
    index = dd.fields.index(f)
    return index > 0 and dd.fields[index - 1].matrix == f.matrix


# ---------------------------------------------------------------------------
# XLSForm
# ---------------------------------------------------------------------------

XLS_TYPES = {
    "integer": "Whole number", "decimal": "Decimal number", "text": "Text", "date": "Date",
    "time": "Time", "datetime": "Date and time", "select_one": "Choose one",
    "select_multiple": "Choose all that apply", "rank": "Rank", "range": "Range / slider",
    "note": "Note (no data)", "calculate": "Calculated", "geopoint": "GPS point",
    "geotrace": "GPS line", "geoshape": "GPS area", "image": "Photo", "audio": "Audio",
    "video": "Video", "file": "File", "barcode": "Barcode", "acknowledge": "Acknowledge (OK)",
    "start": "Start time (automatic)", "end": "End time (automatic)", "today": "Date (automatic)",
    "deviceid": "Device ID (automatic)", "username": "User name (automatic)", "hidden": "Hidden",
}


def describe_xpath(expression: str, names: dict, lists: dict, form) -> str:
    """Rough plain-English version of an XLSForm expression."""
    if not expression:
        return ""

    def choice_label(name, code):
        row = names.get(name)
        if row is None:
            return code
        _, list_name = _base(row.get("type", ""))
        for choice in lists.get(list_name, []):
            if choice.get("name") == code:
                return form.text(choice, "label")
        return code

    text = expression
    text = re.sub(r"selected\(\s*\$\{(\w+)\}\s*,\s*'([^']*)'\s*\)",
                  lambda m: f'{m.group(1)} includes "{choice_label(m.group(1), m.group(2))}"', text)
    text = re.sub(r"\$\{(\w+)\}\s*(!=|=)\s*'([^']*)'",
                  lambda m: (f"{m.group(1)} is {'not ' if m.group(2) == '!=' else ''}blank" if m.group(3) == "" else
                             f'{m.group(1)} {"is not" if m.group(2) == "!=" else "is"} '
                             f'"{choice_label(m.group(1), m.group(3))}"'), text)
    text = re.sub(r"\$\{(\w+)\}", r"\1", text)
    text = re.sub(r"\band\b", "AND", text)
    text = re.sub(r"\bor\b", "OR", text)
    return text


def _base(type_text):
    import xlsform_reader as xr
    return xr.base_type(type_text)


def codebook_from_xlsform(path, title, when="") -> Codebook:
    import xlsform_reader as xr
    form = xr.read_xlsform(path)
    settings = form.settings
    language = settings.get("default_language") or (form.languages[0] if form.languages else None)
    lists = form.choice_lists()
    names = {row.get("name"): row for row in form.survey if row.get("name")}
    sections: list = []
    current = Section("Form start", "", [])
    sections.append(current)
    stack = []
    variables = 0
    for row in form.survey:
        kind, list_name = xr.base_type(row.get("type", ""))
        name = row.get("name", "")
        label = re.sub(r"\*\*|__", "", form.text(row, "label", language))
        if kind in ("begin_group", "begin_repeat"):
            stack.append((kind, name, label or humanise(name)))
            path_titles = " > ".join(t for _, _, t in stack)
            intro = []
            if kind == "begin_repeat":
                intro.append("Repeat group: can be filled in several times")
            if row.get("relevant"):
                intro.append("Shown when: " + describe_xpath(row["relevant"], names, lists, form))
            if row.get("appearance"):
                intro.append(f"Appearance: {row['appearance']}")
            current = Section(path_titles + (" (repeat)" if kind == "begin_repeat" else ""), ". ".join(intro), [])
            sections.append(current)
            continue
        if kind in ("end_group", "end_repeat"):
            if stack:
                stack.pop()
            path_titles = " > ".join(t for _, _, t in stack) or "Form (continued)"
            current = Section(path_titles + (" (continued)" if stack else ""), "", [])
            sections.append(current)
            continue
        if not kind:
            continue
        flags = []
        if row.get("required", "").lower() in ("yes", "true", "true()"):
            flags.append("Required")
        if kind not in ("note",):
            variables += 1
        values = []
        if kind in ("select_one", "select_multiple", "rank") and list_name:
            if kind == "select_multiple":
                values.append("Several answers, stored space-separated:")
            values += [f"{c.get('name')} = {form.text(c, 'label', language)}" for c in lists.get(list_name, [])]
        if row.get("constraint"):
            values.append("Rule: " + describe_xpath(row["constraint"], names, lists, form))
        if row.get("calculation"):
            values.append("Formula: " + row["calculation"])
        if kind == "range" and row.get("parameters"):
            values.append("Parameters: " + row["parameters"])
        notes = [form.text(row, "hint", language)] if form.text(row, "hint", language) else []
        if row.get("appearance"):
            notes.append(f"Appearance: {row['appearance']}")
        relevant = row.get("relevant", "")
        current.rows.append(Row("field", name, flags, label or "(no label)", " | ".join(notes),
                                XLS_TYPES.get(kind, kind) + (f" ({list_name})" if list_name else ""), values,
                                describe_xpath(relevant, names, lists, form) if relevant else "Always (within its group)",
                                relevant, muted=kind == "note"))
    sections = [s for s in sections if s.rows]
    facts = [("Form id", settings.get("form_id", "(not set)")),
             ("Version", settings.get("version", "(not set)")), ("Variables", str(variables)),
             ("Languages", ", ".join(form.languages) or "one")]
    reading = [
        "Variable is the question name. Exports may add group names in front (e.g. section_a/dob), "
        "depending on the export settings.",
        "Choose-all-that-apply questions store the chosen codes separated by spaces; KoboToolbox can also "
        "export one 0/1 column per option.",
        "Shown when: the question only appears when this is true (its group's condition applies too).",
        "Dates are stored as YYYY-MM-DD.",
    ]
    overview_rows = [[s.title, str(sum(1 for r in s.rows if r.kind == "field" and not r.muted))] for s in sections]
    return Codebook(title or settings.get("form_title") or Path(path).stem, Path(path).name, when, facts,
                    reading, ["Section", "Variables"], overview_rows, sections, [])


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

CSS = """
:root { --ink:#1d2733; --muted:#5b6672; --line:#d5dbe1; --head:#e8eef4; --accent:#1f3a5f; --sect:#f3f6f9; }
* { box-sizing: border-box; }
body { margin:0; background:#ffffff; color:var(--ink); font:15px/1.5 -apple-system, "Segoe UI", Roboto, Arial, sans-serif; }
main { max-width:1180px; margin:0 auto; padding:32px 16px 64px; }
h1 { color:var(--accent); font-size:28px; margin:0 0 4px; }
h2 { color:var(--accent); font-size:20px; margin:40px 0 6px; padding-top:8px; border-top:2px solid var(--accent); }
.sub { color:var(--muted); margin:0 0 20px; }
.facts { display:flex; flex-wrap:wrap; gap:10px; margin:16px 0 20px; padding:0; list-style:none; }
.facts li { background:var(--sect); border:1px solid var(--line); border-radius:8px; padding:8px 12px; }
.facts b { display:block; font-size:18px; }
.facts span { color:var(--muted); font-size:13px; }
.box { background:var(--sect); border:1px solid var(--line); border-radius:8px; padding:12px 18px; margin:12px 0; }
.box ul { margin:6px 0; padding-left:20px; }
.intro { color:var(--muted); margin:0 0 10px; }
.wrap { overflow-x:auto; }
table { border-collapse:collapse; width:100%; font-size:13.5px; }
th, td { border:1px solid var(--line); padding:6px 8px; vertical-align:top; text-align:left; }
th { background:var(--head); }
td.var { font-family:Consolas, Menlo, monospace; font-weight:600; white-space:nowrap; }
.flag { display:inline-block; font-family:-apple-system, "Segoe UI", Roboto, Arial, sans-serif; font-weight:500;
        font-size:11.5px; border:1px solid var(--line); border-radius:10px; padding:0 6px; margin:2px 2px 0 0; color:var(--muted); }
.flag.id { border-color:#b4472b; color:#b4472b; }
.flag.req { border-color:var(--accent); color:var(--accent); }
.note { color:var(--muted); font-style:italic; font-size:12.5px; }
.raw { color:var(--muted); font-family:Consolas, Menlo, monospace; font-size:11.5px; }
tr.header td { background:var(--sect); font-weight:600; }
tr.muted td { color:var(--muted); }
.values div { white-space:normal; }
footer { color:var(--muted); font-size:12.5px; margin-top:40px; }
@media print { body { font-size:11px; } main { padding:0; } h2 { break-before:page; } table { font-size:10px; }
  tr { break-inside:avoid; } .wrap { overflow:visible; } }
"""


def esc(text) -> str:
    return html.escape(str(text or ""))


def render_html(book: Codebook, target: Path) -> None:
    out = ["<!doctype html>", '<html lang="en">', "<head>", '<meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width, initial-scale=1">',
           f"<title>Codebook - {esc(book.title)}</title>", f"<style>{CSS}</style>", "</head>", "<body>", "<main>"]
    out.append(f"<h1>Codebook: {esc(book.title)}</h1>")
    out.append(f'<p class="sub">Generated from {esc(book.source)}' + (f" on {esc(book.generated)}" if book.generated else "") + ".</p>")
    out.append('<ul class="facts">' + "".join(f"<li><b>{esc(v)}</b><span>{esc(k)}</span></li>" for k, v in book.facts) + "</ul>")
    out.append('<div class="box"><b>How to read this codebook</b><ul>' +
               "".join(f"<li>{esc(n)}</li>" for n in book.reading_notes) + "</ul></div>")
    if book.identifiers:
        out.append('<div class="box"><b>Identifiers (remove or protect before sharing data)</b><ul>' +
                   "".join(f"<li>{esc(i)}</li>" for i in book.identifiers) + "</ul></div>")
    out.append('<h2 id="overview">Overview</h2><div class="wrap"><table><thead><tr>' +
               "".join(f'<th scope="col">{esc(h)}</th>' for h in book.overview_head) + "</tr></thead><tbody>")
    for i, row in enumerate(book.overview_rows):
        cells = [f'<a href="#s{i}">{esc(row[0])}</a>'] + [esc(c) for c in row[1:]]
        out.append("<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>")
    out.append("</tbody></table></div>")
    for i, section in enumerate(book.sections):
        out.append(f'<h2 id="s{i}">{esc(section.title)}</h2>')
        if section.intro:
            out.append(f'<p class="intro">{esc(section.intro)}</p>')
        out.append('<div class="wrap"><table><thead><tr><th scope="col">Variable</th><th scope="col">Question / label</th>'
                   '<th scope="col">Type</th><th scope="col">Values and rules</th><th scope="col">Shown when</th></tr></thead><tbody>')
        for row in section.rows:
            if row.kind == "header":
                out.append(f'<tr class="header"><td colspan="5">{esc(row.label)}</td></tr>')
                continue
            flags = "".join(
                f'<span class="flag{" id" if f == "Identifier" else (" req" if f == "Required" else "")}">{esc(f)}</span>'
                for f in row.flags)
            label = esc(row.label) + (f'<div class="note">{esc(row.note)}</div>' if row.note else "")
            values = "".join(f"<div>{esc(v)}</div>" for v in row.values)
            shown = esc(row.shown) + (f'<div class="raw">{esc(row.shown_raw)}</div>' if row.shown_raw else "")
            css = ' class="muted"' if row.muted else ""
            out.append(f'<tr{css}><td class="var">{esc(row.variable)}<div>{flags}</div></td><td>{label}</td>'
                       f'<td>{esc(row.type)}</td><td class="values">{values}</td><td>{shown}</td></tr>')
        out.append("</tbody></table></div>")
    out.append("<footer>Made with the redcap-xlsform starter kit. Check every path in a test project before "
               "collecting data.</footer>")
    out += ["</main>", "</body>", "</html>"]
    target.write_text("\n".join(out) + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------
# Word (.docx)
# ---------------------------------------------------------------------------

def render_docx(book: Codebook, target: Path, page: str = "a4") -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    accent = RGBColor(0x1F, 0x3A, 0x5F)
    doc = Document()
    section = doc.sections[0]
    width, height = (Cm(29.7), Cm(21.0)) if page == "a4" else (Cm(27.94), Cm(21.59))
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = width, height
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))
    usable = (width - Cm(3.0)) / 360000  # in cm

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = accent

    def shade(cell, color):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tc_pr.append(shd)

    def repeat_header(row):
        tr_pr = row._tr.get_or_add_trPr()
        element = OxmlElement("w:tblHeader")
        element.set(qn("w:val"), "true")
        tr_pr.append(element)

    def keep_row_together(row):
        tr_pr = row._tr.get_or_add_trPr()
        element = OxmlElement("w:cantSplit")
        element.set(qn("w:val"), "true")
        tr_pr.append(element)

    def widths(table, cm_list):
        table.autofit = False
        for i, w in enumerate(cm_list):
            table.columns[i].width = Cm(w)
            for cell in table.columns[i].cells:
                cell.width = Cm(w)

    def write(cell, lines, size=8.5, bold_first=False, mono_first=False, italic=False, color=None):
        cell.text = ""
        first = True
        for line in lines:
            if line is None:
                continue
            paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(line))
            run.font.size = Pt(size)
            if first and bold_first:
                run.bold = True
            if first and mono_first:
                run.font.name = "Consolas"
            if italic:
                run.italic = True
            if color is not None:
                run.font.color.rgb = color
            first = False

    def add_field(paragraph, instruction):
        run = paragraph.add_run()
        run.font.size = Pt(8)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        text = OxmlElement("w:instrText")
        text.set(qn("xml:space"), "preserve")
        text.text = instruction
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(text)
        run._r.append(end)

    footer = section.footer.paragraphs[0]
    footer.add_run(f"Codebook: {book.title}    Page ").font.size = Pt(8)
    add_field(footer, "PAGE")
    footer.add_run(" of ").font.size = Pt(8)
    add_field(footer, "NUMPAGES")

    doc.add_paragraph(f"Codebook: {book.title}", style="Title")
    sub = doc.add_paragraph(f"Generated from {book.source}" + (f" on {book.generated}" if book.generated else "") + ".")
    sub.runs[0].italic = True
    facts = doc.add_table(rows=2, cols=len(book.facts))
    facts.style = "Table Grid"
    for i, (key, value) in enumerate(book.facts):
        write(facts.rows[0].cells[i], [value], size=12, bold_first=True)
        write(facts.rows[1].cells[i], [key], size=8.5, color=RGBColor(0x5B, 0x66, 0x72))
        shade(facts.rows[0].cells[i], "F3F6F9")
        shade(facts.rows[1].cells[i], "F3F6F9")
    widths(facts, [usable / len(book.facts)] * len(book.facts))

    doc.add_heading("How to read this codebook", level=2)
    for note in book.reading_notes:
        doc.add_paragraph(note, style="List Bullet").paragraph_format.space_after = Pt(1)
    if book.identifiers:
        doc.add_heading("Identifiers (remove or protect before sharing data)", level=2)
        for item in book.identifiers:
            doc.add_paragraph(item, style="List Bullet").paragraph_format.space_after = Pt(1)

    doc.add_heading("Overview", level=1)
    overview = doc.add_table(rows=1 + len(book.overview_rows), cols=len(book.overview_head))
    overview.style = "Table Grid"
    overview.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, head in enumerate(book.overview_head):
        write(overview.rows[0].cells[i], [head], bold_first=True, size=9)
        shade(overview.rows[0].cells[i], "E8EEF4")
    for r, values in enumerate(book.overview_rows, start=1):
        for c, value in enumerate(values):
            write(overview.rows[r].cells[c], [value], size=9)
    first_col = 8.0
    rest = (usable - first_col) / max(1, len(book.overview_head) - 1)
    widths(overview, [first_col] + [min(rest, 3.2)] * (len(book.overview_head) - 1))

    columns = [4.8, 7.4, 3.4, 6.0, usable - 4.8 - 7.4 - 3.4 - 6.0]
    heads = ["Variable", "Question / label", "Type", "Values and rules", "Shown when"]
    grey = RGBColor(0x5B, 0x66, 0x72)
    for section_model in book.sections:
        doc.add_page_break()
        doc.add_heading(section_model.title, level=1)
        if section_model.intro:
            p = doc.add_paragraph(section_model.intro)
            p.runs[0].italic = True
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for i, head in enumerate(heads):
            write(table.rows[0].cells[i], [head], bold_first=True, size=9)
            shade(table.rows[0].cells[i], "E8EEF4")
        repeat_header(table.rows[0])
        for row in section_model.rows:
            cells = table.add_row().cells
            keep_row_together(table.rows[-1])
            if row.kind == "header":
                merged = cells[0].merge(cells[4])
                write(merged, [row.label], bold_first=True, size=9)
                shade(merged, "F3F6F9")
                continue
            write(cells[0], [row.variable] + row.flags, size=8.5, bold_first=True, mono_first=True)
            for paragraph in cells[0].paragraphs[1:]:
                for run in paragraph.runs:
                    run.font.color.rgb = grey
                    run.font.size = Pt(7.5)
            write(cells[1], [row.label], size=8.5, color=grey if row.muted else None)
            if row.note:
                p = cells[1].add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(row.note)
                run.italic = True
                run.font.size = Pt(7.5)
                run.font.color.rgb = grey
            write(cells[2], [row.type], size=8.5)
            write(cells[3], row.values or [""], size=8)
            write(cells[4], [row.shown], size=8)
            if row.shown_raw:
                p = cells[4].add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(row.shown_raw)
                run.font.size = Pt(7)
                run.font.name = "Consolas"
                run.font.color.rgb = grey
        widths(table, columns)
    doc.save(target)


# ---------------------------------------------------------------------------

def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Make a codebook (.docx and .html) from a REDCap data "
                                                 "dictionary and/or an XLSForm.")
    parser.add_argument("--dictionary", help="REDCap data dictionary CSV")
    parser.add_argument("--events", help="instrument-event mapping CSV (longitudinal projects)")
    parser.add_argument("--repeating", help="repeating instruments CSV (event_name, form_name, custom_form_label)")
    parser.add_argument("--xlsform", help="XLSForm .xlsx")
    parser.add_argument("--out", required=True, help="output path without extension, e.g. codebook/my_study_codebook")
    parser.add_argument("--title", help="study or form title shown on the codebook")
    parser.add_argument("--formats", default="docx,html", help="docx,html (default) or just one of them")
    parser.add_argument("--page", choices=["a4", "letter"], default="a4", help="paper size for the .docx")
    parser.add_argument("--date", help="date printed on the codebook (default: today; 'none' to leave it out)")
    args = parser.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")  # never crash on unusual characters
    if not args.dictionary and not args.xlsform:
        parser.error("give --dictionary and/or --xlsform")
    formats = {f.strip().lower() for f in args.formats.split(",") if f.strip()}
    when = "" if (args.date or "").lower() == "none" else (args.date or dt.date.today().isoformat())
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    books = []
    if args.dictionary:
        dd, issues = rd.load_dictionary(args.dictionary)
        errors = [i for i in issues if i.level == "ERROR"]
        if dd is None or not dd.header_ok or errors:
            for issue in errors or issues:
                print(f"{issue.level}: {issue.message}")
            print("Fix the dictionary first (run validate_redcap.py).")
            return 1
        event_map = rd.load_event_map(args.events)[0] if args.events else None
        repeating = rd.load_repeating(args.repeating)[0] if args.repeating else None
        book = codebook_from_redcap(dd, args.title or out.stem, event_map, repeating, when)
        books.append(("redcap", book))
    if args.xlsform:
        book = codebook_from_xlsform(args.xlsform, args.title, when)
        books.append(("xlsform", book))
    for kind, book in books:
        stem = out if len(books) == 1 else out.with_name(f"{out.name}_{kind}")
        if "html" in formats:
            render_html(book, stem.with_suffix(".html"))
            print(f"Wrote {stem.with_suffix('.html')}")
        if "docx" in formats:
            render_docx(book, stem.with_suffix(".docx"), args.page)
            print(f"Wrote {stem.with_suffix('.docx')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
