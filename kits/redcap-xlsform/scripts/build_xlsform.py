#!/usr/bin/env python3
"""
Build an XLSForm .xlsx from plain CSV files, one per sheet - or split an .xlsx back into CSVs.

Keeping an XLSForm as CSV files (survey.csv, choices.csv, settings.csv) makes it easy
to review, compare and edit with a text editor or a coding assistant. This script puts
them together into the .xlsx that KoboToolbox and ODK need, with tidy column widths.

Usage (from the kit folder):
    # CSV folder -> .xlsx
    python scripts/build_xlsform.py samples/03_xlsform/source -o samples/03_xlsform/lakeside_health_check.xlsx

    # .xlsx -> CSV folder (to edit a client's existing form)
    python scripts/build_xlsform.py --split client_form.xlsx -o client_form_source

Needs openpyxl (pip install openpyxl).
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

SHEET_ORDER = ["survey", "choices", "settings", "entities", "external_choices"]
WRAP_COLUMNS = ("label", "hint", "constraint_message", "required_message", "guidance_hint")


def build(source: Path, output: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    csv_files = {p.stem.lower(): p for p in source.glob("*.csv")}
    if "survey" not in csv_files:
        raise SystemExit(f"ERROR: {source} has no survey.csv")
    names = [n for n in SHEET_ORDER if n in csv_files] + sorted(n for n in csv_files if n not in SHEET_ORDER)
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for name in names:
        with open(csv_files[name], newline="", encoding="utf-8-sig") as fh:
            rows = [row for row in csv.reader(fh)]
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
        if not rows:
            continue
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        sheet.freeze_panes = "A2"
        for index, title in enumerate(rows[0], start=1):
            letter = sheet.cell(row=1, column=index).column_letter
            longest = max((len(r[index - 1]) for r in rows if len(r) >= index), default=8)
            wrap = title.split("::")[0].lower() in WRAP_COLUMNS
            sheet.column_dimensions[letter].width = min(max(10, longest + 2), 60 if wrap else 45)
            if wrap:
                for cell in sheet[letter][1:]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"Built {output} from {', '.join(n + '.csv' for n in names)}")


def split(xlsx: Path, output: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    output.mkdir(parents=True, exist_ok=True)
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
                      for v in row]
            while values and values[-1] == "":
                values.pop()
            rows.append(values)
        while rows and not any(rows[-1]):
            rows.pop()
        width = max((len(r) for r in rows), default=0)
        target = output / f"{sheet.title.strip().lower()}.csv"
        with open(target, "w", newline="", encoding="utf-8") as fh:
            csv.writer(fh).writerows([r + [""] * (width - len(r)) for r in rows])
        print(f"Wrote {target} ({max(len(rows) - 1, 0)} rows)")


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Build an XLSForm .xlsx from CSV sheets, or split one into CSVs.")
    parser.add_argument("source", help="folder of CSV files (survey.csv, choices.csv, settings.csv), "
                                       "or an .xlsx file with --split")
    parser.add_argument("-o", "--output", required=True, help="the .xlsx to write (or the folder, with --split)")
    parser.add_argument("--split", action="store_true", help="split an .xlsx into one CSV per sheet")
    args = parser.parse_args(argv)
    source, output = Path(args.source), Path(args.output)
    if not source.exists():
        print(f"ERROR: not found: {source}")
        return 2
    if args.split:
        split(source, output)
    else:
        build(source, output)
    return 0


if __name__ == "__main__":
    sys.exit(main())
