"""Shared helpers for the pdf-binder kit scripts.

You don't run this file yourself. build_binder.py, extract_sds.py,
verify_binder.py, make_templates.py and make_samples.py import it.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import re
import unicodedata
from pathlib import Path

KIT_DIR = Path(__file__).resolve().parent.parent

DEFAULT_ACCENT = "#1F3A5F"


class KitError(Exception):
    """A problem the user can fix (bad file name, wrong column, and so on)."""


# --------------------------------------------------------------------------
# Text for PDF pages
# --------------------------------------------------------------------------
# The built-in PDF fonts (Helvetica) cover Western European text only.
# pdf_safe() swaps or strips anything they can't draw, so a stray character
# in a product name never turns into a black box on the page.
_SWAPS = {
    " ": " ", " ": " ", " ": " ", " ": " ",
    "​": "", "‌": "", "‍": "", "﻿": "",
    "‐": "-", "‑": "-", "‒": "-", "−": "-",
    "≤": "<=", "≥": ">=", "→": "->", "←": "<-",
    "✓": "v", "✔": "v", "✗": "x", "×": "x",
}


def pdf_safe(text) -> str:
    if text is None:
        return ""
    out = []
    for ch in str(text):
        ch = _SWAPS.get(ch, ch)
        if not ch:
            continue
        try:
            ch.encode("cp1252")
            out.append(ch)
        except UnicodeEncodeError:
            plain = unicodedata.normalize("NFKD", ch).encode("ascii", "ignore").decode()
            out.append(plain or "?")
    return "".join(out).replace("\t", " ")


def clean_cell(value) -> str:
    """Turn a spreadsheet or CSV cell into tidy text."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, dt.date):
        return value.isoformat()
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return text.strip()


def long_date(day: dt.date) -> str:
    return f"{day:%B} {day.day}, {day.year}"


def short_date(day: dt.date) -> str:
    return f"{day:%b} {day.day}, {day.year}"


def hex_to_rgb(value: str, fallback: str = DEFAULT_ACCENT) -> tuple[float, float, float]:
    value = (value or fallback).strip().lstrip("#")
    if not re.fullmatch(r"[0-9a-fA-F]{6}", value):
        value = fallback.lstrip("#")
    return tuple(int(value[i:i + 2], 16) / 255 for i in (0, 2, 4))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def natural_key(text: str):
    """Sort key that puts 2 before 10 and 1.2 before 1.10."""
    parts = re.findall(r"\d+|[A-Za-z]+", text or "")
    key = []
    for part in parts:
        if part.isdigit():
            key.append((0, int(part), ""))
        else:
            key.append((1, 0, part.lower()))
    return tuple(key)


def is_example_row(values) -> bool:
    """Template rows that start with EXAMPLE are ignored by the scripts."""
    for value in values:
        if value:
            return value.strip().upper().startswith("EXAMPLE")
    return False


# --------------------------------------------------------------------------
# Reading CSV and XLSX tables
# --------------------------------------------------------------------------
def _norm_header(text: str) -> str:
    text = clean_cell(text).lower()
    text = re.sub(r"\(.*?\)", " ", text)          # drop "(optional)" and similar
    text = text.replace("*", " ").replace(":", " ").replace("_", " ")
    return re.sub(r"\s+", " ", text).strip()


def _match_columns(header_cells, columns):
    """Map canonical column names to positions using each name's synonyms."""
    normalized = [_norm_header(c) for c in header_cells]
    found = {}
    for canonical, synonyms in columns.items():
        for synonym in synonyms:
            if synonym in normalized:
                index = normalized.index(synonym)
                if index not in found.values():
                    found[canonical] = index
                    break
    return found


def _read_csv_rows(path: Path):
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx_rows(path: Path, sheet: str | None, prefer_sheets, columns, required):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - only without requirements
        raise KitError("Reading .xlsx files needs openpyxl. Run: pip install -r requirements.txt") from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise KitError(f"Could not open {path.name} as a spreadsheet ({exc}). "
                       "Close it in Excel if it is open, or save it again as .xlsx.") from exc
    try:
        names = workbook.sheetnames
        if sheet:
            match = [n for n in names if n.lower() == sheet.lower()]
            if not match:
                raise KitError(f"{path.name} has no sheet called '{sheet}'. Sheets: {', '.join(names)}")
            order = match
        else:
            preferred = [n for p in prefer_sheets for n in names if n.lower() == p.lower()]
            order = preferred + [n for n in names if n not in preferred]
        for name in order:
            rows = [[clean_cell(v) for v in row] for row in workbook[name].iter_rows(values_only=True)]
            if _find_header(rows, columns, required) is not None:
                return rows, name
        if sheet:
            return [[clean_cell(v) for v in row] for row in workbook[order[0]].iter_rows(values_only=True)], order[0]
        raise KitError(f"No sheet in {path.name} has the columns {', '.join(required)}.")
    finally:
        workbook.close()


def _find_header(rows, columns, required, scan=20):
    for index, row in enumerate(rows[:scan]):
        found = _match_columns(row, columns)
        if all(r in found for r in required):
            return index, found
    return None


def read_table(path, columns: dict, required=(), sheet: str | None = None, prefer_sheets=()):
    """Read a CSV or XLSX table.

    columns maps a canonical name to a list of accepted header spellings
    (lower case). Returns (rows, info). Each row is a dict of canonical
    name -> text plus "_row", the row number a person sees in Excel.
    """
    path = Path(path)
    if not path.is_file():
        raise KitError(f"File not found: {path}")
    suffix = path.suffix.lower()
    sheet_name = None
    if suffix in (".csv", ".txt", ".tsv"):
        rows = _read_csv_rows(path)
    elif suffix in (".xlsx", ".xlsm"):
        rows, sheet_name = _read_xlsx_rows(path, sheet, prefer_sheets, columns, required)
    elif suffix == ".xls":
        raise KitError(f"{path.name} is an old .xls file. Open it and save it as .xlsx or .csv first.")
    else:
        raise KitError(f"{path.name}: use a .csv or .xlsx file.")
    located = _find_header(rows, columns, required)
    if located is None:
        wanted = ", ".join(required)
        raise KitError(f"Could not find a header row with the columns {wanted} in {path.name}"
                       + (f" (sheet '{sheet_name}')" if sheet_name else "") + ".")
    header_index, positions = located
    records = []
    for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        values = [clean_cell(v) for v in row]
        if not any(values):
            continue
        if is_example_row(values):
            continue
        record = {"_row": offset}
        for canonical in columns:
            pos = positions.get(canonical)
            record[canonical] = values[pos] if pos is not None and pos < len(values) else ""
        records.append(record)
    info = {"sheet": sheet_name, "header_row": header_index + 1, "columns": sorted(positions)}
    return records, info


# --------------------------------------------------------------------------
# Cover files
# --------------------------------------------------------------------------
def load_cover(path) -> dict:
    """Load a cover file (.toml or .json) into a dict."""
    path = Path(path)
    if not path.is_file():
        raise KitError(f"Cover file not found: {path}")
    text = path.read_text(encoding="utf-8-sig")
    if path.suffix.lower() == ".json":
        try:
            data = json.loads(text)
        except json.JSONDecodeError as exc:
            raise KitError(f"{path.name} is not valid JSON: {exc}") from exc
    else:
        try:
            import tomllib
        except ImportError as exc:  # Python 3.10 or older
            raise KitError("Cover files in .toml format need Python 3.11 or newer.") from exc
        try:
            data = tomllib.loads(text)
        except tomllib.TOMLDecodeError as exc:
            raise KitError(f"{path.name} has a formatting mistake: {exc}. "
                           "Check that every value is inside double quotes.") from exc
    if not isinstance(data, dict):
        raise KitError(f"{path.name} should hold settings like title = \"...\"")
    data["_folder"] = str(path.parent)
    return data


def resolve_today(value: str, today: dt.date | None = None) -> str:
    """Replace the word 'today' with today's date written out."""
    if isinstance(value, str) and value.strip().lower() == "today":
        return long_date(today or dt.date.today())
    return value


# --------------------------------------------------------------------------
# XLSX writing helpers (XlsxWriter)
# --------------------------------------------------------------------------
# Text that starts with "=" or looks like a web address stays plain text.
# Formulas are only written on purpose with write_formula(). This also stops
# text pulled out of a PDF from turning into a live formula.
XLSX_OPTIONS = {"strings_to_formulas": False, "strings_to_urls": False}

def xlsx_formats(workbook, accent: str = DEFAULT_ACCENT) -> dict:
    """Common cell formats so every spreadsheet in the kit looks the same."""
    base = {"font_name": "Arial", "font_size": 10, "valign": "top"}

    def fmt(**extra):
        options = dict(base)
        options.update(extra)
        return workbook.add_format(options)

    return {
        "title": fmt(bold=True, font_size=16, font_color=accent, valign="vcenter"),
        "subtitle": fmt(italic=True, font_color="#555555"),
        "label": fmt(bold=True, font_color="#333333"),
        "header": fmt(bold=True, font_color="#FFFFFF", bg_color=accent, border=1,
                      border_color="#BFBFBF", text_wrap=True, valign="vcenter"),
        "text": fmt(border=1, border_color="#D9D9D9", text_wrap=True),
        "text_nowrap": fmt(border=1, border_color="#D9D9D9"),
        "bold": fmt(bold=True, border=1, border_color="#D9D9D9", text_wrap=True),
        "center": fmt(border=1, border_color="#D9D9D9", align="center"),
        "int": fmt(border=1, border_color="#D9D9D9", align="center", num_format="0"),
        "num1": fmt(border=1, border_color="#D9D9D9", align="center", num_format="0.0"),
        "date": fmt(border=1, border_color="#D9D9D9", align="center", num_format="yyyy-mm-dd"),
        "input": fmt(border=1, border_color="#BFBFBF", bg_color="#FFF9DB"),
        "input_date": fmt(border=1, border_color="#BFBFBF", bg_color="#FFF9DB", num_format="yyyy-mm-dd",
                          align="left"),
        "example": fmt(border=1, border_color="#D9D9D9", italic=True, font_color="#7F7F7F", text_wrap=True),
        "example_date": fmt(border=1, border_color="#D9D9D9", italic=True, font_color="#7F7F7F",
                            num_format="yyyy-mm-dd", align="center"),
        "uncertain": fmt(border=1, border_color="#D9D9D9", bg_color="#FFE9A8", text_wrap=True),
        "uncertain_date": fmt(border=1, border_color="#D9D9D9", bg_color="#FFE9A8", align="center",
                              num_format="yyyy-mm-dd"),
        "missing": fmt(border=1, border_color="#D9D9D9", bg_color="#F8CBAD", text_wrap=True),
        "checked": fmt(border=1, border_color="#D9D9D9", bg_color="#E2EFDA", text_wrap=True),
        "checked_date": fmt(border=1, border_color="#D9D9D9", bg_color="#E2EFDA", align="center",
                            num_format="yyyy-mm-dd"),
        "muted": fmt(border=1, border_color="#D9D9D9", font_color="#7F7F7F", italic=True, text_wrap=True),
        "section_row": fmt(bold=True, bg_color="#E9EEF5", border=1, border_color="#D9D9D9", text_wrap=True),
        "section_int": fmt(bold=True, bg_color="#E9EEF5", border=1, border_color="#D9D9D9", align="center"),
        "wrap": fmt(text_wrap=True),
        "note": fmt(text_wrap=True, font_color="#444444"),
        "h2": fmt(bold=True, font_size=12, font_color=accent),
        # Status colours used by conditional formatting
        "st_ok": workbook.add_format({"bg_color": "#C6EFCE", "font_color": "#006100", "bold": True}),
        "st_review": workbook.add_format({"bg_color": "#FFEB9C", "font_color": "#7F6000", "bold": True}),
        "st_check": workbook.add_format({"bg_color": "#FCE4D6", "font_color": "#843C0C", "bold": True}),
        "st_bad": workbook.add_format({"bg_color": "#FFC7CE", "font_color": "#9C0006", "bold": True}),
        "st_info": workbook.add_format({"bg_color": "#DDEBF7", "font_color": "#1F3A5F", "bold": True}),
        "st_grey": workbook.add_format({"bg_color": "#EDEDED", "font_color": "#404040"}),
    }


def write_readme_sheet(workbook, formats, title: str, paragraphs: list[str], name: str = "Read me"):
    """A plain sheet of instructions: one paragraph per row."""
    sheet = workbook.add_worksheet(name)
    sheet.hide_gridlines(2)
    sheet.set_column(0, 0, 110)
    sheet.write(0, 0, title, formats["title"])
    row = 2
    for text in paragraphs:
        if text.startswith("## "):
            row += 1
            sheet.write(row, 0, text[3:], formats["h2"])
        else:
            sheet.write(row, 0, text, formats["note"])
            lines = max(1, len(text) // 105 + 1)
            sheet.set_row(row, 14 * lines)
        row += 1
    return sheet
