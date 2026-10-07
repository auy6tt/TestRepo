#!/usr/bin/env python3
"""Run a replication package in a clean copy and compare its outputs with the originals.

What it does:
  1. Copies the package to a new temporary folder. The original is never touched.
  2. Deletes the files in the output folders of the copy, so every output has
     to be made again by the code.
  3. Runs the master script (Python or R) with a time limit, and saves
     everything it prints to run_log.txt.
  4. Compares each original output with the regenerated one:
       CSV/TSV tables      cell by cell; numbers must agree within a tolerance,
                           text (stars, brackets, labels) must agree exactly
       Excel files         the same, sheet by sheet
       .tex/.txt/.md/.html numbers within the tolerance, the rest exactly
       Images              identical bytes (SHA-256); if not, size, dimensions
                           and the share of pixels that differ
       PDF/EPS/SVG         identical bytes, or identical apart from embedded dates
       Anything else       identical bytes
  5. Writes comparison_report.md, plus comparison_details.csv when cells differ.

Usage:
  python run_and_compare.py PACKAGE --master main.py
  python run_and_compare.py PACKAGE --master main.R --outputs output results --timeout 7200
  python run_and_compare.py PACKAGE --master main.py --originals ../as-received/output
  python run_and_compare.py --compare-only --originals OLD_FOLDER --regenerated NEW_FOLDER

Stata, MATLAB and SAS need licences, so they are not run here. Ask the author to
run the master script on their computer, then compare their new output folder
with --compare-only.

Exit codes: 0 all outputs match, 1 differences or missing outputs, 2 the run failed.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import difflib
import hashlib
import io
import os
import platform
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import time
from pathlib import Path

VERSION = "1.0"

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}
VECTOR_EXTS = {".pdf", ".eps", ".ps", ".svg"}
TABLE_EXTS = {".csv", ".tsv"}
EXCEL_EXTS = {".xlsx", ".xlsm"}
TEXT_EXTS = {".tex", ".txt", ".md", ".html", ".htm", ".rtf", ".json"}
LOG_EXTS = {".log", ".smcl", ".rout"}
JUNK = {".DS_Store", "Thumbs.db", "desktop.ini"}
MISSING_TOKENS = {"", "na", "nan", ".", "n/a", "null", "none", "<na>"}
NUM_RE = re.compile(r"[-+\u2212]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d*)?(?:[eE][-+]?\d+)?"
                    r"|[-+\u2212]?\.\d+(?:[eE][-+]?\d+)?")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def list_files(folder: Path):
    """All files under folder, as {relative path: Path}, skipping system junk."""
    out = {}
    if not folder.is_dir():
        return out
    for path in sorted(folder.rglob("*")):
        if path.is_file() and path.name not in JUNK and "__pycache__" not in path.parts:
            out[path.relative_to(folder).as_posix()] = path
    return out


def snapshot(folder: Path):
    """Size and modification time of every file, to spot changes made by the run."""
    state = {}
    for rel, path in list_files(folder).items():
        st = path.stat()
        state[rel] = (st.st_size, st.st_mtime_ns)
    return state


def fmt_num(x: float) -> str:
    if x == 0:
        return "0"
    return f"{x:.3g}" if abs(x) < 1e-3 or abs(x) >= 1e6 else f"{x:.6g}"


def show_path(path: Path) -> str:
    """Path relative to the current folder when that is shorter and clearer."""
    try:
        rel = os.path.relpath(path)
    except ValueError:  # different drive on Windows
        return str(path)
    return rel if not rel.startswith("..") else str(path)


def md_cell(text) -> str:
    return str(text).replace("|", "\\|").replace("\n", "<br>")


# ---------------------------------------------------------------------------
# Comparing cells and tables
# ---------------------------------------------------------------------------
def split_numbers(text: str):
    """'0.257***' -> ('{}***', [0.257]). Used to compare numbers apart from text."""
    parts, nums, last = [], [], 0
    for m in NUM_RE.finditer(text):
        parts.append(text[last:m.start()])
        nums.append(float(m.group().replace(",", "").replace("\u2212", "-")))
        last = m.end()
    parts.append(text[last:])
    return "{}".join(parts), nums


class Tolerance:
    def __init__(self, rtol: float, atol: float):
        self.rtol, self.atol = rtol, atol

    def close(self, a: float, b: float) -> bool:
        return abs(a - b) <= self.atol + self.rtol * abs(a)

    def describe(self) -> str:
        return f"|original - regenerated| <= {self.atol:g} + {self.rtol:g} x |original|"


def compare_cell(a: str, b: str, tol: Tolerance):
    """Return (same?, largest absolute difference or None, reason)."""
    a, b = (a or "").strip(), (b or "").strip()
    if a == b:
        return True, 0.0, ""
    if a.lower() in MISSING_TOKENS and b.lower() in MISSING_TOKENS:
        return True, 0.0, ""
    ta, na = split_numbers(a)
    tb, nb = split_numbers(b)
    if ta != tb or len(na) != len(nb) or not na:
        return False, None, "text differs"
    worst, same = 0.0, True
    for x, y in zip(na, nb):
        d = abs(x - y)
        worst = max(worst, d)
        if not tol.close(x, y):
            same = False
    return same, worst, "" if same else "number differs"


def compare_grids(orig, new, tol: Tolerance, where_prefix=""):
    """Compare two lists of rows. Returns a result dict."""
    diffs, cells, worst = [], 0, 0.0
    header = orig[0] if orig else []
    for i in range(max(len(orig), len(new))):
        row_o = orig[i] if i < len(orig) else None
        row_n = new[i] if i < len(new) else None
        if row_o is None or row_n is None:
            which = "regenerated" if row_o is None else "original"
            text = ", ".join(row_n if row_o is None else row_o)
            diffs.append((f"{where_prefix}row {i + 1}", "" if row_o is None else text,
                          text if row_o is None else "", f"row only in the {which} file"))
            continue
        label = row_o[0].strip() if row_o and row_o[0].strip() else ""
        for j in range(max(len(row_o), len(row_n))):
            a = row_o[j] if j < len(row_o) else ""
            b = row_n[j] if j < len(row_n) else ""
            cells += 1
            same, d, reason = compare_cell(a, b, tol)
            if d is not None:
                worst = max(worst, d)
            if not same:
                col = header[j].strip() if i > 0 and j < len(header) and header[j].strip() else f"column {j + 1}"
                where = f"{where_prefix}row {i + 1}" + (f" ({label})" if label and j > 0 else "") + f", {col}"
                shown = reason if d is None else fmt_num(d)
                diffs.append((where, a, b, shown))
    return {"cells": cells, "diffs": diffs, "worst": worst,
            "shape": (len(orig), max((len(r) for r in orig), default=0))}


def read_delimited(path: Path):
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    if path.suffix.lower() == ".tsv":
        delim = "\t"
    else:
        first = text.split("\n", 1)[0]
        delim = max([",", ";", "\t", "|"], key=first.count) if first else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def read_excel(path: Path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else (repr(v) if isinstance(v, float) else str(v)) for v in row])
        while rows and not any(c.strip() for c in rows[-1]):
            rows.pop()
        sheets[ws.title] = rows
    wb.close()
    return sheets


def read_stata(path: Path):
    import pandas as pd
    df = pd.read_stata(path, convert_categoricals=False)
    rows = [list(map(str, df.columns))]
    rows += [["" if pd.isna(v) else (repr(v) if isinstance(v, float) else str(v)) for v in rec]
             for rec in df.itertuples(index=False)]
    return rows


# ---------------------------------------------------------------------------
# Comparing files
# ---------------------------------------------------------------------------
def compare_images(orig: Path, new: Path):
    size_o, size_n = orig.stat().st_size, new.stat().st_size
    detail = f"Bytes differ (file size {size_o:,} vs {size_n:,})."
    try:
        from PIL import Image, ImageChops
    except ImportError:
        return detail + " Look at both images side by side."
    try:
        with Image.open(orig) as a, Image.open(new) as b:
            if a.size != b.size:
                return detail + f" Dimensions differ: {a.size[0]}x{a.size[1]} vs {b.size[0]}x{b.size[1]}."
            diff = ImageChops.difference(a.convert("RGB"), b.convert("RGB")).convert("L")
            changed = sum(diff.point(lambda v: 255 if v > 0 else 0).histogram()[255:])
            total = a.size[0] * a.size[1]
            if changed == 0:
                return "Pixels identical; only file metadata differs."
            return (detail + f" Same dimensions ({a.size[0]}x{a.size[1]}); {changed / total:.2%} of pixels "
                    "differ. Look at both images side by side.")
    except Exception as exc:  # unreadable or unusual image format
        return detail + f" Could not open the images to compare pixels ({exc.__class__.__name__})."


VOLATILE = [
    rb"/CreationDate\s*\(D:[^)]*\)", rb"/ModDate\s*\(D:[^)]*\)",
    rb"<xmp:(CreateDate|ModifyDate|MetadataDate)>[^<]*</xmp:\1>",
    rb"<xmpMM:(DocumentID|InstanceID)>[^<]*</xmpMM:\1>",
    rb"/ID\s*\[\s*<[0-9A-Fa-f]*>\s*<[0-9A-Fa-f]*>\s*\]",
    rb"%%CreationDate:[^\n]*", rb"<dc:date>[^<]*</dc:date>",
    rb'\b(id|clip-path|xlink:href|href)="[^"]*"', rb"url\(#[^)]*\)",
]


def strip_volatile(data: bytes) -> bytes:
    for pattern in VOLATILE:
        data = re.sub(pattern, b"", data)
    return data


def compare_file(orig: Path, new: Path, tol: Tolerance, max_diffs: int):
    """Return dict(result, kind, details, diffs)."""
    ext = orig.suffix.lower()
    kind = ("Table" if ext in TABLE_EXTS | EXCEL_EXTS | {".dta"} else "Image" if ext in IMAGE_EXTS
            else "Vector image" if ext in VECTOR_EXTS else "Text" if ext in TEXT_EXTS else "File")
    h_o, h_n = sha256(orig), sha256(new)
    if h_o == h_n:
        return {"result": "Identical", "kind": kind, "details": f"Same bytes (SHA-256 {h_o[:12]})", "diffs": []}
    try:
        if ext in TABLE_EXTS or ext in EXCEL_EXTS or ext == ".dta":
            if ext in TABLE_EXTS:
                sheets_o, sheets_n = {"": read_delimited(orig)}, {"": read_delimited(new)}
            elif ext in EXCEL_EXTS:
                sheets_o, sheets_n = read_excel(orig), read_excel(new)
            else:
                sheets_o, sheets_n = {"": read_stata(orig)}, {"": read_stata(new)}
            diffs, cells, worst = [], 0, 0.0
            for name in sorted(set(sheets_o) | set(sheets_n)):
                if name not in sheets_o or name not in sheets_n:
                    side = "regenerated" if name not in sheets_o else "original"
                    diffs.append((f"sheet {name}", "", "", f"sheet only in the {side} file"))
                    continue
                prefix = f"sheet {name}, " if name else ""
                res = compare_grids(sheets_o[name], sheets_n[name], tol, prefix)
                diffs += res["diffs"]
                cells += res["cells"]
                worst = max(worst, res["worst"])
            if not diffs:
                note = (f"All {cells} cells match; largest number difference {fmt_num(worst)}."
                        if worst else f"All {cells} cells match (formatting differs only).")
                return {"result": "Match", "kind": kind, "details": note, "diffs": []}
            found = f"{len(diffs)} difference{'s' if len(diffs) != 1 else ''}"
            return {"result": "Differs", "kind": kind,
                    "details": f"{found} found ({cells} cells compared).", "diffs": diffs}
        if ext in IMAGE_EXTS:
            detail = compare_images(orig, new)
            result = "Match" if detail.startswith("Pixels identical") else "Differs"
            return {"result": result, "kind": kind, "details": detail, "diffs": []}
        if ext in VECTOR_EXTS:
            if strip_volatile(orig.read_bytes()) == strip_volatile(new.read_bytes()):
                return {"result": "Match", "kind": kind,
                        "details": "Identical apart from embedded dates and internal IDs.", "diffs": []}
            return {"result": "Differs", "kind": kind,
                    "details": f"Content differs (file size {orig.stat().st_size:,} vs {new.stat().st_size:,}). "
                               "Open both and compare by eye.", "diffs": []}
        if ext in TEXT_EXTS:
            lines_o = orig.read_text(encoding="utf-8", errors="replace").splitlines()
            lines_n = new.read_text(encoding="utf-8", errors="replace").splitlines()
            if len(lines_o) == len(lines_n):
                diffs, worst = [], 0.0
                for i, (a, b) in enumerate(zip(lines_o, lines_n), 1):
                    same, d, reason = compare_cell(a, b, tol)
                    worst = max(worst, d or 0.0)
                    if not same:
                        diffs.append((f"line {i}", a.strip(), b.strip(), reason if d is None else fmt_num(d)))
                if not diffs:
                    return {"result": "Match", "kind": kind,
                            "details": f"All {len(lines_o)} lines match; largest number difference {fmt_num(worst)}.",
                            "diffs": []}
                return {"result": "Differs", "kind": kind,
                        "details": f"{len(diffs)} of {len(lines_o)} lines differ.", "diffs": diffs}  # same line count
            udiff = list(difflib.unified_diff(lines_o, lines_n, "original", "regenerated", lineterm="", n=0))
            diffs = [(f"diff line {i}", line, "", "") for i, line in enumerate(udiff[2:max_diffs + 2], 1)]
            return {"result": "Differs", "kind": kind,
                    "details": f"Line count differs ({len(lines_o)} vs {len(lines_n)}).", "diffs": diffs}
    except Exception as exc:  # unreadable file: fall back to a byte comparison
        return {"result": "Differs", "kind": kind,
                "details": f"Bytes differ, and the file could not be read for a closer look ({exc}).", "diffs": []}
    return {"result": "Differs", "kind": kind,
            "details": f"Bytes differ (file size {orig.stat().st_size:,} vs {new.stat().st_size:,}).", "diffs": []}


def match_outputs(originals: dict, regenerated: dict):
    """Pair original and regenerated files: same relative path first, then same unique file name."""
    pairs, used = [], set()
    by_name = {}
    for key, path in regenerated.items():
        by_name.setdefault(path.name.lower(), []).append(key)
    for key, path in originals.items():
        if key in regenerated:
            pairs.append((key, path, regenerated[key]))
            used.add(key)
            continue
        candidates = [k for k in by_name.get(path.name.lower(), []) if k not in used]
        if len(candidates) == 1:
            pairs.append((key, path, regenerated[candidates[0]]))
            used.add(candidates[0])
        else:
            pairs.append((key, path, None))
    new_only = [k for k in regenerated if k not in used]
    return pairs, new_only


# ---------------------------------------------------------------------------
# Running the master script
# ---------------------------------------------------------------------------
def program_version(exe: str) -> str:
    """First line of `exe --version` (Rscript prints it on stderr)."""
    try:
        out = subprocess.run([exe, "--version"], capture_output=True, text=True, timeout=60)
        return (out.stdout or out.stderr).strip().splitlines()[0]
    except (OSError, subprocess.SubprocessError, IndexError):
        return exe


def build_command(master: Path, args):
    # The command runs inside the clean copy, so a relative path such as .venv/bin/python
    # must become absolute here. abspath, not resolve(): resolve() would follow a virtual
    # environment's python symlink out of the environment.
    ext = master.suffix.lower()
    if ext == ".py":
        python = os.path.abspath(shutil.which(args.python) or args.python)
        if not Path(python).exists():
            raise FileNotFoundError(f"Python not found: {args.python}")
        return [python, master.name], program_version(python)
    if ext == ".r":
        rscript = shutil.which(args.rscript)
        if not rscript:
            raise FileNotFoundError(
                f"'{args.rscript}' was not found. Install R (on Ubuntu: sudo apt-get install r-base) or pass "
                "--rscript /path/to/Rscript.")
        rscript = os.path.abspath(rscript)
        return [rscript, master.name], program_version(rscript)
    raise ValueError(f"{master.name}: only Python (.py) and R (.R) master scripts can be run. For Stata, MATLAB or "
                     "SAS, have the author run the package and use --compare-only.")


def run_master(cmd, cwd: Path, log_path: Path, timeout: float, quiet: bool):
    env = dict(os.environ, MPLBACKEND="Agg", PYTHONUNBUFFERED="1", PYTHONDONTWRITEBYTECODE="1")
    posix = os.name == "posix"
    started = time.monotonic()
    with open(log_path, "a", encoding="utf-8") as log:
        proc = subprocess.Popen(cmd, cwd=cwd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, env=env,
                                start_new_session=posix)

        def pump():
            for raw in iter(proc.stdout.readline, b""):
                text = raw.decode("utf-8", errors="replace")
                log.write(text)
                log.flush()
                if not quiet:
                    sys.stdout.write(text)
                    sys.stdout.flush()

        reader = threading.Thread(target=pump, daemon=True)
        reader.start()
        timed_out = False
        try:
            code = proc.wait(timeout=timeout)
        except subprocess.TimeoutExpired:
            timed_out = True
            if posix:
                for sig in (signal.SIGTERM, signal.SIGKILL):
                    try:
                        os.killpg(proc.pid, sig)
                    except ProcessLookupError:
                        break
                    try:
                        proc.wait(timeout=10)
                        break
                    except subprocess.TimeoutExpired:
                        continue
            else:
                proc.kill()
            code = proc.wait()
        reader.join(timeout=10)
    return code, timed_out, time.monotonic() - started


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------
def write_report(path: Path, info: dict, rows: list, warnings: list, max_diffs: int, details_csv: Path | None):
    n_orig = sum(1 for r in rows if r["result"] not in ("New", "Skipped"))
    n_ok = sum(1 for r in rows if r["result"] in ("Identical", "Match"))
    n_diff = sum(1 for r in rows if r["result"] == "Differs")
    n_missing = sum(1 for r in rows if r["result"] == "Not regenerated")
    n_skipped = sum(1 for r in rows if r["result"] == "Skipped")
    skipped = f" ({n_skipped} log file{'s' if n_skipped != 1 else ''} skipped.)" if n_skipped else ""
    if info.get("run_failed"):
        verdict = f"**Result: RUN FAILED.** {info['run_failed']} See run_log.txt."
    elif n_orig == 0:
        verdict = "**Result: NOTHING TO COMPARE.** No original output files were found."
    elif n_ok == n_orig:
        verdict = (f"**Result: REPRODUCED.** All {n_orig} original outputs were regenerated by the code "
                   f"and match.{skipped}")
    else:
        verdict = (f"**Result: DIFFERENCES FOUND.** {n_ok} of {n_orig} outputs match; {n_diff} "
                   f"{'differs' if n_diff == 1 else 'differ'}; {n_missing} "
                   f"{'was' if n_missing == 1 else 'were'} not regenerated.{skipped}")
    lines = [f"# Reproduction check: {info['package_name']}", "", verdict, ""]
    lines += ["## Run", "", "| Item | Value |", "|---|---|"]
    for key, value in info["table"]:
        lines.append(f"| {key} | {md_cell(value)} |")
    lines += ["", "## Outputs", "", "| Output | Type | Result | Details |", "|---|---|---|---|"]
    for r in rows:
        lines.append(f"| {md_cell(r['name'])} | {r['kind']} | {r['result']} | {md_cell(r['details'])} |")
    with_diffs = [r for r in rows if r["diffs"]]
    if with_diffs:
        lines += ["", "## Differences", ""]
        for r in with_diffs:
            lines += [f"### {r['name']}", "", "| Where | Original | Regenerated | Difference |", "|---|---|---|---|"]
            for where, a, b, d in r["diffs"][:max_diffs]:
                lines.append(f"| {md_cell(where)} | {md_cell(a)} | {md_cell(b)} | {md_cell(d)} |")
            if len(r["diffs"]) > max_diffs:
                lines.append(f"| ... | {len(r['diffs']) - max_diffs} more, see comparison_details.csv | | |")
            lines.append("")
    if warnings:
        lines += ["", "## Warnings", ""] + [f"- {w}" for w in warnings]
    lines += ["", "## What to do with this report", ""]
    if n_orig and n_ok == n_orig and not info.get("run_failed"):
        lines += [
            "- Keep this report with the delivery. It shows that the code, run from a clean copy, recreates "
            "every original output on this computer with this software.",
            "- The journal's data editor will run the package on their own computer, so keep the README's "
            "software versions and run time accurate.",
            "",
        ]
    else:
        lines += [
            "- Do not edit code or outputs to make numbers match.",
            "- Send every difference to the author with this report, and let the author decide what to do.",
            "- Small differences can come from software versions, random seeds or rounding. Name the likely cause "
            "if you know it, and say so if you do not.",
            "",
        ]
    if details_csv:
        lines.insert(-1, f"- Every cell-level difference is listed in {details_csv.name}.")
    path.write_text("\n".join(lines), encoding="utf-8")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main(argv=None):
    p = argparse.ArgumentParser(description="Run a replication package in a clean copy and compare outputs.")
    p.add_argument("package", nargs="?", help="Folder that holds the replication package")
    p.add_argument("--master", help="Master script, relative to the package (for example main.py or main.R)")
    p.add_argument("--outputs", nargs="+", default=["output"],
                   help="Output folders, relative to the package (default: output). They are emptied in the "
                        "clean copy before the run, then compared.")
    p.add_argument("--clean", nargs="*", default=[],
                   help="Other folders to empty before the run (for example data/derived). Not compared.")
    p.add_argument("--originals", help="Folder with the original outputs (default: the package's own output folders)")
    p.add_argument("--regenerated", help="With --compare-only: folder with the regenerated outputs")
    p.add_argument("--compare-only", action="store_true", help="Do not run anything; just compare two folders")
    p.add_argument("--timeout", type=float, default=3600, help="Time limit for the run in seconds (default 3600)")
    p.add_argument("--rtol", type=float, default=1e-6, help="Relative tolerance for numbers (default 1e-6)")
    p.add_argument("--atol", type=float, default=1e-9, help="Absolute tolerance for numbers (default 1e-9)")
    p.add_argument("--report-dir", help="Folder for the report (default: <package name>_run_report)")
    p.add_argument("--workdir", help="Where to make the clean copy (default: the system temp folder)")
    p.add_argument("--keep-temp", action="store_true", help="Keep the clean copy after the run")
    p.add_argument("--python", default=sys.executable, help="Python to run .py master scripts with")
    p.add_argument("--rscript", default="Rscript", help="Rscript to run .R master scripts with")
    p.add_argument("--max-diffs", type=int, default=20, help="Differences to show per file in the report")
    p.add_argument("--quiet", action="store_true", help="Do not echo the run's output to the screen")
    args = p.parse_args(argv)
    tol = Tolerance(args.rtol, args.atol)
    started_at = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if args.compare_only:
        if not (args.originals and args.regenerated):
            p.error("--compare-only needs --originals and --regenerated")
        originals_dir, regen_dir = Path(args.originals).resolve(), Path(args.regenerated).resolve()
        for d in (originals_dir, regen_dir):
            if not d.is_dir():
                p.error(f"not a folder: {d}")
        name = regen_dir.name
        report_dir = Path(args.report_dir) if args.report_dir else Path.cwd() / f"{name}_compare_report"
        report_dir.mkdir(parents=True, exist_ok=True)
        originals, regenerated = list_files(originals_dir), list_files(regen_dir)
        info = {"package_name": name, "table": [
            ("Mode", "Compare only (nothing was run)"),
            ("Original outputs", show_path(originals_dir)),
            ("Regenerated outputs", show_path(regen_dir)),
            ("Compared on", started_at),
            ("Tolerance", f"numbers match when {tol.describe()}"),
            ("Tool", f"run_and_compare.py {VERSION}"),
        ]}
        warnings = []
    else:
        if not args.package or not args.master:
            p.error("give the package folder and --master (or use --compare-only)")
        package = Path(args.package).resolve()
        master_rel = Path(args.master)
        if not (package / master_rel).is_file():
            p.error(f"master script not found: {package / master_rel}")
        report_dir = Path(args.report_dir) if args.report_dir else Path.cwd() / f"{package.name}_run_report"
        report_dir.mkdir(parents=True, exist_ok=True)
        if report_dir.resolve() == package or package in report_dir.resolve().parents:
            p.error("put the report outside the package folder")
        log_path = report_dir / "run_log.txt"
        warnings = []

        before = snapshot(package)
        work_root = Path(tempfile.mkdtemp(prefix="repcheck_", dir=args.workdir))
        copy = work_root / package.name
        shutil.copytree(package, copy, ignore=shutil.ignore_patterns(".git", "__pycache__", ".ipynb_checkpoints",
                                                                      ".Rproj.user", ".DS_Store"))
        emptied = []
        for rel in args.outputs + args.clean:
            folder = copy / rel
            if not folder.is_dir():
                emptied.append(f"{rel} (not in the package; the code must create it)")
                continue
            files = [f for f in folder.rglob("*") if f.is_file()]
            for f in files:
                f.unlink()
            emptied.append(f"{rel} ({len(files)} file{'s' if len(files) != 1 else ''} removed)")

        master = copy / master_rel
        run_failed = None
        code, timed_out, seconds, interpreter, cmd = None, False, 0.0, "", []
        log_path.write_text(f"run_and_compare.py {VERSION}\nStarted {started_at}\nPackage: {show_path(package)}\n"
                            f"Clean copy: {copy if args.keep_temp else 'a temporary folder, deleted after the run'}\n",
                            encoding="utf-8")
        try:
            cmd, interpreter = build_command(master, args)
            with open(log_path, "a", encoding="utf-8") as log:
                log.write(f"Command: {Path(cmd[0]).name} {' '.join(cmd[1:])}  (run in the master script's folder)\n"
                          f"Interpreter: {interpreter}\n"
                          + "-" * 70 + "\n")
            code, timed_out, seconds = run_master(cmd, master.parent, log_path, args.timeout, args.quiet)
            with open(log_path, "a", encoding="utf-8") as log:
                log.write("-" * 70 + f"\nExit code: {code}{' (time limit reached)' if timed_out else ''}\n"
                          f"Run time: {seconds:.1f} s\n")
            if timed_out:
                run_failed = f"The run hit the time limit of {args.timeout:g} seconds and was stopped."
            elif code != 0:
                run_failed = f"The master script stopped with exit code {code} after {seconds:.1f} s."
        except (FileNotFoundError, ValueError) as exc:
            run_failed = f"The master script was not run: {exc}"
            with open(log_path, "a", encoding="utf-8") as log:
                log.write(f"NOT RUN: {exc}\n")

        after = snapshot(package)
        changed = sorted(k for k in set(before) | set(after) if before.get(k) != after.get(k))
        if changed:
            warnings.append("The run changed files in the ORIGINAL package folder, outside the clean copy: "
                            + ", ".join(changed[:10]) + (" ..." if len(changed) > 10 else "")
                            + ". The code probably uses absolute paths.")

        # Keep any log files the run wrote, before the clean copy is deleted.
        kept_logs = []
        for rel, path in list_files(copy).items():
            in_logs = any(part.lower() in ("log", "logs") for part in Path(rel).parts[:-1])
            if (path.suffix.lower() in LOG_EXTS or in_logs) and before.get(rel) != (path.stat().st_size, path.stat().st_mtime_ns):
                target = report_dir / "logs_from_run" / rel
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(path, target)
                kept_logs.append(rel)

        if args.originals:
            originals_dir = Path(args.originals).resolve()
            originals = list_files(originals_dir)
            if len(args.outputs) == 1:
                regenerated = list_files(copy / args.outputs[0])
            else:
                regenerated = {f"{rel}/{k}": v for rel in args.outputs for k, v in list_files(copy / rel).items()}
            originals_label = show_path(originals_dir)
        else:
            originals = {f"{rel}/{k}": v for rel in args.outputs for k, v in list_files(package / rel).items()}
            regenerated = {f"{rel}/{k}": v for rel in args.outputs for k, v in list_files(copy / rel).items()}
            originals_label = ", ".join(args.outputs) + " (in the package as received)"

        info = {"package_name": package.name, "run_failed": run_failed, "table": [
            ("Package", show_path(package)),
            ("Master script", str(master_rel)),
            ("Command", f"{' '.join(Path(c).name if i == 0 else c for i, c in enumerate(cmd))}  (run in its own folder)" if cmd else "not run"),
            ("Interpreter", interpreter or "not found"),
            ("Operating system", f"{platform.system()} {platform.release()}"),
            ("Started", started_at),
            ("Run time", f"{seconds:.1f} s" if cmd else "-"),
            ("Exit code", "time limit reached" if timed_out else (str(code) if code is not None else "-")),
            ("Time limit", f"{args.timeout:g} s"),
            ("Clean copy", "fresh temporary folder" + (f" (kept: {copy})" if args.keep_temp else " (deleted afterwards)")),
            ("Emptied before the run", "; ".join(emptied) if emptied else "nothing"),
            ("Original outputs", originals_label),
            ("Tolerance", f"numbers match when {tol.describe()}"),
            ("Logs kept from the run", ", ".join(kept_logs) if kept_logs else "none"),
            ("Tool", f"run_and_compare.py {VERSION}"),
        ]}

    rows = []
    pairs, new_only = match_outputs(originals, regenerated)
    for key, orig, new in pairs:
        if orig.suffix.lower() in LOG_EXTS:
            rows.append({"name": key, "kind": "Log", "result": "Skipped",
                         "details": "Log files change with every run, so they are not compared.", "diffs": []})
            continue
        if new is None:
            rows.append({"name": key, "kind": "-", "result": "Not regenerated",
                         "details": "The run did not create this file.", "diffs": []})
            continue
        res = compare_file(orig, new, tol, args.max_diffs)
        rows.append({"name": key, **res})
    for key in new_only:
        rows.append({"name": key, "kind": "-", "result": "New",
                     "details": "Created by the run; no original to compare with.", "diffs": []})

    details_csv = None
    if any(r["diffs"] for r in rows):
        details_csv = report_dir / "comparison_details.csv"
        with open(details_csv, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["output", "where", "original", "regenerated", "difference"])
            for r in rows:
                for d in r["diffs"]:
                    w.writerow([r["name"], *d])
    elif (report_dir / "comparison_details.csv").exists():
        (report_dir / "comparison_details.csv").unlink()

    report_path = report_dir / "comparison_report.md"
    write_report(report_path, info, rows, warnings, args.max_diffs, details_csv)
    if not args.compare_only and not args.keep_temp:
        shutil.rmtree(work_root, ignore_errors=True)

    n_bad = sum(1 for r in rows if r["result"] in ("Differs", "Not regenerated"))
    print(f"\nReport: {report_path}")
    if info.get("run_failed"):
        print("RUN FAILED: " + info["run_failed"])
        return 2
    summary = ", ".join(f"{sum(1 for r in rows if r['result'] == k)} {k.lower()}"
                        for k in ("Identical", "Match", "Differs", "Not regenerated", "New", "Skipped")
                        if any(r["result"] == k for r in rows))
    print(f"Outputs: {summary}")
    return 1 if n_bad else 0


if __name__ == "__main__":
    sys.exit(main())
