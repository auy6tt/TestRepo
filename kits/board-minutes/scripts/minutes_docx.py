"""Word (.docx) output for the board-minutes kit.

build_minutes.py and make_template.py import this file. You don't run it directly.

How it works: every minutes document is a filled-in template. Without a client
template, the kit builds its own template (the "house" template) in memory and
fills it, so the default output and templates/minutes-template.docx always match.
A client's own .docx works the same way once you type the placeholders into it
(see "Using a client's template" in the README).

Placeholders
  {{name}}   short text, replaced inside the paragraph (keeps the paragraph's formatting)
  {{NAME}}   a block (tables, lists, motions). It must be alone in its own paragraph.
"""

from __future__ import annotations

import copy
import datetime as dt
import re
from pathlib import Path
from typing import Callable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.parts.numbering import NumberingPart
from docx.shared import Emu, Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from minutes_common import (
    UNCLEAR_RE,
    board_name,
    format_date,
    format_time,
    group_label,
    is_confidential,
    item_heading,
    iter_motions,
    locale_of,
    motion_label,
    next_meeting_text,
    parse_date,
    people_list,
    person_text,
    quorum_text,
    result_label,
    roll_call_text,
    auto_item_text,
    smart_quotes,
    status_line,
    status_of,
    vote_counts_text,
    when_where_line,
    word,
)

# --------------------------------------------------------------------------
# House style
# --------------------------------------------------------------------------

BODY_FONT = "Cambria"
UI_FONT = "Calibri"
NAVY = "1F3A5F"
INK = "1F2328"
MUTED = "5A6470"
RULE = "C8D1DB"
BOX_FILL = "EEF3F8"
DRAFT_FILL = "FFF6E8"
DRAFT_ACCENT = "B7791F"
SAMPLE_FILL = "F3F4F6"
ZEBRA = "F6F8FA"
RESULT_COLORS = {"carried": "1A7F37", "failed": "B42318"}
OTHER_RESULT_COLOR = "8A5A00"

S_ORG = "Minutes Org"
S_TITLE = "Minutes Title"
S_SUBTITLE = "Minutes Subtitle"
S_H1 = "Minutes Heading 1"
S_H2 = "Minutes Heading 2"
S_BODY = "Minutes Body"
S_BULLET = "Minutes Bullet"
S_PRESENTER = "Minutes Presenter"
S_BOX_LABEL = "Minutes Box Label"
S_BOX_TEXT = "Minutes Box Text"
S_BOX_DETAIL = "Minutes Box Detail"
S_LABEL = "Minutes Label"
S_VALUE = "Minutes Value"
S_TABLE = "Minutes Table Text"
S_TABLE_HEAD = "Minutes Table Heading"
S_SPACER = "Minutes Spacer"
S_SMALL = "Minutes Small"
S_HEADER = "Minutes Header"
S_SIGN = "Minutes Signature"

KIT_STYLES = [S_ORG, S_TITLE, S_SUBTITLE, S_H1, S_H2, S_BODY, S_BULLET, S_PRESENTER, S_BOX_LABEL,
              S_BOX_TEXT, S_BOX_DETAIL, S_LABEL, S_VALUE, S_TABLE, S_TABLE_HEAD, S_SPACER, S_SMALL,
              S_HEADER, S_SIGN]

SCALAR_NAMES = [
    "organization", "organization_short", "board_name", "meeting_title", "meeting_date",
    "meeting_date_short", "scheduled_time", "called_to_order", "adjourned", "location",
    "meeting_when_where", "presiding", "recording_secretary", "directors_present",
    "directors_absent", "others_present", "members_present", "quorum", "next_meeting", "status",
    "status_line", "footer_note", "prepared_by", "prepared_on", "secretary", "notice",
]
BLOCK_NAMES = ["NOTICE", "MEETING_DETAILS", "ATTENDANCE", "ITEMS", "ITEM", "ACTION_ITEMS",
               "MOTIONS_TABLE", "ATTACHMENTS", "SIGNATURES"]
PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z_]+)(?:\s+([^{}]*?))?\s*\}\}")


class TemplateError(Exception):
    """The template has a problem the user must fix (unknown placeholder, and so on)."""


# --------------------------------------------------------------------------
# Low-level OOXML helpers
# --------------------------------------------------------------------------

PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
             "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
             "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
             "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
               "tblLook", "tblCaption", "tblDescription", "tblPrChange"]
TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight",
              "tblHeader", "tblCellSpacing", "jc", "hidden", "ins", "del", "trPrChange"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge",
              "tcPrChange"]
BORDER_SIDES = ["top", "left", "bottom", "right", "insideH", "insideV"]


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def set_child(parent, tag: str, order: list[str]):
    """Return parent's <w:tag> child, creating it in the position the Word schema requires."""
    existing = parent.find(qn(f"w:{tag}"))
    if existing is not None:
        return existing
    element = OxmlElement(f"w:{tag}")
    rank = order.index(tag)
    for child in parent.iterchildren():
        name = _local(child.tag)
        if name in order and order.index(name) > rank:
            child.addprevious(element)
            return element
    parent.append(element)
    return element


def _border(parent, side: str, val: str = "single", size: int = 4, color: str = RULE, space: int = 0):
    element = parent.find(qn(f"w:{side}"))
    if element is None:
        element = OxmlElement(f"w:{side}")
        rank = BORDER_SIDES.index(side) if side in BORDER_SIDES else 99
        for child in parent.iterchildren():
            if _local(child.tag) in BORDER_SIDES and BORDER_SIDES.index(_local(child.tag)) > rank:
                child.addprevious(element)
                break
        else:
            parent.append(element)
    element.set(qn("w:val"), val)
    if val in ("nil", "none"):
        for attr in ("w:sz", "w:space", "w:color"):
            element.attrib.pop(qn(attr), None)
    else:
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), str(space))
        element.set(qn("w:color"), color)


def paragraph_border(ppr, side: str, size: int = 6, color: str = NAVY, space: int = 4) -> None:
    pbdr = set_child(ppr, "pBdr", PPR_ORDER)
    _border(pbdr, side, "single", size, color, space)


def set_outline_level(ppr, level: int) -> None:
    set_child(ppr, "outlineLvl", PPR_ORDER).set(qn("w:val"), str(level))


def set_char_spacing(rpr, twentieths: int) -> None:
    set_child(rpr, "spacing", RPR_ORDER).set(qn("w:val"), str(twentieths))


def set_lang(rpr, lang: str) -> None:
    set_child(rpr, "lang", RPR_ORDER).set(qn("w:val"), lang)


def table_props(table: Table):
    return table._tbl.tblPr


def set_table_width(table: Table, width: int) -> None:
    tblw = set_child(table_props(table), "tblW", TBLPR_ORDER)
    tblw.set(qn("w:w"), str(width))
    tblw.set(qn("w:type"), "dxa")
    set_child(table_props(table), "tblLayout", TBLPR_ORDER).set(qn("w:type"), "fixed")


def set_table_indent(table: Table, twips: int) -> None:
    """Indent a table so its left border lines up with the text margin (Word 2010 layout rules)."""
    indent = set_child(table_props(table), "tblInd", TBLPR_ORDER)
    indent.set(qn("w:w"), str(twips))
    indent.set(qn("w:type"), "dxa")


def set_column_widths(table: Table, widths: list[int]) -> None:
    for index, width in enumerate(widths):
        table.columns[index].width = Emu(width * 635)
        for row in table.rows:
            row.cells[index].width = Emu(width * 635)


def set_table_borders(table: Table, **sides) -> None:
    """sides: top=(size, color) or None for no border. Missing sides get no border."""
    borders = set_child(table_props(table), "tblBorders", TBLPR_ORDER)
    for side in BORDER_SIDES:
        spec = sides.get(side)
        if spec:
            _border(borders, side, "single", spec[0], spec[1])
        else:
            _border(borders, side, "nil")


def set_cell_margins(table: Table, top: int = 60, left: int = 100, bottom: int = 60, right: int = 100) -> None:
    margins = set_child(table_props(table), "tblCellMar", TBLPR_ORDER)
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = set_child(tcpr, "shd", TCPR_ORDER)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def cell_border(cell, side: str, size: int, color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = set_child(tcpr, "tcBorders", TCPR_ORDER)
    _border(borders, side, "single", size, color)


def row_flags(row, cant_split: bool = True, header: bool = False) -> None:
    trpr = row._tr.get_or_add_trPr()
    if cant_split:
        set_child(trpr, "cantSplit", TRPR_ORDER)
    if header:
        set_child(trpr, "tblHeader", TRPR_ORDER)


def text_width_twips(doc) -> int:
    section = doc.sections[-1]
    return int((section.page_width - section.left_margin - section.right_margin) / 635)


def add_field(paragraph: Paragraph, instruction: str, placeholder: str = "1") -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = placeholder
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def get_numbering_element(doc):
    """The <w:numbering> element, creating the numbering part if the document has none."""
    part = doc.part
    try:
        return part.part_related_by(RT.NUMBERING).element
    except KeyError:
        numbering_part = NumberingPart(PackURI("/word/numbering.xml"), CT.WML_NUMBERING,
                                       parse_xml(f"<w:numbering {nsdecls('w')}/>"), part.package)
        part.relate_to(numbering_part, RT.NUMBERING)
        return numbering_part.element


def ensure_bullet_numbering(doc) -> int:
    """Add (once) a simple bullet list definition and return its numId."""
    numbering = get_numbering_element(doc)
    for abstract in numbering.findall(qn("w:abstractNum")):
        name = abstract.find(qn("w:name"))
        if name is not None and name.get(qn("w:val")) == "MinutesBullet":
            abstract_id = abstract.get(qn("w:abstractNumId"))
            for num in numbering.findall(qn("w:num")):
                ref = num.find(qn("w:abstractNumId"))
                if ref is not None and ref.get(qn("w:val")) == abstract_id:
                    return int(num.get(qn("w:numId")))
    abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abstract_id, num_id = max(abstract_ids + [0]) + 1, max(num_ids + [0]) + 1
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        '<w:multiLevelType w:val="singleLevel"/><w:name w:val="MinutesBullet"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
        '<w:lvlText w:val="•"/><w:lvlJc w:val="left"/>'
        '<w:pPr><w:ind w:left="360" w:hanging="230"/></w:pPr>'
        '<w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Calibri" w:hint="default"/></w:rPr>'
        "</w:lvl></w:abstractNum>"
    )
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/></w:num>')
    cleanup = numbering.find(qn("w:numIdMacAtCleanup"))
    if cleanup is not None:
        cleanup.addprevious(num)
    else:
        numbering.append(num)
    return num_id


# --------------------------------------------------------------------------
# Styles
# --------------------------------------------------------------------------

def _get_or_add_style(doc, name: str, base: str = "Normal"):
    styles = doc.styles
    try:
        return styles[name], False
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        try:
            style.base_style = styles[base]
        except KeyError:
            style.base_style = styles["Normal"]
        style.quick_style = name in (S_H1, S_H2, S_BODY)
        return style, True


def _apply(style, *, font=None, size=None, bold=None, italic=None, color=None, caps=None, spacing=None,
           align=None, before=None, after=None, line=None, line_exact=None, keep_next=None, outline=None,
           left=None, hanging=None, border_bottom=None, next_style=None):
    f = style.font
    if font:
        f.name = font
        rfonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        if rfonts is not None:
            rfonts.set(qn("w:cs"), font)
            rfonts.set(qn("w:eastAsia"), font)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color:
        f.color.rgb = RGBColor.from_string(color)
    if caps is not None:
        f.all_caps = caps
    if spacing is not None:
        set_char_spacing(style.element.get_or_add_rPr(), int(spacing * 20))
    pf = style.paragraph_format
    if align:
        pf.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT,
                        "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if line_exact is not None:
        pf.line_spacing = Pt(line_exact)
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if left is not None:
        pf.left_indent = Inches(left)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)
    if outline is not None:
        set_outline_level(style.element.get_or_add_pPr(), outline)
    if border_bottom:
        paragraph_border(style.element.get_or_add_pPr(), "bottom", border_bottom[0], border_bottom[1], 6)
    if next_style is not None:
        style.next_paragraph_style = next_style


def define_house_styles(doc, locale: str = "en-US") -> None:
    """The kit's own look: Cambria body text, Calibri labels, navy headings."""
    normal = doc.styles["Normal"]
    _apply(normal, font=BODY_FONT, size=11, color=INK, after=6, line=1.12)
    set_lang(normal.element.get_or_add_rPr(), locale)
    specs = {
        S_ORG: dict(font=UI_FONT, size=10.5, bold=True, color=NAVY, caps=True, spacing=1.5, align="center",
                    before=0, after=3),
        S_TITLE: dict(font=BODY_FONT, size=19, color=INK, align="center", before=0, after=3, line=1.0),
        S_SUBTITLE: dict(font=UI_FONT, size=10.5, color=MUTED, align="center", before=0, after=14,
                         border_bottom=(8, NAVY)),
        S_H1: dict(font=UI_FONT, size=12.5, bold=True, color=NAVY, before=15, after=4, keep_next=True, outline=0,
                   line=1.0),
        S_H2: dict(font=UI_FONT, size=11.5, bold=True, color=INK, before=10, after=3, keep_next=True, outline=1,
                   line=1.0),
        S_BODY: dict(after=6),
        S_BULLET: dict(after=3, left=0.25, hanging=0.16),
        S_PRESENTER: dict(font=UI_FONT, size=9.5, italic=True, color=MUTED, after=4, keep_next=True),
        S_BOX_LABEL: dict(font=UI_FONT, size=8.5, bold=True, color=NAVY, caps=True, spacing=1, after=2,
                          keep_next=True, line=1.0),
        S_BOX_TEXT: dict(font=BODY_FONT, size=11, bold=True, color=INK, after=4, keep_next=True, line=1.08),
        S_BOX_DETAIL: dict(font=UI_FONT, size=9.5, color=INK, after=1, line=1.05),
        S_LABEL: dict(font=UI_FONT, size=8.5, bold=True, color=MUTED, caps=True, spacing=0.5, after=0, line=1.0),
        S_VALUE: dict(font=BODY_FONT, size=10.5, color=INK, after=0, line=1.05),
        S_TABLE: dict(font=UI_FONT, size=9, color=INK, after=0, line=1.0),
        S_TABLE_HEAD: dict(font=UI_FONT, size=9, bold=True, color="FFFFFF", after=0, line=1.0),
        S_SPACER: dict(size=4, before=0, after=0, line_exact=6),
        S_SMALL: dict(font=UI_FONT, size=9, color=MUTED, after=3),
        S_HEADER: dict(font=UI_FONT, size=8.5, color=MUTED, after=0, line=1.0),
        S_SIGN: dict(font=UI_FONT, size=10, color=INK, after=0, line=1.0),
    }
    for name in KIT_STYLES:
        base = "Normal" if name not in (S_BULLET,) else S_BODY
        style, _ = _get_or_add_style(doc, name, base)
        _apply(style, **specs[name])


def ensure_kit_styles(doc) -> None:
    """Add any missing kit styles to a client template, using the template's own fonts."""
    def has(name):
        try:
            doc.styles[name]
            return True
        except KeyError:
            return False

    minimal = {
        S_ORG: dict(bold=True, caps=True, align="center", after=3),
        S_TITLE: dict(size=16, align="center", after=3),
        S_SUBTITLE: dict(size=10, align="center", after=12),
        S_BODY: dict(after=6),
        S_BULLET: dict(after=3, left=0.25, hanging=0.16),
        S_PRESENTER: dict(italic=True, size=9.5, after=4, keep_next=True),
        S_BOX_LABEL: dict(bold=True, caps=True, size=8.5, after=2, keep_next=True),
        S_BOX_TEXT: dict(bold=True, after=4, keep_next=True),
        S_BOX_DETAIL: dict(size=9.5, after=1),
        S_LABEL: dict(bold=True, caps=True, size=8.5, after=0),
        S_VALUE: dict(size=10.5, after=0),
        S_TABLE: dict(size=9.5, after=0, line=1.0),
        S_TABLE_HEAD: dict(size=9, bold=True, color="FFFFFF", after=0, line=1.0),
        S_SPACER: dict(size=4, before=0, after=0, line_exact=6),
        S_SMALL: dict(size=9, after=3),
        S_HEADER: dict(size=8.5, after=0),
        S_SIGN: dict(size=10, after=0),
    }
    for name, base, level in ((S_H1, "Heading 1", 0), (S_H2, "Heading 2", 1)):
        if has(name):
            continue
        style, _ = _get_or_add_style(doc, name, base if has(base) else "Normal")
        if has(base):
            _apply(style, keep_next=True, outline=level)
        else:
            _apply(style, bold=True, size=12.5 if level == 0 else 11.5, before=14 if level == 0 else 10,
                   after=4, keep_next=True, outline=level)
    for name, spec in minimal.items():
        if not has(name):
            style, _ = _get_or_add_style(doc, name, S_BODY if name == S_BULLET and has(S_BODY) else "Normal")
            _apply(style, **spec)


# --------------------------------------------------------------------------
# The house template
# --------------------------------------------------------------------------

def _plain_paragraph(container, text: str, style: str):
    paragraph = container.add_paragraph(style=style)
    if text:
        paragraph.add_run(text)
    return paragraph


def build_house_template(locale: str = "en-US", appendices: bool = True):
    """Create the kit's generic minutes template (a python-docx Document)."""
    doc = Document()
    section = doc.sections[0]
    if locale == "en-GB":
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin, section.bottom_margin = Inches(0.85), Inches(0.8)
    section.header_distance = section.footer_distance = Inches(0.4)
    define_house_styles(doc, locale)
    width = text_width_twips(doc)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")  # python-docx's default omits this required attribute

    # Header on pages 2+, footer on every page.
    section.different_first_page_header_footer = True
    header_p = section.header.paragraphs[0]
    header_p.style = doc.styles[S_HEADER]
    header_p.paragraph_format.tab_stops.add_tab_stop(Emu(width * 635), WD_TAB_ALIGNMENT.RIGHT)
    header_p.add_run("{{organization_short}}\tMinutes · {{meeting_date_short}}")
    paragraph_border(header_p._p.get_or_add_pPr(), "bottom", 4, RULE, 4)
    for footer in (section.footer, section.first_page_footer):
        footer_p = footer.paragraphs[0]
        footer_p.style = doc.styles[S_HEADER]
        footer_p.paragraph_format.tab_stops.add_tab_stop(Emu(width * 635), WD_TAB_ALIGNMENT.RIGHT)
        footer_p.add_run("{{footer_note}}\tPage ")
        add_field(footer_p, "PAGE", "1")
        footer_p.add_run(" of ")
        add_field(footer_p, "NUMPAGES", "1")
    section.first_page_header.paragraphs[0].style = doc.styles[S_HEADER]

    _plain_paragraph(doc, "{{organization}}", S_ORG)
    _plain_paragraph(doc, "Minutes of the {{meeting_title}}", S_TITLE)
    _plain_paragraph(doc, "{{meeting_when_where}}", S_SUBTITLE)
    for block in ("NOTICE", "MEETING_DETAILS", "ITEMS", "ATTACHMENTS", "SIGNATURES"):
        _plain_paragraph(doc, "{{" + block + "}}", S_BODY)
    if appendices:
        # "Page break before" (not a page-break character) so a full last page never leaves a blank page.
        _plain_paragraph(doc, "Appendix A: Action Items", S_H1).paragraph_format.page_break_before = True
        _plain_paragraph(doc, "{{ACTION_ITEMS}}", S_BODY)
        _plain_paragraph(doc, "Appendix B: Summary of Motions", S_H1)
        _plain_paragraph(doc, "{{MOTIONS_TABLE}}", S_BODY)

    props = doc.core_properties
    props.title = "Board meeting minutes template"
    props.author = ""
    props.last_modified_by = ""
    props.comments = "Placeholders in {{double braces}} are filled by build_minutes.py. See the kit README."
    props.category = "Meeting minutes"
    props.keywords = "minutes"
    props.revision = 1
    now = dt.datetime.now(dt.timezone.utc).replace(microsecond=0, tzinfo=None)
    props.created = props.modified = now
    return doc


# --------------------------------------------------------------------------
# Rendering helpers (everything is added at the end of the document body,
# then moved into place by the placeholder engine)
# --------------------------------------------------------------------------

class Ctx:
    def __init__(self, doc, m: dict):
        self.doc = doc
        self.m = m
        self.locale = locale_of(m)
        self.width = text_width_twips(doc)
        self.bullet_id = ensure_bullet_numbering(doc)
        self.motion_numbers = {id(motion): number for number, _, motion in iter_motions(m)}


def add_text(paragraph: Paragraph, text: str, *, bold=None, italic=None, color=None, size=None,
             smart: bool = True) -> None:
    """Add text as runs; [UNCLEAR] markers get a yellow highlight so nobody misses them."""
    if not text:
        return
    text = smart_quotes(text) if smart else text
    position = 0
    pieces = []
    for match in UNCLEAR_RE.finditer(text):
        if match.start() > position:
            pieces.append((text[position:match.start()], False))
        pieces.append((match.group(0), True))
        position = match.end()
    if position < len(text):
        pieces.append((text[position:], False))
    for piece, unclear in pieces:
        run = paragraph.add_run(piece)
        if bold is not None or unclear:
            run.bold = True if unclear else bold
        if italic is not None:
            run.italic = italic
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if size:
            run.font.size = Pt(size)
        if unclear:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def para(ctx: Ctx, text: str = "", style: str = S_BODY, container=None, **fmt) -> Paragraph:
    target = container if container is not None else ctx.doc
    paragraph = target.add_paragraph(style=style)
    add_text(paragraph, text, **fmt)
    return paragraph


def bullet(ctx: Ctx, text: str, container=None) -> Paragraph:
    paragraph = para(ctx, "", S_BULLET, container)
    numpr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numpr.get_or_add_ilvl().val = 0
    numpr.get_or_add_numId().val = ctx.bullet_id
    add_text(paragraph, text)
    return paragraph


def spacer(ctx: Ctx) -> Paragraph:
    return ctx.doc.add_paragraph(style=S_SPACER)


def box(ctx: Ctx, fill: str, accent: str):
    """A one-cell shaded table with a coloured left edge, used for motions and notices."""
    table = ctx.doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table, ctx.width)
    set_column_widths(table, [ctx.width])
    set_table_borders(table, left=(24, accent))
    set_cell_margins(table, top=90, left=170, bottom=90, right=150)
    set_table_indent(table, 170)
    shade_cell(table.cell(0, 0), fill)
    row_flags(table.rows[0], cant_split=True)
    return table.cell(0, 0)


def grid_table(ctx: Ctx, headers: list[str], widths: list[int], rows: list[list], zebra: bool = True):
    """A data table: navy header row that repeats on each page, light rules, optional zebra stripes."""
    table = ctx.doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    scale = ctx.width / total if total > ctx.width else 1
    widths = [int(w * scale) for w in widths]
    set_table_width(table, sum(widths))
    set_column_widths(table, widths)
    set_table_borders(table, top=(4, NAVY), bottom=(4, RULE), insideH=(4, RULE))
    set_cell_margins(table, top=36, left=80, bottom=36, right=80)
    set_table_indent(table, 80)
    head = table.rows[0]
    row_flags(head, cant_split=True, header=True)
    for index, label in enumerate(headers):
        cell = head.cells[index]
        shade_cell(cell, NAVY)
        cell.paragraphs[0].style = ctx.doc.styles[S_TABLE_HEAD]
        add_text(cell.paragraphs[0], label)
    for row_index, values in enumerate(rows, start=1):
        row = table.rows[row_index]
        row_flags(row, cant_split=True)
        for col, value in enumerate(values):
            cell = row.cells[col]
            if zebra and row_index % 2 == 0:
                shade_cell(cell, ZEBRA)
            paragraph = cell.paragraphs[0]
            paragraph.style = ctx.doc.styles[S_TABLE]
            if isinstance(value, tuple):  # (text, color, bold)
                add_text(paragraph, value[0], color=value[1], bold=value[2])
            elif isinstance(value, list):  # one line per entry
                for position, line in enumerate(value):
                    target = paragraph if position == 0 else cell.add_paragraph(style=S_TABLE)
                    add_text(target, line)
            else:
                add_text(paragraph, str(value) if value is not None else "")
    return table


def details_table(ctx: Ctx, rows: list[tuple[str, list[str]]]):
    table = ctx.doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    label_width = int(ctx.width * 0.25)
    set_table_width(table, ctx.width)
    set_column_widths(table, [label_width, ctx.width - label_width])
    set_table_borders(table, top=(4, RULE), bottom=(4, RULE), insideH=(4, RULE))
    set_cell_margins(table, top=55, left=0, bottom=55, right=110)
    for index, (label, values) in enumerate(rows):
        row = table.rows[index]
        row_flags(row, cant_split=True)
        label_p = row.cells[0].paragraphs[0]
        label_p.style = ctx.doc.styles[S_LABEL]
        label_p.paragraph_format.space_before = Pt(1.5)
        add_text(label_p, label)
        cell = row.cells[1]
        for position, value in enumerate(values or [""]):
            paragraph = cell.paragraphs[0] if position == 0 else cell.add_paragraph()
            paragraph.style = ctx.doc.styles[S_VALUE]
            add_text(paragraph, value)
    return table


def short_date(value, locale: str) -> str:
    d = parse_date(value)
    if d is None:
        return "" if value is None else str(value)
    month = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"][d.month - 1]
    return f"{d.day} {month} {d.year}" if locale == "en-GB" else f"{month} {d.day}, {d.year}"


def vote_short(vote: dict | None) -> str:
    if not vote:
        return "—"
    numbers = [vote.get(k) for k in ("in_favor", "opposed", "abstained")]
    if numbers[0] is not None and numbers[1] is not None:
        shown = [str(n) for n in numbers if n is not None]
        return "–".join(shown)
    if vote.get("unanimous"):
        return "Unanimous"
    method = vote.get("method")
    return {"voice": "Voice", "roll_call": "Roll call", "show_of_hands": "Hands", "ballot": "Ballot",
            "unanimous_consent": "Consent", "electronic": "Electronic"}.get(method, "—")


# --------------------------------------------------------------------------
# Blocks
# --------------------------------------------------------------------------

def render_notice(ctx: Ctx, _arg=None) -> None:
    m = ctx.m
    notice = (m.get("document") or {}).get("notice")
    if status_of(m) == "draft":
        cell = box(ctx, DRAFT_FILL, DRAFT_ACCENT)
        first = cell.paragraphs[0]
        first.style = ctx.doc.styles[S_BOX_LABEL]
        add_text(first, "Draft", color=DRAFT_ACCENT)
        para(ctx, f"These minutes have not yet been approved by the {board_name(m)}.", S_BOX_DETAIL, cell)
        if notice:
            para(ctx, notice, S_BOX_DETAIL, cell, italic=True)
        spacer(ctx)
    elif notice:
        cell = box(ctx, SAMPLE_FILL, MUTED)
        cell.paragraphs[0].style = ctx.doc.styles[S_BOX_DETAIL]
        add_text(cell.paragraphs[0], notice, italic=True)
        spacer(ctx)


def _attendance_rows(ctx: Ctx) -> list[tuple[str, list[str]]]:
    m, locale = ctx.m, ctx.locale
    attendance = m["attendance"]
    label = group_label(m)
    rows = [(f"{label} present", [person_text(p, locale) for p in attendance["directors_present"]]),
            (f"{label} absent", [person_text(p, locale) for p in attendance.get("directors_absent") or []]
             or ["None"])]
    if attendance.get("others_present"):
        rows.append(("Also present", [person_text(p, locale) for p in attendance["others_present"]]))
    if attendance.get("members_present"):
        rows.append(("Members present", [attendance["members_present"]]))
    rows.append(("Quorum", [quorum_text(m)]))
    return rows


def render_meeting_details(ctx: Ctx, _arg=None) -> None:
    meeting, locale = ctx.m["meeting"], ctx.locale
    rows = [("Presiding", [person_text(meeting["presiding"], locale)])]
    if meeting.get("recording_secretary"):
        rows.append(("Recording secretary", [person_text(meeting["recording_secretary"], locale)]))
    rows.append(("Called to order", [format_time(meeting["called_to_order"], locale)]))
    rows.append(("Adjourned", [format_time(meeting["adjourned"], locale)]))
    details_table(ctx, rows + _attendance_rows(ctx))
    spacer(ctx)


def render_attendance(ctx: Ctx, _arg=None) -> None:
    details_table(ctx, _attendance_rows(ctx))
    spacer(ctx)


def render_motion(ctx: Ctx, motion: dict) -> None:
    locale = ctx.locale
    number = ctx.motion_numbers.get(id(motion), 0)
    cell = box(ctx, BOX_FILL, NAVY)
    label = cell.paragraphs[0]
    label.style = ctx.doc.styles[S_BOX_LABEL]
    add_text(label, motion_label(number, motion))
    para(ctx, motion["text"], S_BOX_TEXT, cell)

    parts = []
    if motion.get("moved_by"):
        mover = motion["moved_by"]
        if motion.get("on_behalf_of"):
            mover += f", on behalf of the {motion['on_behalf_of']}"
        parts.append(("Moved by ", mover))
    elif motion.get("on_behalf_of"):
        parts.append(("Moved by ", f"the {motion['on_behalf_of']}"))
    if motion.get("seconded_by"):
        parts.append(("Seconded by ", motion["seconded_by"]))
    elif motion.get("second_note"):
        parts.append(("", motion["second_note"]))
    if parts:
        line = para(ctx, "", S_BOX_DETAIL, cell)
        for index, (lead, value) in enumerate(parts):
            if index:
                add_text(line, "  ·  ", color=MUTED)
            if lead:
                add_text(line, lead, color=MUTED)
            add_text(line, value)

    for amendment in motion.get("amendments") or []:
        line = para(ctx, "", S_BOX_DETAIL, cell)
        add_text(line, "Amendment: ", bold=True)
        add_text(line, amendment["text"])
        details = []
        if amendment.get("moved_by"):
            details.append(f"moved by {amendment['moved_by']}")
        if amendment.get("seconded_by"):
            details.append(f"seconded by {amendment['seconded_by']}")
        elif amendment.get("second_note"):
            details.append(amendment["second_note"])
        votes = vote_counts_text(amendment.get("vote"), locale)
        if votes:
            details.append(votes[0].lower() + votes[1:])
        details.append(result_label(amendment.get("result")).lower())
        add_text(line, f" ({'; '.join(details)})", color=MUTED)

    result = motion.get("result")
    line = para(ctx, "", S_BOX_DETAIL, cell)
    line.paragraph_format.space_before = Pt(2)
    color = RESULT_COLORS.get(result, OTHER_RESULT_COLOR)
    add_text(line, result_label(result).upper() if result != "[UNCLEAR]" else "[UNCLEAR]", bold=True, color=color)
    votes = vote_counts_text(motion.get("vote"), locale)
    if votes:
        add_text(line, "  —  ", color=MUTED)
        add_text(line, votes)
    vote = motion.get("vote") or {}
    roll = roll_call_text(vote)
    if roll:
        line = para(ctx, "", S_BOX_DETAIL, cell)
        add_text(line, "Roll call: ", color=MUTED)
        add_text(line, roll)
    if vote.get("recused"):
        line = para(ctx, "", S_BOX_DETAIL, cell)
        add_text(line, "Recused: ", color=MUTED)
        add_text(line, ", ".join(vote["recused"]))
    for note in (vote.get("note"), motion.get("notes")):
        if note:
            para(ctx, note, S_BOX_DETAIL, cell, italic=True, color=MUTED)
    spacer(ctx)


def render_item_body(ctx: Ctx, item: dict) -> None:
    kind = item.get("kind", "business")
    if kind == "heading":
        return
    if item.get("presenter"):
        para(ctx, f"Presenter: {item['presenter']}", S_PRESENTER)
    summary = item.get("summary") or []
    if not summary:
        for sentence in auto_item_text(ctx.m, item):
            para(ctx, sentence)
    for entry in summary:
        if entry.startswith("- "):
            bullet(ctx, entry[2:].strip())
        else:
            paragraph = para(ctx, entry)
            if entry.rstrip().endswith(":"):  # a lead-in line stays with the list below it
                paragraph.paragraph_format.keep_with_next = True
    if kind == "executive_session" and item.get("outcome"):
        para(ctx, item["outcome"])
    for motion in item.get("motions") or []:
        render_motion(ctx, motion)


def _item_key(value) -> str:
    return str(value or "").strip().lower().rstrip(".):")


def render_items(ctx: Ctx, skip: set[str] | None = None) -> None:
    skip = skip or set()
    for item in ctx.m["items"]:
        if item.get("number") and _item_key(item["number"]) in skip:
            continue
        style = S_H2 if item.get("level") == 2 else S_H1
        para(ctx, item_heading(item), style)
        render_item_body(ctx, item)


def render_single_item(ctx: Ctx, number: str) -> None:
    for item in ctx.m["items"]:
        if _item_key(item.get("number")) == _item_key(number):
            render_item_body(ctx, item)
            return
    raise TemplateError(f"{{{{ITEM {number}}}}}: no agenda item has number {number!r} in the minutes JSON.")


def render_action_items(ctx: Ctx, _arg=None) -> None:
    actions = ctx.m.get("action_items") or []
    if not actions:
        para(ctx, "No action items were recorded.")
        return
    rows = []
    for index, action in enumerate(actions, start=1):
        rows.append([str(index), action["task"], action["owner"], short_date(action["due"], ctx.locale),
                     action.get("item", "")])
    grid_table(ctx, ["#", "Action", "Owner", "Due", "Item"], [430, 4570, 1830, 1430, 1100], rows)
    spacer(ctx)


def render_motions_table(ctx: Ctx, _arg=None) -> None:
    motions = list(iter_motions(ctx.m))
    if not motions:
        para(ctx, "No motions were made.")
        return
    note = para(ctx, f"Votes show in {word(ctx.locale, 'favor', 'favour')}–opposed–abstained where counts were "
                     "recorded.", S_SMALL)
    note.paragraph_format.keep_with_next = True
    rows = []
    for number, item, motion in motions:
        moved = motion.get("moved_by") or motion.get("on_behalf_of") or "—"
        second = motion.get("seconded_by") or "—"
        result = motion.get("result")
        rows.append([str(number), item.get("number", ""), motion["text"], [moved, second],
                     vote_short(motion.get("vote")),
                     (result_label(result), RESULT_COLORS.get(result, OTHER_RESULT_COLOR), True)])
    grid_table(ctx, ["#", "Item", "Motion", "Moved / seconded", "Vote", "Result"],
               [400, 620, 3820, 2120, 1200, 1200], rows)
    spacer(ctx)


def render_attachments(ctx: Ctx, _arg=None) -> None:
    attachments = ctx.m.get("attachments") or []
    if not attachments:
        return
    para(ctx, "Attachments", S_H1)
    for attachment in attachments:
        bullet(ctx, attachment)


def render_signatures(ctx: Ctx, _arg=None) -> None:
    m, locale = ctx.m, ctx.locale
    signatures = m.get("signatures") or {}
    secretary = signatures.get("secretary") or m["meeting"].get("recording_secretary")
    label = signatures.get("submitted_label", "Respectfully submitted")
    intro = para(ctx, f"{label},", S_SIGN)
    intro.paragraph_format.space_before = Pt(18)
    intro.paragraph_format.keep_with_next = True
    gap = para(ctx, "", S_SIGN)
    gap.paragraph_format.space_before = Pt(20)
    gap.paragraph_format.keep_with_next = True
    table = ctx.doc.add_table(rows=1, cols=3)
    widths = [int(ctx.width * 0.52), int(ctx.width * 0.1), int(ctx.width * 0.3)]
    set_table_width(table, sum(widths))
    set_column_widths(table, widths)
    set_table_borders(table)
    set_cell_margins(table, top=40, left=0, bottom=0, right=0)
    row_flags(table.rows[0], cant_split=True)
    left, right = table.cell(0, 0), table.cell(0, 2)
    cell_border(left, "top", 6, INK)
    cell_border(right, "top", 6, INK)
    left.paragraphs[0].style = ctx.doc.styles[S_SIGN]
    add_text(left.paragraphs[0], person_text(secretary, locale, details=False) if secretary else "Secretary")
    right.paragraphs[0].style = ctx.doc.styles[S_SIGN]
    add_text(right.paragraphs[0], "Date")
    table.cell(0, 1).paragraphs[0].style = ctx.doc.styles[S_SIGN]
    if signatures.get("show_approval_line", True):
        if status_of(m) == "approved" and (m.get("document") or {}).get("approved_on"):
            text = (f"Approved by the {board_name(m)} on "
                    f"{format_date(m['document']['approved_on'], locale, weekday=False)}.")
        else:
            text = f"Approved by the {board_name(m)} on: ______________________________"
        line = para(ctx, text, S_SIGN)
        line.paragraph_format.space_before = Pt(22)


BLOCK_RENDERERS: dict[str, Callable] = {
    "NOTICE": render_notice,
    "MEETING_DETAILS": render_meeting_details,
    "ATTENDANCE": render_attendance,
    "ACTION_ITEMS": render_action_items,
    "MOTIONS_TABLE": render_motions_table,
    "ATTACHMENTS": render_attachments,
    "SIGNATURES": render_signatures,
}


# --------------------------------------------------------------------------
# Placeholder engine
# --------------------------------------------------------------------------

def scalar_values(m: dict) -> dict[str, str]:
    locale = locale_of(m)
    meeting, organization, attendance = m["meeting"], m["organization"], m["attendance"]
    document = m.get("document") or {}
    signatures = m.get("signatures") or {}
    secretary = signatures.get("secretary") or meeting.get("recording_secretary")
    footer = status_line(m) + (". Confidential." if is_confidential(m) else ".")
    return {
        "organization": organization["name"],
        "organization_short": organization.get("short_name") or organization["name"],
        "board_name": board_name(m),
        "meeting_title": meeting["title"],
        "meeting_date": format_date(meeting["date"], locale),
        "meeting_date_short": format_date(meeting["date"], locale, weekday=False),
        "scheduled_time": format_time(meeting.get("scheduled_time"), locale),
        "called_to_order": format_time(meeting.get("called_to_order"), locale),
        "adjourned": format_time(meeting.get("adjourned"), locale),
        "location": meeting.get("location", ""),
        "meeting_when_where": when_where_line(m),
        "presiding": person_text(meeting["presiding"], locale, details=False),
        "recording_secretary": person_text(meeting.get("recording_secretary"), locale, details=False),
        "directors_present": people_list(attendance["directors_present"], locale),
        "directors_absent": people_list(attendance.get("directors_absent"), locale),
        "others_present": people_list(attendance.get("others_present"), locale),
        "members_present": attendance.get("members_present", ""),
        "quorum": quorum_text(m),
        "next_meeting": next_meeting_text(m),
        "status": "DRAFT" if status_of(m) == "draft" else "APPROVED",
        "status_line": status_line(m),
        "footer_note": footer,
        "prepared_by": document.get("prepared_by", ""),
        "prepared_on": format_date(document.get("prepared_on"), locale, weekday=False),
        "secretary": person_text(secretary, locale, details=False) if secretary else "",
        "notice": document.get("notice", ""),
    }


def _story_roots(doc):
    """(root element, parent object) for the body and every header and footer in use."""
    roots = [(doc.element.body, doc._body)]
    for section in doc.sections:
        for part in (section.header, section.first_page_header, section.even_page_header,
                     section.footer, section.first_page_footer, section.even_page_footer):
            if not part.is_linked_to_previous:
                roots.append((part._element, part))
    return roots


def _paragraphs(root, parent) -> list[Paragraph]:
    return [Paragraph(p, parent) for p in root.iter(qn("w:p"))]


def _replace_scalars(paragraph: Paragraph, values: dict, problems: list[str]) -> None:
    runs = paragraph.runs
    if not runs:
        return
    full = "".join(run.text for run in runs)
    if "{{" not in full:
        return
    bounds, position = [], 0
    for run in runs:
        bounds.append((position, position + len(run.text)))
        position += len(run.text)
    for match in reversed(list(PLACEHOLDER_RE.finditer(full))):
        name = match.group(1)
        if name in BLOCK_NAMES:
            problems.append(f"{match.group(0)} must be alone in its own paragraph (line).")
            continue
        key = name if name in values else name.lower()  # {{ORGANIZATION}} works like {{organization}}
        if key not in values:
            problems.append(f"Unknown placeholder {match.group(0)}.")
            continue
        start, end = match.span()
        first = next(i for i, (a, b) in enumerate(bounds) if a <= start < b)
        last = next(i for i, (a, b) in enumerate(bounds) if a < end <= b)
        value = smart_quotes(values[key])
        offset = bounds[first][0]
        if first == last:
            text = runs[first].text
            runs[first].text = text[:start - offset] + value + text[end - offset:]
        else:
            runs[first].text = runs[first].text[:start - offset] + value
            for index in range(first + 1, last):
                runs[index].text = ""
            runs[last].text = runs[last].text[end - bounds[last][0]:]


def _highlight_unclear_runs(paragraph: Paragraph) -> None:
    for run in list(paragraph.runs):
        text = run.text
        if not UNCLEAR_RE.search(text) or run.font.highlight_color == WD_COLOR_INDEX.YELLOW:
            continue
        pieces = [piece for piece in re.split(f"({UNCLEAR_RE.pattern})", text) if piece]
        run.text = pieces[0]
        if UNCLEAR_RE.fullmatch(pieces[0]):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        anchor = run._r
        for piece in pieces[1:]:
            new_r = copy.deepcopy(run._r)
            anchor.addnext(new_r)
            anchor = new_r
            new_run = Run(new_r, paragraph)
            new_run.text = piece
            new_run.font.highlight_color = WD_COLOR_INDEX.YELLOW if UNCLEAR_RE.fullmatch(piece) else None


def _replace_block(doc, anchor: Paragraph, render: Callable[[], None]) -> None:
    body = doc.element.body
    before = set(body.iterchildren())
    render()
    new_elements = [el for el in body.iterchildren() if el not in before]
    for element in new_elements:
        anchor._p.addprevious(element)
    parent = anchor._p.getparent()
    parent.remove(anchor._p)
    if _local(parent.tag) == "tc":
        children = [c for c in parent.iterchildren() if _local(c.tag) in ("p", "tbl")]
        if not children or _local(children[-1].tag) != "p":
            parent.append(OxmlElement("w:p"))


def fill_document(doc, m: dict, house: bool) -> None:
    """Replace every placeholder in `doc` with content from the minutes `m`."""
    if not house:
        ensure_kit_styles(doc)
    ctx = Ctx(doc, m)
    problems: list[str] = []

    # Block placeholders (body only).
    anchors = []
    for paragraph in _paragraphs(doc.element.body, doc._body):
        text = paragraph.text.strip()
        match = PLACEHOLDER_RE.fullmatch(text)
        if match and match.group(1).isupper():
            name, arg = match.group(1), (match.group(2) or "").strip()
            if name not in BLOCK_NAMES:
                if name.lower() in SCALAR_NAMES and not arg:
                    continue  # a short placeholder typed in capitals; filled below
                problems.append(f"Unknown block placeholder {match.group(0)}. "
                                f"Blocks: {', '.join('{{' + b + '}}' for b in BLOCK_NAMES)}.")
                continue
            if name == "ITEM" and not arg:
                problems.append("{{ITEM}} needs an item number, for example {{ITEM 5}}.")
                continue
            anchors.append((paragraph, name, arg))
    for root, parent in _story_roots(doc)[1:]:
        for paragraph in _paragraphs(root, parent):
            for match in PLACEHOLDER_RE.finditer(paragraph.text):
                if match.group(1) in BLOCK_NAMES:
                    problems.append(f"{match.group(0)} is in a header or footer; blocks only work in the main text.")
    if problems:
        raise TemplateError("\n".join(problems))

    placed_items = {_item_key(arg) for _, name, arg in anchors if name == "ITEM"}
    for paragraph, name, arg in anchors:
        if name == "ITEMS":
            renderer = lambda: render_items(ctx, skip=placed_items)  # noqa: E731
        elif name == "ITEM":
            renderer = lambda arg=arg: render_single_item(ctx, arg)  # noqa: E731
        else:
            renderer = lambda name=name: BLOCK_RENDERERS[name](ctx)  # noqa: E731
        _replace_block(doc, paragraph, renderer)

    # Short placeholders everywhere (body, tables, headers, footers).
    values = scalar_values(m)
    for root, parent in _story_roots(doc):
        for paragraph in _paragraphs(root, parent):
            _replace_scalars(paragraph, values, problems)
            _highlight_unclear_runs(paragraph)
    if problems:
        known = ", ".join("{{" + n + "}}" for n in SCALAR_NAMES)
        raise TemplateError("\n".join(sorted(set(problems))) + f"\nShort placeholders you can use: {known}.")


def set_properties(doc, m: dict) -> None:
    document = m.get("document") or {}
    meeting, organization = m["meeting"], m["organization"]
    props = doc.core_properties
    props.title = f"Minutes – {organization.get('short_name') or organization['name']} – {meeting['date']}"
    props.subject = meeting["title"]
    author = document.get("prepared_by") or organization["name"]
    props.author = author
    props.last_modified_by = author
    props.category = "Meeting minutes"
    props.keywords = "minutes, " + ("draft" if status_of(m) == "draft" else "approved")
    props.comments = ""
    props.revision = 1
    now = dt.datetime.now(dt.timezone.utc).replace(microsecond=0, tzinfo=None)
    props.created = props.modified = now


def render_docx(m: dict, out_path: str | Path, template_path: str | Path | None = None,
                appendices: bool = True) -> Path:
    """Write the minutes as a .docx. Uses the house template unless template_path is given."""
    out_path = Path(out_path)
    if template_path:
        template_path = Path(template_path)
        if template_path.suffix.lower() != ".docx":
            raise TemplateError("The template must be a .docx file. Open it in Word or LibreOffice and "
                                "save it as a Word document (.docx) first.")
        if not template_path.exists():
            raise TemplateError(f"Template not found: {template_path}")
        try:
            doc = Document(str(template_path))
        except Exception as exc:  # python-docx raises several error types for damaged files
            raise TemplateError(f"Could not open the template {template_path.name}: {exc}") from None
        house = False
    else:
        doc = build_house_template(locale_of(m), appendices=appendices)
        house = True
    fill_document(doc, m, house=house)
    set_properties(doc, m)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
