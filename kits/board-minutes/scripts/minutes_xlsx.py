"""Excel (.xlsx) output for the board-minutes kit.

build_minutes.py imports this file. You don't run it directly.

Sheets: Motions, Action Items, Questions (with an empty Answer column for the
secretary) and Meeting (key facts and counts).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from minutes_common import (
    UNCLEAR_RE,
    METHOD_LABELS,
    find_unclear,
    format_date,
    format_time,
    group_label,
    iter_motions,
    locale_of,
    motion_label,
    next_meeting_text,
    parse_date,
    people_list,
    person_text,
    quorum_text,
    result_label,
    roll_call_text,
    status_of,
    word,
)

NAVY = "1F3A5F"
INK = "1F2328"
MUTED = "5A6470"
RULE = "D0D7DE"
ZEBRA = "F6F8FA"
UNCLEAR_FILL = "FFF2CC"
RESULT_FILLS = {"carried": ("E6F4EA", "1A7F37"), "failed": ("FDECEA", "B42318")}
OTHER_RESULT = ("FFF4E5", "8A5A00")

HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BODY_FONT = Font(name="Calibri", size=10, color=INK)
THIN = Side(style="thin", color=RULE)
HEADER_ROW = 4


def _title(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Calibri", size=10, color=MUTED)
    ws.row_dimensions[1].height = 20


def _table(ws, columns: list[tuple[str, float]], rows: list[list], date_format: str,
           center: tuple[int, ...] = ()) -> None:
    for index, (label, width) in enumerate(columns, start=1):
        cell = ws.cell(row=HEADER_ROW, column=index, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[HEADER_ROW].height = 30
    for offset, values in enumerate(rows, start=1):
        row_number = HEADER_ROW + offset
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=row_number, column=index, value=value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top",
                                       horizontal="center" if index in center else None)
            cell.border = Border(bottom=THIN)
            if offset % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=ZEBRA)
            if hasattr(value, "year"):
                cell.number_format = date_format
                cell.alignment = Alignment(horizontal="left", vertical="top")
            if isinstance(value, str) and UNCLEAR_RE.search(value):
                cell.fill = PatternFill("solid", fgColor=UNCLEAR_FILL)
                cell.font = Font(name="Calibri", size=10, bold=True, color=INK)
    last_col = get_column_letter(len(columns))
    last_row = HEADER_ROW + max(len(rows), 1)
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{last_row}"
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.oddFooter.left.text = "&8&A"
    ws.oddFooter.right.text = "&8Page &P of &N"


def _first_upper(text: str) -> str:
    return text[:1].upper() + text[1:]


def _due_value(value):
    parsed = parse_date(value)
    return parsed if parsed else value


def build_workbook(m: dict) -> Workbook:
    locale = locale_of(m)
    organization = m["organization"]
    meeting = m["meeting"]
    org_name = organization.get("short_name") or organization["name"]
    subtitle = f"{meeting['title']} · {format_date(meeting['date'], locale)}"
    if status_of(m) == "draft":
        subtitle += " · DRAFT, not yet approved"
    date_format = "d mmm yyyy" if locale == "en-GB" else "mmm d, yyyy"

    wb = Workbook()

    # Motions -----------------------------------------------------------
    ws = wb.active
    ws.title = "Motions"
    _title(ws, f"{org_name}: motions and votes", subtitle)
    rows = []
    for number, item, motion in iter_motions(m):
        vote = motion.get("vote") or {}
        details = []
        if motion.get("on_behalf_of"):
            details.append(f"On behalf of the {motion['on_behalf_of']}.")
        if motion.get("second_note"):
            details.append(motion["second_note"])
        if vote.get("unanimous"):
            details.append("Unanimous.")
        if vote.get("roll_call"):
            details.append("Roll call: " + roll_call_text(vote))
        if vote.get("recused"):
            details.append("Recused: " + ", ".join(vote["recused"]))
        for amendment in motion.get("amendments") or []:
            details.append(f"Amendment ({result_label(amendment.get('result')).lower()}): {amendment['text']}")
        for note in (vote.get("note"), motion.get("notes")):
            if note:
                details.append(note)
        rows.append([
            motion_label(number, motion), item.get("number", ""), motion["text"],
            motion.get("moved_by") or "", motion.get("seconded_by") or "",
            _first_upper(METHOD_LABELS.get(vote.get("method"), vote.get("method") or "")),
            vote.get("in_favor"), vote.get("opposed"), vote.get("abstained"),
            result_label(motion.get("result")), " ".join(details), motion.get("source", ""),
        ])
    columns = [("Motion", 11), ("Item", 7), ("Motion text", 50), ("Moved by", 18), ("Seconded by", 20),
               ("Vote method", 14), (f"In {word(locale, 'favor', 'favour')}", 9), ("Opposed", 9),
               ("Abstained", 10), ("Result", 12), ("Details", 42), ("Recording time", 11)]
    _table(ws, columns, rows, date_format, center=(2, 7, 8, 9))
    for offset, (_, _, motion) in enumerate(iter_motions(m), start=1):
        cell = ws.cell(row=HEADER_ROW + offset, column=10)
        fill, color = RESULT_FILLS.get(motion.get("result"), OTHER_RESULT)
        if motion.get("result") == "[UNCLEAR]":
            fill, color = UNCLEAR_FILL, INK
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Calibri", size=10, bold=True, color=color)
    if not rows:
        ws.cell(row=HEADER_ROW + 1, column=1, value="No motions were made.").font = BODY_FONT

    # Action items ------------------------------------------------------
    ws = wb.create_sheet("Action Items")
    _title(ws, f"{org_name}: action items", subtitle)
    actions = m.get("action_items") or []
    rows = [[index, action["task"], action["owner"], _due_value(action["due"]), action.get("item", ""), "Open",
             action.get("source", "")] for index, action in enumerate(actions, start=1)]
    _table(ws, [("#", 5), ("Action", 60), ("Owner", 22), ("Due", 15), ("Item", 7), ("Status", 12),
                ("Recording time", 11)], rows, date_format, center=(1, 5))
    if rows:
        status = DataValidation(type="list", formula1='"Open,In progress,Done"', allow_blank=True)
        ws.add_data_validation(status)
        status.add(f"F{HEADER_ROW + 1}:F{HEADER_ROW + len(rows)}")
    else:
        ws.cell(row=HEADER_ROW + 1, column=2, value="No action items were recorded.").font = BODY_FONT

    # Questions ---------------------------------------------------------
    ws = wb.create_sheet("Questions")
    _title(ws, f"{org_name}: questions for the secretary", subtitle)
    questions = m.get("questions") or []
    rows = [[index, q["question"], q.get("where", ""), q.get("source", ""), ""]
            for index, q in enumerate(questions, start=1)]
    _table(ws, [("#", 5), ("Question", 70), ("Where in the minutes", 26), ("Recording time", 11),
                ("Answer", 40)], rows, date_format, center=(1,))
    if not rows:
        ws.cell(row=HEADER_ROW + 1, column=2, value="No open questions.").font = BODY_FONT

    # Meeting facts -----------------------------------------------------
    ws = wb.create_sheet("Meeting")
    _title(ws, f"{org_name}: meeting facts", subtitle)
    attendance = m["attendance"]
    label = group_label(m)
    motions = list(iter_motions(m))
    facts = [
        ["Organization", organization["name"]],
        ["Meeting", meeting["title"]],
        ["Date", format_date(meeting["date"], locale)],
        ["Scheduled time", format_time(meeting.get("scheduled_time"), locale)],
        ["Called to order", format_time(meeting.get("called_to_order"), locale)],
        ["Adjourned", format_time(meeting.get("adjourned"), locale)],
        ["Location", meeting.get("location", "")],
        ["Presiding", person_text(meeting.get("presiding"), locale)],
        ["Recording secretary", person_text(meeting.get("recording_secretary"), locale)],
        [f"{label} present", people_list(attendance["directors_present"], locale)],
        [f"{label} absent", people_list(attendance.get("directors_absent"), locale)],
        ["Also present", people_list(attendance.get("others_present"), locale)],
        ["Members present", attendance.get("members_present", "")],
        ["Quorum", quorum_text(m)],
        ["Next meeting", next_meeting_text(m)],
        ["Status", "Draft (not yet approved)" if status_of(m) == "draft" else "Approved"],
        ["Motions", len(motions)],
        ["Motions carried", sum(1 for _, _, mo in motions if mo.get("result") == "carried")],
        ["Action items", len(actions)],
        ["Open questions", len(questions)],
        ["[UNCLEAR] markers", len(find_unclear(m))],
    ]
    _table(ws, [("Field", 24), ("Value", 90)], facts, date_format)
    ws.auto_filter.ref = None
    for row in range(HEADER_ROW + 1, HEADER_ROW + len(facts) + 1):
        ws.cell(row=row, column=1).font = Font(name="Calibri", size=10, bold=True, color=INK)
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

    document = m.get("document") or {}
    wb.properties.creator = document.get("prepared_by") or organization["name"]
    wb.properties.lastModifiedBy = wb.properties.creator
    wb.properties.title = f"Motions – {org_name} – {meeting['date']}"
    return wb


def render_xlsx(m: dict, out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(m).save(str(out_path))
    return out_path
