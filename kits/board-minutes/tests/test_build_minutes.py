"""Tests for build_minutes.py, make_template.py and the Word/Excel modules."""

import datetime as dt
import shutil

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.text.run import Run
from openpyxl import load_workbook

import build_minutes
import make_template
from conftest import SAMPLE_JSON, SAMPLES_DIR, TEMPLATES_DIR
from minutes_docx import TemplateError, render_docx


def all_text(path) -> str:
    """Every paragraph in the body (tables included), headers and footers, one per line."""
    doc = Document(str(path))
    roots = [doc.element.body]
    for section in doc.sections:
        for part in (section.header, section.first_page_header, section.footer, section.first_page_footer):
            if not part.is_linked_to_previous:
                roots.append(part._element)
    lines = []
    for root in roots:
        for p in root.iter(qn("w:p")):
            lines.append("".join(t.text or "" for t in p.iter(qn("w:t"))))
    return "\n".join(lines)


def body_lines(path) -> list[str]:
    doc = Document(str(path))
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p"))]


def test_sample_builds_docx_and_xlsx(tmp_path, sample):
    code = build_minutes.main([str(SAMPLE_JSON), "--out-dir", str(tmp_path)])
    assert code == 0
    docx_path = tmp_path / "wrenfield-commons-2026-09-16-minutes.docx"
    xlsx_path = tmp_path / "wrenfield-commons-2026-09-16-motions.xlsx"
    assert docx_path.exists() and xlsx_path.exists()
    text = all_text(docx_path)
    assert "{{" not in text and "}}" not in text
    assert "Wrenfield Commons Homeowners Association" in text
    assert "Minutes of the Regular Meeting of the Board of Directors" in text
    assert "Wednesday, September 16, 2026 · 7:00 p.m. · By Zoom videoconference" in text
    assert "Draft minutes, subject to approval by the Board of Directors. Confidential." in text
    for item in sample["items"]:
        for motion in item.get("motions", []):
            assert motion["text"].replace("'", "’") in text
    assert "Roll call: Diane Okafor: yes; Raymond Castillo: yes; Grace Lindqvist: no" in text
    assert "Appendix A: Action Items" in text and "Appendix B: Summary of Motions" in text
    assert "account 1147" in text and "4,100" not in text  # executive session detail stays out


def test_unclear_markers_are_highlighted(tmp_path):
    out = tmp_path / "m.docx"
    render_docx(build_minutes.load_minutes(SAMPLE_JSON), out)
    doc = Document(str(out))
    highlighted = [r.text for p in doc.element.body.iter(qn("w:p"))
                   for r in [Run(x, None) for x in p.iter(qn("w:r"))]
                   if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert "[UNCLEAR: $5,600 or $5,650]" in highlighted
    assert "[UNCLEAR: Raymond Castillo or Grace Lindqvist]" in highlighted
    assert "[UNCLEAR]" in highlighted


def test_document_properties(tmp_path, sample, write_json):
    sample["document"]["prepared_by"] = "Sam Writer Minutes Service"
    path = write_json(sample)
    assert build_minutes.main([str(path), "--out-dir", str(tmp_path), "--no-xlsx"]) == 0
    props = Document(str(tmp_path / "wrenfield-commons-hoa-2026-09-16-minutes.docx")).core_properties
    assert props.author == "Sam Writer Minutes Service"
    assert props.title == "Minutes – Wrenfield Commons HOA – 2026-09-16"
    assert "python-docx" not in (props.author + props.last_modified_by + props.comments)


def test_output_names(tmp_path, write_json, sample):
    path = write_json(sample, "acme-2026-10-01.json")
    assert build_minutes.main([str(path), "--out-dir", str(tmp_path / "out"), "--name", "acme"]) == 0
    assert (tmp_path / "out" / "acme-minutes.docx").exists()
    assert (tmp_path / "out" / "acme-motions.xlsx").exists()
    assert build_minutes.output_base(path, None) == "acme-2026-10-01"
    assert build_minutes.output_base(path.with_name("minutes.json"), None, sample) == "wrenfield-commons-hoa-2026-09-16"
    assert build_minutes.output_base(path.with_name("acme-minutes.json"), None, sample) == "acme"


def test_committed_sample_docx_matches_a_fresh_build(tmp_path):
    out = tmp_path / "fresh.docx"
    render_docx(build_minutes.load_minutes(SAMPLE_JSON), out)
    assert body_lines(out) == body_lines(SAMPLES_DIR / "wrenfield-commons-2026-09-16-minutes.docx")


def test_xlsx_contents(tmp_path):
    assert build_minutes.main([str(SAMPLE_JSON), "--out-dir", str(tmp_path)]) == 0
    wb = load_workbook(tmp_path / "wrenfield-commons-2026-09-16-motions.xlsx")
    assert wb.sheetnames == ["Motions", "Action Items", "Questions", "Meeting"]
    motions = wb["Motions"]
    assert motions["A4"].value == "Motion" and motions.freeze_panes == "A5"
    rows = list(motions.iter_rows(min_row=5, values_only=True))
    assert len(rows) == 8
    assert rows[3][:10] == ("Motion 4", "9a", "To renew the contract with Pinecrest Grounds Maintenance for two "
                            "years at $5,150 per month, starting November 1.", "Raymond Castillo", "Helen Whitaker",
                            "Roll call vote", 4, 1, None, "Carried")
    actions = wb["Action Items"]
    assert actions.max_row == 4 + 9
    assert actions["D5"].value == dt.datetime(2026, 10, 21)
    assert actions["D7"].value == "Budget workshop"
    assert actions["C6"].fill.fgColor.rgb.endswith("FFF2CC")  # [UNCLEAR] owner is highlighted
    questions = wb["Questions"]
    assert questions.max_row == 4 + 4 and questions["E4"].value == "Answer"


def test_schema_errors_stop_the_build(tmp_path, sample, write_json, capsys):
    del sample["meeting"]["date"]
    sample["items"][4]["motions"][0]["seconder"] = "Raymond Castillo"
    sample["meeting"]["called_to_order"] = "7:02 pm"
    path = write_json(sample)
    assert build_minutes.main([str(path), "--out-dir", str(tmp_path)]) == 1
    err = capsys.readouterr().err
    assert "'date' is a required property" in err
    assert "$.items[4].motions[0]: unknown field(s) 'seconder'" in err
    assert "$.meeting.called_to_order" in err and "24-hour HH:MM" in err
    assert not list(tmp_path.glob("*.docx"))


def test_bad_json_file(tmp_path, capsys):
    path = tmp_path / "broken.json"
    path.write_text('{"schema_version": "1.0",', encoding="utf-8")
    assert build_minutes.main([str(path)]) == 1
    assert "not valid JSON" in capsys.readouterr().err


def test_generic_template_gives_the_same_document(tmp_path):
    default = tmp_path / "default.docx"
    templated = tmp_path / "templated.docx"
    minutes = build_minutes.load_minutes(SAMPLE_JSON)
    render_docx(minutes, default)
    render_docx(minutes, templated, template_path=TEMPLATES_DIR / "minutes-template.docx")
    assert body_lines(default) == body_lines(templated)
    assert all_text(default) == all_text(templated)


def make_client_template(path, drop_numbering=False):
    doc = Document()
    if drop_numbering:
        rid = next(rel.rId for rel in doc.part.rels.values() if rel.reltype == RT.NUMBERING)
        doc.part.drop_rel(rid)
    section = doc.sections[0]
    section.header.paragraphs[0].text = "{{organization}} | Board minutes"
    section.footer.paragraphs[0].text = "{{status_line}}"
    split = doc.add_paragraph()
    split.add_run("Association: {{organ")
    split.add_run("ization}}").bold = True
    doc.add_paragraph("Date: {{meeting_date}}, called to order at {{called_to_order}}")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Directors present"
    table.cell(0, 1).text = "{{directors_present}}"
    doc.add_paragraph("{{ORGANIZATION}}")
    doc.add_heading("APPROVAL OF MINUTES", level=1)
    doc.add_paragraph("{{ITEM 5}}")
    doc.add_heading("OTHER BUSINESS", level=1)
    doc.add_paragraph("{{ITEMS}}")
    doc.add_paragraph("{{ACTION_ITEMS}}")
    doc.save(str(path))


@pytest.mark.parametrize("drop_numbering", [False, True])
def test_client_template(tmp_path, drop_numbering):
    template = tmp_path / "client.docx"
    make_client_template(template, drop_numbering)
    out = tmp_path / "out.docx"
    render_docx(build_minutes.load_minutes(SAMPLE_JSON), out, template_path=template)
    text = all_text(out)
    assert "{{" not in text
    assert "Wrenfield Commons Homeowners Association | Board minutes" in text
    assert "Association: Wrenfield Commons Homeowners Association" in text
    assert "Date: Wednesday, September 16, 2026, called to order at 7:02 p.m." in text
    assert "Diane Okafor, President; Raymond Castillo, Vice President;" in text
    assert "Draft minutes, subject to approval by the Board of Directors" in text
    lines = body_lines(out)
    # Item 5 sits under the client's own heading, and {{ITEMS}} does not repeat it.
    assert lines.index("APPROVAL OF MINUTES") < lines.index("To approve the minutes of the August 19, 2026 "
                                                            "regular meeting as corrected.")
    assert "5. Approval of Minutes: August 19, 2026 Regular Meeting" not in lines
    assert "6. Treasurer’s Report: August 2026 Financial Statements" in lines
    doc = Document(str(out))
    assert "Minutes Body" in [s.name for s in doc.styles]
    bullets = [p for p in doc.paragraphs if p.style.name == "Minutes Bullet"]
    assert bullets and all(p._p.pPr.numPr is not None for p in bullets)


def test_template_errors(tmp_path):
    minutes = build_minutes.load_minutes(SAMPLE_JSON)

    def build_with(paragraphs, header=None):
        doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
        if header:
            doc.sections[0].header.paragraphs[0].text = header
        path = tmp_path / "t.docx"
        doc.save(str(path))
        render_docx(minutes, tmp_path / "o.docx", template_path=path)

    with pytest.raises(TemplateError, match="Unknown placeholder"):
        build_with(["Date: {{meeting_dat}}"])
    with pytest.raises(TemplateError, match="must be alone in its own paragraph"):
        build_with(["See {{ITEMS}} below"])
    with pytest.raises(TemplateError, match="Unknown block placeholder"):
        build_with(["{{MOTIONS}}"])
    with pytest.raises(TemplateError, match="header or footer"):
        build_with(["{{ITEMS}}"], header="{{ATTENDANCE}}")
    with pytest.raises(TemplateError, match="no agenda item has number"):
        build_with(["{{ITEM 99}}"])
    with pytest.raises(TemplateError, match="must be a .docx"):
        render_docx(minutes, tmp_path / "o.docx", template_path=tmp_path / "missing.doc")


def test_template_error_exit_code(tmp_path, capsys):
    doc = Document()
    doc.add_paragraph("{{nonsense}}")
    template = tmp_path / "bad.docx"
    doc.save(str(template))
    assert build_minutes.main([str(SAMPLE_JSON), "--out-dir", str(tmp_path), "--template", str(template)]) == 2
    assert "Unknown placeholder {{nonsense}}" in capsys.readouterr().err


def test_uk_locale(tmp_path, sample, write_json):
    sample["document"]["locale"] = "en-GB"
    path = write_json(sample)
    assert build_minutes.main([str(path), "--out-dir", str(tmp_path)]) == 0
    out = tmp_path / "wrenfield-commons-hoa-2026-09-16-minutes.docx"
    doc = Document(str(out))
    assert round(doc.sections[0].page_height.inches, 2) == 11.69  # A4
    text = all_text(out)
    assert "Wednesday 16 September 2026 · 7.00pm" in text
    assert "called the meeting to order at 7:02 p.m." in text  # the writer's own words are not changed
    assert "4 in favour, 0 opposed" in text


def test_approved_minutes_have_no_draft_notice(tmp_path, sample, write_json):
    sample["document"].update({"status": "approved", "approved_on": "2026-10-21"})
    sample["document"].pop("notice")
    path = write_json(sample)
    assert build_minutes.main([str(path), "--out-dir", str(tmp_path), "--no-appendix"]) == 0
    text = all_text(tmp_path / "wrenfield-commons-hoa-2026-09-16-minutes.docx")
    assert "not yet been approved" not in text
    assert "Approved by the Board of Directors on October 21, 2026." in text
    assert "Confidential" not in text
    assert "Appendix A" not in text


def test_minimal_minutes(tmp_path, write_json):
    minimal = {
        "schema_version": "1.0",
        "organization": {"name": "St. Brigid Parish Council"},
        "meeting": {"title": "Meeting of the Parish Council", "date": "2026-05-04", "called_to_order": "18:30",
                    "adjourned": "[UNCLEAR]", "presiding": {"name": "Ann Lee"}},
        "attendance": {"directors_present": [{"name": "Ann Lee"}], "directors_absent": [],
                       "quorum": {"met": None}, "group_label": "Council members"},
        "items": [{"title": "Adjournment", "kind": "adjournment"}],
    }
    path = write_json(minimal)
    assert build_minutes.main([str(path), "--out-dir", str(tmp_path)]) == 0
    text = all_text(tmp_path / "st-brigid-parish-council-2026-05-04-minutes.docx")
    assert "Council members present" in text
    assert "The meeting was adjourned at [UNCLEAR]." in text
    assert "No motions were made." in text and "No action items were recorded." in text


def test_make_template(tmp_path):
    out = tmp_path / "t.docx"
    assert make_template.main(["-o", str(out)]) == 0
    text = all_text(out)
    for placeholder in ("{{organization}}", "{{meeting_title}}", "{{NOTICE}}", "{{MEETING_DETAILS}}", "{{ITEMS}}",
                        "{{SIGNATURES}}", "{{ACTION_ITEMS}}", "{{MOTIONS_TABLE}}", "{{footer_note}}"):
        assert placeholder in text
    a4 = tmp_path / "a4.docx"
    assert make_template.main(["-o", str(a4), "--locale", "en-GB", "--no-appendix"]) == 0
    assert round(Document(str(a4)).sections[0].page_width.inches, 2) == 8.27
    assert "{{ACTION_ITEMS}}" not in all_text(a4)


def test_committed_template_is_current(tmp_path):
    out = tmp_path / "t.docx"
    make_template.main(["-o", str(out)])
    assert body_lines(out) == body_lines(TEMPLATES_DIR / "minutes-template.docx")


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed")
def test_pdf_option(tmp_path):
    assert build_minutes.main([str(SAMPLE_JSON), "--out-dir", str(tmp_path), "--pdf", "--no-xlsx"]) == 0
    pdf = tmp_path / "wrenfield-commons-2026-09-16-minutes.pdf"
    assert pdf.exists() and pdf.stat().st_size > 20_000
    assert pdf.read_bytes()[:5] == b"%PDF-"


def test_skeleton_is_valid_json_but_must_be_filled_in(capsys):
    skeleton = TEMPLATES_DIR / "minutes-skeleton.json"
    data = build_minutes.load_minutes(skeleton)
    errors = build_minutes.schema_errors(data)
    assert errors and all("YYYY-MM-DD" in e or "HH:MM" in e for e in errors)
    assert build_minutes.main([str(skeleton), "--out-dir", "/nonexistent-folder"]) == 1
    assert "does not match the minutes format" in capsys.readouterr().err
