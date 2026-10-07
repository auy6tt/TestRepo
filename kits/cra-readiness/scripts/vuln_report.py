#!/usr/bin/env python3
"""Check the components in a CycloneDX SBOM for known vulnerabilities.

  * Python packages (purl pkg:pypi/...) are checked with pip-audit.
  * Node packages (purl pkg:npm/...) are checked with `npm audit --json` in the
    project folders you name with --node-project. Any npm packages not covered
    that way are looked up in the same npm advisory database through its bulk
    advisory service, so a stored SBOM can be re-checked without the code.
  * Other ecosystems (Go, Rust, Java, ...) are checked with the OSV API, but only
    if api.osv.dev can be reached.
  * Hand-listed components (firmware libraries, SDKs) are listed for a manual check.

Extra details (severity, CVE numbers, fixed versions) come from OSV records
(API, or OSV's public data files when the API is blocked). Every CVE is also
checked against CISA's catalogue of known exploited vulnerabilities (KEV).

Output: a client-facing report (.md, .docx, .xlsx), findings.json for next
month's comparison, and the raw tool output for your records.

Example
  python vuln_report.py deliverables/sbom/sprout-s1-2.3.0.cdx.json \\
      --node-project ../client/companion-app --client client.json --out deliverables
  python vuln_report.py sbom.cdx.json --previous last-month/vulnerability-data/findings.json --out this-month
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import urllib.parse
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import kitlib as K  # noqa: E402

PART_PROP = f"{K.KIT_NAME}:part"
SOURCE_PROP = f"{K.KIT_NAME}:source"
OSV_API = "https://api.osv.dev/v1"
OSV_BUCKET = "https://osv-vulnerabilities.storage.googleapis.com"
NPM_BULK = "https://registry.npmjs.org/-/npm/v1/security/advisories/bulk"
KEV_URLS = [
    "https://www.cisa.gov/sites/default/files/feeds/known_exploited_vulnerabilities.json",
    "https://raw.githubusercontent.com/cisagov/kev-data/main/known_exploited_vulnerabilities.json",
]
OSV_ECOSYSTEMS = {"pypi": "PyPI", "npm": "npm", "maven": "Maven", "golang": "Go", "cargo": "crates.io",
                  "nuget": "NuGet", "gem": "RubyGems", "composer": "Packagist", "pub": "Pub", "hex": "Hex",
                  "swift": "SwiftURL", "cocoapods": "CocoaPods"}


# ---------------------------------------------------------------------------
# Reading the SBOM
# ---------------------------------------------------------------------------

def parse_purl(purl: str) -> dict | None:
    m = re.match(r"^pkg:([a-zA-Z0-9.+-]+)/(.+)$", purl or "")
    if not m:
        return None
    ptype, rest = m.group(1).lower(), m.group(2)
    rest = rest.split("#", 1)[0].split("?", 1)[0]
    version = None
    if "@" in rest:
        rest, version = rest.rsplit("@", 1)
        version = urllib.parse.unquote(version)
    pieces = [urllib.parse.unquote(p) for p in rest.split("/") if p]
    name = pieces[-1] if pieces else ""
    namespace = "/".join(pieces[:-1])
    full = f"{namespace}/{name}" if namespace else name
    return {"type": ptype, "namespace": namespace, "name": name, "full_name": full, "version": version}


def get_prop(comp: dict, name: str) -> str:
    for prop in comp.get("properties", []) or []:
        if prop.get("name") == name:
            return prop.get("value", "")
    return ""


class Sbom:
    """The parts of an SBOM this report needs: components, parts and who depends on whom."""

    def __init__(self, path: Path):
        self.path = path
        self.data = K.load_json(path)
        if self.data.get("bomFormat") != "CycloneDX":
            K.fail(f"{path} is not a CycloneDX JSON SBOM.")
        self.sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
        self.root = self.data.get("metadata", {}).get("component", {}) or {}
        comps = []
        self._flatten(self.data.get("components", []), comps)
        self.components = {c.get("bom-ref") or f"{c.get('name')}@{c.get('version')}": c for c in comps}
        self.deps = {d.get("ref"): list(d.get("dependsOn", []) or []) for d in self.data.get("dependencies", []) or []}
        root_ref = self.root.get("bom-ref")
        # Parts exist only in product SBOMs merged by make_sbom.py (they carry the part property).
        self.part_roots = [r for r in self.deps.get(root_ref, [])
                           if r in self.components and get_prop(self.components[r], PART_PROP)]
        self.single = not self.part_roots
        if self.single:
            self.part_roots = [root_ref] if root_ref else []
        self.part_name = {}
        for ref in self.part_roots:
            comp = self.components.get(ref) or self.root
            self.part_name[ref] = get_prop(comp, PART_PROP) or comp.get("name", "product")
        self._reach_cache: dict[str, set] = {}

    def _flatten(self, items, out):
        for comp in items or []:
            out.append(comp)
            self._flatten(comp.get("components"), out)

    def product_name(self) -> str:
        return self.root.get("name", "product")

    def product_version(self) -> str:
        return self.root.get("version", "")

    def parts_of(self, ref: str) -> list[str]:
        comp = self.components.get(ref, {})
        value = get_prop(comp, PART_PROP)
        if value:
            return [v.strip() for v in value.split(",") if v.strip()]
        return [self.part_name[r] for r in self.part_roots] if self.single else []

    def _reach(self, ref: str) -> set:
        if ref in self._reach_cache:
            return self._reach_cache[ref]
        seen, stack = set(), [ref]
        while stack:
            cur = stack.pop()
            for nxt in self.deps.get(cur, []):
                if nxt not in seen:
                    seen.add(nxt)
                    stack.append(nxt)
        self._reach_cache[ref] = seen
        return seen

    def direct_and_parents(self, ref: str) -> tuple[bool, list[str]]:
        """Is the component a direct dependency? If not, which direct dependencies bring it in?"""
        direct, parents = False, []
        for root_ref in self.part_roots:
            for d in self.deps.get(root_ref, []):
                if d == ref:
                    direct = True
                elif ref in self._reach(d):
                    comp = self.components.get(d, {})
                    parents.append(display_name(comp))
        return direct, sorted(set(parents))


def display_name(comp: dict) -> str:
    parsed = parse_purl(comp.get("purl", ""))
    if parsed and parsed["type"] == "npm":
        return parsed["full_name"]
    group = comp.get("group")
    return f"{group}/{comp.get('name')}" if group else comp.get("name", "")


# ---------------------------------------------------------------------------
# OSV records (severity, aliases, fixed versions)
# ---------------------------------------------------------------------------

class Osv:
    def __init__(self, offline: bool):
        self.offline = offline
        self.api_ok = False
        self.bucket_ok = not offline
        self.cache: dict[str, dict | None] = {}
        if not offline:
            data, err = K.http_json(f"{OSV_API}/query", {"package": {"name": "jinja2", "ecosystem": "PyPI"},
                                                           "version": "2.4.1"}, timeout=8)
            self.api_ok = data is not None and err is None
        self.status = ("OSV API (api.osv.dev)" if self.api_ok else
                       "OSV data files (osv-vulnerabilities.storage.googleapis.com); OSV API not reachable"
                       if not offline else "skipped (--offline)")

    def record(self, vuln_id: str, ecosystem: str) -> dict | None:
        key = f"{ecosystem}/{vuln_id}"
        if key in self.cache:
            return self.cache[key]
        rec = None
        if self.api_ok:
            rec, _err = K.http_json(f"{OSV_API}/vulns/{urllib.parse.quote(vuln_id)}", timeout=20)
        if rec is None and self.bucket_ok:
            rec, err = K.http_json(f"{OSV_BUCKET}/{ecosystem}/{urllib.parse.quote(vuln_id)}.json", timeout=20)
            if err and not err.startswith("HTTP"):
                self.bucket_ok = False  # network problem: stop trying
                self.status = "not reachable; severity may be missing"
        self.cache[key] = rec
        return rec

    def best_record(self, ids: list[str], ecosystem: str) -> dict | None:
        ordered = sorted(ids, key=lambda i: (not i.startswith("GHSA-"), not i.startswith("PYSEC-"), i))
        for vuln_id in ordered:
            if vuln_id.startswith("CVE-") and len(ordered) > 1:
                continue
            rec = self.record(vuln_id, ecosystem)
            if rec:
                return rec
        return None

    def query_batch(self, queries: list[dict]) -> list[list[str]]:
        """OSV API querybatch: returns vulnerability ids per query (API only)."""
        if not self.api_ok or not queries:
            return [[] for _ in queries]
        data, _err = K.http_json(f"{OSV_API}/querybatch", {"queries": queries}, timeout=60)
        results = (data or {}).get("results", [])
        return [[v["id"] for v in (r.get("vulns") or [])] for r in results] + [[]] * (len(queries) - len(results))


def osv_details(rec: dict | None, ecosystem: str, name: str, installed: str) -> dict:
    """Pull summary, severity, CVSS, aliases and the fixed version from an OSV record."""
    out = {"summary": "", "severity": None, "cvss_vector": "", "cvss_score": None, "aliases": [], "fixed": None,
           "published": ""}
    if not rec:
        return out
    out["summary"] = (rec.get("summary") or "").strip()
    out["aliases"] = [rec.get("id", "")] + list(rec.get("aliases", []) or [])
    out["published"] = (rec.get("published") or "")[:10]
    db_sev = (rec.get("database_specific") or {}).get("severity")
    if db_sev:
        out["severity"] = K.normalise_severity(db_sev)
    for sev in rec.get("severity", []) or []:
        if sev.get("type") == "CVSS_V3":
            out["cvss_vector"] = sev.get("score", "")
            out["cvss_score"] = K.cvss3_base_score(out["cvss_vector"])
        elif sev.get("type") == "CVSS_V4" and not out["cvss_vector"]:
            out["cvss_vector"] = sev.get("score", "")
    if not out["severity"] and out["cvss_score"] is not None:
        out["severity"] = K.severity_from_score(out["cvss_score"])
    wanted = name.lower().replace("_", "-")
    for aff in rec.get("affected", []) or []:
        pkg = aff.get("package", {}) or {}
        if pkg.get("ecosystem") != ecosystem or pkg.get("name", "").lower().replace("_", "-") != wanted:
            continue
        for rng in aff.get("ranges", []) or []:
            if rng.get("type") not in ("ECOSYSTEM", "SEMVER"):
                continue
            introduced, fixed = None, None
            for event in rng.get("events", []):
                if "introduced" in event:
                    introduced, fixed = event["introduced"], None
                if "fixed" in event:
                    fixed = event["fixed"]
                    if (introduced in (None, "0") or not K.version_lt(installed, introduced)) and K.version_lt(installed, fixed):
                        out["fixed"] = fixed
                        return out
    return out


def clean_summary(text: str, limit: int = 180) -> str:
    text = re.sub(r"#+\s*\w+\s*", " ", text or "")
    text = re.sub(r"[`*]", "", text)  # keep underscores: they are part of names like set_key
    text = re.sub(r"\s+", " ", text).strip()
    first = re.split(r"(?<=[.!?])\s", text, maxsplit=1)[0]
    if len(first) > limit:
        first = first[: limit - 3].rsplit(" ", 1)[0] + "..."
    return first


# ---------------------------------------------------------------------------
# Python: pip-audit
# ---------------------------------------------------------------------------

def check_python(sbom: Sbom, refs: list[str], raw_dir: Path, sources: list, not_checked: list) -> list[dict]:
    tool = K.find_python_tool("pip-audit")
    if not tool:
        K.fail("pip-audit not found. Run scripts/setup.sh and use the kit's Python.")
    pins = {}
    for ref in refs:
        comp = sbom.components[ref]
        parsed = parse_purl(comp.get("purl", ""))
        version = comp.get("version") or (parsed or {}).get("version")
        name = (parsed or {}).get("name") or comp.get("name")
        if not version:
            not_checked.append(nc_row(sbom, ref, "No version in the SBOM, so it cannot be checked."))
            continue
        pins.setdefault((name.lower(), version), []).append(ref)
    req = raw_dir / "pinned-python-components.txt"
    K.write_text(req, "".join(f"{n}=={v}\n" for n, v in sorted(pins)))
    out_file = raw_dir / "pip-audit.json"
    K.info(f"Checking {len(pins)} Python packages with pip-audit (PyPI advisory data)")
    proc = K.run([tool, "-r", str(req), "--no-deps", "--disable-pip", "-f", "json", "--desc", "on",
                  "--aliases", "on", "--progress-spinner", "off", "-o", str(out_file)], check=False)
    if not out_file.exists():
        K.fail(f"pip-audit did not produce results:\n{proc.stderr[-1500:]}")
    data = K.load_json(out_file)
    sources.append({"source": "pip-audit (PyPI vulnerability data)",
                    "detail": f"pip-audit {K.tool_version([tool, '--version'])}; {len(pins)} packages"})
    findings = []
    for dep in data.get("dependencies", []):
        key = (dep["name"].lower(), dep.get("version"))
        dep_refs = pins.get(key) or [r for (n, v), rs in pins.items() if n.replace("_", "-") == key[0].replace("_", "-")
                                     and v == key[1] for r in rs]
        if "skip_reason" in dep:
            for ref in dep_refs:
                not_checked.append(nc_row(sbom, ref, f"pip-audit skipped it: {dep['skip_reason']}"))
            continue
        groups = group_ids(dep.get("vulns", []))
        for group in groups:
            for ref in dep_refs:
                findings.append({
                    "ref": ref, "ecosystem": "PyPI", "name": dep["name"], "installed": dep["version"],
                    "ids": group["ids"], "summary": group["summary"], "fix_versions": group["fix_versions"],
                    "severity": None, "cvss_score": None, "cvss_vector": "", "link": "",
                    "source": "pip-audit", "fix_hint": None,
                })
    return findings


def group_ids(vulns: list[dict]) -> list[dict]:
    """pip-audit can list one problem several times under different IDs. Merge them."""
    groups: list[dict] = []
    for v in vulns:
        ids = {v["id"], *(v.get("aliases") or [])}
        match = [g for g in groups if g["ids"] & ids]
        merged = {"ids": set(ids), "fix_versions": set(v.get("fix_versions") or []),
                  "summary": v.get("description") or ""}
        for g in match:
            merged["ids"] |= g["ids"]
            merged["fix_versions"] |= g["fix_versions"]
            merged["summary"] = merged["summary"] or g["summary"]
            groups.remove(g)
        groups.append(merged)
    for g in groups:
        g["ids"] = sorted(g["ids"])
        g["fix_versions"] = sorted(g["fix_versions"], key=K._version_key)
    return groups


# ---------------------------------------------------------------------------
# Node: npm audit and the npm bulk advisory service
# ---------------------------------------------------------------------------

def npm_range_match(rng: str, version: str) -> bool:
    try:
        import semantic_version as sv
        return sv.NpmSpec(rng).match(sv.Version.coerce(version))
    except Exception:  # noqa: BLE001 - if a range cannot be read, assume it applies
        return True


def npm_fixed_from_range(rng: str, version: str) -> str | None:
    """Best guess at the first fixed version from an npm range such as '>=1.0.0 <1.15.1'."""
    for sub in rng.split("||"):
        sub = sub.strip()
        if not npm_range_match(sub, version):
            continue
        upper = re.findall(r"<\s*v?(\d[\w.+-]*)", sub)
        upper = [u for u in upper if not re.search(r"<=\s*v?" + re.escape(u), sub)]
        if upper:
            return upper[-1]
        incl = re.findall(r"<=\s*v?(\d[\w.+-]*)", sub) or re.findall(r"\s-\s*v?(\d[\w.+-]*)", sub)
        if incl:
            return f"> {incl[-1]}"
    return None


def advisory_to_ids(adv: dict) -> list[str]:
    url = adv.get("url", "")
    m = re.search(r"(GHSA-[0-9a-z]{4}-[0-9a-z]{4}-[0-9a-z]{4})", url)
    ids = [m.group(1)] if m else [f"npm-{adv.get('source') or adv.get('id')}"]
    return ids


def run_npm_audit(folder: Path, raw_dir: Path, include_dev: bool, sources: list) -> tuple[dict, dict]:
    """Return ({package: [advisories]}, {package: fixAvailable}) from npm audit --json."""
    if not ((folder / "package-lock.json").exists() or (folder / "npm-shrinkwrap.json").exists()):
        K.warn(f"{folder} has no package-lock.json; npm audit needs one. Using the npm advisory service instead.")
        return {}, {}
    cmd = ["npm", "audit", "--json", "--package-lock-only"]
    if not include_dev:
        cmd.append("--omit=dev")
    K.info(f"Running npm audit in {folder}")
    proc = K.run(cmd, cwd=folder, check=False)
    try:
        data = json.loads(proc.stdout or "{}")
    except ValueError:
        K.warn(f"npm audit returned something unexpected in {folder}:\n{(proc.stderr or proc.stdout)[-800:]}")
        return {}, {}
    if "error" in data and "vulnerabilities" not in data:
        K.warn(f"npm audit failed in {folder}: {data['error'].get('summary', data['error'])}")
        return {}, {}
    K.write_json(raw_dir / f"npm-audit-{K.slugify(folder.name)}.json", data)
    sources.append({"source": f"npm audit in {folder.name}",
                    "detail": f"npm {K.tool_version(['npm', '--version'])}; "
                              f"{data.get('metadata', {}).get('dependencies', {}).get('prod', '?')} production packages"})
    advisories, fixes = {}, {}
    for name, entry in (data.get("vulnerabilities") or {}).items():
        advs = [v for v in entry.get("via", []) if isinstance(v, dict)]
        if advs:
            advisories[name] = [{"url": a.get("url", ""), "title": a.get("title", ""), "severity": a.get("severity"),
                                 "range": a.get("range", "*"), "cvss": a.get("cvss") or {}, "source": a.get("source")}
                                for a in advs]
        fixes[name] = entry.get("fixAvailable")
    return advisories, fixes


def npm_bulk(packages: dict[str, set], raw_dir: Path, sources: list) -> dict | None:
    payload = {name: sorted(vers) for name, vers in packages.items()}
    K.info(f"Checking {sum(len(v) for v in payload.values())} npm packages with the npm advisory service")
    data, err = K.http_json(NPM_BULK, payload, timeout=60)
    if data is None:
        K.warn(f"npm advisory service not reachable ({err}).")
        return None
    K.write_json(raw_dir / "npm-advisories.json", data)
    sources.append({"source": "npm advisory database (bulk advisory service, same data as npm audit)",
                    "detail": f"{len(payload)} package names"})
    return {name: [{"url": a.get("url", ""), "title": a.get("title", ""), "severity": a.get("severity"),
                    "range": a.get("vulnerable_versions", "*"), "cvss": a.get("cvss") or {}, "source": a.get("id")}
                   for a in advs] for name, advs in data.items()}


def check_node(sbom: Sbom, refs: list[str], projects: list[Path], raw_dir: Path, include_dev: bool,
               sources: list, not_checked: list) -> list[dict]:
    by_part: dict[str, list[str]] = {}
    for ref in refs:
        for part in sbom.parts_of(ref) or [""]:
            by_part.setdefault(part, []).append(ref)
    findings = []
    covered = set()
    for folder in projects:
        pkg_name = ""
        if (folder / "package.json").exists():
            pkg_name = K.load_json(folder / "package.json").get("name", "")
        part = next((p for p in by_part if p in (folder.name, pkg_name)), None)
        if part is None and len(by_part) == 1:
            part = next(iter(by_part))
        if part is None:
            K.warn(f"Could not match {folder} to a part of the SBOM; skipping npm audit there.")
            continue
        advisories, fixes = run_npm_audit(folder, raw_dir, include_dev, sources)
        if not advisories and not fixes and not (raw_dir / f"npm-audit-{K.slugify(folder.name)}.json").exists():
            continue
        for ref in by_part[part]:
            findings += npm_findings_for(sbom, ref, advisories, fixes, "npm audit")
            covered.add(ref)

    rest = [r for r in refs if r not in covered]
    if rest:
        packages: dict[str, set] = {}
        for ref in rest:
            parsed = parse_purl(sbom.components[ref].get("purl", ""))
            version = sbom.components[ref].get("version") or parsed.get("version")
            if version:
                packages.setdefault(parsed["full_name"], set()).add(version)
            else:
                not_checked.append(nc_row(sbom, ref, "No version in the SBOM, so it cannot be checked."))
        advisories = npm_bulk(packages, raw_dir, sources) if packages else {}
        if advisories is None:
            for ref in rest:
                not_checked.append(nc_row(sbom, ref, "npm advisory service not reachable."))
        else:
            for ref in rest:
                findings += npm_findings_for(sbom, ref, advisories, {}, "npm advisory database")
    return findings


def npm_findings_for(sbom: Sbom, ref: str, advisories: dict, fixes: dict, source: str) -> list[dict]:
    comp = sbom.components[ref]
    parsed = parse_purl(comp.get("purl", ""))
    name = parsed["full_name"]
    version = comp.get("version") or parsed.get("version") or ""
    out = []
    for adv in advisories.get(name, []):
        if not npm_range_match(adv["range"], version):
            continue
        cvss = adv.get("cvss") or {}
        out.append({
            "ref": ref, "ecosystem": "npm", "name": name, "installed": version,
            "ids": advisory_to_ids(adv), "summary": adv.get("title", ""), "fix_versions": [],
            "range": adv["range"], "severity": K.normalise_severity(adv.get("severity")),
            "cvss_score": cvss.get("score") or None, "cvss_vector": cvss.get("vectorString") or "",
            "link": adv.get("url", ""), "source": source, "fix_hint": fixes.get(name),
        })
    return out


# ---------------------------------------------------------------------------
# Other ecosystems through the OSV API
# ---------------------------------------------------------------------------

def check_other(sbom: Sbom, refs: list[str], osv: Osv, sources: list, not_checked: list) -> list[dict]:
    if not refs:
        return []
    if not osv.api_ok:
        for ref in refs:
            not_checked.append(nc_row(sbom, ref, "Needs the OSV API (api.osv.dev), which could not be reached."))
        return []
    queries = [{"package": {"purl": sbom.components[r]["purl"]}} for r in refs]
    results = osv.query_batch(queries)
    sources.append({"source": "OSV API querybatch", "detail": f"{len(refs)} components"})
    findings = []
    for ref, ids in zip(refs, results):
        comp = sbom.components[ref]
        parsed = parse_purl(comp["purl"])
        for vuln_id in ids:
            findings.append({
                "ref": ref, "ecosystem": OSV_ECOSYSTEMS.get(parsed["type"], parsed["type"]),
                "name": parsed["full_name"], "installed": comp.get("version") or parsed.get("version") or "",
                "ids": [vuln_id], "summary": "", "fix_versions": [], "severity": None, "cvss_score": None,
                "cvss_vector": "", "link": "", "source": "OSV API", "fix_hint": None,
            })
    return findings


# ---------------------------------------------------------------------------
# Manual checks
# ---------------------------------------------------------------------------

def nc_row(sbom: Sbom, ref: str, reason: str, where: str = "") -> dict:
    comp = sbom.components.get(ref, {})
    return {"component": display_name(comp), "version": comp.get("version", ""),
            "part": ", ".join(sbom.parts_of(ref)), "reason": reason, "where": where,
            "purl": comp.get("purl", ""), "cpe": comp.get("cpe", "")}


def manual_where(comp: dict) -> str:
    places = []
    parsed = parse_purl(comp.get("purl", ""))
    if parsed and parsed["type"] == "github" and parsed["namespace"]:
        places.append(f"https://github.com/{parsed['namespace']}/{parsed['name']}/security/advisories")
    for ref in comp.get("externalReferences", []) or []:
        if ref.get("type") in ("website", "advisories", "vcs") and ref.get("url") and ref["url"] not in places:
            places.append(ref["url"])
    term = urllib.parse.quote_plus(comp.get("name", ""))
    places.append(f"https://osv.dev/list?q={term}")
    places.append(f"https://nvd.nist.gov/vuln/search/results?form_type=Basic&results_type=overview&query={term}&search_type=all")
    places.append("https://euvd.enisa.europa.eu/")
    return " ; ".join(places)


# ---------------------------------------------------------------------------
# KEV: known exploited vulnerabilities
# ---------------------------------------------------------------------------

def load_kev(kev_file: str | None, offline: bool) -> tuple[dict, str]:
    data, origin = None, ""
    if kev_file:
        data, origin = K.load_json(kev_file), f"local file {Path(kev_file).name}"
    elif not offline:
        for url in KEV_URLS:
            data, _err = K.http_json(url, timeout=30)
            if data and "vulnerabilities" in data:
                origin = url
                break
            data = None
    if not data:
        return {}, "not checked (catalogue not reachable; pass --kev-file to use a downloaded copy)"
    entries = {v.get("cveID"): v for v in data.get("vulnerabilities", []) if v.get("cveID")}
    label = (f"CISA KEV catalogue version {data.get('catalogVersion', '?')} "
             f"({data.get('count', len(entries))} entries, released {str(data.get('dateReleased', ''))[:10]}) from {origin}")
    return entries, label


# ---------------------------------------------------------------------------
# Putting findings together
# ---------------------------------------------------------------------------

def primary_id(ids: list[str]) -> str:
    for prefix in ("CVE-", "GHSA-", "PYSEC-"):
        for i in sorted(ids):
            if i.startswith(prefix):
                return i
    return sorted(ids)[0] if ids else "unknown"


def enrich(findings: list[dict], osv: Osv, kev: dict) -> None:
    eco_dirs = {"PyPI": "PyPI", "npm": "npm"}

    def work(f):
        rec = osv.best_record(f["ids"], eco_dirs.get(f["ecosystem"], f["ecosystem"]))
        return f, osv_details(rec, f["ecosystem"], f["name"], f["installed"])

    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(pool.map(work, findings))
    for f, det in results:
        ids = set(f["ids"]) | {a for a in det["aliases"] if a}
        f["ids"] = sorted(ids)
        f["id"] = primary_id(f["ids"])
        f["aliases"] = [i for i in f["ids"] if i != f["id"]]
        if det["summary"]:
            f["summary"] = det["summary"]
        f["summary"] = clean_summary(f["summary"]) or "See the advisory."
        if not f.get("severity"):
            f["severity"] = det["severity"] or K.severity_from_score(det["cvss_score"])
        if not f.get("cvss_score") and det["cvss_score"] is not None:
            f["cvss_score"] = det["cvss_score"]
        if not f.get("cvss_vector"):
            f["cvss_vector"] = det["cvss_vector"]
        f["published"] = det["published"]
        fixed = det["fixed"]
        if not fixed and f["fix_versions"]:
            above = [v for v in f["fix_versions"] if K.version_lt(f["installed"], v)]
            fixed = K.min_version(above) if above else K.min_version(f["fix_versions"])
        if not fixed and f.get("range"):
            fixed = npm_fixed_from_range(f["range"], f["installed"])
        f["fixed"] = fixed
        if not f["link"]:
            ghsa = next((i for i in f["ids"] if i.startswith("GHSA-")), None)
            f["link"] = f"https://github.com/advisories/{ghsa}" if ghsa else f"https://osv.dev/vulnerability/{f['id']}"
        cves = [i for i in f["ids"] if i.startswith("CVE-")]
        hits = [kev[c] for c in cves if c in kev]
        f["kev"] = ({"cve": hits[0]["cveID"], "date_added": hits[0].get("dateAdded", ""),
                     "ransomware": hits[0].get("knownRansomwareCampaignUse", "")} if hits else None)


def merge_duplicates(findings: list[dict]) -> list[dict]:
    """Merge findings for the same component whose IDs overlap (aliases of one problem)."""
    merged: list[dict] = []
    for f in findings:
        same = [m for m in merged if m["ref"] == f["ref"] and set(m["ids"]) & set(f["ids"])]
        if not same:
            merged.append(f)
            continue
        m = same[0]
        worst = min((m, f), key=lambda x: K.severity_rank(x["severity"]))
        m["ids"] = sorted(set(m["ids"]) | set(f["ids"]))
        m["severity"] = worst["severity"]
        m["summary"], m["link"] = worst["summary"], worst["link"]
        m["cvss_score"] = max((x for x in (m.get("cvss_score"), f.get("cvss_score")) if x is not None), default=None)
        exact = [x for x in (m.get("fixed"), f.get("fixed")) if x and not str(x).startswith(">")]
        m["fixed"] = K.max_version(exact) if exact else (m.get("fixed") or f.get("fixed"))
        m["kev"] = m.get("kev") or f.get("kev")
        if not isinstance(m.get("fix_hint"), dict) and f.get("fix_hint") is not None:
            m["fix_hint"] = f["fix_hint"]
        m["source"] = ", ".join(sorted(set(m["source"].split(", ")) | {f["source"]}))
    for m in merged:
        m["id"] = primary_id(m["ids"])
        m["aliases"] = [i for i in m["ids"] if i != m["id"]]
    return merged


def assign_actions(findings: list[dict], sbom: Sbom) -> list[dict]:
    """Give each finding a recommended action, then group findings that share one."""
    upgrades: dict[tuple, list[dict]] = {}
    npm_targets: dict[tuple, tuple] = {}
    for f in findings:
        direct, parents = sbom.direct_and_parents(f["ref"])
        f["direct"] = direct
        f["via"] = parents
        f["part"] = ", ".join(sbom.parts_of(f["ref"])) or sbom.product_name()
        hint = f.get("fix_hint")
        if isinstance(hint, dict) and hint.get("name"):
            key = ("npm-fix", f["part"], hint["name"], hint.get("version"))
            major = " This is a major version change, so test carefully." if hint.get("isSemVerMajor") else ""
            current = next((c.get("version") for c in sbom.components.values()
                            if display_name(c) == hint["name"] and f["part"] in sbom.parts_of(c.get("bom-ref", ""))), None)
            frm = f" from {current}" if current else ""
            f["action_key"] = key
            f["action"] = f"Update {hint['name']}{frm} to {hint.get('version')} in {f['part']}.{major}"
            npm_targets[(f["part"], hint["name"])] = (key, f["action"])

    for f in findings:
        if "action_key" in f:
            continue
        hint = f.get("fix_hint")
        if hint is True:
            joined = next((npm_targets[(f["part"], p)] for p in f["via"] if (f["part"], p) in npm_targets), None)
            if joined:
                f["action_key"], f["action"] = joined
            else:
                f["action_key"] = ("npm-audit-fix", f["part"])
                f["action"] = (f"Run `npm audit fix` in {f['part']}, test, and ship the updated package-lock.json.")
            continue
        if not f.get("fixed"):
            f["action_key"] = ("no-fix", f["part"], f["name"], f["id"])
            f["action"] = (f"No fixed version of {f['name']} is listed yet for {f['id']}. Read the advisory for "
                           "a workaround, check whether your product uses the affected feature, and watch for a fix.")
            continue
        upgrades.setdefault((f["part"], f["ecosystem"], f["name"], f["installed"]), []).append(f)

    npm_parents: dict[tuple, list] = {}
    for (part, eco, name, installed), group in upgrades.items():
        exact = [g["fixed"] for g in group if not str(g["fixed"]).startswith(">")]
        target = K.max_version(exact) if exact else group[0]["fixed"]
        major = (" This is a major version change, so test carefully."
                 if exact and K.major_of(target) != K.major_of(installed) else "")
        direct = group[0]["direct"]
        via = group[0]["via"]
        version_text = target if not str(target).startswith(">") else f"a version newer than {target[1:].strip()}"
        if direct or not via:
            text = f"Update {name} from {installed} to {version_text} or later in {part}.{major}"
        elif eco == "PyPI":
            text = (f"Raise {name} from {installed} to {version_text} or later in {part}. It comes in through "
                    f"{', '.join(via)}: add `{name}>={target}` to requirements.txt, or update "
                    f"{', '.join(via)}.{major}")
        else:
            # npm packages that come in through the same parent: one action for the parent.
            npm_parents.setdefault((part, tuple(via)), []).append((name, installed, version_text, group))
            continue
        for g in group:
            g["action_key"] = ("upgrade", part, eco, name, installed)
            g["action"] = text

    for (part, via), items in npm_parents.items():
        wanted = ", ".join(f"{n} {vt} or later (now {inst})" for n, inst, vt, _ in items)
        names = ", ".join(n for n, *_ in items)
        text = (f"Update {', '.join(via)} in {part} to a release that pulls in {wanted}. If there is no such "
                f"release yet, add overrides for {names} in package.json and test.")
        for *_, group in items:
            for g in group:
                g["action_key"] = ("npm-parent", part, via)
                g["action"] = text

    actions: dict[tuple, dict] = {}
    for f in findings:
        act = actions.setdefault(f["action_key"], {"action": f["action"], "part": f["part"], "findings": [],
                                                   "components": set(), "kev": False})
        act["findings"].append(f)
        act["components"].add(f["name"])
        act["kev"] = act["kev"] or bool(f.get("kev"))
    ordered = []
    for act in actions.values():
        act["severity"] = min((f["severity"] for f in act["findings"]), key=K.severity_rank)
        act["ids"] = [f["id"] for f in sorted(act["findings"], key=lambda x: K.severity_rank(x["severity"]))]
        act["components"] = sorted(act["components"])
        ordered.append(act)
    ordered.sort(key=lambda a: (not a["kev"], K.severity_rank(a["severity"]), -len(a["findings"]), a["action"]))
    return ordered


def compare(findings: list[dict], previous: dict) -> dict:
    prev = previous.get("findings", [])

    def same(a, b):
        return (a["ecosystem"] == b["ecosystem"] and a["name"].lower() == b["name"].lower()
                and set(a.get("ids") or [a["id"]]) & set(b.get("ids") or [b["id"]]))

    new = [f for f in findings if not any(same(f, p) for p in prev)]
    gone = [p for p in prev if not any(same(f, p) for f in findings)]
    for f in findings:
        f["status"] = "New" if f in new else "Still open"
    return {"date": previous.get("generated", "the last check"), "new": new, "fixed": gone}


# ---------------------------------------------------------------------------
# Writing the report
# ---------------------------------------------------------------------------

def severity_counts(findings: list[dict]) -> dict:
    counts = {s: 0 for s in K.SEVERITY_ORDER}
    for f in findings:
        counts[f["severity"]] = counts.get(f["severity"], 0) + 1
    return counts


def write_markdown(ctx: dict) -> str:
    f_all, actions, nc = ctx["findings"], ctx["actions"], ctx["not_checked"]
    counts = severity_counts(f_all)
    kev_hits = [f for f in f_all if f.get("kev")]
    affected = sorted({(f["name"], f["installed"], f["part"]) for f in f_all})
    md = [f"# Known-vulnerability report: {ctx['product']} {ctx['version']}", ""]
    md += [f"- **Prepared for:** {ctx['client_name']}",
           f"- **Prepared by:** {ctx['prepared_by']}",
           f"- **Date of check:** {ctx['date']}",
           f"- **SBOM checked:** `{ctx['sbom_name']}` (SHA-256 `{ctx['sbom_sha'][:16]}...`)", ""]
    md += ["> This report lists publicly known vulnerabilities in the third-party components of your product. "
           f"It matches your SBOM against public vulnerability databases as of {ctx['date']}. It does not test "
           "your own code, and a listed vulnerability is not always exploitable in your product. "
           "This is technical information, not legal advice.", ""]
    md += ["## Summary", "", "| Item | Result |", "|---|---|",
           f"| Components in the SBOM | {ctx['total_components']} |",
           f"| Checked automatically | {ctx['checked']} |",
           f"| Need a manual check | {len(nc)} |",
           f"| Components with known vulnerabilities | {len(affected)} |",
           f"| Known vulnerabilities | {len(f_all)} |",
           "| By severity | " + " · ".join(f"{s} {counts[s]}" for s in K.SEVERITY_ORDER) + " |",
           f"| On CISA's known-exploited list (KEV) | {len(kev_hits) if ctx['kev_checked'] else 'not checked'} |"]
    if ctx.get("comparison"):
        cmp_ = ctx["comparison"]
        md.append(f"| Changes since {cmp_['date']} | {len(cmp_['new'])} new · {len(cmp_['fixed'])} no longer found |")
    md.append("")
    if kev_hits:
        md += ["> **Act today.** These vulnerabilities are on CISA's list of vulnerabilities known to be "
               "exploited in real attacks: " + ", ".join(f"{f['id']} in {f['name']} {f['installed']}" for f in kev_hits)
               + ". Check straight away whether your product is affected. If it is, and the vulnerability is "
               "being exploited, the CRA reporting clock may already be running: open the reporting runbook "
               "and call your lawyer.", ""]
    if actions:
        md += ["## What to do first", "",
               "Each action below fixes one or more findings. They are sorted by the most serious finding they fix.", ""]
        for n, act in enumerate(actions, start=1):
            ids = ", ".join(act["ids"][:6]) + (f" and {len(act['ids']) - 6} more" if len(act["ids"]) > 6 else "")
            flag = " **Known exploited.**" if act["kev"] else ""
            md.append(f"{n}. **{act['action']}**{flag} Fixes {len(act['findings'])} finding(s), highest severity "
                      f"{act['severity']}, in {', '.join(act['components'])} ({ids}).")
        md.append("")
    elif ctx["checked"] == 0:
        md += ["## What to do first", "", "None of the components could be checked automatically. Check the "
               "components listed under 'Components that need a manual check'.", ""]
    else:
        md += ["## What to do first", "", "No known vulnerabilities were found in the components that could be "
               "checked automatically. Keep checking every month and before every release.", ""]
    md += ["## Before your next release", "",
           "- From 11 December 2027 the CRA expects products to go on sale without known exploitable "
           "vulnerabilities (verify current dates). For each finding, either fix it or write down why it "
           "cannot be exploited in your product (for example, the affected feature is never used). Use the "
           "'Your assessment' column in the spreadsheet.",
           "- Re-run the SBOM and this check for every release, and keep both files with the release.",
           "- The components in 'Need a manual check' are not covered by the automatic tools. Check them by hand.", ""]
    md += ["## Is anything here reportable?", ""]
    if kev_hits:
        md.append("Some findings are on the known-exploited list (see the warning above). Being on that list "
                  "does not prove your product is being attacked, but it means you must assess it now.")
    elif ctx["kev_checked"]:
        md.append(f"None of the findings is on CISA's known-exploited list ({ctx['kev_label']}).")
    else:
        md.append("The known-exploited list could not be checked this time (see 'How this check was done').")
    md += ["", "Since 11 September 2026 manufacturers must report actively exploited vulnerabilities in their "
           "products, and severe incidents, to ENISA's single reporting platform: an early warning within 24 "
           "hours, a notification within 72 hours and a final report later (verify current dates). If you learn "
           "that any vulnerability in your product is being exploited, follow your reporting runbook straight "
           "away and involve your lawyer.", ""]

    action_no = {id(f): n for n, act in enumerate(actions, start=1) for f in act["findings"]}
    md += ["## All findings", "", "The Action column points to the numbered list in 'What to do first'.", "",
           "| Severity | ID | Component | Fixed in | Action | Part | What it is |",
           "|---|---|---|---|---|---|---|"]
    for f in sorted(f_all, key=lambda x: (K.severity_rank(x["severity"]), x["name"].lower(), x["id"])):
        idtext = f"[{f['id']}]({f['link']})"
        if f.get("kev"):
            idtext += " (KEV)"
        if f.get("status") == "New":
            idtext += " (new)"
        md.append(f"| {f['severity']} | {idtext} | {K.md_escape(f['name'])} {K.md_escape(f['installed'])} | "
                  f"{K.md_escape(f['fixed'] or 'no fix yet')} | {action_no.get(id(f), '')} | "
                  f"{K.md_escape(f['part'])} | {K.md_escape(f['summary'])} |")
    if not f_all:
        md.append("| - | - | - | - | - | - | No known vulnerabilities found |")
    md.append("")
    if ctx.get("comparison") and ctx["comparison"]["fixed"]:
        md += [f"## No longer found since {ctx['comparison']['date']}", "",
               "| ID | Component | Was installed |", "|---|---|---|"]
        for p in ctx["comparison"]["fixed"]:
            md.append(f"| {p.get('id')} | {K.md_escape(p.get('name'))} | {K.md_escape(p.get('installed'))} |")
        md.append("")
    md += ["## Components that need a manual check", ""]
    if nc:
        md += ["| Component | Version | Part | Why | Where to look |", "|---|---|---|---|---|"]
        for row in nc:
            where = row["where"].split(" ; ")[0] if row["where"] else ""
            where_md = f"[{where}]({where})" if where.startswith("http") else where
            md.append(f"| {K.md_escape(row['component'])} | {K.md_escape(row['version'])} | {K.md_escape(row['part'])} | "
                      f"{K.md_escape(row['reason'])} | {where_md} |")
        md += ["", "More places to look are in the spreadsheet: the vendor's security advisories, OSV "
               "(osv.dev), the US National Vulnerability Database and the EU Vulnerability Database (EUVD).", ""]
    else:
        md += ["None.", ""]
    md += ["## How this check was done", ""]
    for src in ctx["sources"]:
        md.append(f"- **{src['source']}**: {src['detail']}")
    md += [f"- **Severity and CVE details:** {ctx['osv_status']}",
           f"- **Known exploited check:** {ctx['kev_label']}",
           "", "Severity follows the rating in each advisory (GitHub's rating where there is one, otherwise "
           "the CVSS v3 base score). 'Fixed in' is the first release that fixes the problem for your version. "
           "New vulnerabilities are published every day, so this report is a snapshot.", ""]
    return "\n".join(md)


def write_xlsx(ctx: dict, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    row = K.xlsx_title(ws, f"Known-vulnerability report: {ctx['product']} {ctx['version']}",
                       f"Checked on {ctx['date']} against SBOM {ctx['sbom_name']}. Not legal advice.")
    counts = severity_counts(ctx["findings"])
    facts = [("Prepared for", ctx["client_name"]), ("Prepared by", ctx["prepared_by"]),
             ("Components in the SBOM", ctx["total_components"]), ("Checked automatically", ctx["checked"]),
             ("Need a manual check", len(ctx["not_checked"])), ("Known vulnerabilities", len(ctx["findings"]))]
    facts += [(f"Severity: {s}", counts[s]) for s in K.SEVERITY_ORDER]
    facts += [("On CISA KEV list", sum(1 for f in ctx["findings"] if f.get("kev")) if ctx["kev_checked"] else "not checked")]
    if ctx.get("comparison"):
        facts += [(f"New since {ctx['comparison']['date']}", len(ctx["comparison"]["new"])),
                  (f"No longer found since {ctx['comparison']['date']}", len(ctx["comparison"]["fixed"]))]
    for label, value in facts:
        ws.cell(row=row, column=1, value=label).font = K.xlsx_font(bold=True)
        ws.cell(row=row, column=2, value=value).font = K.xlsx_font()
        row += 1
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 60

    act = wb.create_sheet("Actions")
    headers = ["#", "Action", "Part", "Findings fixed", "Highest severity", "Known exploited?", "Components", "IDs",
               "Owner", "Target date", "Done?"]
    K.xlsx_header(act, 1, headers, [5, 70, 18, 10, 12, 11, 30, 40, 18, 13, 10])
    yellow = PatternFill("solid", fgColor=K.INPUT_FILL)
    for n, a in enumerate(ctx["actions"], start=1):
        values = [n, a["action"], a["part"], len(a["findings"]), a["severity"], "YES" if a["kev"] else "no",
                  ", ".join(a["components"]), ", ".join(a["ids"]), "", "", ""]
        for c, v in enumerate(values, start=1):
            act.cell(row=n + 1, column=c, value=v)
    K.xlsx_body_style(act, 2, len(ctx["actions"]) + 1, len(headers))
    for r in range(2, len(ctx["actions"]) + 2):
        for c in (9, 10, 11):
            act.cell(row=r, column=c).fill = yellow
    act.freeze_panes = "C2"

    fs = wb.create_sheet("Findings")
    headers = ["Severity", "ID", "Other IDs", "Component", "Installed", "Fixed in", "Part", "Direct?",
               "Comes in through", "What it is", "CVSS score", "Known exploited (KEV)", "Status", "Advisory link",
               "Recommended action", "Your assessment", "Notes / evidence"]
    K.xlsx_header(fs, 1, headers, [11, 22, 30, 20, 11, 12, 16, 8, 22, 55, 8, 14, 10, 45, 55, 26, 40])
    rows = sorted(ctx["findings"], key=lambda x: (K.severity_rank(x["severity"]), x["name"].lower(), x["id"]))
    for r, f in enumerate(rows, start=2):
        values = [f["severity"], f["id"], ", ".join(f["aliases"]), f["name"], f["installed"], f["fixed"] or "no fix yet",
                  f["part"], "Yes" if f["direct"] else "No", ", ".join(f["via"]), f["summary"], f.get("cvss_score"),
                  f"YES ({f['kev']['date_added']})" if f.get("kev") else "no", f.get("status", ""), f["link"],
                  f["action"], "Not assessed yet", ""]
        for c, v in enumerate(values, start=1):
            fs.cell(row=r, column=c, value=v)
        fs.cell(row=r, column=14).hyperlink = f["link"]
    last = len(rows) + 1
    K.xlsx_body_style(fs, 2, last, len(headers))
    choice = DataValidation(type="list", allow_blank=True, formula1=(
        '"Not assessed yet,Affected - fix planned,Fixed,Not affected - code not used,'
        'Not affected - mitigated,Not affected - other (see notes)"'))
    fs.add_data_validation(choice)
    for r in range(2, last + 1):
        fs.cell(row=r, column=16).fill = yellow
        fs.cell(row=r, column=17).fill = yellow
        choice.add(fs.cell(row=r, column=16))
    colours = {"Critical": "C00000", "High": "F4B084", "Medium": "FFE699", "Low": "DDEBF7"}
    for sev, colour in colours.items():
        fs.conditional_formatting.add(f"A2:A{max(last, 2)}", FormulaRule(formula=[f'$A2="{sev}"'],
                                      fill=PatternFill("solid", fgColor=colour)))
    fs.freeze_panes = "C2"
    fs.auto_filter.ref = f"A1:Q{max(last, 2)}"

    mc = wb.create_sheet("Manual checks")
    headers = ["Component", "Version", "Part", "Why it needs a manual check", "Where to look", "Checked on",
               "Result (IDs found or 'none found')", "Checked by"]
    K.xlsx_header(mc, 1, headers, [30, 12, 16, 45, 80, 12, 34, 16])
    for r, nrow in enumerate(ctx["not_checked"], start=2):
        values = [nrow["component"], nrow["version"], nrow["part"], nrow["reason"],
                  nrow["where"].replace(" ; ", "\n"), "", "", ""]
        for c, v in enumerate(values, start=1):
            mc.cell(row=r, column=c, value=v)
    K.xlsx_body_style(mc, 2, len(ctx["not_checked"]) + 1, len(headers))
    for r in range(2, len(ctx["not_checked"]) + 2):
        for c in (6, 7, 8):
            mc.cell(row=r, column=c).fill = yellow

    src = wb.create_sheet("Sources")
    K.xlsx_header(src, 1, ["Source", "Detail"], [60, 90])
    rows_src = [(s["source"], s["detail"]) for s in ctx["sources"]]
    rows_src += [("Severity and CVE details", ctx["osv_status"]), ("Known exploited check", ctx["kev_label"]),
                 ("SBOM file", ctx["sbom_name"]), ("SBOM SHA-256", ctx["sbom_sha"])]
    for r, (a, b) in enumerate(rows_src, start=2):
        src.cell(row=r, column=1, value=a)
        src.cell(row=r, column=2, value=b)
    K.xlsx_body_style(src, 2, len(rows_src) + 1, 2)
    wb.save(path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("sbom", help="CycloneDX JSON SBOM (from make_sbom.py or the client)")
    ap.add_argument("--out", required=True, help="output folder")
    ap.add_argument("--node-project", action="append", default=[], help="Node project folder to run npm audit in")
    ap.add_argument("--client", help="client.json (for names in the report)")
    ap.add_argument("--prepared-by", help="your name and business")
    ap.add_argument("--previous", help="findings.json from the last check, to show what changed")
    ap.add_argument("--kev-file", help="a downloaded copy of CISA's known_exploited_vulnerabilities.json")
    ap.add_argument("--include-dev", action="store_true", help="npm audit: include development dependencies")
    ap.add_argument("--offline", action="store_true", help="skip OSV and KEV look-ups (tools still need their databases)")
    args = ap.parse_args()

    sbom = Sbom(Path(args.sbom))
    client = K.load_json(args.client) if args.client else {}
    out_dir = Path(args.out)
    raw_dir = out_dir / "vulnerability-data"
    raw_dir.mkdir(parents=True, exist_ok=True)

    python_refs, node_refs, other_refs, not_checked = [], [], [], []
    for ref, comp in sbom.components.items():
        if ref in sbom.part_roots:
            continue
        parsed = parse_purl(comp.get("purl", ""))
        if parsed and parsed["type"] == "pypi":
            python_refs.append(ref)
        elif parsed and parsed["type"] == "npm":
            node_refs.append(ref)
        elif parsed and parsed["type"] in OSV_ECOSYSTEMS:
            other_refs.append(ref)
        else:
            reason = ("Listed by hand: no automatic database covers it reliably."
                      if get_prop(comp, SOURCE_PROP) else "No package identifier (purl) that the tools can look up.")
            not_checked.append(nc_row(sbom, ref, reason, manual_where(comp)))

    sources: list[dict] = []
    osv = Osv(args.offline)
    findings = []
    if python_refs:
        findings += check_python(sbom, python_refs, raw_dir, sources, not_checked)
    if node_refs:
        projects = [Path(p).expanduser().resolve() for p in args.node_project]
        findings += check_node(sbom, node_refs, projects, raw_dir, args.include_dev, sources, not_checked)
    findings += check_other(sbom, other_refs, osv, sources, not_checked)

    K.info(f"Looking up details for {len(findings)} findings ({osv.status})")
    kev, kev_label = load_kev(args.kev_file, args.offline)
    enrich(findings, osv, kev)
    findings = merge_duplicates(findings)
    actions = assign_actions(findings, sbom)
    comparison = compare(findings, K.load_json(args.previous)) if args.previous else None
    if comparison is None:
        for f in findings:
            f["status"] = "First check"

    checked = len(python_refs) + len(node_refs) + len(other_refs) - sum(
        1 for row in not_checked if not row["reason"].startswith("Listed by hand")
        and not row["reason"].startswith("No package identifier"))
    ctx = {
        "product": sbom.product_name(), "version": sbom.product_version(),
        "client_name": K.dig(client, "company.legal_name") or (sbom.root.get("supplier") or {}).get("name", ""),
        "prepared_by": args.prepared_by or K.dig(client, "engagement.consultant_name") or "",
        "date": K.today(), "sbom_name": sbom.path.name, "sbom_sha": sbom.sha256,
        "total_components": len(sbom.components) - (0 if sbom.single else len(sbom.part_roots)),
        "checked": checked, "findings": findings, "actions": actions, "not_checked": not_checked,
        "sources": sources, "osv_status": osv.status, "kev_label": kev_label, "kev_checked": bool(kev),
        "comparison": comparison,
    }

    md = write_markdown(ctx)
    md_path = out_dir / "vulnerability-report.md"
    K.write_text(md_path, md)
    footer = f"Prepared by {ctx['prepared_by']}  |  Not legal advice" if ctx["prepared_by"] else "Not legal advice"
    K.md_to_docx(md, out_dir / "vulnerability-report.docx", footer=footer)
    write_xlsx(ctx, out_dir / "vulnerability-report.xlsx")
    K.write_json(raw_dir / "findings.json", {
        "generated": ctx["date"], "sbom": ctx["sbom_name"], "sbom_sha256": ctx["sbom_sha"],
        "product": {"name": ctx["product"], "version": ctx["version"]},
        "findings": [{k: v for k, v in f.items() if k not in ("action_key", "fix_hint", "ref")} for f in findings],
        "not_checked": not_checked, "sources": sources, "kev": kev_label, "osv": osv.status,
    })
    counts = severity_counts(findings)
    K.info("")
    K.info(f"Report: {md_path}, .docx and .xlsx")
    K.info(f"Found {len(findings)} known vulnerabilities: " + ", ".join(f"{s} {counts[s]}" for s in K.SEVERITY_ORDER))
    K.info(f"{len(actions)} recommended actions; {len(not_checked)} components need a manual check")
    if any(f.get("kev") for f in findings):
        K.warn("At least one finding is on CISA's known-exploited list. Read the report today.")


if __name__ == "__main__":
    main()
