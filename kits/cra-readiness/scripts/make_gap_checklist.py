#!/usr/bin/env python3
"""Build the CRA gap checklist spreadsheet (.xlsx).

The requirements come from templates/gap-checklist-items.csv: a plain-English
paraphrase of the CRA's essential requirements and related duties, written for
small makers. It is not the legal text.

  python make_gap_checklist.py --out ../templates/cra-gap-checklist.xlsx      # blank checklist
  python make_gap_checklist.py --client client.json --status gap-status.csv \\
      --out clients/acme/2026-10/cra-gap-checklist.xlsx                      # filled in

gap-status.csv columns: id, status, evidence, owner, target_date, notes
Status values: Not started, In progress, Done, Needs lawyer, Not applicable
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import kitlib as K  # noqa: E402

STATUSES = ["Not started", "In progress", "Done", "Needs lawyer", "Not applicable"]
STATUS_FILLS = {"Done": "C6EFCE", "In progress": "FFEB9C", "Not started": "FFC7CE", "Needs lawyer": "E4DFEC",
                "Not applicable": "EDEDED"}
LAST_ROW = 500  # formulas look this far down, so added rows are counted too


def read_csv(path: Path) -> list[dict]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return [{k.strip(): (v or "").strip() for k, v in row.items() if k} for row in csv.DictReader(fh)]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", required=True, help="output .xlsx file")
    ap.add_argument("--items", default=str(K.TEMPLATES_DIR / "gap-checklist-items.csv"))
    ap.add_argument("--status", help="CSV with the client's status per item (optional)")
    ap.add_argument("--client", help="client.json (for the title page)")
    ap.add_argument("--product", help="product name (default: first product in client.json)")
    args = ap.parse_args()

    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    items = read_csv(Path(args.items))
    status = {row["id"]: row for row in read_csv(Path(args.status))} if args.status else {}
    unknown = sorted(set(status) - {i["id"] for i in items})
    if unknown:
        K.warn(f"IDs in {args.status} that are not in the checklist: {', '.join(unknown)}")
    client = K.load_json(args.client) if args.client else {}
    company = K.dig(client, "company.legal_name") or ""
    product = args.product or ((client.get("products") or [{}])[0].get("name") or "")
    prepared_by = K.dig(client, "engagement.consultant_name") or ""
    yellow = PatternFill("solid", fgColor=K.INPUT_FILL)

    wb = Workbook()
    # ---- How to use -------------------------------------------------------
    how = wb.active
    how.title = "How to use"
    title = "CRA gap checklist" + (f": {product}" if product else "")
    row = K.xlsx_title(how, title, "Plain-English summary of the EU Cyber Resilience Act (Regulation (EU) 2024/2847) "
                                   "for small makers. Not the legal text and not legal advice.")
    facts = [("Client", company or "(fill in)"), ("Product", product or "(fill in)"),
             ("Prepared by", prepared_by or "(fill in)"), ("Status as of", K.today())]
    for label, value in facts:
        how.cell(row=row, column=1, value=label).font = K.xlsx_font(bold=True)
        how.cell(row=row, column=2, value=value).font = K.xlsx_font()
        row += 1
    blocks = [
        ("How to fill it in", [
            "1. Work through the 'Checklist' sheet with the client, one area at a time (about 60 to 90 minutes).",
            "2. Fill in only the yellow columns: Status, Evidence, Owner, Target date and Notes / gaps.",
            "3. Status: pick from the list. Use 'Needs lawyer' for anything about scope, category, role or conformity.",
            "4. Evidence: the file name or link that proves the item (for example 'vulnerability-report.xlsx, 2026-10-07').",
            "5. The 'Summary' sheet counts progress per area by itself.",
        ]),
        ("What the statuses mean", [
            "Not started: nothing in place yet.",
            "In progress: partly done, or done but not yet approved or published.",
            "Done: in place, with evidence you could show an authority.",
            "Needs lawyer: waiting for a legal answer. Do not guess.",
            "Not applicable: does not apply to this product; say why in Notes.",
        ]),
        ("Key dates (verify current dates)", [
            "11 September 2026: reporting of actively exploited vulnerabilities and severe incidents to ENISA's single "
            "reporting platform applies (early warning within 24 hours, notification within 72 hours, then a final report).",
            "11 December 2027: all other CRA obligations apply.",
        ]),
        ("Ground rules", [
            "This checklist shows gaps. Completing it does not make a product 'CRA compliant'.",
            "Scope, product category, role and conformity assessment questions go to the client's lawyer.",
            "The manufacturer signs the EU declaration of conformity, never the consultant.",
            "References point to Regulation (EU) 2024/2847. Check the official text for the exact wording.",
        ]),
    ]
    for heading, lines in blocks:
        row += 1
        how.cell(row=row, column=1, value=heading).font = K.xlsx_font(bold=True, size=11, color=K.HEADER_FILL)
        row += 1
        for line in lines:
            cell = how.cell(row=row, column=1, value=line)
            cell.font = K.xlsx_font()
            how.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            how.row_dimensions[row].height = 28 if len(line) > 110 else 15
            row += 1
    row += 1
    how.cell(row=row, column=1, value="Example of a filled-in row").font = K.xlsx_font(bold=True, size=11, color=K.HEADER_FILL)
    row += 1
    example_head = ["ID", "Status", "Evidence (file or link)", "Owner", "Target date", "Notes / gaps"]
    example = ["P2", "In progress", "vulnerability-report.xlsx (2026-10-07)", "Firmware lead", "2026-11-30",
               "6 updates planned for release 2.4; 1 finding marked 'not affected'."]
    K.xlsx_header(how, row, example_head)
    for c, value in enumerate(example, start=1):
        cell = how.cell(row=row + 1, column=c, value=value)
        cell.fill = yellow
    K.xlsx_body_style(how, row + 1, row + 1, len(example))
    for col, width in zip("ABCDEF", [24, 16, 36, 18, 13, 50]):
        how.column_dimensions[col].width = width

    # ---- Checklist ------------------------------------------------------------
    ws = wb.create_sheet("Checklist")
    headers = ["ID", "Area", "Requirement (plain English)", "CRA reference", "What good evidence looks like",
               "Quick win for small makers", "Status", "Evidence (file or link)", "Owner", "Target date",
               "Notes / gaps"]
    K.xlsx_header(ws, 1, headers, [6, 20, 58, 16, 34, 40, 14, 30, 18, 12, 40])
    for r, item in enumerate(items, start=2):
        st = status.get(item["id"], {})
        values = [item["id"], item["area"], item["requirement"], item["reference"], item["evidence"],
                  item["quick_win"], st.get("status") or "Not started", st.get("evidence", ""),
                  st.get("owner") or "", st.get("target_date", ""), st.get("notes", "")]
        if values[6] not in STATUSES:
            K.warn(f"{item['id']}: unknown status '{values[6]}', using 'Not started'.")
            values[6] = "Not started"
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)
    last = len(items) + 1
    K.xlsx_body_style(ws, 2, last, len(headers))
    for r in range(2, last + 1):
        for c in range(7, 12):
            ws.cell(row=r, column=c).fill = yellow
        ws.cell(row=r, column=1).font = K.xlsx_font(bold=True)
    choice = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False,
                            showErrorMessage=True, errorTitle="Status", error="Pick a status from the list.")
    ws.add_data_validation(choice)
    choice.add(f"G2:G{LAST_ROW}")
    for value, colour in STATUS_FILLS.items():
        ws.conditional_formatting.add(f"G2:G{LAST_ROW}", CellIsRule(operator="equal", formula=[f'"{value}"'],
                                      fill=PatternFill("solid", fgColor=colour)))
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:K{last}"
    ws.print_options.gridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    # ---- Summary ----------------------------------------------------------------
    sm = wb.create_sheet("Summary", 1)
    nxt = K.xlsx_title(sm, "Progress by area", "Counts update automatically from the Checklist sheet.")
    head = ["Area", "Items", "Done", "In progress", "Not started", "Needs lawyer", "Not applicable", "Progress"]
    K.xlsx_header(sm, nxt, head, [44, 9, 9, 12, 12, 13, 15, 11])
    areas = []
    for item in items:
        if item["area"] not in areas:
            areas.append(item["area"])
    cached: dict[str, object] = {}
    rng_area = f"Checklist!$B$2:$B${LAST_ROW}"
    rng_status = f"Checklist!$G$2:$G${LAST_ROW}"
    first = nxt + 1
    for i, area in enumerate(areas):
        r = first + i
        sm.cell(row=r, column=1, value=area)
        sm.cell(row=r, column=2, value=f"=COUNTIF({rng_area},$A{r})")
        area_items = [it for it in items if it["area"] == area]
        cached[f"B{r}"] = len(area_items)
        for c, st_name in zip(range(3, 8), ["Done", "In progress", "Not started", "Needs lawyer", "Not applicable"]):
            col = "CDEFG"[c - 3]
            sm.cell(row=r, column=c, value=f'=COUNTIFS({rng_area},$A{r},{rng_status},"{st_name}")')
            cached[f"{col}{r}"] = sum(1 for it in area_items
                                      if (status.get(it["id"], {}).get("status") or "Not started") == st_name)
        sm.cell(row=r, column=8, value=f"=IFERROR(C{r}/(B{r}-G{r}),0)")
        applicable = cached[f"B{r}"] - cached[f"G{r}"]
        cached[f"H{r}"] = round(cached[f"C{r}"] / applicable, 6) if applicable else 0
    total = first + len(areas)
    sm.cell(row=total, column=1, value="All areas")
    for c in range(2, 8):
        col = "BCDEFG"[c - 2]
        sm.cell(row=total, column=c, value=f"=SUM({col}{first}:{col}{total - 1})")
        cached[f"{col}{total}"] = sum(cached[f"{col}{r}"] for r in range(first, total))
    sm.cell(row=total, column=8, value=f"=IFERROR(C{total}/(B{total}-G{total}),0)")
    applicable = cached[f"B{total}"] - cached[f"G{total}"]
    cached[f"H{total}"] = round(cached[f"C{total}"] / applicable, 6) if applicable else 0
    K.xlsx_body_style(sm, first, total, len(head))
    for c in range(1, 9):
        sm.cell(row=total, column=c).font = K.xlsx_font(bold=True)
    for r in range(first, total + 1):
        sm.cell(row=r, column=8).number_format = "0%"
    note_row = total + 2
    sm.cell(row=note_row, column=1, value=("Progress = Done ÷ (Items − Not applicable). 100% means every applicable "
                                           "item has evidence. It does not mean the product is CRA compliant.")).font = \
        K.xlsx_font(italic=True, color="595959")
    sm.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
    sm.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True)
    sm.row_dimensions[note_row].height = 28

    wb.calculation.fullCalcOnLoad = True
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    K.xlsx_inject_cached_values(out, {"Summary": cached})
    done = cached[f"C{total}"]
    K.info(f"Wrote {out}: {len(items)} items in {len(areas)} areas; {done} done, "
           f"{cached[f'D{total}']} in progress, {cached[f'F{total}']} need the lawyer.")


if __name__ == "__main__":
    main()
