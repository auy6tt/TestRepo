#!/usr/bin/env python3
"""Make a CycloneDX SBOM (software bill of materials) for one product release.

Give it one or more project folders and/or existing CycloneDX JSON files:

  * Python folders (requirements.txt, pyproject.toml, poetry.lock or Pipfile.lock)
    are scanned with cyclonedx-py.
  * Node folders (package.json with package-lock.json, or no lockfile) are
    scanned with cyclonedx-npm. yarn.lock and pnpm-lock.yaml need cdxgen.
  * Existing CycloneDX JSON files (for example from the client's own build)
    are read as they are.
  * Components that no tool can see, such as C libraries inside firmware,
    come from a CSV file (--extra-components).

Everything is merged into one product SBOM, checked against the CycloneDX 1.6
schema, and listed in a readable table (.md and .xlsx).

Examples
  python make_sbom.py ../client/backend ../client/app \\
      --product "Sprout S1" --product-version 2.3.0 --product-type device \\
      --supplier "Mossbyte Labs BV" --extra-components firmware-components.csv \\
      --out deliverables/sbom

  python make_sbom.py client-sbom.cdx.json --out deliverables/sbom   # list an existing SBOM
"""

from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import json
import re
import shutil
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import kitlib as K  # noqa: E402

PART_PROP = f"{K.KIT_NAME}:part"
SOURCE_PROP = f"{K.KIT_NAME}:source"
NOTES_PROP = f"{K.KIT_NAME}:notes"
CYCLONEDX_TYPES = {"application", "framework", "library", "container", "platform", "operating-system",
                   "device", "device-driver", "firmware", "file", "machine-learning-model", "data",
                   "cryptographic-asset"}


# ---------------------------------------------------------------------------
# Detecting what a folder contains
# ---------------------------------------------------------------------------

def detect(folder: Path) -> list[str]:
    kinds = []
    if (folder / "package.json").exists():
        kinds.append("node")
    if (folder / "poetry.lock").exists():
        kinds.append("python-poetry")
    elif (folder / "Pipfile.lock").exists():
        kinds.append("python-pipenv")
    elif (folder / "requirements.txt").exists():
        kinds.append("python-requirements")
    elif (folder / "pyproject.toml").exists() and pyproject_dependencies(folder / "pyproject.toml") is not None:
        kinds.append("python-pyproject")
    return kinds


def pyproject_data(path: Path) -> dict:
    try:
        import tomllib
    except ModuleNotFoundError:  # Python 3.10
        import tomli as tomllib  # type: ignore
    with open(path, "rb") as fh:
        return tomllib.load(fh)


def pyproject_dependencies(path: Path):
    project = pyproject_data(path).get("project", {})
    if "dependencies" not in project:
        return None
    return list(project.get("dependencies") or [])


def normalise_pkg(name: str) -> str:
    return re.sub(r"[-_.]+", "-", name).lower()


def requirement_names(req_file: Path, seen: set | None = None) -> list[str]:
    """Top-level package names in a requirements file (follows -r includes)."""
    seen = seen or set()
    if req_file in seen or not req_file.exists():
        return []
    seen.add(req_file)
    names = []
    for raw in req_file.read_text(encoding="utf-8").splitlines():
        line = raw.split(" #")[0].strip()
        if not line or line.startswith("#"):
            continue
        include = re.match(r"^(-r|--requirement)\s+(\S+)", line)
        if include:
            names += requirement_names((req_file.parent / include.group(2)).resolve(), seen)
            continue
        if line.startswith("-"):
            continue
        match = re.match(r"^([A-Za-z0-9][A-Za-z0-9._-]*)", line)
        if match:
            names.append(match.group(1))
    return names


def unpinned_requirements(req_file: Path) -> list[str]:
    loose = []
    for raw in req_file.read_text(encoding="utf-8").splitlines():
        line = raw.split(" #")[0].strip()
        if not line or line.startswith(("#", "-")):
            continue
        if "==" not in line and "@" not in line:
            loose.append(line)
    return loose


# ---------------------------------------------------------------------------
# Running the SBOM tools
# ---------------------------------------------------------------------------

def install_and_scan(tool: str, folder: Path, req_file: Path, common: list, pyproject_opt: list,
                     no_deps: bool = False) -> None:
    """Install the requirements into a throwaway environment (without pip itself), then list what is there."""
    env_dir = K.make_temp_dir("cra-pyenv-")
    try:
        K.run([sys.executable, "-m", "venv", "--without-pip", str(env_dir)])
        target_python = env_dir / "bin" / "python"
        cmd = [sys.executable, "-m", "pip", "--python", str(target_python), "install",
               "--disable-pip-version-check", "--no-input", "-q"]
        if no_deps:
            cmd.append("--no-deps")
        K.run([*cmd, "-r", str(req_file)], cwd=folder)
        K.run([tool, "environment", *common, *pyproject_opt, str(target_python)])
    finally:
        shutil.rmtree(env_dir, ignore_errors=True)


def poetry_lock_requirements(folder: Path, include_dev: bool) -> list[str]:
    """Pinned requirement lines from poetry.lock (main group unless include_dev)."""
    try:
        import tomllib
    except ModuleNotFoundError:  # Python 3.10
        import tomli as tomllib  # type: ignore
    with open(folder / "poetry.lock", "rb") as fh:
        data = tomllib.load(fh)
    lines = []
    for pkg in data.get("package", []):
        groups = pkg.get("groups") or [pkg.get("category", "main")]
        if not include_dev and "main" not in groups:
            continue
        marker = pkg.get("markers")
        if isinstance(marker, dict):
            marker = marker.get("main") or next(iter(marker.values()), None)
        line = f"{pkg['name']}=={pkg['version']}"
        lines.append(f"{line} ; {marker}" if marker else line)
    return lines


def poetry_direct_names(pyproject: Path) -> list[str]:
    if not pyproject.exists():
        return []
    data = pyproject_data(pyproject)
    names = []
    for dep in data.get("project", {}).get("dependencies", []) or []:
        match = re.match(r"^\s*([A-Za-z0-9][A-Za-z0-9._-]*)", dep)
        if match:
            names.append(match.group(1))
    names += [n for n in (data.get("tool", {}).get("poetry", {}).get("dependencies", {}) or {}) if n.lower() != "python"]
    return names


def sbom_python(folder: Path, kind: str, out_file: Path, mode: str, include_dev: bool, notes: list[str]) -> dict:
    tool = K.find_python_tool("cyclonedx-py")
    if not tool:
        K.fail("cyclonedx-py not found. Run scripts/setup.sh and use the kit's Python "
               "(see README: 'Set up the tools').")
    common = ["--sv", "1.6", "--of", "JSON", "-o", str(out_file)]
    pyproject = folder / "pyproject.toml"
    project_meta = pyproject_data(pyproject).get("project", {}) if pyproject.exists() else {}
    poetry_meta = pyproject_data(pyproject).get("tool", {}).get("poetry") if pyproject.exists() else None
    # cyclonedx-py reads the root component from pyproject.toml, but stops with an error when a
    # [tool.poetry] table has no name (newer Poetry projects). Then we describe the root ourselves.
    use_pyproject = bool(project_meta) and (poetry_meta is None or "name" in poetry_meta)
    pyproject_opt = ["--pyproject", str(pyproject)] if use_pyproject else []
    direct_names: list[str] | None = None

    if kind == "python-poetry":
        lock_lines = poetry_lock_requirements(folder, include_dev)
        done = False
        try:
            K.run([tool, "poetry", *common, *([] if include_dev else ["--no-dev"]), str(folder)])
            done = bool(K.load_json(out_file).get("components")) or not lock_lines
        except (RuntimeError, ValueError):
            pass
        if not done:
            # Newer Poetry projects (dependencies under [project]) are not read by cyclonedx-py's
            # poetry mode, so install exactly what poetry.lock pins and scan that instead.
            K.info("  cyclonedx-py could not read this Poetry project; using the versions pinned in poetry.lock")
            req_file = K.make_temp_dir("cra-req-") / "requirements.txt"
            req_file.write_text("\n".join(lock_lines) + "\n", encoding="utf-8")
            direct_names = poetry_direct_names(pyproject)
            try:
                install_and_scan(tool, folder, req_file, common, pyproject_opt, no_deps=True)
            except RuntimeError as exc:
                K.warn(f"Could not install {folder.name}'s locked packages for scanning; listing them from "
                       f"poetry.lock only.\n{exc}")
                K.run([tool, "requirements", *common, *pyproject_opt, str(req_file)])
                notes.append(f"{folder.name}: packages listed from poetry.lock without installing them, so "
                             "the SBOM shows which packages are used but not which one pulls in which.")
    elif kind == "python-pipenv":
        K.run([tool, "pipenv", *common, *pyproject_opt, *(["--dev"] if include_dev else []), str(folder)])
        if (folder / "Pipfile").exists():
            pipfile = pyproject_data(folder / "Pipfile")
            direct_names = list(pipfile.get("packages", {}) or {})
            if include_dev:
                direct_names += list(pipfile.get("dev-packages", {}) or {})
    else:
        if kind == "python-requirements":
            req_file = folder / "requirements.txt"
            loose = unpinned_requirements(req_file)
            if loose:
                notes.append(f"{folder.name}: requirements.txt has unpinned lines ({', '.join(loose[:5])}). "
                             "Versions were resolved today and may differ from what shipped. "
                             "Ask the client for 'pip freeze' output from the release build.")
        else:
            deps = pyproject_dependencies(pyproject) or []
            req_file = K.make_temp_dir("cra-req-") / "requirements.txt"
            req_file.write_text("\n".join(deps) + "\n", encoding="utf-8")
            notes.append(f"{folder.name}: no lockfile, so versions come from today's resolution of "
                         "pyproject.toml. Ask the client which versions shipped.")
        direct_names = requirement_names(req_file)

        chosen_mode = mode
        if mode == "environment":
            try:
                install_and_scan(tool, folder, req_file, common, pyproject_opt)
            except RuntimeError as exc:
                K.warn(f"Could not install {folder.name}'s requirements into a test environment, so "
                       f"falling back to reading the file only.\n{exc}")
                notes.append(f"{folder.name}: dependencies could not be installed for scanning (private "
                             "packages or build errors?), so only the packages named in requirements.txt "
                             "are listed. Indirect dependencies are missing.")
                mode = "requirements"
        if mode == "requirements":
            K.run([tool, "requirements", *common, *pyproject_opt, str(req_file)])
            if chosen_mode == "requirements":
                notes.append(f"{folder.name}: 'requirements' mode lists only what the file names. "
                             "Use the default mode to include indirect dependencies.")

    sbom = K.load_json(out_file)
    meta = sbom.setdefault("metadata", {})
    if not meta.get("component"):
        # Describe the project from pyproject.toml [project], or by its folder name.
        root = {"type": "application", "name": project_meta.get("name") or folder.name, "bom-ref": "root-component"}
        for key in ("version", "description"):
            if isinstance(project_meta.get(key), str):
                root[key] = project_meta[key]
        meta["component"] = root
    if direct_names is not None:
        set_python_root_dependencies(sbom, direct_names)
    return sbom


def set_python_root_dependencies(sbom: dict, direct_names: list[str]) -> None:
    """cyclonedx-py does not know which packages the project asks for directly; add that."""
    wanted = {normalise_pkg(n) for n in direct_names}
    root = sbom.setdefault("metadata", {}).get("component")
    if not root:
        return
    root_ref = root.get("bom-ref") or "root-component"
    root["bom-ref"] = root_ref
    refs = [c["bom-ref"] for c in sbom.get("components", []) if normalise_pkg(c.get("name", "")) in wanted]
    deps = sbom.setdefault("dependencies", [])
    for entry in deps:
        if entry.get("ref") == root_ref:
            entry["dependsOn"] = sorted(set(entry.get("dependsOn", [])) | set(refs))
            return
    deps.append({"ref": root_ref, "dependsOn": sorted(refs)})


def sbom_node(folder: Path, out_file: Path, include_dev: bool, node_tool: str, notes: list[str]) -> tuple[dict, str]:
    has_npm_lock = (folder / "package-lock.json").exists() or (folder / "npm-shrinkwrap.json").exists()
    has_other_lock = (folder / "yarn.lock").exists() or (folder / "pnpm-lock.yaml").exists()
    use_cdxgen = node_tool == "cdxgen" or (node_tool == "auto" and has_other_lock and not has_npm_lock)

    if use_cdxgen:
        cdxgen = K.find_npm_tool("cdxgen")
        if not cdxgen:
            K.fail(f"{folder} uses yarn or pnpm. Install cdxgen with scripts/setup.sh, or ask the client "
                   "for an npm lockfile or a CycloneDX SBOM from their build.")
        cmd = [cdxgen, "-t", "js", "--spec-version", "1.6", "-o", str(out_file), "--no-recurse",
               "--no-install-deps", "--json-pretty", "--no-validate"]
        if not include_dev:
            cmd += ["--required-only", "--no-auto-compositions"]
        K.run([*cmd, str(folder)], cwd=folder)
        return K.load_json(out_file), "cdxgen"

    tool = K.find_npm_tool("cyclonedx-npm")
    npx = shutil.which("npx")
    if tool:
        base = [tool]
    elif npx:
        K.warn("cyclonedx-npm is not installed locally; using npx to fetch it (slower).")
        base = [npx, "--yes", "@cyclonedx/cyclonedx-npm"]
    else:
        K.fail("cyclonedx-npm not found. Run scripts/setup.sh first.")

    work = folder
    lock_only = not (folder / "node_modules").exists()
    if not has_npm_lock:
        work = K.make_temp_dir("cra-npm-")
        shutil.copy(folder / "package.json", work / "package.json")
        if (folder / ".npmrc").exists():
            shutil.copy(folder / ".npmrc", work / ".npmrc")
        K.run(["npm", "install", "--package-lock-only", "--ignore-scripts", "--no-audit", "--no-fund"], cwd=work)
        lock_only = True
        notes.append(f"{folder.name}: no package-lock.json, so versions come from today's resolution of "
                     "package.json. Ask the client for the lockfile of the shipped release.")

    cmd = [*base, "--output-format", "JSON", "--spec-version", "1.6", "--flatten-components",
           "--mc-type", "application", "--output-file", str(out_file)]
    if lock_only:
        cmd.append("--package-lock-only")
    if not include_dev:
        cmd += ["--omit", "dev"]
    cmd += ["--", str(work / "package.json")]
    try:
        K.run(cmd, cwd=work)
    except RuntimeError as exc:
        K.warn(f"npm reported problems in {folder.name}; retrying with --ignore-npm-errors.\n{exc}")
        K.run([*cmd[:-2], "--ignore-npm-errors", *cmd[-2:]], cwd=work)
        notes.append(f"{folder.name}: npm reported problems with the dependency tree "
                     "(missing or invalid packages). Check the lockfile with the client.")
    if work != folder:
        shutil.rmtree(work, ignore_errors=True)
    return K.load_json(out_file), "cyclonedx-npm"


# ---------------------------------------------------------------------------
# Components listed by hand (firmware libraries, SDKs, RTOS ...)
# ---------------------------------------------------------------------------

def licence_entry(text: str) -> list[dict]:
    text = (text or "").strip()
    if not text:
        return []
    try:
        from cyclonedx.spdx import is_supported_id
    except ImportError:  # pragma: no cover
        def is_supported_id(_value):
            return False
    if is_supported_id(text):
        return [{"license": {"id": text}}]
    if re.search(r"\s(AND|OR|WITH)\s", text):
        return [{"expression": text}]
    return [{"license": {"name": text}}]


def read_extra_components(csv_path: Path) -> dict[str, dict]:
    """Return {part_name: {"root": component, "components": [...]}} from a CSV file."""
    parts: dict[str, dict] = {}
    with open(csv_path, newline="", encoding="utf-8-sig") as fh:
        for row in csv.DictReader(fh):
            row = {k.strip().lower(): (v or "").strip() for k, v in row.items() if k}
            if not row.get("name"):
                continue
            part = row.get("part") or csv_path.stem
            ctype = (row.get("type") or "library").lower()
            if ctype not in CYCLONEDX_TYPES:
                K.warn(f"{csv_path.name}: unknown type '{ctype}' for {row['name']}; using 'library'.")
                ctype = "library"
            comp = {"type": ctype, "name": row["name"]}
            if row.get("version"):
                comp["version"] = row["version"]
            comp["bom-ref"] = f"{part}|{row['name']}@{row.get('version') or 'unknown'}"
            if row.get("supplier"):
                comp["supplier"] = {"name": row["supplier"]}
            lic = licence_entry(row.get("licence") or row.get("license", ""))
            if lic:
                comp["licenses"] = lic
            if row.get("purl"):
                comp["purl"] = row["purl"]
            if row.get("cpe"):
                comp["cpe"] = row["cpe"]
            if row.get("website"):
                comp["externalReferences"] = [{"type": "website", "url": row["website"]}]
            props = [{"name": SOURCE_PROP, "value": f"listed by hand in {csv_path.name}"}]
            if row.get("notes"):
                props.append({"name": NOTES_PROP, "value": row["notes"]})
            comp["properties"] = props
            entry = parts.setdefault(part, {"root": None, "components": []})
            if row["name"] == part:
                entry["root"] = comp
            else:
                entry["components"].append(comp)
    return parts


# ---------------------------------------------------------------------------
# Merging
# ---------------------------------------------------------------------------

def flatten_components(components: list[dict]) -> list[dict]:
    out = []
    for comp in components or []:
        nested = comp.pop("components", None)
        out.append(comp)
        if nested:
            out.extend(flatten_components(nested))
    return out


def set_prop(comp: dict, name: str, value: str) -> None:
    props = comp.setdefault("properties", [])
    for prop in props:
        if prop.get("name") == name:
            if value not in prop["value"].split(", "):
                prop["value"] = f"{prop['value']}, {value}"
            return
    props.append({"name": name, "value": value})


def get_prop(comp: dict, name: str) -> str:
    for prop in comp.get("properties", []) or []:
        if prop.get("name") == name:
            return prop.get("value", "")
    return ""


def merge_parts(parts: list[dict], product: dict, tools: list[dict], prepared_by: str | None) -> dict:
    """parts: [{"name", "sbom"}] -> one product SBOM."""
    components: dict[str, dict] = {}
    deps: dict[str, set] = {}
    part_roots = []

    def add_dep(ref, depends_on):
        deps.setdefault(ref, set()).update(depends_on)

    for part in parts:
        sbom = copy.deepcopy(part["sbom"])
        name = part["name"]
        root = sbom.get("metadata", {}).get("component") or {"type": "application", "name": name}
        root.pop("components", None)
        root.setdefault("version", product.get("version") or "unversioned")
        old_root_ref = root.get("bom-ref")
        new_root_ref = f"part:{name}@{root.get('version')}"
        n = 2
        while new_root_ref in components:
            new_root_ref = f"part:{name}@{root.get('version')}#{n}"
            n += 1
        ref_map = {}
        if old_root_ref:
            ref_map[old_root_ref] = new_root_ref
        root["bom-ref"] = new_root_ref
        if root.get("type") not in CYCLONEDX_TYPES:
            root["type"] = "application"
        set_prop(root, PART_PROP, name)
        components[new_root_ref] = root
        part_roots.append(new_root_ref)

        for comp in flatten_components(sbom.get("components", [])):
            old_ref = comp.get("bom-ref") or f"{comp.get('name')}@{comp.get('version')}"
            existing = components.get(old_ref)
            if existing is not None and existing.get("purl") and existing.get("purl") == comp.get("purl"):
                ref_map[old_ref] = old_ref          # same package used by two parts: list it once
                set_prop(existing, PART_PROP, name)
                continue
            new_ref = old_ref
            if existing is not None:
                new_ref = f"{name}|{old_ref}"
            ref_map[old_ref] = new_ref
            comp["bom-ref"] = new_ref
            set_prop(comp, PART_PROP, name)
            components[new_ref] = comp

        for entry in sbom.get("dependencies", []) or []:
            ref = ref_map.get(entry.get("ref"), entry.get("ref"))
            add_dep(ref, {ref_map.get(d, d) for d in entry.get("dependsOn", []) or []})
        deps.setdefault(new_root_ref, set())

    product_ref = "product"
    product_component = {
        "type": product["type"],
        "bom-ref": product_ref,
        "name": product["name"],
        "version": product.get("version") or "unversioned",
    }
    if product.get("supplier"):
        product_component["supplier"] = {"name": product["supplier"]}
        product_component["manufacturer"] = {"name": product["supplier"]}
    if product.get("description"):
        product_component["description"] = product["description"]
    add_dep(product_ref, set(part_roots))

    known = set(components) | {product_ref}
    dependencies = []
    for ref in sorted(deps):
        if ref not in known:
            continue
        dependencies.append({"ref": ref, "dependsOn": sorted(d for d in deps[ref] if d in known and d != ref)})

    metadata = {
        "timestamp": K.now_utc_iso(),
        "tools": {"components": tools},
        "component": product_component,
    }
    if product.get("supplier"):
        metadata["supplier"] = {"name": product["supplier"]}
        metadata["manufacturer"] = {"name": product["supplier"]}
    if prepared_by:
        metadata["authors"] = [{"name": prepared_by}]
    metadata["properties"] = [{
        "name": f"{K.KIT_NAME}:confirmation",
        "value": "Draft until the manufacturer confirms this list matches the shipped release.",
    }]
    return {
        "$schema": "http://cyclonedx.org/schema/bom-1.6.schema.json",
        "bomFormat": "CycloneDX",
        "specVersion": "1.6",
        "serialNumber": f"urn:uuid:{uuid.uuid4()}",
        "version": 1,
        "metadata": metadata,
        "components": list(components.values()),
        "dependencies": dependencies,
    }


def validate_sbom(text: str, spec: str) -> str | None:
    try:
        from cyclonedx.schema import SchemaVersion
        from cyclonedx.validation.json import JsonStrictValidator
    except ImportError:
        return "skipped (cyclonedx-python-lib not installed)"
    version = getattr(SchemaVersion, "V" + spec.replace(".", "_"), None)
    if version is None:
        return f"skipped (unknown spec version {spec})"
    error = JsonStrictValidator(version).validate_str(text)
    return None if error is None else str(error)


# ---------------------------------------------------------------------------
# Readable component table
# ---------------------------------------------------------------------------

ECOSYSTEMS = {"pypi": "PyPI (Python)", "npm": "npm (Node)", "github": "GitHub source", "generic": "Other",
              "maven": "Maven (Java)", "golang": "Go", "cargo": "Cargo (Rust)", "nuget": "NuGet (.NET)",
              "gem": "RubyGems", "composer": "Composer (PHP)", "conan": "Conan (C/C++)", "deb": "Debian package"}


def licence_text(comp: dict) -> str:
    names = []
    for entry in comp.get("licenses", []) or []:
        if "expression" in entry:
            names.append(entry["expression"])
            continue
        lic = entry.get("license", {})
        value = lic.get("id") or lic.get("name") or ""
        value = value.replace("License :: OSI Approved :: ", "").replace("License :: ", "")
        if value == "OSI Approved":
            value = "OSI-approved (name not stated)"
        if value:
            names.append(value)
    spdx_like = [n for n in names if re.match(r"^[A-Za-z0-9.+-]+$", n) and " " not in n]
    chosen = spdx_like or names
    seen = []
    for n in chosen:
        if n not in seen:
            seen.append(n)
    return " / ".join(seen) if seen else "Not stated"


def supplier_text(comp: dict) -> str:
    for key in ("supplier", "manufacturer"):
        if isinstance(comp.get(key), dict) and comp[key].get("name"):
            return comp[key]["name"]
    for key in ("author", "publisher"):
        if comp.get(key):
            return comp[key]
    authors = comp.get("authors") or []
    if authors and isinstance(authors[0], dict):
        return authors[0].get("name", "")
    return ""


def ecosystem_of(comp: dict) -> str:
    purl = comp.get("purl") or ""
    m = re.match(r"^pkg:([a-z0-9.+-]+)/", purl)
    if m:
        return ECOSYSTEMS.get(m.group(1), m.group(1))
    if get_prop(comp, SOURCE_PROP).startswith("listed by hand"):
        return "Listed by hand"
    return "Not identified"


def display_name(comp: dict) -> str:
    group = comp.get("group")
    return f"{group}/{comp['name']}" if group else comp.get("name", "")


def table_rows(bom: dict) -> list[dict]:
    comps = {c["bom-ref"]: c for c in bom.get("components", []) if c.get("bom-ref")}
    deps = {d["ref"]: set(d.get("dependsOn", [])) for d in bom.get("dependencies", [])}
    root = bom.get("metadata", {}).get("component", {}) or {}
    product_ref = root.get("bom-ref")
    part_refs = [r for r in deps.get(product_ref, set()) if r in comps and get_prop(comps[r], PART_PROP)]
    single_project = not part_refs
    if single_project:  # a single-project SBOM: its root is the metadata component
        part_refs = [product_ref] if product_ref else []
    direct = set()
    for ref in part_refs:
        direct |= deps.get(ref, set())
    rows = []
    for ref, comp in comps.items():
        if ref in part_refs:
            continue
        part = get_prop(comp, PART_PROP) or (root.get("name", "") if single_project else "")
        rows.append({
            "part": part,
            "name": display_name(comp),
            "version": comp.get("version", ""),
            "type": comp.get("type", ""),
            "ecosystem": ecosystem_of(comp),
            "direct": "Yes" if (ref in direct or get_prop(comp, SOURCE_PROP)) else "No",
            "licence": licence_text(comp),
            "supplier": supplier_text(comp),
            "purl": comp.get("purl", ""),
            "notes": get_prop(comp, NOTES_PROP),
        })
    rows.sort(key=lambda r: (r["part"], r["direct"] != "Yes", r["name"].lower()))
    return rows


def part_summary(bom: dict) -> list[dict]:
    comps = {c["bom-ref"]: c for c in bom.get("components", []) if c.get("bom-ref")}
    deps = {d["ref"]: d.get("dependsOn", []) for d in bom.get("dependencies", [])}
    product_ref = bom.get("metadata", {}).get("component", {}).get("bom-ref")
    out = []
    for ref in deps.get(product_ref, []):
        comp = comps.get(ref)
        if comp and get_prop(comp, PART_PROP):
            out.append({"name": comp.get("name"), "version": comp.get("version", ""), "type": comp.get("type", "")})
    return out


COPYLEFT = re.compile(r"\b(A?GPL|LGPL|MPL|EPL|CDDL|EUPL|OSL|CC-BY-SA)", re.I)


def write_tables(bom: dict, rows: list[dict], sbom_path: Path, out_dir: Path, base: str,
                 notes: list[str], prepared_by: str | None) -> tuple[Path, Path]:
    meta = bom.get("metadata", {})
    product = meta.get("component", {})
    sha256 = hashlib.sha256(sbom_path.read_bytes()).hexdigest()
    parts = part_summary(bom)
    by_part: dict[str, list[dict]] = {}
    for row in rows:
        by_part.setdefault(row["part"] or "Components", []).append(row)
    licence_counts: dict[str, int] = {}
    for row in rows:
        licence_counts[row["licence"]] = licence_counts.get(row["licence"], 0) + 1
    unknown = licence_counts.get("Not stated", 0)
    copyleft = sorted({r["licence"] for r in rows if COPYLEFT.search(r["licence"])})

    # ---- Markdown -------------------------------------------------------
    md = [f"# Software bill of materials: {product.get('name', '')} {product.get('version', '')}", ""]
    md += [
        f"- **Product:** {product.get('name', '')} (version {product.get('version', '')})",
        f"- **Manufacturer:** {meta.get('manufacturer', {}).get('name', 'not stated')}",
        f"- **SBOM file:** `{sbom_path.name}` (CycloneDX {bom.get('specVersion')}, JSON)",
        f"- **SHA-256 of the SBOM file:** `{sha256}`",
        f"- **Generated:** {K.today()}" + (f" by {prepared_by}" if prepared_by else ""),
        f"- **Components listed:** {len(rows)}",
        "",
        "> **Please confirm.** This list was produced from the dependency files you provided. "
        "Check that it matches the release you actually ship, including firmware libraries and "
        "anything copied into the code by hand. Tell us about anything missing or wrong, then "
        "confirm in writing. The SBOM stays a draft until you do.",
        "",
    ]
    if parts:
        md += ["## Parts of the product", "", "| Part | Version | Type | Components |", "|---|---|---|---|"]
        for p in parts:
            md.append(f"| {K.md_escape(p['name'])} | {K.md_escape(p['version'])} | {p['type']} | "
                      f"{len(by_part.get(p['name'], []))} |")
        md.append("")
    md += ["## Licences at a glance", "",
           "Licence names come from package metadata. They can be incomplete or wrong, and this "
           "is not a licence review.", "", "| Licence | Components |", "|---|---|"]
    for lic, count in sorted(licence_counts.items(), key=lambda kv: (-kv[1], kv[0])):
        md.append(f"| {K.md_escape(lic)} | {count} |")
    md.append("")
    if unknown or copyleft:
        md.append("Worth a look:")
        md.append("")
        if unknown:
            md.append(f"- {unknown} component(s) have no licence in their metadata.")
        if copyleft:
            md.append(f"- Copyleft-style licences found ({', '.join(copyleft)}). These can come with "
                      "conditions when you distribute the product. Ask your lawyer if unsure.")
        md.append("")
    if notes:
        md += ["## Notes from the scan", ""] + [f"- {n}" for n in notes] + [""]
    md += ["## Components", ""]
    for part, prow in by_part.items():
        md += [f"### {part}", "", "| # | Component | Version | Type | Direct? | Licence | Package URL (purl) |",
               "|---|---|---|---|---|---|---|"]
        for n, row in enumerate(prow, start=1):
            md.append(f"| {n} | {K.md_escape(row['name'])} | {K.md_escape(row['version'])} | {row['type']} | "
                      f"{row['direct']} | {K.md_escape(row['licence'])} | {K.md_escape(row['purl'])} |")
        md.append("")
    md += ["## How to read this", "",
           "- **Direct?** Yes means your project asks for this component itself. No means it comes in "
           "through another component.",
           "- **Package URL (purl)** is the standard name vulnerability databases use to look a component up.",
           "- The full machine-readable version is the `.cdx.json` file. Keep one per release, together "
           "with the release itself.", ""]
    md_path = out_dir / f"{base}-components.md"
    K.write_text(md_path, "\n".join(md))

    # ---- Excel -----------------------------------------------------------
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Components"
    headers = ["Part", "Component", "Version", "Type", "Ecosystem", "Direct?", "Licence",
               "Supplier / author", "Package URL (purl)", "Notes"]
    K.xlsx_header(ws, 1, headers, [18, 30, 12, 13, 16, 9, 22, 26, 46, 40])
    for r, row in enumerate(rows, start=2):
        values = [row["part"], row["name"], row["version"], row["type"], row["ecosystem"], row["direct"],
                  row["licence"], row["supplier"], row["purl"], row["notes"]]
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)
    K.xlsx_body_style(ws, 2, len(rows) + 1, len(headers))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:J{len(rows) + 1}"

    info = wb.create_sheet("About", 0)
    nxt = K.xlsx_title(info, f"SBOM: {product.get('name', '')} {product.get('version', '')}",
                       "Readable copy of the CycloneDX SBOM. The .cdx.json file is the official version.")
    facts = [
        ("Product", product.get("name", "")),
        ("Version", product.get("version", "")),
        ("Manufacturer", meta.get("manufacturer", {}).get("name", "")),
        ("SBOM file", sbom_path.name),
        ("SBOM format", f"CycloneDX {bom.get('specVersion')} JSON"),
        ("SHA-256 of SBOM file", sha256),
        ("Generated", K.today()),
        ("Prepared by", prepared_by or ""),
        ("Components listed", len(rows)),
        ("Components without licence data", unknown),
    ]
    for label, value in facts:
        info.cell(row=nxt, column=1, value=label).font = K.xlsx_font(bold=True)
        info.cell(row=nxt, column=2, value=value).font = K.xlsx_font()
        nxt += 1
    nxt += 1
    info.cell(row=nxt, column=1, value="Parts").font = K.xlsx_font(bold=True, size=11, color=K.HEADER_FILL)
    nxt += 1
    for p in parts:
        info.cell(row=nxt, column=1, value=p["name"]).font = K.xlsx_font()
        info.cell(row=nxt, column=2, value=f"version {p['version']}, {p['type']}, "
                  f"{len(by_part.get(p['name'], []))} components").font = K.xlsx_font()
        nxt += 1
    nxt += 1
    info.cell(row=nxt, column=1, value="Client confirmation").font = K.xlsx_font(bold=True, size=11, color=K.HEADER_FILL)
    nxt += 1
    info.cell(row=nxt, column=1, value=("Fill in the yellow cells when you have checked that this list matches "
                                        "the release you ship.")).font = K.xlsx_font(italic=True)
    nxt += 1
    yellow = PatternFill("solid", fgColor=K.INPUT_FILL)
    choice = DataValidation(type="list", formula1='"Yes - matches what ships,No - see notes"', allow_blank=True)
    info.add_data_validation(choice)
    for label in ("Matches the shipped release?", "Checked by (name, role)", "Date", "Missing or wrong items"):
        info.cell(row=nxt, column=1, value=label).font = K.xlsx_font(bold=True)
        cell = info.cell(row=nxt, column=2)
        cell.fill = yellow
        cell.font = K.xlsx_font()
        if label.startswith("Matches"):
            choice.add(cell)
        nxt += 1
    if notes:
        nxt += 1
        info.cell(row=nxt, column=1, value="Notes from the scan").font = K.xlsx_font(bold=True, size=11, color=K.HEADER_FILL)
        nxt += 1
        for note in notes:
            info.cell(row=nxt, column=1, value=note).font = K.xlsx_font()
            nxt += 1
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 80

    lic_ws = wb.create_sheet("Licences")
    K.xlsx_header(lic_ws, 1, ["Licence (as stated in metadata)", "Components"], [44, 14])
    for r, (lic, count) in enumerate(sorted(licence_counts.items(), key=lambda kv: (-kv[1], kv[0])), start=2):
        lic_ws.cell(row=r, column=1, value=lic)
        lic_ws.cell(row=r, column=2, value=count)
    K.xlsx_body_style(lic_ws, 2, len(licence_counts) + 1, 2)
    xlsx_path = out_dir / f"{base}-components.xlsx"
    wb.save(xlsx_path)
    return md_path, xlsx_path


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="*", help="project folders and/or CycloneDX JSON files")
    ap.add_argument("--out", required=True, help="output folder")
    ap.add_argument("--product", help="product name (default: from --client or the single project)")
    ap.add_argument("--product-version", help="release version of the product")
    ap.add_argument("--product-type", choices=["device", "application", "firmware"],
                    help="device = hardware product (default when firmware parts are listed)")
    ap.add_argument("--supplier", help="manufacturer / supplier name")
    ap.add_argument("--prepared-by", help="who prepared the SBOM (you)")
    ap.add_argument("--client", help="client.json to read defaults from")
    ap.add_argument("--extra-components", action="append", default=[],
                    help="CSV of components no tool can see (repeatable). See templates/extra-components.csv")
    ap.add_argument("--include-dev", action="store_true", help="also list development-only dependencies")
    ap.add_argument("--python-mode", choices=["environment", "requirements"], default="environment",
                    help="environment (default) installs into a throwaway folder to find indirect "
                         "dependencies; requirements only reads the file")
    ap.add_argument("--node-tool", choices=["auto", "cyclonedx-npm", "cdxgen"], default="auto")
    args = ap.parse_args()
    if not args.inputs and not args.extra_components:
        ap.error("give at least one project folder, SBOM file or --extra-components CSV")

    client = K.load_json(args.client) if args.client else {}
    first_product = (client.get("products") or [{}])[0]
    out_dir = Path(args.out)
    parts_dir = out_dir / "parts"
    parts_dir.mkdir(parents=True, exist_ok=True)
    notes: list[str] = []
    parts: list[dict] = []
    tools_used: dict[str, str] = {}

    for raw in args.inputs:
        path = Path(raw).expanduser().resolve()
        if path.is_file():
            if path.suffix.lower() != ".json":
                K.fail(f"{raw}: only CycloneDX JSON files can be read. Convert XML to JSON first.")
            sbom = K.load_json(path)
            if sbom.get("bomFormat") != "CycloneDX":
                K.fail(f"{raw} is not a CycloneDX SBOM (bomFormat is missing).")
            name = (sbom.get("metadata", {}).get("component", {}) or {}).get("name") or path.stem
            parts.append({"name": name, "sbom": sbom, "source": str(path)})
            K.info(f"Read existing SBOM {path.name} ({len(flatten_components(copy.deepcopy(sbom.get('components', []))))} components)")
            continue
        if not path.is_dir():
            K.fail(f"{raw} does not exist.")
        kinds = detect(path)
        if not kinds:
            K.fail(f"{raw}: no package.json, requirements.txt, pyproject.toml, poetry.lock or Pipfile.lock found.")
        for kind in kinds:
            name = path.name if len(kinds) == 1 else f"{path.name}-{kind.split('-')[0]}"
            out_file = parts_dir / f"{slugify_part(name)}.cdx.json"
            K.info(f"Scanning {path} as a {kind.replace('-', ' ')} project")
            if kind == "node":
                sbom, tool = sbom_node(path, out_file, args.include_dev, args.node_tool, notes)
                tools_used[tool] = tools_used.get(tool) or (
                    K.tool_version([K.find_npm_tool(tool) or tool, "--version"]))
            else:
                sbom = sbom_python(path, kind, out_file, args.python_mode, args.include_dev, notes)
                tools_used["cyclonedx-py"] = K.tool_version([K.find_python_tool("cyclonedx-py"), "--version"])
            K.write_json(out_file, sbom)
            parts.append({"name": name, "sbom": sbom, "source": str(path)})
            K.info(f"  -> {out_file} ({len(sbom.get('components', []))} components)")

    for csv_file in args.extra_components:
        extra = read_extra_components(Path(csv_file))
        for part_name, entry in extra.items():
            root = entry["root"] or {"type": "firmware", "name": part_name,
                                     "version": args.product_version or first_product.get("current_version")}
            root = dict(root)
            root.setdefault("bom-ref", f"{part_name}-root")
            components = entry["components"]
            refs = [c["bom-ref"] for c in components]
            sbom = {"metadata": {"component": root}, "components": components,
                    "dependencies": [{"ref": root["bom-ref"], "dependsOn": refs}]}
            parts.append({"name": part_name, "sbom": sbom, "source": csv_file})
            K.info(f"Added {len(components)} hand-listed components for {part_name} from {csv_file}")

    product_name = args.product or first_product.get("name")
    single = len(parts) == 1 and not product_name
    supplier = args.supplier or K.dig(client, "company.legal_name")
    prepared_by = args.prepared_by or K.dig(client, "engagement.consultant_name")
    version = args.product_version or first_product.get("current_version")
    has_firmware = any((p["sbom"].get("metadata", {}).get("component") or {}).get("type") == "firmware" for p in parts)
    product = {
        "name": product_name or (parts[0]["name"] if parts else "product"),
        "version": version,
        "type": args.product_type or ("device" if has_firmware else "application"),
        "supplier": supplier,
        "description": first_product.get("description") if product_name == first_product.get("name") else None,
    }

    tools = [{"type": "application", "name": K.KIT_NAME, "version": K.KIT_VERSION}]
    for name, ver in tools_used.items():
        tools.append({"type": "application", "name": name, "version": str(ver)[:60]})

    if single:
        bom = parts[0]["sbom"]
        root_version = (bom.get("metadata", {}).get("component") or {}).get("version", "")
        base = slugify_part(f"{parts[0]['name']}-{root_version}" if root_version else parts[0]["name"])
    else:
        bom = merge_parts(parts, product, tools, prepared_by)
        base = slugify_part(f"{product['name']}-{product['version'] or 'unversioned'}")

    sbom_path = out_dir / f"{base}.cdx.json"
    text = json.dumps(bom, indent=2, ensure_ascii=False) + "\n"
    sbom_path.write_text(text, encoding="utf-8")
    problem = validate_sbom(text, str(bom.get("specVersion", "1.6")))
    if problem:
        if problem.startswith("skipped"):
            K.warn(f"Schema check {problem}")
        else:
            K.warn(f"The SBOM does not pass the CycloneDX schema check:\n{problem[:1500]}")
    else:
        K.info(f"Schema check: {sbom_path.name} is valid CycloneDX {bom.get('specVersion')}")

    rows = table_rows(bom)
    md_path, xlsx_path = write_tables(bom, rows, sbom_path, out_dir, base, notes, prepared_by)
    K.info("")
    K.info(f"SBOM:   {sbom_path}")
    K.info(f"Table:  {md_path}")
    K.info(f"        {xlsx_path}")
    K.info(f"Parts:  {', '.join(p['name'] for p in parts)}; {len(rows)} components in total")
    for note in notes:
        K.warn(note)
    K.info("Next: ask the client to confirm the list matches what ships (see the About sheet).")


def slugify_part(text: str) -> str:
    return K.slugify(text).strip(".-") or "sbom"


if __name__ == "__main__":
    main()
