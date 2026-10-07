"""Shared helpers for the CRA readiness kit scripts.

Finding tools, running commands, simple web requests, version sorting,
Markdown-to-Word conversion and spreadsheet styling all live here so the
main scripts stay short and readable.
"""

from __future__ import annotations

import datetime as _dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

KIT_DIR = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = KIT_DIR / "templates"
KIT_NAME = "cra-readiness-kit"
KIT_VERSION = "1.0"
USER_AGENT = f"{KIT_NAME}/{KIT_VERSION}"


# --------------------------------------------------------------------------
# Small utilities
# --------------------------------------------------------------------------

def today() -> str:
    """Today's date (YYYY-MM-DD). Set CRA_TODAY to pin it for repeatable output."""
    return os.environ.get("CRA_TODAY") or _dt.date.today().isoformat()


def now_utc_iso() -> str:
    return _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def slugify(text: str) -> str:
    text = re.sub(r"[^A-Za-z0-9.]+", "-", str(text)).strip("-").lower()
    return re.sub(r"-{2,}", "-", text) or "product"


def info(msg: str) -> None:
    print(msg, flush=True)


def warn(msg: str) -> None:
    print(f"WARNING: {msg}", file=sys.stderr, flush=True)


def fail(msg: str, code: int = 1) -> None:
    print(f"ERROR: {msg}", file=sys.stderr, flush=True)
    sys.exit(code)


def load_json(path: str | Path) -> dict:
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def write_json(path: str | Path, data) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2, ensure_ascii=False)
        fh.write("\n")


def write_text(path: str | Path, text: str) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def dig(data: dict, dotted: str, default=None):
    """Read a nested value: dig(client, 'company.legal_name')."""
    cur = data
    for part in dotted.split("."):
        if isinstance(cur, dict) and part in cur:
            cur = cur[part]
        else:
            return default
    return cur


def md_escape(text) -> str:
    """Make text safe for a Markdown table cell."""
    text = "" if text is None else str(text)
    return text.replace("|", "\\|").replace("\n", " ").strip()


# --------------------------------------------------------------------------
# Finding and running tools
# --------------------------------------------------------------------------

def find_python_tool(name: str) -> str | None:
    """Find a Python command-line tool (cyclonedx-py, pip-audit).

    Looks next to the Python running this script first (the kit's virtualenv),
    then on PATH.
    """
    candidate = Path(sys.executable).parent / name
    if candidate.exists():
        return str(candidate)
    return shutil.which(name)


def npm_tool_dirs() -> list[Path]:
    dirs = []
    env_dir = os.environ.get("CRA_NPM_TOOLS")
    if env_dir:
        dirs.append(Path(env_dir).expanduser())
    # setup.sh puts the npm tools in a folder called "npm" next to the virtualenv.
    dirs.append(Path(sys.prefix).parent / "npm")
    dirs.append(KIT_DIR)
    return dirs


def find_npm_tool(bin_name: str) -> str | None:
    """Find a locally installed npm command (cyclonedx-npm, cdxgen)."""
    for base in npm_tool_dirs():
        candidate = base / "node_modules" / ".bin" / bin_name
        if candidate.exists():
            return str(candidate)
    return shutil.which(bin_name)


def run(cmd: list[str], cwd: str | Path | None = None, env: dict | None = None,
        check: bool = True, timeout: int = 900) -> subprocess.CompletedProcess:
    """Run a command and capture its output. Raises RuntimeError when check fails."""
    shown = " ".join(str(c) for c in cmd)
    info(f"  $ {shown}")
    proc = subprocess.run([str(c) for c in cmd], cwd=str(cwd) if cwd else None,
                          env=env, capture_output=True, text=True, timeout=timeout)
    if check and proc.returncode != 0:
        tail = (proc.stderr or proc.stdout or "").strip()[-2000:]
        raise RuntimeError(f"Command failed ({proc.returncode}): {shown}\n{tail}")
    return proc


def tool_version(cmd: list[str]) -> str:
    """Return the first version number a tool prints for --version."""
    try:
        out = subprocess.run([str(c) for c in cmd], capture_output=True, text=True, timeout=60)
        text = re.sub(r"\x1b\[[0-9;]*m", "", (out.stdout or "") + "\n" + (out.stderr or ""))
        match = re.search(r"\b(\d+\.\d+(?:\.\d+)?)\b", text)
        return match.group(1) if match else "unknown"
    except Exception:  # noqa: BLE001 - version lookups must never stop a run
        return "unknown"


# --------------------------------------------------------------------------
# Web requests (standard library only; honours HTTPS_PROXY and SSL_CERT_FILE)
# --------------------------------------------------------------------------

def http_json(url: str, payload=None, timeout: int = 30):
    """GET (or POST when payload is given) JSON. Returns (data, error_text)."""
    data = None
    headers = {"User-Agent": USER_AGENT, "Accept": "application/json"}
    if payload is not None:
        data = json.dumps(payload).encode("utf-8")
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(url, data=data, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return json.loads(resp.read().decode("utf-8")), None
    except urllib.error.HTTPError as exc:
        return None, f"HTTP {exc.code}"
    except (urllib.error.URLError, TimeoutError, OSError, ValueError) as exc:
        return None, str(getattr(exc, "reason", exc))


# --------------------------------------------------------------------------
# Versions
# --------------------------------------------------------------------------

def _version_key(text: str):
    try:
        from packaging.version import Version
        return (1, Version(str(text)))
    except Exception:  # noqa: BLE001 - fall back to a rough numeric sort
        parts = re.findall(r"\d+|[A-Za-z]+", str(text))
        return (0, tuple((0, int(p)) if p.isdigit() else (1, p) for p in parts))


def version_lt(a: str, b: str) -> bool:
    ka, kb = _version_key(a), _version_key(b)
    if ka[0] != kb[0]:
        # One side could not be parsed: compare the rough form of both.
        ka, kb = _version_key_rough(a), _version_key_rough(b)
        return ka < kb
    return ka[1] < kb[1]


def _version_key_rough(text: str):
    parts = re.findall(r"\d+|[A-Za-z]+", str(text))
    return tuple((0, int(p)) if p.isdigit() else (1, p) for p in parts)


def max_version(versions):
    best = None
    for v in versions:
        if v and (best is None or version_lt(best, v)):
            best = v
    return best


def min_version(versions):
    best = None
    for v in versions:
        if v and (best is None or version_lt(v, best)):
            best = v
    return best


def major_of(version: str) -> str:
    m = re.match(r"\D*(\d+)", str(version or ""))
    return m.group(1) if m else ""


# --------------------------------------------------------------------------
# CVSS v3 base score (used when an advisory only gives a vector string)
# --------------------------------------------------------------------------

_CVSS3_WEIGHTS = {
    "AV": {"N": 0.85, "A": 0.62, "L": 0.55, "P": 0.2},
    "AC": {"L": 0.77, "H": 0.44},
    "UI": {"N": 0.85, "R": 0.62},
    "C": {"H": 0.56, "L": 0.22, "N": 0.0},
    "I": {"H": 0.56, "L": 0.22, "N": 0.0},
    "A": {"H": 0.56, "L": 0.22, "N": 0.0},
}


def _roundup(value: float) -> float:
    as_int = round(value * 100000)
    if as_int % 10000 == 0:
        return as_int / 100000.0
    return (as_int // 10000 + 1) / 10.0


def cvss3_base_score(vector: str) -> float | None:
    """Return the CVSS 3.x base score for a vector like CVSS:3.1/AV:N/AC:L/..."""
    if not vector or not vector.startswith("CVSS:3"):
        return None
    try:
        metrics = dict(part.split(":", 1) for part in vector.split("/")[1:])
        scope_changed = metrics["S"] == "C"
        pr_weights = {"N": 0.85, "L": 0.68 if scope_changed else 0.62, "H": 0.5 if scope_changed else 0.27}
        iss = 1 - ((1 - _CVSS3_WEIGHTS["C"][metrics["C"]]) * (1 - _CVSS3_WEIGHTS["I"][metrics["I"]])
                   * (1 - _CVSS3_WEIGHTS["A"][metrics["A"]]))
        if scope_changed:
            impact = 7.52 * (iss - 0.029) - 3.25 * (iss - 0.02) ** 15
        else:
            impact = 6.42 * iss
        exploitability = (8.22 * _CVSS3_WEIGHTS["AV"][metrics["AV"]] * _CVSS3_WEIGHTS["AC"][metrics["AC"]]
                          * pr_weights[metrics["PR"]] * _CVSS3_WEIGHTS["UI"][metrics["UI"]])
        if impact <= 0:
            return 0.0
        if scope_changed:
            return _roundup(min(1.08 * (impact + exploitability), 10))
        return _roundup(min(impact + exploitability, 10))
    except (KeyError, ValueError):
        return None


def severity_from_score(score: float | None) -> str:
    if score is None:
        return "Not rated"
    if score >= 9.0:
        return "Critical"
    if score >= 7.0:
        return "High"
    if score >= 4.0:
        return "Medium"
    if score > 0:
        return "Low"
    return "Not rated"


SEVERITY_ORDER = ["Critical", "High", "Medium", "Low", "Not rated"]


def normalise_severity(text: str | None) -> str:
    mapping = {"critical": "Critical", "high": "High", "moderate": "Medium", "medium": "Medium",
               "low": "Low", "info": "Low", "none": "Not rated"}
    return mapping.get(str(text or "").strip().lower(), "Not rated")


def severity_rank(sev: str) -> int:
    return SEVERITY_ORDER.index(sev) if sev in SEVERITY_ORDER else len(SEVERITY_ORDER)


# --------------------------------------------------------------------------
# Markdown to Word (.docx)
# --------------------------------------------------------------------------

FONT = "Arial"
HEADING_COLOUR = "1F3864"
TABLE_HEADER_FILL = "D9E2F3"


def _set_cell_fill(cell, hex_fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _paragraph_border(paragraph, side: str = "bottom", colour: str = "8EA9DB", size: int = 8) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), "4")
    edge.set(qn("w:color"), colour)
    borders.append(edge)
    p_pr.append(borders)


def _paragraph_shading(paragraph, hex_fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    p_pr.append(shd)


def _add_hyperlink(paragraph, url: str, text: str, bold: bool = False, size=None) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    r_pr.append(fonts)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "0563C1")
    r_pr.append(colour)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        r_pr.append(sz)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


_INLINE = re.compile(
    r"(\*\*[^*]+\*\*"           # **bold**
    r"|`[^`]+`"                  # `code`
    r"|\[[^\]]+\]\([^)\s]+\)"    # [text](url)
    r"|<(?:https?://|mailto:)[^>\s]+>"  # <https://...>
    r"|\*[^*\s][^*]*\*)"         # *italic*
)


def _add_inline(paragraph, text: str, size=None, bold_all: bool = False) -> None:
    from docx.shared import Pt
    pos = 0
    for match in _INLINE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            run.bold = bold_all or None
            if size:
                run.font.size = Pt(size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if size:
                run.font.size = Pt(size)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt((size or 10.5) - 1)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)\s]+)\)", token).groups()
            if re.match(r"^(https?://|mailto:)", url):
                _add_hyperlink(paragraph, url, label, bold=bold_all, size=size)
            else:
                run = paragraph.add_run(label)
                run.bold = bold_all or None
        elif token.startswith("<"):
            url = token[1:-1]
            _add_hyperlink(paragraph, url, url.replace("mailto:", ""), bold=bold_all, size=size)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.bold = bold_all or None
            if size:
                run.font.size = Pt(size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = bold_all or None
        if size:
            run.font.size = Pt(size)


def _split_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    cells = re.split(r"(?<!\\)\|", line)
    return [c.strip().replace("\\|", "|") for c in cells]


def _column_widths(header: list[str], rows: list[list[str]], total_mm: float) -> list[float]:
    """Share the page width between columns: every column gets room for its longest
    word, and the rest goes to the columns with the most text."""
    def mm(chars: float) -> float:          # 9 pt Arial plus cell padding
        return 1.95 * chars + 3.2

    def tokens(text: str) -> list[str]:
        out = []
        for word in text.split():
            out += [word] if re.match(r"^(CVE|GHSA|PYSEC)-", word) else word.split("-")
        return out

    mins, wants = [], []
    for c in range(len(header)):
        body = [re.sub(r"[*`]", "", re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", r[c] if c < len(r) else ""))
                for r in rows] or [""]
        head = re.sub(r"[*`]", "", header[c])
        longest_word = max((len(w) for t in body for w in tokens(t)), default=3)
        head_word = max((len(w) for w in tokens(head)), default=3)
        # Body words get room up to a cap (very long words may wrap); header words are bold, so wider.
        lo = max(min(max(mm(longest_word), 12), 36), 2.15 * head_word + 3.2)
        avg = sum(len(t) for t in body) / len(body)
        want = max(lo, min(mm(0.5 * avg + 0.5 * min(max(len(t) for t in body), 70)), 95))
        if not any(t.strip() for t in body):
            want = max(lo, total_mm * 0.4)  # an empty column is for answers: leave room to write
        mins.append(lo)
        wants.append(want)
    if sum(mins) >= total_mm:
        return [total_mm * m / sum(mins) for m in mins]
    extra = total_mm - sum(mins)
    need = [w - m for w, m in zip(wants, mins)]
    if sum(need) <= 0:
        return [m + extra / len(mins) for m in mins]
    return [m + extra * n / sum(need) for m, n in zip(mins, need)]


def _add_page_number_field(run) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _set_repeat_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def md_to_docx(md_text: str, out_path: str | Path, footer: str | None = None) -> Path:
    """Convert the simple Markdown used in this kit into a tidy Word document.

    Supports headings, paragraphs, bullet and numbered lists, task boxes
    (- [ ] / - [x]), tables, block quotes, horizontal rules, code blocks,
    **bold**, *italic*, `code` and links. HTML comments are dropped.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    md_text = re.sub(r"<!--.*?-->", "", md_text, flags=re.S)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(20)
    section.top_margin = section.bottom_margin = Mm(18)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5), ("Heading 4", 10.5)):
        st = styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(HEADING_COLOUR)
        rpr = st.element.get_or_add_rPr()
        rpr.rFonts.set(qn("w:ascii"), FONT)
        rpr.rFonts.set(qn("w:hAnsi"), FONT)
        rpr.rFonts.set(qn("w:eastAsia"), FONT)
        for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rpr.rFonts.attrib.pop(qn(theme_attr), None)  # theme fonts would override Arial
        st.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        st.paragraph_format.space_after = Pt(6)
    for name in ("List Bullet", "List Bullet 2", "List Bullet 3"):
        styles[name].font.name = FONT

    if footer:
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run(f"{footer}  |  Page ")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        num = fp.add_run()
        num.font.size = Pt(8)
        _add_page_number_field(num)

    lines = md_text.splitlines()
    i = 0
    first_block = True
    para_buffer: list[str] = []

    def flush_paragraph():
        nonlocal para_buffer
        if para_buffer:
            p = doc.add_paragraph()
            _add_inline(p, " ".join(s.strip() for s in para_buffer))
            para_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Mm(4)
                _paragraph_shading(p, "F2F2F2")
                run = p.add_run(lines[i] if lines[i] else " ")
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                i += 1
            i += 1
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and first_block:
                p = doc.add_paragraph(style="Title")
                _add_inline(p, text)
                _paragraph_border(p)
            else:
                p = doc.add_heading(level=level)
                _add_inline(p, text)
            first_block = False
            i += 1
            continue
        first_block = False

        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            flush_paragraph()
            p = doc.add_paragraph()
            _paragraph_border(p, colour="BFBFBF", size=6)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            flush_paragraph()
            header = _split_table_row(stripped)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_table_row(lines[i]))
                i += 1
            ncols = len(header)
            table = doc.add_table(rows=1, cols=ncols)
            table.style = "Table Grid"
            table.autofit = False
            widths = _column_widths(header, rows, total_mm=170)
            for c, text in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                _add_inline(cell.paragraphs[0], text, size=9, bold_all=True)
                _set_cell_fill(cell, TABLE_HEADER_FILL)
            _set_repeat_header(table.rows[0])
            for row in rows:
                cells = table.add_row().cells
                for c in range(ncols):
                    text = row[c] if c < len(row) else ""
                    parts = re.split(r"<br\s*/?>", text)
                    para = cells[c].paragraphs[0]
                    para.paragraph_format.space_after = Pt(0)
                    for n, part in enumerate(parts):
                        if n:
                            para = cells[c].add_paragraph()
                            para.paragraph_format.space_after = Pt(0)
                        _add_inline(para, part, size=9)
            for c, width in enumerate(widths):
                table.columns[c].width = Mm(width)
                for row_cells in table.rows:
                    row_cells.cells[c].width = Mm(width)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip()[1:].strip())
                i += 1
            for chunk in "\n".join(quote_lines).split("\n\n"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(5)
                p.paragraph_format.right_indent = Mm(5)
                _paragraph_shading(p, "EAF1FB")
                _paragraph_border(p, side="left", colour="2F5597", size=18)
                _add_inline(p, " ".join(s for s in chunk.splitlines() if s))
            continue

        item = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", line)
        if item:
            flush_paragraph()
            indent = len(item.group(1).replace("\t", "    "))
            level = 1 if indent < 2 else (2 if indent < 5 else 3)
            marker, text = item.group(2), item.group(3)
            # Join wrapped continuation lines of the same item.
            while (i + 1 < len(lines) and lines[i + 1].strip()
                   and not re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[i + 1])
                   and lines[i + 1].startswith("  ") and not lines[i + 1].strip().startswith("|")):
                i += 1
                text += " " + lines[i].strip()
            box = re.match(r"^\[( |x|X)\]\s+(.*)$", text)
            if box:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(6 * level)
                p.paragraph_format.first_line_indent = Mm(-6)
                p.paragraph_format.space_after = Pt(3)
                p.add_run("\u2611  " if box.group(1).lower() == "x" else "\u2610  ")
                _add_inline(p, box.group(2))
            elif marker[0].isdigit():
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(7 * level)
                p.paragraph_format.first_line_indent = Mm(-6)
                p.paragraph_format.space_after = Pt(3)
                p.add_run(f"{marker.rstrip(')').rstrip('.')}.\t")
                _add_inline(p, text)
            else:
                style = {1: "List Bullet", 2: "List Bullet 2", 3: "List Bullet 3"}[level]
                p = doc.add_paragraph(style=style)
                p.paragraph_format.space_after = Pt(3)
                _add_inline(p, text)
            i += 1
            continue

        para_buffer.append(stripped)
        i += 1

    flush_paragraph()
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.author = os.environ.get("CRA_AUTHOR", "")
    doc.save(str(out_path))
    return out_path


# --------------------------------------------------------------------------
# Spreadsheets (.xlsx)
# --------------------------------------------------------------------------

XLSX_FONT = "Arial"
HEADER_FILL = "1F3864"
INPUT_FILL = "FFF2CC"  # pale yellow: cells the reader fills in


def xlsx_header(ws, row: int, headers: list[str], widths: list[float] | None = None) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="BFBFBF")
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = Font(name=XLSX_FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    if widths:
        from openpyxl.utils import get_column_letter
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width


def xlsx_body_style(ws, first_row: int, last_row: int, ncols: int, wrap: bool = True) -> None:
    """Arial 10, thin grey borders, text wrapped and aligned to the top."""
    from openpyxl.styles import Alignment, Border, Font, Side
    thin = Side(style="thin", color="D9D9D9")
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, max_col=ncols):
        for cell in row:
            old = cell.font
            cell.font = Font(name=XLSX_FONT, size=10, bold=old.bold, italic=old.italic,
                             underline=old.underline, color=old.color)
            cell.alignment = Alignment(wrap_text=wrap, vertical="top")
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)


def xlsx_font(bold: bool = False, italic: bool = False, size: float = 10, color: str | None = None):
    from openpyxl.styles import Font
    return Font(name=XLSX_FONT, bold=bold, italic=italic, size=size, color=color)


def xlsx_title(ws, text: str, subtitle: str | None = None) -> int:
    """Write a title (and optional subtitle) at the top of a sheet. Returns next free row."""
    from openpyxl.styles import Font
    ws["A1"] = text
    ws["A1"].font = Font(name=XLSX_FONT, bold=True, size=14, color=HEADER_FILL)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(name=XLSX_FONT, italic=True, size=10, color="595959")
        return 4
    return 3


def xlsx_inject_cached_values(path: str | Path, values: dict[str, dict[str, object]]) -> None:
    """Store pre-computed results for formula cells.

    openpyxl saves formulas without results, so viewers that do not
    recalculate (file previews, some web viewers) would show empty cells.
    This writes our own computed results next to each formula. Excel and
    LibreOffice still recalculate when the file is opened.
    values = {"Sheet name": {"B4": 12, "C4": 0.5, ...}}
    """
    import xml.etree.ElementTree as ET

    path = Path(path)
    ns_main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(path) as zin:
        items = {name: zin.read(name) for name in zin.namelist()}
    workbook = ET.fromstring(items["xl/workbook.xml"])
    rels = ET.fromstring(items["xl/_rels/workbook.xml.rels"])
    rel_targets = {r.get("Id"): r.get("Target") for r in rels}
    sheet_files = {}
    for sheet in workbook.iter(f"{{{ns_main}}}sheet"):
        target = rel_targets[sheet.get(f"{{{ns_rel}}}id")]
        target = target.lstrip("/")
        sheet_files[sheet.get("name")] = target if target.startswith("xl/") else f"xl/{target}"

    for sheet_name, cells in values.items():
        member = sheet_files.get(sheet_name)
        if not member:
            continue
        xml = items[member].decode("utf-8")
        for ref, value in cells.items():
            pattern = re.compile(r'<c r="%s"([^>]*)>(<f>.*?</f>)(?:<v>.*?</v>|<v\s*/>)?</c>' % re.escape(ref), re.S)

            def repl(match, value=value):
                attrs = re.sub(r'\s*t="[^"]*"', "", match.group(1))
                if isinstance(value, str):
                    safe = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    return f'<c r="{ref}"{attrs} t="str">{match.group(2)}<v>{safe}</v></c>'
                return f'<c r="{ref}"{attrs}>{match.group(2)}<v>{value}</v></c>'

            xml = pattern.sub(repl, xml, count=1)
        items[member] = xml.encode("utf-8")

    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items.items():
            zout.writestr(name, data)
    tmp.replace(path)


def make_temp_dir(prefix: str) -> Path:
    return Path(tempfile.mkdtemp(prefix=prefix))
