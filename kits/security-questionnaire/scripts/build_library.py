#!/usr/bin/env python3
"""Seed a draft answer library from a client's policy documents.

What it does
  1. Reads every .md, .txt, .docx and .pdf file in a folder (and subfolders)
     and splits each one into sections and short passages.
  2. For each seed question (templates/question_bank.csv by default) it finds
     the passages that use the same important words (TF-IDF cosine score).
  3. Writes a draft library: every row gets the best candidate source
     (document, section, excerpt) and Status "Draft", or Status
     "NEEDS CLIENT INPUT" when nothing in the documents matches.

What it never does
  It never writes an answer. A person (with Claude's help) reads each
  candidate source and writes the answer using only what the source says.

Usage (from kits/security-questionnaire/; client files go in clients/<client>/,
which git ignores)
  python scripts/build_library.py --docs clients/<client>/docs --out clients/<client>/answer_library.xlsx
         [--questions templates/question_bank.csv] [--client "Client name"]
         [--owner "Name, Title"] [--top 3] [--min-score 0.12]
         [--gaps clients/<client>/questions_for_client.md]
  python scripts/build_library.py --blank templates/answer_library.xlsx
  python scripts/build_library.py --check clients/<client>/answer_library.xlsx
         [--gaps clients/<client>/questions_for_client.md]

Run with --help for all options.
"""
from __future__ import annotations

import argparse
import csv
import re
import signal
import sys
import warnings
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import sqkit  # noqa: E402

SUPPORTED = {".md": "markdown", ".markdown": "markdown", ".txt": "text", ".docx": "docx", ".pdf": "pdf"}
CONVERT_HINT = {".doc": "docx", ".odt": "docx", ".rtf": "docx", ".pages": "pdf"}
EXCERPT_CHARS = 450


@dataclass
class Section:
    ref: str                 # what we cite: "4.2 Multi-factor authentication" or "page 3"
    context: str             # headings above the passage, used for matching
    paragraphs: list[str] = field(default_factory=list)


@dataclass
class Document:
    path: Path
    kind: str
    title: str
    sections: list[Section]
    notes: list[str] = field(default_factory=list)

    @property
    def chars(self) -> int:
        return sum(len(p) for s in self.sections for p in s.paragraphs)


@dataclass
class Passage:
    doc: Document
    section: Section
    text: str


# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------

def strip_markdown(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"(\*\*|__)(.+?)\1", r"\2", text)
    text = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", text)
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", text).strip()


# Bullets, including the private-use glyphs Word's Symbol font leaves in PDFs.
LIST_ITEM_RE = re.compile("^\\s*(?:[-*+•‣⁃∙▪▫●◦·–—"
                          "]|\\d{1,2}[.)]|[a-z][.)])\\s+")


def parse_markdown(text: str) -> tuple[str | None, list[Section]]:
    title = None
    sections: list[Section] = []
    stack: list[tuple[int, str]] = []
    current = Section(ref="(start of document)", context="")
    para: list[str] = []

    def flush():
        if para:
            current.paragraphs.append(" ".join(para))
            para.clear()

    for raw in text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        heading = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", stripped)
        if heading:
            flush()
            level, text_h = len(heading.group(1)), strip_markdown(heading.group(2))
            if level == 1 and title is None and not sections and not current.paragraphs:
                title = text_h
                continue
            if current.paragraphs:
                sections.append(current)
            stack = [(lv, h) for lv, h in stack if lv < level] + [(level, text_h)]
            current = Section(ref=text_h, context=" ".join(h for _lv, h in stack))
            continue
        if not stripped:
            flush()
            continue
        if re.match(r"^\|?\s*:?-{3,}", stripped) or re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            flush()
            continue
        if stripped.startswith("|"):
            flush()
            cells = [strip_markdown(c) for c in stripped.strip("|").split("|")]
            row = ": ".join(c for c in cells if c)
            if row:
                current.paragraphs.append(row)
            continue
        if stripped.startswith(">"):
            stripped = stripped.lstrip("> ").strip()
        if LIST_ITEM_RE.match(stripped):
            flush()
            para.append(strip_markdown(LIST_ITEM_RE.sub("", stripped, count=1)))
            continue
        para.append(strip_markdown(stripped))
    flush()
    if current.paragraphs:
        sections.append(current)
    return title, sections


def iter_docx_blocks(document):
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, document)
        elif tag == "tbl":
            yield Table(child, document)


def parse_docx(path: Path) -> tuple[str | None, list[Section]]:
    import docx
    from docx.table import Table

    document = docx.Document(str(path))
    title = (document.core_properties.title or "").strip() or None
    has_title_style = any((p.style is not None and (p.style.name or "").lower() == "title") for p in document.paragraphs)
    sections: list[Section] = []
    stack: list[tuple[int, str]] = []
    current = Section(ref="(start of document)", context="")
    seen_body = False

    for block in iter_docx_blocks(document):
        if isinstance(block, Table):
            for row in block.rows:
                cells = []
                for cell in row.cells:
                    t = re.sub(r"\s+", " ", cell.text).strip()
                    if t and t not in cells:
                        cells.append(t)
                if cells:
                    current.paragraphs.append(": ".join(cells))
                    seen_body = True
            continue
        text = re.sub(r"\s+", " ", block.text).strip()
        if not text:
            continue
        style = (block.style.name or "").lower() if block.style is not None else ""
        if style == "title":
            title = title or text
            continue
        m = re.match(r"heading\s*(\d)", style)
        if m:
            level = int(m.group(1))
            if level == 1 and not has_title_style and not seen_body and not sections and not current.paragraphs:
                title = title or text
                has_title_style = True
                continue
            if current.paragraphs:
                sections.append(current)
            stack = [(lv, h) for lv, h in stack if lv < level] + [(level, text)]
            current = Section(ref=text, context=" ".join(h for _lv, h in stack))
            continue
        current.paragraphs.append(LIST_ITEM_RE.sub("", text, count=1) if style.startswith("list") else text)
        seen_body = True
    if current.paragraphs:
        sections.append(current)
    return title, sections


NUMBERED_HEADING_RE = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){0,3})\.?\s+([A-Z][^\n]{1,80})$")


def _is_heading_line(line: str) -> bool:
    if len(line) > 90 or line.endswith((".", ",", ";")):
        return False
    if NUMBERED_HEADING_RE.match(line):
        words = line.split()
        return len(words) <= 10
    letters = re.sub(r"[^A-Za-z]", "", line)
    return len(letters) >= 4 and letters.isupper() and len(line.split()) <= 8


def _lines_to_paragraphs(lines: list[str]) -> list[str]:
    """Rebuild paragraphs from layout lines (PDF/plain text)."""
    paragraphs: list[str] = []
    buf = ""
    for line in lines:
        if not line:
            if buf:
                paragraphs.append(buf)
                buf = ""
            continue
        if LIST_ITEM_RE.match(line):
            if buf:
                paragraphs.append(buf)
            buf = LIST_ITEM_RE.sub("", line, count=1)
            continue
        if buf.endswith("-") and not buf.endswith(" -"):
            buf = buf + line        # "us-" + "west-2": keep the hyphen, drop the line break
        elif buf:
            buf = buf + " " + line
        else:
            buf = line
    if buf:
        paragraphs.append(buf)
    # Split very long paragraphs into chunks of whole sentences.
    out = []
    for p in paragraphs:
        p = re.sub(r"\s+", " ", p).strip()
        if len(p) <= 600:
            out.append(p)
            continue
        sentences = re.split(r"(?<=[.!?])\s+(?=[A-Z0-9])", p)
        chunk = ""
        for s in sentences:
            if chunk and len(chunk) + len(s) > 400:
                out.append(chunk)
                chunk = s
            else:
                chunk = f"{chunk} {s}".strip()
        if chunk:
            out.append(chunk)
    return out


def parse_lines(pages: list[str]) -> list[Section]:
    sections: list[Section] = []
    current_ref, current_lines, current_page = "(start of document)", [], 1
    for page_no, page_text in enumerate(pages, start=1):
        for raw in page_text.splitlines():
            line = re.sub(r"\s+", " ", raw).strip()
            if line and _is_heading_line(line):
                paras = _lines_to_paragraphs(current_lines)
                if paras:
                    ref = current_ref if current_ref.startswith("(") else f"{current_ref} (page {current_page})"
                    sections.append(Section(ref=ref, context=current_ref.strip("()"), paragraphs=paras))
                current_ref, current_lines, current_page = line, [], page_no
            else:
                current_lines.append(line)
        current_lines.append("")
    paras = _lines_to_paragraphs(current_lines)
    if paras:
        ref = current_ref if current_ref.startswith("(") else f"{current_ref} (page {current_page})"
        sections.append(Section(ref=ref, context=current_ref.strip("()"), paragraphs=paras))
    return sections


def parse_pdf(path: Path) -> tuple[str | None, list[Section], list[str]]:
    from pypdf import PdfReader

    notes = []
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:  # noqa: BLE001
            return None, [], ["Encrypted PDF: ask the client for an unlocked copy."]
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001
            pages.append("")
            notes.append(f"Could not read a page: {exc}")
    text_chars = sum(len(p.strip()) for p in pages)
    if text_chars < 40 * max(1, len(pages)):
        notes.append("Very little text found. This may be a scanned PDF: ask for the original file or run OCR first.")
    title = None
    try:
        meta_title = (reader.metadata.title if reader.metadata else None) or ""
        if meta_title.strip() and not meta_title.lower().endswith((".docx", ".doc", ".pdf")):
            title = meta_title.strip()
    except Exception:  # noqa: BLE001
        pass
    if not title:
        for line in (pages[0] if pages else "").splitlines():
            if line.strip():
                title = line.strip()
                break
    sections = parse_lines(pages)
    # Drop the title line if it became its own "section" text.
    if sections and title and sections[0].ref.startswith("(") and sections[0].paragraphs[:1] == [title]:
        sections[0].paragraphs = sections[0].paragraphs[1:]
        if not sections[0].paragraphs:
            sections = sections[1:]
    return title, sections, notes


def add_version(title: str, sections: list[Section]) -> str:
    """'Access Control Policy' -> 'Access Control Policy v1.3' if a version is stated near the top."""
    head = " ".join(p for s in sections[:2] for p in s.paragraphs[:6])[:1500]
    m = re.search(r"\bversion\s*:?\s*v?(\d+(?:\.\d+){0,2})\b", head, re.I)
    if m and not re.search(r"\bv\d", title):
        return f"{title} v{m.group(1)}"
    return title


def load_documents(folder: Path) -> tuple[list[Document], list[str]]:
    docs: list[Document] = []
    skipped: list[str] = []
    for path in sorted(p for p in folder.rglob("*") if p.is_file()):
        if path.name.startswith(("~$", ".")):
            continue
        ext = path.suffix.lower()
        kind = SUPPORTED.get(ext)
        if not kind:
            hint = CONVERT_HINT.get(ext)
            msg = f"{path.relative_to(folder)}: skipped (unsupported type)"
            if hint:
                msg += f". Convert it first: soffice --headless --convert-to {hint} \"{path.name}\""
            skipped.append(msg)
            continue
        notes: list[str] = []
        try:
            if kind == "markdown":
                title, sections = parse_markdown(path.read_text(encoding="utf-8", errors="replace"))
            elif kind == "text":
                text = path.read_text(encoding="utf-8", errors="replace")
                title = next((ln.strip() for ln in text.splitlines() if ln.strip()), None)
                sections = parse_lines([text])
            elif kind == "docx":
                title, sections = parse_docx(path)
            else:
                title, sections, notes = parse_pdf(path)
        except Exception as exc:  # noqa: BLE001
            skipped.append(f"{path.relative_to(folder)}: could not be read ({exc})")
            continue
        title = title or path.stem.replace("_", " ").replace("-", " ").title()
        title = add_version(title, sections)
        if not sections:
            notes.append("No text found.")
        docs.append(Document(path=path, kind=kind, title=title, sections=sections, notes=notes))
    return docs, skipped


WHOLE_DOCUMENT = "(whole document)"


def make_passages(docs: list[Document]) -> list[Passage]:
    """Passages at three sizes, so both narrow and broad questions find a match:
    each paragraph, each whole section (bullets stay with their lead-in), and
    one passage for the document itself (title + purpose) for questions like
    'Do you have a documented incident response plan?'."""
    passages = []
    for doc in docs:
        body = [s for s in doc.sections if not s.ref.startswith("(")]
        if body:
            whole = Section(ref=WHOLE_DOCUMENT, context=doc.title, paragraphs=[body[0].paragraphs[0]])
            passages.append(Passage(doc=doc, section=whole, text=f"{doc.title}. {body[0].paragraphs[0]}"))
        for section in doc.sections:
            paras = [p for p in section.paragraphs if len(p) >= 3]
            for p in paras:
                passages.append(Passage(doc=doc, section=section, text=p))
            joined = " • ".join(paras)
            if len(paras) > 1 and len(joined) <= 1500:
                passages.append(Passage(doc=doc, section=section, text=joined))
    return passages


def excerpt(text: str, limit: int = EXCERPT_CHARS) -> str:
    if len(text) <= limit:
        return text
    cut = text[:limit].rsplit(" ", 1)[0]
    return cut + " ..."


# --------------------------------------------------------------------------
# Matching seed questions to passages
# --------------------------------------------------------------------------

def find_evidence(questions: list[dict], passages: list[Passage], top: int, floor: float):
    """For each question, the best passage per section, best first, down to `floor`.

    The query is the canonical question plus its alternate phrasings and
    search terms. Scores are TF-IDF cosine similarity (0 to 1)."""
    passage_feats = [sqkit.features(sqkit.tokens(f"{p.section.context} {p.text}", keep_phrase_words=False))
                     for p in passages]
    index = sqkit.TfidfIndex(passage_feats)
    vectors = [index.vector(f) for f in passage_feats]
    results = []
    for q in questions:
        query = " ".join([q["question"], *q.get("alternates", []), *q.get("search_terms", [])])
        qv = index.vector(sqkit.features(sqkit.tokens(query, keep_phrase_words=False)))
        scored = sorted(((index.cosine(qv, v), i) for i, v in enumerate(vectors)), reverse=True)
        picked, seen = [], set()
        for score, i in scored:
            if score < floor or len(picked) >= top:
                break
            p = passages[i]
            key = (p.doc.path, p.section.ref)
            if key in seen:
                continue
            seen.add(key)
            picked.append((score, p))
        results.append(picked)
    return results


def display_text(p: Passage) -> str:
    """A short bullet such as 'Critical: 7 days' means little alone: show its whole section."""
    if len(p.text) < 80 and len(p.section.paragraphs) > 1:
        return " • ".join(p.section.paragraphs)
    return p.text


# --------------------------------------------------------------------------
# Commands
# --------------------------------------------------------------------------

def cmd_blank(out: Path) -> int:
    sqkit.write_library(out, [], active="Readme")
    print(f"Wrote blank library template: {out}")
    return 0


def cmd_check(path: Path, stale_days: int, gaps: str | None = None, client: str = "") -> int:
    entries = sqkit.read_library(path)
    if gaps:
        write_gaps(Path(gaps), entries, client)
        n = sum(1 for e in entries if e["status"] == sqkit.NEEDS_INPUT)
        print(f"Wrote {n} NEEDS CLIENT INPUT question(s) for the client: {gaps}")
    issues = sqkit.check_library(entries, stale_days=stale_days)
    counts = {s: sum(1 for e in entries if e.get("status") == s) for s in sqkit.STATUSES}
    blank = sum(1 for e in entries if not e.get("status"))
    print(f"{path}: {len(entries)} entries | " + " | ".join(f"{k}: {v}" for k, v in counts.items())
          + (f" | no status: {blank}" if blank else ""))
    if not issues:
        print("No problems found.")
        return 0
    for level in ("ERROR", "WARNING"):
        rows = [i for i in issues if i[0] == level]
        if rows:
            print(f"\n{level}S ({len(rows)})")
            for _lvl, eid, msg in rows:
                print(f"  {eid:10} {msg}")
    errors = sum(1 for i in issues if i[0] == "ERROR")
    print(f"\n{errors} error(s), {len(issues) - errors} warning(s).")
    return 1 if errors else 0


def write_gaps(path: Path, entries: list[dict], client: str) -> None:
    gaps = [e for e in entries if e["status"] == sqkit.NEEDS_INPUT]
    if path.suffix.lower() == ".csv":
        with open(path, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["ID", "Category", "Question", "Client answer", "Supporting document (if any)"])
            for e in gaps:
                w.writerow([e["id"], e["category"], e["question"], "", ""])
        return
    lines = [f"# Questions for {client or 'the client'}", "",
             "Your documents don't cover these topics. A short factual answer for each is enough. "
             "If something is not in place, just say so: an honest \"No\" or \"Not yet\" is fine. "
             "If a document covers it, send the document instead."]
    current = None
    for e in gaps:
        if e["category"] != current:
            current = e["category"]
            lines += ["", f"## {current}", ""]
        lines.append(f"- [ ] **{e['id']}** {e['question']}")
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def cmd_build(args) -> int:
    docs_dir = Path(args.docs)
    if not docs_dir.is_dir():
        print(f"Not a folder: {docs_dir}", file=sys.stderr)
        return 2
    out = Path(args.out)
    if out.suffix.lower() != ".xlsx":
        print("--out must end in .xlsx", file=sys.stderr)
        return 2
    docs, skipped = load_documents(docs_dir)
    if not docs:
        print(f"No readable documents in {docs_dir}. Supported: {', '.join(SUPPORTED)}", file=sys.stderr)
        for s in skipped:
            print("  " + s, file=sys.stderr)
        return 2
    passages = make_passages(docs)
    questions = sqkit.read_question_bank(args.questions)
    sqkit.assign_ids(questions)
    if args.weak_score > args.min_score:
        args.weak_score = args.min_score
    results = find_evidence(questions, passages, args.top, args.weak_score)

    entries, evidence = [], []
    for q, all_picked in zip(questions, results):
        picked = [(s, p) for s, p in all_picked if s >= args.min_score]
        weak = [(s, p) for s, p in all_picked if s < args.min_score]
        best = picked[0] if picked else None
        if best:
            score, p = best
            note = (f"Possible source found automatically (score {score:.2f}). Read the section; write an answer only "
                    f"if it really supports one. Other candidates: Evidence sheet.")
        elif weak:
            score, p = weak[0]
            note = (f"Documents seem silent. Only a weak match (score {score:.2f}): {p.doc.title}, {p.section.ref}. "
                    f"Check it on the Evidence sheet, then ask the client.")
        else:
            note = "No matching text found in the documents provided. Ask the client (or for a document that covers it)."
        entries.append({
            "id": q["id"], "category": q["category"], "question": q["question"],
            "alternates": "\n".join(q["alternates"]), "answer": "", "short": "",
            "source_doc": best[1].doc.title if best else "",
            "source_section": best[1].section.ref if best else "",
            "excerpt": excerpt(display_text(best[1])) if best else "",
            "owner": args.owner or "", "reviewed": None, "confidence": "",
            "status": sqkit.STATUS_DRAFT if best else sqkit.NEEDS_INPUT, "notes": note,
        })
        for rank, (score, p) in enumerate(all_picked, start=1):
            label = rank if score >= args.min_score else f"{rank} (weak)"
            evidence.append([q["id"], q["question"], label, round(score, 2), p.doc.title, p.section.ref,
                             str(p.doc.path.relative_to(docs_dir)), excerpt(display_text(p), 700)])

    documents = [[str(d.path.relative_to(docs_dir)), d.title, d.kind, len(d.sections), d.chars, " ".join(d.notes)]
                 for d in docs]
    documents += [[s.split(":")[0], "", "", 0, 0, s.split(":", 1)[1].strip()] for s in skipped]
    out.parent.mkdir(parents=True, exist_ok=True)
    sqkit.write_library(out, entries, client=args.client or "", evidence=evidence, documents=documents)

    # ---- Summary
    print(f"Read {len(docs)} document(s) from {docs_dir} ({len(passages)} passages):")
    for d in docs:
        print(f"  {str(d.path.relative_to(docs_dir)):42} {d.title[:44]:44} {len(d.sections):3} sections"
              + (f"  NOTE: {' '.join(d.notes)}" if d.notes else ""))
    for s in skipped:
        print(f"  {s}")
    found = sum(1 for e in entries if e["status"] == sqkit.STATUS_DRAFT)
    print(f"\nSeed questions: {len(entries)} (from {args.questions})")
    print(f"  Possible source found: {found}")
    print(f"  NEEDS CLIENT INPUT:    {len(entries) - found}")
    by_cat: dict[str, list[int]] = {}
    for e in entries:
        c = by_cat.setdefault(e["category"], [0, 0])
        c[0 if e["status"] == sqkit.STATUS_DRAFT else 1] += 1
    print("\n  Category                        source found / needs input")
    for cat, (f, n) in by_cat.items():
        print(f"  {cat:32} {f:3} / {n}")
    print(f"\nWrote draft library: {out}")
    if args.gaps:
        write_gaps(Path(args.gaps), entries, args.client)
        print(f"Wrote questions for the client: {args.gaps}")
    print("\nNext: for each Draft row, read the cited section and write the answer from it (facts only), "
          "fix the source if a better one exists, and set Short Answer and Confidence. "
          "Ask the client about every NEEDS CLIENT INPUT row. No answer is used until its Status is Approved.")
    return 0


def main(argv=None) -> int:
    if hasattr(signal, "SIGPIPE"):  # exit quietly when output is piped to head/less
        signal.signal(signal.SIGPIPE, signal.SIG_DFL)
    parser = argparse.ArgumentParser(
        description="Seed a draft answer library from client documents, write a blank template, or check a library.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Examples (run from kits/security-questionnaire/; client files go in clients/<client>/, which git ignores):\n"
               "  python scripts/build_library.py --docs clients/acme/docs --out clients/acme/answer_library.xlsx "
               "--client \"Acme\" --gaps clients/acme/questions_for_client.md\n"
               "  python scripts/build_library.py --check clients/acme/answer_library.xlsx "
               "--gaps clients/acme/questions_for_client.md\n"
               "  python scripts/build_library.py --docs samples/policies --out /tmp/draft.xlsx --client \"Quarterhour\"\n"
               "  python scripts/build_library.py --check samples/library/answer_library.xlsx\n"
               "  python scripts/build_library.py --blank templates/answer_library.xlsx")
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--docs", help="folder of client documents (.md .txt .docx .pdf), searched recursively")
    mode.add_argument("--blank", metavar="OUT.xlsx", help="write an empty library template and stop")
    mode.add_argument("--check", metavar="LIBRARY.xlsx", help="check a library for missing sources, bad values and risky wording")
    parser.add_argument("--out", help="where to write the draft library (.xlsx); required with --docs")
    parser.add_argument("--questions", default=str(sqkit.DEFAULT_QUESTION_BANK),
                        help="seed questions: a CSV like templates/question_bank.csv, or an existing library .xlsx "
                             "(default: %(default)s)")
    parser.add_argument("--client", default="", help="client name, shown on the Readme sheet")
    parser.add_argument("--owner", default="", help="default Owner for every row, e.g. 'Dana Whitfield, CTO'")
    parser.add_argument("--top", type=int, default=3, help="candidate passages to keep per question (default: 3)")
    parser.add_argument("--min-score", type=float, default=0.12,
                        help="lowest similarity (0-1) that counts as a possible source (default: 0.12)")
    parser.add_argument("--weak-score", type=float, default=0.07,
                        help="matches between this and --min-score are listed as 'weak' on the Evidence sheet "
                             "but the row stays NEEDS CLIENT INPUT (default: 0.07)")
    parser.add_argument("--gaps", metavar="FILE.md|FILE.csv",
                        help="also write the NEEDS CLIENT INPUT questions as a checklist to send to the client "
                             "(with --docs: from the new draft; with --check: from the reviewed library)")
    parser.add_argument("--stale-days", type=int, default=365, help="--check: warn when Last Reviewed is older (default: 365)")
    args = parser.parse_args(argv)

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        if args.blank:
            return cmd_blank(Path(args.blank))
        if args.check:
            return cmd_check(Path(args.check), args.stale_days, args.gaps, args.client)
        if not args.out:
            parser.error("--out is required with --docs")
        return cmd_build(args)


if __name__ == "__main__":
    sys.exit(main())
