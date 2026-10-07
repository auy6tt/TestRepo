#!/usr/bin/env python3
"""Fill a buyer's security questionnaire from the client's answer library.

What it does
  1. Opens the buyer's questionnaire (.xlsx, .xlsm or .csv; .xls and .ods are
     converted with LibreOffice first) and finds, on each sheet, the header row
     and the question / answer / Yes-No / comment / evidence columns. You can
     set any of these yourself.
  2. Compares every question with the library's canonical questions and
     alternate phrasings. Score (0 to 1) = 70% TF-IDF keyword similarity +
     30% fuzzy text similarity (rapidfuzz). The matched library ID, score,
     shared keywords and runner-up are all shown.
  3. Writes proposed answers (and sources) into a COPY of the file. Only the
     answer cells change; the buyer's formatting, drop-downs and other sheets
     stay exactly as they were.
  4. Adds a review sheet listing every question, NEEDS CLIENT INPUT and CHECK
     items first, with links to the answer cells.

Only library answers with Status "Approved" are used. Weak or missing
matches are never guessed: they are marked NEEDS CLIENT INPUT.

Usage
  python fill_questionnaire.py BUYER.xlsx --library answer_library.xlsx [--out BUYER_DRAFT.xlsx]
  python fill_questionnaire.py BUYER.xlsx --library answer_library.xlsx --dry-run
  python fill_questionnaire.py --finalize BUYER_DRAFT.xlsx --out BUYER_FINAL.xlsx

Run with --help for every option (columns, thresholds, sources, config file).
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import shutil
import subprocess
import sys
import tempfile
import warnings
from dataclasses import dataclass, field
from pathlib import Path

from rapidfuzz import fuzz

sys.path.insert(0, str(Path(__file__).resolve().parent))
import sqkit  # noqa: E402
from xlsx_patch import XlsxError, XlsxPackage, col_letter, col_number, quote_sheet  # noqa: E402

REVIEW_SHEET = "Review - remove before sending"
NEEDS, CHECK, OK, SKIPPED = sqkit.NEEDS_INPUT, "CHECK", "OK", "SKIPPED"
PRIORITY = {NEEDS: 0, CHECK: 1, OK: 2, SKIPPED: 3}
HEADER_SCAN_ROWS = 30
ROLE_NAMES = {"question": "question", "answer": "answer", "short": "Yes/No answer", "comment": "comments",
              "source": "evidence/source", "id": "question ID", "domain": "domain"}

DEFAULTS = {
    "ok_score": 0.50, "min_score": 0.25, "margin": 0.05, "fuzzy_weight": 0.30, "stale_days": 365,
    "placeholder": sqkit.NEEDS_INPUT, "source_mode": "auto", "include_drafts": False, "overwrite": False,
}

# Cell text that means "not answered yet".
EMPTYISH = re.compile(r"^\s*$|^\s*(-+|—|–|\.+|\?|tbd|tbc|n/?a\?)\s*$|^\s*\[[^\]]*\]\s*$|^\s*<[^>]*>\s*$|"
                      r"^\s*(please\s+)?(select|choose|enter|type|insert|click)\b.{0,40}$", re.I)
# Questions that ask for something a canned answer may not cover.
ASKS_FOR_MORE = re.compile(
    r"\bif (yes|so|not|no)\b|\bplease (attach|provide|share|specify|state|list|upload|name|give|include)\b|"
    r"\b(attach|upload)\b|\bcopy of\b|\b(provide|share|state|specify|list|name|give)\s+(a|an|the|your|any|all|which)?\s*"
    r"(copy|list|summary|evidence|name|names|date|dates|details?|example|frequency|limit|version|report)\b|"
    r"\bwhen (was|were|did)\b|\bwhat (date|year)\b|\bdate of\b|\bcoverage limit\b|\bhow many\b", re.I)
YES_NO_START = re.compile(r"^\s*(is|are|do|does|did|has|have|can|will|would|should|was|were|could|must)\b", re.I)


# --------------------------------------------------------------------------
# Reading questionnaires
# --------------------------------------------------------------------------

class XlsxGrid:
    def __init__(self, wb, ws):
        self.wb, self.ws, self.title = wb, ws, ws.title
        self.max_row, self.max_col = ws.max_row, ws.max_column
        self.hidden = ws.sheet_state != "visible"
        self._merged = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]

    def value(self, r: int, c: int) -> str:
        if r < 1 or c < 1:
            return ""
        return sqkit.clean(self.ws.cell(row=r, column=c).value)

    def merged(self, r: int, c: int):
        for m in self._merged:
            if m[0] <= r <= m[2] and m[1] <= c <= m[3]:
                return m
        return None

    def dropdown(self, c: int, rows: list[int]):
        letter = col_letter(c)
        for dv in self.ws.data_validations.dataValidation:
            if dv.type != "list" or not dv.formula1:
                continue
            if any(f"{letter}{r}" in dv.sqref for r in rows[:60]):
                try:
                    return self._list_values(dv.formula1)
                except Exception:  # noqa: BLE001 - unreadable list: treat as no list
                    return None
        return None

    def _list_values(self, formula: str):
        f = formula.strip().lstrip("=")
        if f.startswith('"'):
            return [v.strip() for v in f.strip('"').split(",") if v.strip()]
        sheet, rng = self.ws.title, f
        if re.fullmatch(r"[A-Za-z_][\w.]*", f) and not re.fullmatch(r"[A-Za-z]{1,3}\d+", f):
            dn = self.wb.defined_names.get(f) or self.ws.defined_names.get(f)
            if dn is None:
                return None
            sheet, rng = next(iter(dn.destinations))
        elif "!" in f:
            sheet, rng = f.rsplit("!", 1)
            sheet = sheet.strip("'").replace("''", "'")
        cells = self.wb[sheet][rng.replace("$", "")]
        if not isinstance(cells, tuple):
            cells = ((cells,),)
        values = []
        for row in cells:
            for cell in (row if isinstance(row, tuple) else (row,)):
                v = sqkit.clean(cell.value)
                if v:
                    values.append(v)
        return values or None


class CsvGrid:
    def __init__(self, rows: list[list[str]], title: str):
        self.rows, self.title, self.hidden = rows, title, False
        self.max_row = len(rows)
        self.max_col = max((len(r) for r in rows), default=0)

    def value(self, r: int, c: int) -> str:
        if 1 <= r <= len(self.rows) and 1 <= c <= len(self.rows[r - 1]):
            return self.rows[r - 1][c - 1].strip()
        return ""

    def merged(self, r, c):
        return None

    def dropdown(self, c, rows):
        return None


def read_csv(path: Path):
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(text.splitlines(), dialect))
    return rows, dialect, encoding


# --------------------------------------------------------------------------
# Finding the header row and columns
# --------------------------------------------------------------------------

def header_roles(text: str) -> dict[str, int]:
    t = re.sub(r"\s+", " ", text.lower()).strip()
    if not t or len(t) > 80:
        return {}
    roles: dict[str, int] = {}
    idish = (re.fullmatch(r"(q(uestion)?\.?\s*)?(#|no\.?|nr\.?|num(ber)?|id|ref(erence)?\.?)(\s*(#|no\.?|id))?", t)
             or re.fullmatch(r"(item|control|requirement|req|question|ref|reference)\.?\s*(#|no\.?|number|id|ref)\.?", t)
             or t in {"s/n", "sn", "s.no", "s.no.", "sr. no.", "sr no", "sl no", "sl. no."})
    if idish:
        roles["id"] = 3
    yes_no = re.search(r"yes\s*/\s*no|\by\s*/\s*n\b|yes or no|yes,\s*no|yes\s*-\s*no", t)
    if not idish:
        if "question" in t:
            roles["question"] = 4
        elif re.search(r"requirement|control (statement|question|description|text)|criteria|criterion|"
                       r"assessment item|query|^item$|^control$", t):
            roles["question"] = 3
        elif t in ("description", "topic", "item description"):
            roles["question"] = 1
    if yes_no:
        roles["short"] = 4
    elif re.search(r"\b(compliant|compliance|implemented|in place|status|choice|select|selection)\b", t) \
            and not re.search(r"comment|explan|detail", t):
        roles["short"] = 2
    if re.search(r"\b(answers?|responses?|reply|replies)\b", t) and not yes_no:
        roles["answer"] = 3
    if re.search(r"comment|explanation|explain|\bdetails?\b|remark|\bnotes?\b|justification|additional info|"
                 r"clarification|elaborat|implementation|narrative", t):
        roles["comment"] = 2
    if re.search(r"evidence|supporting|documentation|artifact|artefact|attachment|\bsource\b|\blink\b|\burl\b|"
                 r"\breferences?\b", t) and "id" not in roles:
        roles["source"] = 2
    if re.search(r"\b(domain|category|section|area|topic|family|chapter)\b", t):
        roles["domain"] = 1
    return roles


@dataclass
class Layout:
    grid: object
    header_row: int
    cols: dict[str, int]
    short_values: list[str] | None = None
    notes: list[str] = field(default_factory=list)

    @property
    def long_col(self):
        return self.cols.get("answer") or self.cols.get("comment")


def _is_yes_no_list(values) -> bool:
    if not values:
        return False
    kinds = {classify_choice(v) for v in values}
    return "Yes" in kinds and "No" in kinds


def resolve_col(spec, grid, header_row: int):
    """A column given as a letter (E), a number (5) or its header text."""
    if spec is None:
        return None
    s = str(spec).strip()
    if s.lower() in ("", "none", "off", "no"):
        return 0
    for c in range(1, grid.max_col + 1):
        if grid.value(header_row, c).lower() == s.lower():
            return c
    if re.fullmatch(r"[A-Za-z]{1,3}", s):
        return col_number(s)
    if s.isdigit():
        return int(s)
    raise SystemExit(f"Column '{spec}' not found in header row {header_row} of sheet '{grid.title}'.")


def detect_layout(grid, args) -> Layout | None:
    rows = [args.header_row] if args.header_row else range(1, min(grid.max_row, HEADER_SCAN_ROWS) + 1)
    best = None
    for r in rows:
        cells = {}
        for c in range(1, grid.max_col + 1):
            roles = header_roles(grid.value(r, c))
            if roles:
                cells[c] = roles
        has_q = any("question" in roles for roles in cells.values())
        has_a = any(k in roles for roles in cells.values() for k in ("answer", "short", "comment"))
        if not args.header_row and not (has_q and has_a):
            continue
        score = sum(max(roles.values()) for roles in cells.values())
        if best is None or score > best[0]:
            best = (score, r, cells)
    if best is None:
        if not args.header_row:
            return None
        best = (0, args.header_row, {})
    _, header_row, cells = best

    cols: dict[str, int] = {}
    used: set[int] = set()
    for role in ("question", "short", "answer", "comment", "source", "id", "domain"):
        options = sorted(((roles[role], -c, c) for c, roles in cells.items() if role in roles and c not in used), reverse=True)
        if options:
            cols[role] = options[0][2]
            used.add(options[0][2])

    data_rows = list(range(header_row + 1, min(grid.max_row, header_row + 60) + 1))
    # A column with a Yes/No drop-down is the short-answer column, whatever its header says.
    if "short" not in cols:
        for c in range(1, grid.max_col + 1):
            if c not in used and grid.value(header_row, c) and _is_yes_no_list(grid.dropdown(c, data_rows)):
                cols["short"] = c
                used.add(c)
                break
    if "answer" in cols and "short" not in cols and _is_yes_no_list(grid.dropdown(cols["answer"], data_rows)):
        cols["short"] = cols.pop("answer")

    for role, spec in (("question", args.question_col), ("answer", args.answer_col), ("short", args.short_col),
                       ("comment", args.comment_col), ("source", args.source_col), ("id", args.id_col)):
        c = resolve_col(spec, grid, header_row)
        if c is None:
            continue
        for other in [k for k, v in cols.items() if v == c and k != role]:
            cols.pop(other)
        if c == 0:
            cols.pop(role, None)
        else:
            cols[role] = c
    if "question" not in cols:
        return None
    layout = Layout(grid=grid, header_row=header_row, cols=cols)
    if "short" in cols:
        layout.short_values = grid.dropdown(cols["short"], data_rows)
    if not layout.long_col and "short" not in cols:
        layout.notes.append("No answer column found: set one with --answer-col.")
    return layout


def describe_layout(layout: Layout) -> str:
    g = layout.grid
    parts = [f"header row {layout.header_row}"]
    for role in ("id", "question", "short", "answer", "comment", "source"):
        if role in layout.cols:
            c = layout.cols[role]
            parts.append(f"{ROLE_NAMES[role]} {col_letter(c)} \"{g.value(layout.header_row, c)[:30]}\"")
    if layout.short_values:
        parts.append(f"drop-down: {', '.join(layout.short_values)}")
    return "; ".join(parts)


# --------------------------------------------------------------------------
# Questions and matching
# --------------------------------------------------------------------------

@dataclass
class Match:
    score: float
    tfidf: float
    fuzzy: float
    entry: dict
    phrasing: str
    tokens: list[str]
    shared: list[str] = field(default_factory=list)


@dataclass
class Item:
    layout: Layout
    row: int
    question: str
    qid: str = ""
    status: str = ""
    reasons: list[str] = field(default_factory=list)
    match: Match | None = None
    runner: Match | None = None
    answer: str = ""
    short: str = ""
    short_written: str = ""
    source: str = ""
    writes: dict[int, str] = field(default_factory=dict)  # column -> text
    forced: str = ""          # library ID chosen by the reviewer (--use-review)
    carry_resolved: str = ""  # carried over from the previous review sheet
    carry_notes: str = ""

    @property
    def sheet(self) -> str:
        return self.layout.grid.title

    def ref(self, col: int) -> str:
        return f"{col_letter(col)}{self.row}"

    @property
    def review_cell(self) -> str:
        """The cell the review sheet links to: the answer cell, else the Yes/No cell."""
        cols = self.layout.cols
        return self.ref(self.layout.long_col or cols.get("short") or cols["question"])


def collect_items(layout: Layout, args) -> list[Item]:
    g, cols = layout.grid, layout.cols
    qcol = cols["question"]
    targets = [cols[k] for k in ("answer", "short", "comment") if k in cols]
    skip = re.compile(args.skip_regex, re.I) if args.skip_regex else None
    items = []
    last_row = args.last_row or g.max_row
    for r in range(max(layout.header_row + 1, args.first_row or 0), last_row + 1):
        text = g.value(r, qcol)
        if not text:
            continue
        words = len(text.split())
        if words < 3 and "?" not in text:
            continue
        m = g.merged(r, qcol)
        if m and any(m[1] <= c <= m[3] for c in targets):
            continue  # a section banner merged across the answer columns
        if "id" in cols and not g.value(r, cols["id"]) and not text.rstrip().endswith("?") and words <= 8:
            continue  # a section heading inside the question column
        if skip and skip.search(text):
            continue
        items.append(Item(layout=layout, row=r, question=text, qid=g.value(r, cols["id"]) if "id" in cols else ""))
    return items


ENTRY_WEIGHT = 0.3  # share of the keyword score taken from all of an entry's phrasings together


class Matcher:
    """Score = (1 - w) x keyword score + w x fuzzy score, where w is --fuzzy-weight (0.3).

    keyword score: TF-IDF cosine similarity of the normalised words and word pairs
                   (0.7 x best single phrasing + 0.3 x all phrasings of the entry together;
                   known security terms such as "mfa" count 1.5 times).
    fuzzy score:   rapidfuzz token-sort ratio against the closest phrasing (catches spelling
                   and word-order differences).
    Both are between 0 and 1, so the score is too.
    """

    def __init__(self, entries: list[dict], fuzzy_weight: float, extra_texts=(), extra_stopwords=()):
        self.stop = set(extra_stopwords)
        self.w = fuzzy_weight
        self.entries = [e for e in entries if e.get("status") != sqkit.STATUS_RETIRED and e.get("question")]
        self.phrasings = []
        for i, e in enumerate(self.entries):
            for text in [e["question"], *sqkit.split_alternates(e.get("alternates", ""))]:
                toks = sqkit.tokens(text, self.stop)
                if toks:
                    self.phrasings.append((i, text, toks))
        corpus = [sqkit.features(t) for _i, _x, t in self.phrasings]
        corpus += [sqkit.features(sqkit.tokens(x, self.stop)) for x in extra_texts]
        self.index = sqkit.TfidfIndex(corpus, concept_boost=sqkit.CONCEPT_BOOST)
        self.vectors = [self.index.vector(sqkit.features(t)) for _i, _x, t in self.phrasings]
        together: dict[int, list[str]] = {}
        for i, _x, toks in self.phrasings:
            together.setdefault(i, []).extend(toks)
        self.entry_vectors = {i: self.index.vector(sqkit.features(t)) for i, t in together.items()}

    def rank(self, question: str, top: int = 2) -> list[Match]:
        q_toks = sqkit.tokens(question, self.stop)
        if not q_toks:
            return []
        qv = self.index.vector(sqkit.features(q_toks))
        q_str = " ".join(q_toks)
        best_cos: dict[int, tuple[float, int]] = {}
        best_fuzzy: dict[int, float] = {}
        for n, ((i, _text, toks), vec) in enumerate(zip(self.phrasings, self.vectors)):
            c = self.index.cosine(qv, vec)
            if i not in best_cos or c > best_cos[i][0]:
                best_cos[i] = (c, n)
            if self.w:
                best_fuzzy[i] = max(best_fuzzy.get(i, 0.0), fuzz.token_sort_ratio(q_str, " ".join(toks)) / 100)
        matches = []
        for i, (c, n) in best_cos.items():
            keyword = (1 - ENTRY_WEIGHT) * c + ENTRY_WEIGHT * self.index.cosine(qv, self.entry_vectors[i])
            fz = best_fuzzy.get(i, 0.0)
            _i, text, toks = self.phrasings[n]
            matches.append(Match((1 - self.w) * keyword + self.w * fz, keyword, fz, self.entries[i], text, toks))
        ranked = sorted(matches, key=lambda m: -m.score)[:top]
        for m in ranked:
            shared = set(q_toks) & set(m.tokens)
            m.shared = sorted(shared, key=lambda t: -self.index.idf.get(t, 0))[:6]
        return ranked


def classify_choice(value: str):
    n = re.sub(r"[^a-z]", "", value.lower())
    if n in ("na", "nap") or n.startswith("notapplicable") or n.startswith("na") and len(n) <= 3:
        return "N/A"
    if n.startswith("partial") or n.startswith("inprogress") or n == "some":
        return "Partial"
    if n in ("y", "true") or n.startswith("yes") or n in ("compliant", "implemented", "inplace") \
            or n.startswith("fullycompliant") or n.startswith("fullyimplemented"):
        return "Yes"
    if n in ("n", "false") or n.startswith("no") or n.startswith("noncompliant") or n.startswith("notimplemented"):
        return "No"
    return None


def map_short(short: str, allowed):
    if not allowed:
        return short
    for value in allowed:
        if classify_choice(value) == short:
            return value
    return None


def decide(item: Item, ranked: list[Match], args, today: dt.date) -> None:
    best = ranked[0] if ranked else None
    item.match = best
    item.runner = ranked[1] if len(ranked) > 1 else None
    forced = bool(item.forced)
    if forced and item.forced.upper() in NO_MATCH_WORDS:
        item.status = NEEDS
        item.reasons.append("The reviewer chose no library answer for this question. Get the answer from the client.")
        return
    if not forced and (best is None or best.score < args.min_score):
        hint = f" (closest: {best.entry['id']}, score {best.score:.2f})" if best else ""
        item.status = NEEDS
        item.reasons.append(f"No library question is similar enough{hint}. Get the answer from the client, "
                            f"add it to the library, then fill it in.")
        return
    e = best.entry
    status = e.get("status", "")
    usable = bool(e.get("answer")) and (status == sqkit.STATUS_APPROVED
                                        or (args.include_drafts and status == sqkit.STATUS_DRAFT))
    if not usable:
        item.status = NEEDS
        if not e.get("answer") or status == sqkit.NEEDS_INPUT:
            item.reasons.append(f"Matches {e['id']} (score {best.score:.2f}), which has no approved answer yet. "
                                f"Ask the client, then update {e['id']} in the library.")
        else:
            item.reasons.append(f"Matches {e['id']} (score {best.score:.2f}), but its status is "
                                f"'{status or 'blank'}'. Approve it in the library or use --include-drafts.")
        return

    item.answer, item.short, item.source = e["answer"], e.get("short", ""), sqkit.source_text(e)
    checks = []
    r2 = item.runner
    if not forced and best.score < args.ok_score:
        checks.append(f"Score {best.score:.2f} is below {args.ok_score:.2f}: confirm {e['id']} really answers this question.")
    if not forced and r2 and r2.score >= best.score - args.margin and r2.entry.get("answer") != e.get("answer"):
        checks.append(f"Close second match {r2.entry['id']} ({r2.score:.2f}): make sure the right answer was picked.")
    if not forced and sqkit.has_negation(item.question) != sqkit.has_negation(best.phrasing):
        checks.append("The question may be worded the opposite way to the library question: check Yes/No.")
    if status == sqkit.STATUS_DRAFT:
        checks.append(f"{e['id']} is a Draft (not yet approved by the client).")
    if e.get("confidence") == "Low":
        checks.append(f"{e['id']} has Low confidence in the library.")
    reviewed = e.get("reviewed")
    if reviewed and (today - reviewed).days > args.stale_days:
        checks.append(f"{e['id']} was last reviewed on {reviewed.isoformat()}, over {args.stale_days} days ago.")
    if not item.source:
        checks.append(f"{e['id']} has no source in the library. Add one before sending.")
    if ASKS_FOR_MORE.search(item.question):
        checks.append("The question asks for specifics (a date, number, list or attachment): make sure the answer covers them.")
    layout = item.layout
    if "short" in layout.cols:
        if item.short:
            mapped = map_short(item.short, layout.short_values)
            if mapped is None:
                checks.append(f"Short answer '{item.short}' is not in this column's drop-down "
                              f"({', '.join(layout.short_values)}): choose one with the client.")
            else:
                item.short_written = mapped
        elif YES_NO_START.match(item.question):
            checks.append(f"{e['id']} has no Short Answer, so the Yes/No cell was left blank.")
    if forced:
        item.reasons.append(f"Library entry {e['id']} was chosen by the reviewer.")
    if checks:
        item.status = CHECK
        item.reasons += checks
    else:
        item.status = OK
        if not forced:
            item.reasons.append("Strong match. Still read the answer before sending.")


def plan_writes(item: Item, args) -> None:
    layout, cols = item.layout, item.layout.cols
    long_col = layout.long_col
    src_line = f"Source: {item.source}" if item.source else ""
    if item.status in (OK, CHECK):
        text = item.answer
        source_col = None
        if args.source_mode in ("auto", "column"):
            if cols.get("source"):
                source_col = cols["source"]
            elif args.source_mode == "auto" and cols.get("comment") and cols["comment"] != long_col:
                source_col = cols["comment"]
        if src_line and long_col and (args.source_mode == "append" or (args.source_mode == "auto" and not source_col)):
            text = f"{text}\n{src_line}"
        if long_col:
            item.writes[long_col] = text
        if cols.get("short") and item.short_written:
            item.writes[cols["short"]] = item.short_written
        if source_col and src_line:
            item.writes[source_col] = src_line
    elif item.status == NEEDS and long_col and args.placeholder:
        item.writes[long_col] = args.placeholder


def already_answered(item: Item) -> bool:
    g, cols = item.layout.grid, item.layout.cols
    for role in ("answer", "short"):
        if role in cols and not EMPTYISH.match(g.value(item.row, cols[role])):
            return True
    if "answer" not in cols and "comment" in cols and not EMPTYISH.match(g.value(item.row, cols["comment"])):
        return True
    return False


# --------------------------------------------------------------------------
# Review sheet
# --------------------------------------------------------------------------

REVIEW_HEADERS = ["Status", "What to do", "Sheet", "Cell", "Ref", "Buyer question", "Proposed Yes/No",
                  "Proposed answer", "Library ID", "Library question matched", "Score", "Shared keywords",
                  "Runner-up", "Source", "Library confidence", "Last reviewed", "Use library ID",
                  "Resolved (Y/N)", "Reviewer notes"]
REVIEW_WIDTHS = [17, 44, 14, 9, 8, 46, 10, 60, 10, 40, 7, 24, 16, 36, 11, 12, 11, 10, 30]
NO_MATCH_WORDS = ("NONE", "-", "NO", "NO MATCH")
STATUS_FILL = {NEEDS: "F8CBAD", CHECK: "FFE699", OK: "C6EFCE", SKIPPED: "D9D9D9"}


def review_table(items: list[Item]) -> list[list]:
    rows = []
    for it in items:
        m, e = it.match, (it.match.entry if it.match else {})
        rows.append({
            "status": it.status, "reasons": " ".join(it.reasons), "sheet": it.sheet, "cell": it.review_cell,
            "ref": it.qid, "question": it.question, "short": it.short_written or (it.short if it.status != NEEDS else ""),
            "answer": it.answer if it.status in (OK, CHECK) else "", "lib_id": e.get("id", "") if m else "",
            "lib_q": m.phrasing if m else "", "score": round(m.score, 2) if m else "",
            "shared": ", ".join(m.shared) if m else "",
            "runner": f"{it.runner.entry['id']} ({it.runner.score:.2f})" if it.runner else "",
            "source": it.source, "confidence": e.get("confidence", "") if m else "",
            "reviewed": e["reviewed"].isoformat() if m and e.get("reviewed") else "",
            "use_id": it.forced, "resolved": it.carry_resolved, "notes": it.carry_notes,
        })
    return rows


def build_review_sheet(pkg: XlsxPackage, items: list[Item], meta: dict) -> None:
    ordered = sorted(items, key=lambda it: (PRIORITY[it.status], it.sheet, it.row))
    counts = {s: sum(1 for it in items if it.status == s) for s in PRIORITY}
    wrap = {"wrap": True}
    title = {"bold": True, "size": 14}
    header = {"bold": True, "fill": "1F3864", "color": "FFFFFF", "wrap": True, "valign": "center"}
    link = {"color": "0563C1", "underline": True}
    rows: list[list] = [
        [(f"Review: {meta['questionnaire']} (remove this sheet before sending to the buyer)", title)],
        [f"Library: {meta['library']} ({meta['approved']} approved answers). Generated {meta['date']}. "
         f"Thresholds: OK at {meta['ok_score']:.2f} or more; CHECK at {meta['min_score']:.2f} or more; below that, no match."],
        [f"{len(items)} questions: {counts[OK]} OK, {counts[CHECK]} CHECK, {counts[NEEDS]} NEEDS CLIENT INPUT"
         + (f", {counts[SKIPPED]} skipped (already answered)" if counts[SKIPPED] else "") + "."],
        [f"Score: how close the buyer's question is to the library question, from 0 to 1 "
         f"({100 - meta['fuzzy_pct']}% TF-IDF keyword match + {meta['fuzzy_pct']}% fuzzy text match). "
         "Shared keywords show why it matched; Runner-up is the next best library entry."],
        ["How to finish: fix every NEEDS CLIENT INPUT and CHECK row in the questionnaire, then type Y in Resolved. "
         "Wrong match? Type the right library ID (or NONE) in 'Use library ID' and rerun with --use-review. "
         "The client's technical owner must read and approve every answer. Then delete this sheet, or run "
         "fill_questionnaire.py --finalize, before sending."],
        [],
        [(h, header) for h in REVIEW_HEADERS],
    ]
    first_data = len(rows) + 1
    links = []
    for i, r in enumerate(review_table(ordered)):
        n = first_data + i
        rows.append([
            (r["status"], {"bold": True, "fill": STATUS_FILL[r["status"]], "wrap": True}), (r["reasons"], wrap),
            (r["sheet"], wrap), (r["cell"], link), r["ref"], (r["question"], wrap), r["short"], (r["answer"], wrap),
            r["lib_id"], (r["lib_q"], wrap), r["score"], (r["shared"], wrap), r["runner"], (r["source"], wrap),
            r["confidence"], r["reviewed"], r["use_id"], r["resolved"], (r["notes"] or None, wrap),
        ])
        links.append((f"D{n}", f"{quote_sheet(r['sheet'])}!{r['cell']}", f"{r['sheet']}!{r['cell']}"))
    last = max(len(rows), first_data)
    header_row = first_data - 1
    end_col = col_letter(len(REVIEW_HEADERS))
    pkg.add_sheet(REVIEW_SHEET, rows, widths=REVIEW_WIDTHS, freeze_rows=header_row,
                  autofilter=f"A{header_row}:{end_col}{last}", links=links,
                  lists=[(f"R{first_data}:R{max(last, first_data)}", ["Y", "N"])], tab_color="C00000")


def write_review_csv(path: Path, items: list[Item]) -> None:
    ordered = sorted(items, key=lambda it: (PRIORITY[it.status], it.row))
    keys = ["status", "reasons", "sheet", "cell", "ref", "question", "short", "answer", "lib_id", "lib_q", "score",
            "shared", "runner", "source", "confidence", "reviewed", "use_id", "resolved", "notes"]
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(REVIEW_HEADERS)
        for r in review_table(ordered):
            w.writerow([r[k] for k in keys])


# --------------------------------------------------------------------------
# Commands
# --------------------------------------------------------------------------

def find_soffice():
    for name in ("soffice", "libreoffice"):
        if shutil.which(name):
            return shutil.which(name)
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def convert_to_xlsx(path: Path) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise SystemExit(f"{path.name}: open it in Excel or LibreOffice and save as .xlsx first (LibreOffice not found).")
    outdir = Path(tempfile.mkdtemp(prefix="sq-convert-"))
    subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), str(path)],
                   check=True, capture_output=True, timeout=300)
    converted = outdir / (path.stem + ".xlsx")
    if not converted.exists():
        raise SystemExit(f"LibreOffice could not convert {path.name} to .xlsx")
    print(f"Converted {path.name} to .xlsx with LibreOffice (the filled copy will be .xlsx).")
    return converted


def load_grids(path: Path):
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    return wb, [XlsxGrid(wb, ws) for ws in wb.worksheets]


def print_items(items: list[Item], limit: int | None = None) -> None:
    for it in items[:limit] if limit else items:
        m = it.match
        match = f"{m.entry['id']:8} {m.score:.2f}" if m else "-"
        print(f"  {it.sheet[:14]:14} {it.ref(it.layout.cols['question']):6} {it.status:18} {match:14} {it.question[:70]}")


def cmd_fill(args) -> int:
    src = Path(args.questionnaire)
    if not src.exists():
        print(f"File not found: {src}", file=sys.stderr)
        return 2
    library_path = Path(args.library)
    entries = sqkit.read_library(library_path)
    approved = [e for e in entries if e.get("status") == sqkit.STATUS_APPROVED and e.get("answer")]
    errors = [i for i in sqkit.check_library(entries, stale_days=args.stale_days) if i[0] == "ERROR"]
    if errors:
        print(f"Warning: the library has {len(errors)} error(s). Run build_library.py --check {library_path}")
    unapproved = sum(1 for e in entries if e.get("answer") and e.get("status") != sqkit.STATUS_APPROVED)
    if unapproved and not args.include_drafts:
        print(f"Note: {unapproved} library answer(s) are not Approved yet and will not be used.")

    ext = src.suffix.lower()
    is_csv = ext in (".csv", ".tsv", ".txt")
    work = src
    if ext in (".xls", ".ods"):
        work, out_ext = convert_to_xlsx(src), ".xlsx"
    elif ext in (".xlsx", ".xlsm"):
        out_ext = ext
    elif is_csv:
        out_ext = ext
    else:
        print(f"Unsupported file type {ext}. Use .xlsx, .xlsm, .csv (or .xls/.ods with LibreOffice installed).",
              file=sys.stderr)
        return 2
    out = Path(args.out) if args.out else src.with_name(f"{src.stem}_DRAFT{out_ext}")
    if out.suffix.lower() != out_ext:
        print(f"--out must end in {out_ext} (same type as the questionnaire).", file=sys.stderr)
        return 2
    if out.resolve() == src.resolve():
        print("--out must be a different file: the buyer's original is never changed.", file=sys.stderr)
        return 2

    if is_csv:
        rows, dialect, encoding = read_csv(work)
        grids = [CsvGrid(rows, src.stem)]
    else:
        wb, grids = load_grids(work)
        if REVIEW_SHEET in wb.sheetnames:
            print(f"This file already has a '{REVIEW_SHEET}' sheet: it is a filled draft. Run on the buyer's original file.",
                  file=sys.stderr)
            return 2

    wanted = set(args.sheet or [])
    if wanted - {g.title for g in grids}:
        print(f"Sheet(s) not found: {', '.join(sorted(wanted - {g.title for g in grids}))}. "
              f"Sheets: {', '.join(g.title for g in grids)}", file=sys.stderr)
        return 2
    layouts = []
    for g in grids:
        if wanted and g.title not in wanted:
            continue
        if not wanted and g.hidden:
            continue
        layout = detect_layout(g, args)
        if layout:
            layouts.append(layout)
        elif wanted:
            print(f"Sheet '{g.title}': no question column found. Set --header-row and --question-col.", file=sys.stderr)
            return 2

    items: list[Item] = []
    print(f"Questionnaire: {src}")
    for layout in layouts:
        found = collect_items(layout, args)
        if not found and not wanted:
            continue
        print(f"  Sheet '{layout.grid.title}': {describe_layout(layout)}")
        for note in layout.notes:
            print(f"    NOTE: {note}")
        print(f"    {len(found)} question(s) found")
        items += found
    if not items:
        print("No questions found. Use --sheet, --header-row and --question-col to point at them (try --dry-run).",
              file=sys.stderr)
        return 2

    matcher = Matcher(entries, args.fuzzy_weight, extra_texts=[it.question for it in items],
                      extra_stopwords=set(w.lower() for w in (args.ignore_words or [])))
    today = dt.date.today()
    for it in items:
        if already_answered(it) and not args.overwrite:
            it.status = SKIPPED
            it.reasons.append("Already has an answer in the buyer's file; left unchanged (use --overwrite to replace).")
            continue
        decide(it, matcher.rank(it.question), args, today)
        plan_writes(it, args)

    counts = {s: sum(1 for it in items if it.status == s) for s in PRIORITY}
    print(f"Library: {library_path} ({len(approved)} approved answers)")
    print(f"Results: {counts[OK]} OK, {counts[CHECK]} CHECK, {counts[NEEDS]} NEEDS CLIENT INPUT"
          + (f", {counts[SKIPPED]} skipped (already answered)" if counts[SKIPPED] else ""))
    if args.dry_run:
        print("\nDry run: nothing written. Per question:")
        print_items(items)
        return 0

    meta = {"questionnaire": src.name, "library": library_path.name, "approved": len(approved),
            "date": today.isoformat(), "ok_score": args.ok_score, "min_score": args.min_score,
            "fuzzy_pct": round(args.fuzzy_weight * 100)}
    out.parent.mkdir(parents=True, exist_ok=True)
    if is_csv:
        for it in items:
            for col, text in it.writes.items():
                row = rows[it.row - 1]
                row.extend([""] * (col - len(row)))
                row[col - 1] = text
        with open(out, "w", newline="", encoding=encoding if encoding != "latin-1" else "cp1252") as fh:
            csv.writer(fh, dialect).writerows(rows)
        review_path = out.with_name(out.stem + "_REVIEW.csv")
        write_review_csv(review_path, items)
        print(f"Wrote {out}\nWrote {review_path} (the review list)")
    else:
        try:
            pkg = XlsxPackage(work)
            by_sheet: dict[str, dict[str, str]] = {}
            for it in items:
                for col, text in it.writes.items():
                    by_sheet.setdefault(it.sheet, {})[it.ref(col)] = text
            for sheet, values in by_sheet.items():
                for ref in pkg.set_cells(sheet, values):
                    for it in items:
                        if it.sheet == sheet and any(it.ref(c) == ref for c in it.writes):
                            it.status = CHECK if it.status == OK else it.status
                            it.reasons.append(f"Cell {ref} holds a formula, so it was not changed: fill it by hand.")
            build_review_sheet(pkg, items, meta)
            pkg.save(out)
        except XlsxError as exc:
            print(f"Cannot write the filled copy: {exc}", file=sys.stderr)
            return 2
        written = verify_output(out, items)
        if pkg.signed:
            print("Note: the original file was digitally signed. The signature does not cover the filled copy.")
        print(f"Wrote {out} ({written} cells filled; review sheet '{REVIEW_SHEET}')")
    print("Next: work through the review list (NEEDS CLIENT INPUT and CHECK first), have the client's technical "
          "owner approve every answer, then finalise.")
    return 0


def verify_output(path: Path, items: list[Item]) -> int:
    """Re-open the filled copy and confirm every planned cell holds its text."""
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    count, missing = 0, []
    for it in items:
        for col, text in it.writes.items():
            value = sqkit.clean(wb[it.sheet][it.ref(col)].value)
            if value == text.strip() or value == text:
                count += 1
            elif not any(r.startswith(f"Cell {it.ref(col)} holds a formula") for r in it.reasons):
                missing.append(f"{it.sheet}!{it.ref(col)}")
    if REVIEW_SHEET not in wb.sheetnames or missing:
        raise SystemExit(f"Check failed on {path}: review sheet present={REVIEW_SHEET in wb.sheetnames}, "
                         f"cells not written: {missing[:10]}")
    return count


def cmd_finalize(args) -> int:
    import openpyxl

    src = Path(args.finalize)
    ext = src.suffix.lower()
    out = Path(args.out) if args.out else src.with_name(src.stem.replace("_DRAFT", "") + "_FINAL" + ext)
    if out.resolve() == src.resolve():
        print("--out must be a different file.", file=sys.stderr)
        return 2
    placeholder = (args.placeholder or sqkit.NEEDS_INPUT).lower()
    problems, unresolved = [], []
    if ext in (".csv", ".tsv", ".txt"):
        rows, _dialect, _enc = read_csv(src)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                if placeholder in value.lower():
                    problems.append(f"{col_letter(c)}{r}")
    elif ext in (".xlsx", ".xlsm"):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(src, data_only=True)
        for ws in wb.worksheets:
            if ws.title == REVIEW_SHEET:
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and placeholder in cell.value.lower():
                        problems.append(f"{ws.title}!{cell.coordinate}")
        if REVIEW_SHEET in wb.sheetnames:
            ws = wb[REVIEW_SHEET]
            header_row, cols = None, {}
            for row in ws.iter_rows(min_row=1, max_row=15):
                names = [sqkit.clean(c.value) for c in row]
                if "Resolved (Y/N)" in names:
                    header_row = row[0].row
                    cols = {n: i for i, n in enumerate(names)}
                    break
            if header_row:
                for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                    status = sqkit.clean(row[cols["Status"]])
                    done = sqkit.clean(row[cols["Resolved (Y/N)"]]).upper() in ("Y", "YES")
                    if status in (NEEDS, CHECK) and not done:
                        unresolved.append(f"{sqkit.clean(row[cols['Sheet']])}!{sqkit.clean(row[cols['Cell']])} "
                                          f"[{status}] {sqkit.clean(row[cols['Buyer question']])[:70]}")
    else:
        print("Finalize works on .xlsx, .xlsm and .csv files.", file=sys.stderr)
        return 2

    if problems:
        print(f"{len(problems)} cell(s) still say '{args.placeholder or sqkit.NEEDS_INPUT}':")
        for p in problems[:40]:
            print(f"  {p}")
    if unresolved:
        print(f"{len(unresolved)} review item(s) not marked Resolved = Y:")
        for u in unresolved[:40]:
            print(f"  {u}")
    if (problems or unresolved) and not args.force:
        print("Not finalised. Fix these (or use --force if the client accepts them as they are).")
        return 1

    if ext in (".csv", ".tsv", ".txt"):
        shutil.copyfile(src, out)
    else:
        try:
            pkg = XlsxPackage(src)
            if REVIEW_SHEET in pkg.sheet_names():
                pkg.remove_sheet(REVIEW_SHEET)
            pkg.save(out)
        except XlsxError as exc:
            print(f"Cannot finalise: {exc}", file=sys.stderr)
            return 2
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            check = openpyxl.load_workbook(out)
        if REVIEW_SHEET in check.sheetnames:
            raise SystemExit("Check failed: the review sheet is still present.")
    print(f"Wrote {out} (review sheet removed). Send it only after the client's technical owner has signed off.")
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Fill a buyer's security questionnaire from an answer library (writes a copy plus a review sheet).",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Examples:\n"
               "  python fill_questionnaire.py ../samples/questionnaire/buyer.xlsx --library ../samples/library/answer_library.xlsx\n"
               "  python fill_questionnaire.py buyer.xlsx --library lib.xlsx --sheet \"Security\" --question-col C --answer-col E\n"
               "  python fill_questionnaire.py buyer.csv --library lib.xlsx --ok-score 0.6 --min-score 0.4\n"
               "  python fill_questionnaire.py --finalize buyer_DRAFT.xlsx --out buyer_FINAL.xlsx\n\n"
               "Config file: --config settings.json with any option names as keys, e.g.\n"
               '  {"sheet": ["Security"], "header_row": 4, "question_col": "C", "short_col": "D", "ok_score": 0.6}')
    p.add_argument("questionnaire", nargs="?", help="the buyer's questionnaire (.xlsx .xlsm .csv; .xls .ods via LibreOffice)")
    p.add_argument("--library", help="the client's answer library (.xlsx)")
    p.add_argument("--out", help="where to write the filled copy (default: <name>_DRAFT next to the original)")
    p.add_argument("--config", help="JSON file with option values; command-line options override it")
    p.add_argument("--finalize", metavar="REVIEWED.xlsx",
                   help="check a reviewed draft (no placeholders left, every review item Resolved = Y) and save a copy "
                        "without the review sheet")
    p.add_argument("--force", action="store_true", help="--finalize even if items are unresolved")
    p.add_argument("--dry-run", action="store_true", help="show detected columns and matches; write nothing")

    g = p.add_argument_group("where the questions are (detected automatically unless set)")
    g.add_argument("--sheet", action="append", help="sheet to fill (repeat for several; default: every visible sheet "
                                                      "with a question column)")
    g.add_argument("--header-row", type=int, help="row number of the column headers")
    g.add_argument("--question-col", help="question column: letter, number or header text")
    g.add_argument("--answer-col", help="column for the full answer")
    g.add_argument("--short-col", help="column for Yes/No/Partial/N/A")
    g.add_argument("--comment-col", help="comments/explanation column")
    g.add_argument("--source-col", help="evidence/reference column ('none' to never use one)")
    g.add_argument("--id-col", help="question ID column")
    g.add_argument("--first-row", type=int, help="first row to read questions from")
    g.add_argument("--last-row", type=int, help="last row to read questions from")
    g.add_argument("--skip-regex", help="skip questions whose text matches this regular expression")

    m = p.add_argument_group("matching")
    m.add_argument("--ok-score", type=float, help=f"score at or above which a match is OK (default {DEFAULTS['ok_score']})")
    m.add_argument("--min-score", type=float,
                   help=f"score below which there is no match: NEEDS CLIENT INPUT (default {DEFAULTS['min_score']})")
    m.add_argument("--margin", type=float,
                   help=f"flag CHECK when the runner-up is this close to the best (default {DEFAULTS['margin']})")
    m.add_argument("--fuzzy-weight", type=float,
                   help=f"share of the fuzzy text score in the total, 0 to 1 (default {DEFAULTS['fuzzy_weight']})")
    m.add_argument("--ignore-words", nargs="*", help="extra words to ignore, e.g. the buyer's company name")
    m.add_argument("--include-drafts", action="store_true", default=None,
                   help="also use Draft library answers (always flagged CHECK)")
    m.add_argument("--stale-days", type=int, help=f"flag answers reviewed longer ago than this (default {DEFAULTS['stale_days']})")

    w = p.add_argument_group("writing")
    w.add_argument("--source-mode", choices=["auto", "column", "append", "review-only"],
                   help="where sources go: auto = evidence column, else comments column, else a 'Source:' line under the "
                        "answer; column = only an evidence column; append = always under the answer; review-only = only "
                        "on the review sheet (default auto)")
    w.add_argument("--placeholder", help=f"text for unanswered questions (default '{DEFAULTS['placeholder']}'; '' leaves them blank)")
    w.add_argument("--overwrite", action="store_true", default=None, help="replace answers already in the buyer's file")
    return p


def main(argv=None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    if args.config:
        with open(args.config, encoding="utf-8") as fh:
            config = json.load(fh)
        known = {a.dest for a in parser._actions}
        unknown = set(config) - known
        if unknown:
            parser.error(f"unknown option(s) in {args.config}: {', '.join(sorted(unknown))}")
        if isinstance(config.get("sheet"), str):
            config["sheet"] = [config["sheet"]]
        parser.set_defaults(**config)
        args = parser.parse_args(argv)
    for key, value in DEFAULTS.items():
        if getattr(args, key, None) is None:
            setattr(args, key, value)
    if not 0 <= args.fuzzy_weight <= 1:
        parser.error("--fuzzy-weight must be between 0 and 1")
    if args.min_score > args.ok_score:
        parser.error("--min-score must not be higher than --ok-score")
    if args.finalize:
        return cmd_finalize(args)
    if not args.questionnaire or not args.library:
        parser.error("give the questionnaire file and --library (or use --finalize)")
    return cmd_fill(args)


if __name__ == "__main__":
    sys.exit(main())
