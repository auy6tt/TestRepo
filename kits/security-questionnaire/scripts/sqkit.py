"""Shared code for the security questionnaire kit.

This module holds three things that both scripts need:

1. The answer library format: its columns, allowed values, and how to read
   and write the library workbook (with drop-down lists and a Readme sheet).
2. Text normalisation: lower-casing, turning common acronyms and synonyms
   into one word (MFA, 2FA and "two-factor" all become "mfa"), dropping
   filler words and reducing words to their stem.
3. A small TF-IDF index. Questions and passages are compared with cosine
   similarity, a well-known and easy-to-explain method: 0 means no words in
   common, 1 means the same important words.

Nothing here calls the internet or any AI service.
"""
from __future__ import annotations

import csv
import datetime as dt
import math
import re
import unicodedata
import warnings
from collections import Counter
from pathlib import Path

import snowballstemmer

# --------------------------------------------------------------------------
# Library format
# --------------------------------------------------------------------------

KIT_DIR = Path(__file__).resolve().parent.parent
DEFAULT_QUESTION_BANK = KIT_DIR / "templates" / "question_bank.csv"

LIBRARY_SHEET = "Library"
NEEDS_INPUT = "NEEDS CLIENT INPUT"

STATUS_APPROVED = "Approved"
STATUS_DRAFT = "Draft"
STATUS_RETIRED = "Retired"
STATUSES = [STATUS_APPROVED, STATUS_DRAFT, NEEDS_INPUT, STATUS_RETIRED]
SHORT_ANSWERS = ["Yes", "No", "Partial", "N/A"]
CONFIDENCE_LEVELS = ["High", "Medium", "Low"]

# key, column header, column width, what goes in the column
COLUMNS = [
    ("id", "ID", 10,
     "Unique ID such as AC-002. Never reuse or renumber an ID: past questionnaires refer to it."),
    ("category", "Category", 20,
     "Topic area. Pick one from the drop-down (the list lives on the Lists sheet)."),
    ("question", "Canonical Question", 46,
     "The question in plain, neutral words. One topic per row."),
    ("alternates", "Alternate Phrasings", 40,
     "Other ways buyers ask the same thing, one per line. The fill script matches against these too. "
     "Keep the same meaning and the same Yes/No direction as the canonical question."),
    ("answer", "Approved Answer", 60,
     "The exact words the client has approved for customers. Facts from the source only. "
     "Leave blank until it can be written from a source; never guess."),
    ("short", "Short Answer", 11,
     "Yes, No, Partial or N/A. Used for Yes/No columns in questionnaires. Leave blank for open questions."),
    ("source_doc", "Source Document", 28,
     "The client document the answer comes from (title and version), or "
     "'Written confirmation from <name>, <date>'."),
    ("source_section", "Source Section", 24,
     "Section number and heading, or page, in the source document."),
    ("excerpt", "Source Excerpt", 50,
     "The sentence(s) in the source that support the answer, copied word for word."),
    ("owner", "Owner", 18,
     "The client person who stands behind this answer, usually the technical owner."),
    ("reviewed", "Last Reviewed", 13,
     "Date the owner last confirmed the answer (YYYY-MM-DD). Review at least yearly and whenever the source changes."),
    ("confidence", "Confidence", 11,
     "High: stated directly in a current source. Medium: needs light interpretation, or the source is over a year old. "
     "Low: only partly supported; confirm before use."),
    ("status", "Status", 20,
     "Approved: client signed off. Draft: written, not yet approved. NEEDS CLIENT INPUT: the documents are silent. "
     "Retired: no longer used."),
    ("notes", "Notes", 40,
     "Internal notes: open questions, caveats, what changed. Never sent to buyers."),
]
COLUMN_KEYS = [c[0] for c in COLUMNS]
HEADER_TO_KEY = {c[1].lower(): c[0] for c in COLUMNS}

CATEGORIES = [
    ("GOV", "Governance & Policies"),
    ("CMP", "Compliance & Certifications"),
    ("HR", "People Security"),
    ("AC", "Access Control"),
    ("CRY", "Encryption & Keys"),
    ("DAT", "Data Protection & Privacy"),
    ("BCK", "Backup & Recovery"),
    ("BCP", "Business Continuity & DR"),
    ("IR", "Incident Response"),
    ("VEN", "Vendor Management"),
    ("SDL", "Secure Development"),
    ("VUL", "Vulnerability Management"),
    ("LOG", "Logging & Monitoring"),
    ("NET", "Network & Infrastructure"),
    ("END", "Endpoints & Assets"),
    ("PHY", "Physical Security"),
    ("OTH", "Other"),
]

# Words that tend to overstate a control. The library check flags them so a
# human confirms the source really supports them.
RISKY_WORDS = [
    "certified", "certification", "compliant", "compliance with", "guarantee", "guaranteed",
    "100%", "always", "never", "fully", "all systems", "fully compliant", "audited", "accredited",
    "military-grade", "unhackable", "bank-level", "best-in-class",
]


def clean(value) -> str:
    """Return a cell value as stripped text ('' for empty)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def split_alternates(text: str) -> list[str]:
    """Alternate phrasings are stored one per line (or separated by ' | ')."""
    parts = re.split(r"\r?\n|\s\|\s", text or "")
    return [p.strip(" -•\t") for p in parts if p.strip(" -•\t")]


def as_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = clean(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d %B %Y", "%d %b %Y", "%B %d, %Y"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def source_text(entry: dict) -> str:
    """'Access Control Policy v1.3, 4.2 Multi-factor authentication'."""
    doc = clean(entry.get("source_doc"))
    section = clean(entry.get("source_section"))
    if doc and section:
        return f"{doc}, {section}"
    return doc or section


# --------------------------------------------------------------------------
# Reading and writing the library workbook
# --------------------------------------------------------------------------

def read_library(path) -> list[dict]:
    """Read the Library sheet into a list of dicts keyed by COLUMN_KEYS."""
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[LIBRARY_SHEET] if LIBRARY_SHEET in wb.sheetnames else None
    header_row = None
    sheets = [ws] if ws is not None else wb.worksheets
    for sheet in sheets:
        for row in sheet.iter_rows(min_row=1, max_row=10):
            texts = [clean(c.value).lower() for c in row]
            if "canonical question" in texts:
                ws, header_row = sheet, row[0].row
                break
        if header_row:
            break
    if not header_row:
        raise SystemExit(f"{path}: no sheet with a 'Canonical Question' header. Is this an answer library?")

    col_keys = {}
    for cell in ws[header_row]:
        key = HEADER_TO_KEY.get(clean(cell.value).lower())
        if key:
            col_keys[cell.column] = key
    missing = [c[1] for c in COLUMNS if c[0] in ("id", "question", "answer", "status") and c[0] not in col_keys.values()]
    if missing:
        raise SystemExit(f"{path}: the Library sheet is missing column(s): {', '.join(missing)}")

    entries = []
    for row in ws.iter_rows(min_row=header_row + 1):
        entry = {key: "" for key in COLUMN_KEYS}
        for cell in row:
            key = col_keys.get(cell.column)
            if key == "reviewed":
                entry[key] = as_date(cell.value)
            elif key:
                entry[key] = clean(cell.value)
        if not entry["id"] and not entry["question"]:
            continue
        entry["_row"] = row[0].row
        entries.append(entry)
    return entries


def check_library(entries: list[dict], stale_days: int = 365, today: dt.date | None = None) -> list[tuple[str, str, str]]:
    """Return a list of (level, id, message). level is ERROR or WARNING."""
    today = today or dt.date.today()
    issues = []
    seen = {}
    for e in entries:
        eid = e.get("id") or f"row {e.get('_row', '?')}"
        status = e.get("status", "")
        answer = e.get("answer", "")
        if not e.get("id"):
            issues.append(("ERROR", eid, "Missing ID."))
        elif e["id"] in seen:
            issues.append(("ERROR", eid, f"Duplicate ID (also on row {seen[e['id']]})."))
        else:
            seen[e["id"]] = e.get("_row", "?")
        if not e.get("question"):
            issues.append(("ERROR", eid, "Missing canonical question."))
        if status and status not in STATUSES:
            issues.append(("WARNING", eid, f"Unknown status '{status}'. Use one of: {', '.join(STATUSES)}."))
        if e.get("short") and e["short"] not in SHORT_ANSWERS:
            issues.append(("ERROR", eid, f"Short answer '{e['short']}' is not one of {', '.join(SHORT_ANSWERS)}."))
        if e.get("confidence") and e["confidence"] not in CONFIDENCE_LEVELS:
            issues.append(("WARNING", eid, f"Confidence '{e['confidence']}' is not one of {', '.join(CONFIDENCE_LEVELS)}."))
        if NEEDS_INPUT.lower() in answer.lower():
            issues.append(("ERROR", eid, "The answer text contains NEEDS CLIENT INPUT. Leave the answer blank and use the Status column."))
        if status == STATUS_APPROVED:
            if not answer:
                issues.append(("ERROR", eid, "Approved but the answer is empty."))
            if not e.get("source_doc"):
                issues.append(("ERROR", eid, "Approved but no source document. Every answer must cite a source."))
            if not e.get("owner"):
                issues.append(("WARNING", eid, "Approved but no owner."))
            if not e.get("reviewed"):
                issues.append(("WARNING", eid, "Approved but no Last Reviewed date."))
            if not e.get("confidence"):
                issues.append(("WARNING", eid, "Approved but no confidence level."))
        elif answer and status in ("", NEEDS_INPUT):
            issues.append(("WARNING", eid, f"Has an answer but Status is '{status or 'blank'}'. Set Draft or Approved."))
        reviewed = e.get("reviewed")
        if status == STATUS_APPROVED and reviewed and (today - reviewed).days > stale_days:
            issues.append(("WARNING", eid, f"Last reviewed {reviewed.isoformat()}, over {stale_days} days ago."))
        if answer:
            low, quoted = answer.lower(), (e.get("excerpt") or "").lower()

            def says(text, word):
                return re.search(r"(?<![a-z])" + re.escape(word) + r"(?![a-z])", text)

            # A strong word is fine when the quoted source uses it too.
            hits = [w for w in RISKY_WORDS if says(low, w) and not says(quoted, w)]
            if hits:
                issues.append(("WARNING", eid, f"Check wording, could overstate: {', '.join(hits)}. "
                                              "Keep only if the source says exactly this (and quote it in Source Excerpt)."))
    return issues


def write_library(path, entries: list[dict], *, client: str = "", evidence: list[list] | None = None,
                  documents: list[list] | None = None, changelog: list[list] | None = None,
                  active: str = LIBRARY_SHEET) -> None:
    """Write an answer library workbook (also used for the blank template)."""
    import openpyxl
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    wb.properties.creator = "Security questionnaire kit"
    wb.properties.title = "Answer library" + (f": {client}" if client else "")
    readme = wb.active
    readme.title = "Readme"
    lib = wb.create_sheet(LIBRARY_SHEET)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")
    wrap_top = Alignment(wrap_text=True, vertical="top")

    # ---- Library sheet
    for col, (key, header, width, _help) in enumerate(COLUMNS, start=1):
        cell = lib.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        lib.column_dimensions[get_column_letter(col)].width = width
    lib.row_dimensions[1].height = 30
    for r, entry in enumerate(entries, start=2):
        for col, key in enumerate(COLUMN_KEYS, start=1):
            value = entry.get(key, "")
            if key == "reviewed":
                value = as_date(value)
            cell = lib.cell(row=r, column=col, value=value if value not in ("", None) else None)
            cell.alignment = wrap_top
            if key == "reviewed":
                cell.number_format = "yyyy-mm-dd"
    last = max(len(entries) + 1, 1000)
    lib.freeze_panes = "C2"
    lib.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(entries) + 1}"

    letter = {key: get_column_letter(i) for i, key in enumerate(COLUMN_KEYS, start=1)}

    def add_list(col_key, formula, strict=True, prompt=""):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                            errorStyle="stop" if strict else "warning")
        dv.error = "Pick a value from the list."
        dv.errorTitle = "Not in the list"
        if prompt:
            dv.prompt = prompt
        dv.add(f"{letter[col_key]}2:{letter[col_key]}{last}")
        lib.add_data_validation(dv)

    n_cat = len(CATEGORIES)
    add_list("category", f"Lists!$A$2:$A${n_cat + 1}", strict=False)
    add_list("short", '"' + ",".join(SHORT_ANSWERS) + '"')
    add_list("confidence", '"' + ",".join(CONFIDENCE_LEVELS) + '"')
    add_list("status", '"' + ",".join(STATUSES) + '"')
    date_dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
    date_dv.error = "Enter a date as YYYY-MM-DD."
    date_dv.errorTitle = "Not a date"
    date_dv.add(f"{letter['reviewed']}2:{letter['reviewed']}{last}")
    lib.add_data_validation(date_dv)
    id_col = letter["id"]
    id_dv = DataValidation(type="custom", formula1=f"COUNTIF(${id_col}$2:${id_col}${last},{id_col}2)=1", allow_blank=True)
    id_dv.error = "This ID is already used. IDs must be unique."
    id_dv.errorTitle = "Duplicate ID"
    id_dv.add(f"{id_col}2:{id_col}{last}")
    lib.add_data_validation(id_dv)

    amber = PatternFill("solid", fgColor="FCE4D6", bgColor="FCE4D6")
    yellow = PatternFill("solid", fgColor="FFF2CC", bgColor="FFF2CC")
    green = PatternFill("solid", fgColor="E2EFDA", bgColor="E2EFDA")
    red = PatternFill("solid", fgColor="F8CBAD", bgColor="F8CBAD")
    st = letter["status"]
    status_range = f"{st}2:{st}{last}"
    answer_range = f"{letter['answer']}2:{letter['answer']}{last}"
    lib.conditional_formatting.add(status_range, FormulaRule(formula=[f'${st}2="{NEEDS_INPUT}"'], fill=amber))
    lib.conditional_formatting.add(answer_range, FormulaRule(formula=[f'${st}2="{NEEDS_INPUT}"'], fill=amber))
    lib.conditional_formatting.add(status_range, FormulaRule(formula=[f'${st}2="{STATUS_DRAFT}"'], fill=yellow))
    lib.conditional_formatting.add(status_range, FormulaRule(formula=[f'${st}2="{STATUS_APPROVED}"'], fill=green))
    src = letter["source_doc"]
    lib.conditional_formatting.add(
        f"{src}2:{src}{last}",
        FormulaRule(formula=[f'AND(${st}2="{STATUS_APPROVED}",LEN(TRIM(${src}2))=0)'], fill=red))
    rv = letter["reviewed"]
    lib.conditional_formatting.add(
        f"{rv}2:{rv}{last}",
        FormulaRule(formula=[f"AND(ISNUMBER({rv}2),TODAY()-{rv}2>365)"], font=Font(color="C00000", bold=True)))

    # ---- Readme sheet
    readme.column_dimensions["A"].width = 26
    readme.column_dimensions["B"].width = 100
    title = "Security questionnaire answer library" + (f": {client}" if client else "")
    readme["A1"] = title
    readme["A1"].font = Font(bold=True, size=14)
    lines = [
        ("What this is", "The client's approved answers to common security questions, each tied to the document it comes from. "
                         "Use it to fill questionnaires faster and keep answers consistent."),
        ("The rules", "1. Answers are statements the client makes to its customers. Write only what the client's documents "
                      "or written confirmations say. Never invent a control, a date, a number or a certification."),
        ("", "2. Every answer cites its source (document and section, or a dated written confirmation). "
             "If the documents are silent, leave the answer blank and set Status to NEEDS CLIENT INPUT."),
        ("", "3. The client's technical owner reviews and approves every answer. Only Approved answers are used to fill questionnaires."),
        ("", "4. This library is not an audit, certification or attestation."),
        ("", "5. Confidential: share only with people covered by the NDA."),
        ("How to add an answer", "Find the supporting text in the source, write the answer in plain words using only those facts, "
                                 "fill Source Document, Source Section and Source Excerpt, set Short Answer and Confidence, "
                                 "and set Status to Draft. When the owner approves it, set Status to Approved, Owner and Last Reviewed."),
        ("Keeping it current", "Review every answer at least once a year and whenever a source document changes. "
                               "Log changes on the Change Log sheet. Retire answers instead of deleting them."),
        ("Matching tips", "When a buyer asks a known question in new words, add their wording to Alternate Phrasings "
                          "(one per line). Keep the same Yes/No direction as the canonical question."),
        ("Example row", "ID: AC-002 | Category: Access Control | Canonical Question: Is multi-factor authentication required for "
                        "all workforce accounts? | Approved Answer: Yes. MFA is required for all staff accounts on email, "
                        "cloud hosting and code hosting. | Short Answer: Yes | Source Document: Access Control Policy v1.3 | "
                        "Source Section: 4.2 Multi-factor authentication | Status: Approved"),
    ]
    r = 3
    for label, text in lines:
        readme.cell(row=r, column=1, value=label).font = Font(bold=True)
        cell = readme.cell(row=r, column=2, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        readme.cell(row=r, column=1).alignment = Alignment(vertical="top")
        r += 1
    r += 1
    readme.cell(row=r, column=1, value="Column").font = Font(bold=True)
    readme.cell(row=r, column=2, value="What to put in it").font = Font(bold=True)
    r += 1
    for _key, header, _w, help_text in COLUMNS:
        readme.cell(row=r, column=1, value=header).alignment = Alignment(vertical="top")
        readme.cell(row=r, column=2, value=help_text).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    # ---- Optional draft-only sheets
    def table_sheet(name, headers, rows, widths):
        ws = wb.create_sheet(name)
        for c, (h, w) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = w
        for i, row in enumerate(rows, start=2):
            for c, value in enumerate(row, start=1):
                ws.cell(row=i, column=c, value=value).alignment = wrap_top
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        return ws

    if evidence is not None:
        table_sheet("Evidence",
                    ["Library ID", "Canonical Question", "Rank", "Score", "Source Document", "Source Section", "File", "Excerpt"],
                    evidence, [10, 40, 6, 7, 28, 26, 26, 80])
    if documents is not None:
        table_sheet("Documents", ["File", "Title used for citations", "Type", "Sections", "Characters", "Notes"],
                    documents, [34, 34, 7, 9, 11, 60])
    table_sheet("Change Log", ["Date", "ID", "What changed", "Changed by", "Approved by"], changelog or [],
                [12, 10, 70, 18, 18])

    lists = wb.create_sheet("Lists")
    for c, (header, values) in enumerate([
        ("Category", [name for _p, name in CATEGORIES]),
        ("ID prefix", [p for p, _n in CATEGORIES]),
        ("Short Answer", SHORT_ANSWERS),
        ("Confidence", CONFIDENCE_LEVELS),
        ("Status", STATUSES),
    ], start=1):
        lists.cell(row=1, column=c, value=header).font = Font(bold=True)
        lists.column_dimensions[get_column_letter(c)].width = 28 if c == 1 else 18
        for i, value in enumerate(values, start=2):
            lists.cell(row=i, column=c, value=value)

    # Exactly one selected tab, otherwise Excel opens the sheets "grouped".
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == active
    wb.active = wb.sheetnames.index(active)
    wb.save(path)
    tidy_xlsx(path)


def tidy_xlsx(path) -> None:
    """Fix openpyxl's non-standard font element order so the file passes strict validators."""
    from xlsx_patch import XlsxPackage

    pkg = XlsxPackage(path)
    pkg.tidy_font_order()
    pkg.save(path)


def read_question_bank(path) -> list[dict]:
    """Read seed questions from a CSV or XLSX (question bank or an existing library)."""
    path = Path(path)
    rows: list[dict] = []
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        for e in read_library(path):
            rows.append({"id": e["id"], "category": e["category"], "question": e["question"],
                         "alternates": split_alternates(e["alternates"]), "search_terms": []})
        return rows
    with open(path, newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        fields = {(f or "").strip().lower(): f for f in reader.fieldnames or []}

        def get(row, *names):
            for n in names:
                if n in fields:
                    return clean(row.get(fields[n]))
            return ""

        for row in reader:
            question = get(row, "canonical question", "question")
            if not question:
                continue
            rows.append({
                "id": get(row, "id"),
                "category": get(row, "category") or "Other",
                "question": question,
                "alternates": split_alternates(get(row, "alternate phrasings", "alternates")),
                "search_terms": [t.strip() for t in re.split(r"[;,]", get(row, "search terms", "keywords")) if t.strip()],
            })
    return rows


def assign_ids(rows: list[dict]) -> None:
    """Give rows without an ID one based on their category prefix."""
    prefix_for = {name.lower(): p for p, name in CATEGORIES}
    used = {r["id"] for r in rows if r.get("id")}
    counters: Counter = Counter()
    for r in rows:
        if r.get("id"):
            continue
        prefix = prefix_for.get(r.get("category", "").lower(), "OTH")
        while True:
            counters[prefix] += 1
            candidate = f"{prefix}-{counters[prefix]:03d}"
            if candidate not in used:
                break
        r["id"] = candidate
        used.add(candidate)


# --------------------------------------------------------------------------
# Text normalisation
# --------------------------------------------------------------------------

# Each group's phrases are replaced by the group's single token before
# matching. Edit freely: add phrases your clients' buyers use.
SYNONYM_GROUPS = {
    "mfa": ["multi factor authentication", "multifactor authentication", "multi factor", "multifactor",
            "two factor authentication", "two factor", "2 factor", "2fa", "two step verification",
            "2 step verification", "two step", "second factor", "hardware security keys", "hardware security key",
            "hardware tokens", "hardware token", "security keys", "security key", "authenticator apps",
            "authenticator app", "one time passwords", "one time password", "totp", "mfa"],
    "sso": ["single sign on", "single signon", "sso"],
    "pentest": ["penetration testing", "penetration tests", "penetration test", "pen testing", "pen tests",
                "pen test", "pentests", "pentesting", "pentest"],
    "soc2": ["service organization control 2", "service organisation control 2", "soc 2", "soc2", "soc ii"],
    "iso27001": ["iso iec 27001", "iso 27001", "iso27001", "27001"],
    "gdpr": ["general data protection regulation", "uk gdpr", "eu gdpr", "gdpr"],
    "dpa": ["data processing agreement", "data processing addendum", "data protection agreement", "dpa"],
    "tls": ["transport layer security", "tls", "ssl", "https"],
    "kms": ["key management service", "key management system", "kms", "hardware security module", "hsm"],
    "rbac": ["role based access control", "role based access", "role based", "rbac"],
    "leastprivilege": ["principle of least privilege", "least privilege", "least privileged", "need to know"],
    "accessreview": ["user access reviews", "user access review", "access reviews", "access review",
                     "access recertification", "recertification", "recertified", "entitlement review"],
    "offboard": ["off boarding", "offboarding", "leavers", "leaver", "terminated employees",
                 "termination of employment", "departing employees", "employee departures", "last working day",
                 "last day", "final day", "leaves the company", "leave the company", "leaves your company",
                 "leave your company", "leaves the organization", "leave the organization"],
    "backgroundcheck": ["background screening checks", "background screening", "background checks", "background check",
                        "pre employment screening", "criminal record checks", "criminal background",
                        "background verification", "screening checks", "screening", "screened", "screen"],
    "securitytraining": ["security awareness training", "security and privacy awareness training", "security awareness",
                         "awareness training", "security training", "phishing training", "phishing simulation"],
    "nda": ["non disclosure agreements", "non disclosure agreement", "nondisclosure agreement",
            "confidentiality agreements", "confidentiality agreement", "ndas", "nda"],
    "subprocessor": ["sub processors", "sub processor", "subprocessors", "subprocessor"],
    "thirdparty": ["third parties", "third party", "3rd party"],
    "vpn": ["virtual private network", "vpn"],
    "waf": ["web application firewall", "waf"],
    "ddos": ["distributed denial of service", "denial of service", "ddos"],
    "antimalware": ["endpoint detection and response", "endpoint protection", "anti malware", "antimalware",
                    "anti virus", "antivirus", "edr", "malware protection"],
    "mdm": ["mobile device management", "device management", "mdm"],
    "diskencrypt": ["full disk encryption", "disk encryption", "fde", "filevault", "bitlocker"],
    "byod": ["bring your own device", "byod", "personal devices", "personal device", "personally owned devices",
             "own devices"],
    "siem": ["security information and event management", "siem"],
    "intrusiondetect": ["intrusion detection systems", "intrusion detection system", "intrusion detection",
                        "intrusion prevention systems", "intrusion prevention system", "intrusion prevention"],
    "sast": ["static application security testing", "static code analysis", "static analysis", "sast"],
    "depscan": ["software composition analysis", "dependency scanning", "dependency checks", "dependency check",
                "dependency scans", "dependabot", "sca"],
    "sdlc": ["secure software development lifecycle", "secure software development life cycle",
             "software development lifecycle", "software development life cycle", "secure development lifecycle",
             "development lifecycle", "ssdlc", "sdlc"],
    "codereview": ["code reviews", "code review", "peer reviewed", "peer review", "pull requests", "pull request",
                   "merge requests", "merge request"],
    "sla": ["service level agreement", "service level", "uptime commitment", "sla"],
    "ai": ["artificial intelligence", "machine learning", "large language models", "large language model",
           "generative ai", "genai", "llms", "llm", "ml models", "ai"],
    "dsar": ["data subject access requests", "data subject access request", "data subject requests",
             "data subject request", "dsars", "dsar", "right to erasure", "right to be forgotten"],
    "personaldata": ["personally identifiable information", "personal information", "personal data", "pii"],
    "multitenant": ["multi tenant", "multitenant", "multi tenancy", "multitenancy"],
    "cyberinsurance": ["cyber liability insurance", "cyber insurance", "cyber liability", "cybersecurity insurance"],
    "vulnscan": ["vulnerability scanning", "vulnerability scans", "vulnerability scan", "vulnerability scanner"],
    "statuspage": ["status pages", "status page", "status site", "status updates"],
    "rto": ["recovery time objectives", "recovery time objective", "rto"],
    "rpo": ["recovery point objectives", "recovery point objective", "maximum tolerable data loss",
            "acceptable data loss", "data loss", "rpo"],
    "dlp": ["data loss prevention", "dlp"],
    "inventory": ["asset inventories", "asset inventory", "asset registers", "asset register", "inventories",
                  "inventory"],
    "disasterrecovery": ["disaster recovery", "dr"],
    "businesscontinuity": ["business continuity", "bcp"],
    "incidentresponse": ["incident response", "incident management", "irp"],
    "riskassessment": ["risk assessments", "risk assessment", "risk register", "risk analysis"],
    "securitylead": ["chief information security officer", "information security officer", "security officer",
                     "head of security", "security lead", "ciso"],
    "secretsmanager": ["secrets manager", "secret manager", "secrets management", "vault"],
    "availabilityzone": ["availability zones", "availability zone", "multi az"],
    "pitr": ["point in time recovery", "point in time restore", "pitr"],
    "datacenter": ["data centres", "data centre", "data centers", "data center", "datacenters", "datacenter",
                   "datacentre"],
    "aws": ["amazon web services", "aws"],
    "customerdata": ["customer data", "client data", "our data", "customers data", "tenant data"],
    "prod": ["production", "prod", "live"],
    "nonprod": ["non production", "nonproduction", "pre production", "staging"],
    "encryptrest": ["encryption at rest", "encrypted at rest", "at rest encryption", "data at rest", "at rest",
                    "while stored", "when stored"],
    "encrypttransit": ["encryption in transit", "encrypted in transit", "data in transit", "in transit", "in flight"],
    "auditlog": ["audit logs", "audit log", "audit trails", "audit trail", "security logs", "event logs",
                 "activity logs", "activity log"],
    "patch": ["patch management", "patching", "patches", "patched", "security updates"],
    # In a buyer's questionnaire "our" means the buyer and "your" means the vendor.
    "customeruser": ["our employees", "our employee", "our users", "our user", "our staff", "our people",
                     "our personnel", "our team members", "customer users", "customers users", "end users"],
    "staff": ["your employees", "your staff", "your personnel", "your people", "employees", "employee", "staff",
              "personnel", "workforce", "team members"],
    "outage": ["becomes unavailable", "became unavailable", "becoming unavailable", "is unavailable", "unavailable",
               "goes down", "went down", "outages", "outage", "downtime", "fails", "failed", "failure"],
    "timeframe": ["timeframes", "timeframe", "time frames", "time frame", "timelines", "timeline", "how quickly",
                  "how soon", "how fast", "within what time"],
    "vendor": ["third party providers", "service providers", "suppliers", "supplier", "vendors", "vendor",
               "subcontractors", "subcontractor"],
    "always": ["24 hours a day", "around the clock", "round the clock", "24x7", "24 x 7", "24 7"],
    "laptop": ["laptops", "laptop", "notebooks", "workstations", "workstation", "desktops", "desktop computers",
               "computers", "computer", "company devices", "company device"],
    "screenlock": ["screen locks", "screen lock", "screen locking", "lock screen", "auto lock", "automatic lock"],
    "newhire": ["new hires", "new hire", "new employees", "new employee", "new staff", "new joiners", "joiners",
                "new starters", "new personnel"],
    "contractend": ["end of the contract", "end of contract", "contract ends", "contract end", "contract termination",
                    "termination of the contract", "terminate the contract", "termination", "cancellation", "cancels",
                    "cancel", "when we leave", "if we leave"],
    "notify": ["notifications", "notification", "notifies", "notified", "notifying", "notify"],
    "retain": ["retention period", "retention", "retained", "retaining", "retains", "retain", "kept for", "kept"],
    "login": ["sign in", "signin", "sign on", "log in", "logon", "log on", "logins", "login"],
    "idp": ["identity providers", "identity provider", "idp"],
    "backup": ["backed up", "back up", "backing up", "back ups", "backups", "backup", "snapshots", "snapshot"],
    "admin": ["administrators", "administrator", "administrative", "administration", "admins", "admin"],
    "separate": ["segregation", "segregated", "segregate", "separation", "separated", "separate", "isolation",
                 "isolated", "isolate"],
    "delete": ["deletion", "deleted", "deletes", "delete", "erasure", "erased", "erase", "purged", "purge",
               "destroyed", "destruction", "destroy", "disposal", "dispose"],
}

# Word endings the stemmer leaves different ("response" -> respons, "respond" -> respond).
STEM_ALIASES = {"respons": "respond", "retent": "retain", "recoveri": "recov", "revoc": "revok", "analyz": "analys",
                "analysi": "analys", "licenc": "licens", "centr": "center", "storag": "store", "notif": "notifi",
                "subscript": "subscrib"}

STOPWORDS = set("""
a about above after again against all also am an and any are aren as at be because been before being below
between both but by can cannot could did didn do does doesn doing don down during each either etc few for from
further had has have having he her here hers him his how i if in into is isn it its itself just may me might
more most must my no nor not of off on once only or other ought our ours out over own per please same shall she
should so some such than that the their theirs them then there these they this those through to too under
until up upon us very via was we were what when where whether which while who whom why will with within
without would yes you your yours e g eg ie i.e n a
company companies organisation organization organisations organizations business firm team
describe explain provide detail details list indicate confirm specify state include includes including
currently applicable relevant appropriate ensure ensures exist exists existing place
use used uses using ever anywhere
""".split())

# Subjects such as "the vendor" or "your company" say who is answering, not
# what the question is about, so they are removed.
_SUBJECT_RE = re.compile(r"\b(?:the|your)\s+(?:vendor|supplier|service provider|provider|company|organisation|"
                         r"organization|business|firm|contractor)\b(?!s)")
_STEMMER = snowballstemmer.stemmer("english")
_CANON = set(SYNONYM_GROUPS)


def _build_synonym_regex():
    pairs = []
    for canon, phrases in SYNONYM_GROUPS.items():
        for p in phrases:
            pairs.append((p, canon))
    pairs.sort(key=lambda x: -len(x[0]))
    lookup = {p: c for p, c in pairs}
    pattern = re.compile(r"(?<![a-z0-9])(" + "|".join(re.escape(p) for p, _ in pairs) + r")(?![a-z0-9])")
    return pattern, lookup


_SYN_RE, _SYN_LOOKUP = _build_synonym_regex()


def _ascii_lower(text: str) -> str:
    text = unicodedata.normalize("NFKD", text or "")
    text = text.encode("ascii", "ignore").decode("ascii")
    return text.lower()


def normalise(text: str, keep_phrase_words: bool = True) -> str:
    """Lower-case, simplify punctuation and replace synonyms with a group token.

    keep_phrase_words: a multi-word phrase such as "incident response" becomes
    "incidentresponse incident response" (True) or just "incidentresponse" (False)."""
    t = _ascii_lower(text)
    t = t.replace("&", " and ")
    # Cloud region names such as us-east-1 become one token ("us" is a stop word).
    t = re.sub(r"\b(us|eu|ap|ca|sa|me|af|il|mx)-(east|west|north|south|central|northeast|southeast|northwest|"
               r"southwest)-(\d)\b", r"\1\2\3", t)
    t = re.sub(r"(\w)'s\b", r"\1", t)          # vendor's -> vendor
    t = re.sub(r"\b(\w{3,})fications?\b", r"\1fy", t)  # classification -> classify (stems then agree)
    t = re.sub(r"\b(\w{3,})is(e|ed|es|ing|ation|ations)\b", r"\1iz\2", t)  # British -ise -> -ize
    t = re.sub(r"(\w)s'(?=\s|$)", r"\1s", t)    # customers' -> customers
    t = re.sub(r"[^a-z0-9@]+", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    t = _SUBJECT_RE.sub(" ", t)

    def swap(m):
        phrase = m.group(1)
        token = _SYN_LOOKUP[phrase]
        return token + (" " + phrase if keep_phrase_words and " " in phrase else "")

    t = _SYN_RE.sub(swap, t)
    return re.sub(r"\s+", " ", t).strip()


def tokens(text: str, extra_stopwords: set[str] | None = None, keep_phrase_words: bool = True) -> list[str]:
    """Normalised, stop-word-free, stemmed tokens."""
    out = []
    stop = STOPWORDS | (extra_stopwords or set())
    for word in normalise(text, keep_phrase_words).split():
        if word in _CANON:
            out.append(word)
        elif word in stop or len(word) < 2:
            continue
        else:
            stem = _STEMMER.stemWord(word)
            out.append(STEM_ALIASES.get(stem, stem))
    return out


def features(toks: list[str]) -> list[str]:
    """Single words plus pairs of neighbouring words ('access review' style)."""
    return toks + [f"{a}_{b}" for a, b in zip(toks, toks[1:])]


_NEGATION_RE = re.compile(r"\b(not|no|never|none|nobody|neither|nor|without|cannot|prohibit\w*|forbid\w*|"
                          r"disallow\w*|ban|banned)\b|n't\b")
_NOT_NEGATION_RE = re.compile(r"\b(no|not) (later|more|less|fewer) than\b|\bwhether or not\b|\bif (not|no)\b|"
                              r"\byes\s*/\s*no\b|\by\s*/\s*n\b|\byes or no\b|\bno\.\s*\d|\bif so\b")


def has_negation(text: str) -> bool:
    """True if the question is worded negatively ('not', 'never', 'prohibited'...)."""
    t = _NOT_NEGATION_RE.sub(" ", (text or "").lower())
    return bool(_NEGATION_RE.search(t))


# --------------------------------------------------------------------------
# TF-IDF similarity
# --------------------------------------------------------------------------

CONCEPT_BOOST = 1.5  # known security terms (the SYNONYM_GROUPS tokens, e.g. "mfa") count 1.5 times


class TfidfIndex:
    """Plain TF-IDF with cosine similarity.

    weight(term) = (1 + log(count in text)) * idf(term), times CONCEPT_BOOST for known security terms
    idf(term)    = log((1 + N) / (1 + texts containing term)) + 1
    Vectors are scaled to length 1, so the dot product is the cosine.
    """

    def __init__(self, corpus: list[list[str]], concept_boost: float = 1.0):
        self.n = len(corpus)
        self.boost = concept_boost
        df: Counter = Counter()
        for feats in corpus:
            df.update(set(feats))
        self.idf = {t: math.log((1 + self.n) / (1 + c)) + 1 for t, c in df.items()}
        self.unseen_idf = math.log(1 + self.n) + 1

    def vector(self, feats: list[str]) -> dict[str, float]:
        counts = Counter(feats)
        vec = {}
        for t, c in counts.items():
            w = (1 + math.log(c)) * self.idf.get(t, self.unseen_idf)
            vec[t] = w * self.boost if t in _CANON else w
        norm = math.sqrt(sum(w * w for w in vec.values())) or 1.0
        return {t: w / norm for t, w in vec.items()}

    @staticmethod
    def cosine(a: dict[str, float], b: dict[str, float]) -> float:
        if len(a) > len(b):
            a, b = b, a
        return sum(w * b.get(t, 0.0) for t, w in a.items())
