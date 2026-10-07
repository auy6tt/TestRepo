#!/usr/bin/env python3
"""Batch-convert a folder of Microsoft Publisher files to searchable PDFs.

For every .pub file in a folder (and all its subfolders) this script:

  1. converts it to PDF with LibreOffice (one file at a time, with a time limit),
  2. saves a PNG picture of the first page,
  3. records the result in index.xlsx (one row per file),
  4. builds contact-sheet.html, a page of thumbnails the client can browse,
     search and tick ("rebuild this one as a template").

The client's files are never changed. Output folders mirror the original
folder structure.

Examples
  # Count what is there before you quote (nothing is converted):
  python scripts/convert_archive.py ~/jobs/stmarys/incoming --survey

  # Convert every .pub file:
  python scripts/convert_archive.py ~/jobs/stmarys/incoming ~/jobs/stmarys/archive \
      --title "St Mary's School newsletters"

  # Also convert Word, PowerPoint and drawing files LibreOffice can open:
  python scripts/convert_archive.py incoming archive --ext all

Run with --help to see every option.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

sys.path.insert(0, str(Path(__file__).resolve().parent))

from office_tools import (  # noqa: E402
    MODULE_BY_EXT,
    MODULE_LABEL,
    PUBLISHER_FILTER,
    OfficeToolsError,
    Soffice,
    find_soffice,
    has_ole_signature,
    pdf_target,
    read_pdf,
    render_page,
    soffice_version,
    soffice_version_text,
)

EXTENSION_GROUPS = {
    "pub": ["pub"],
    "word": ["doc", "docx", "odt", "rtf", "wpd", "wps"],
    "slides": ["ppt", "pptx", "odp"],
    "drawing": ["odg", "vsd", "vsdx", "cdr", "pmd"],
}

# Folders and files that are never documents (system junk, lock files, Mac extras).
SKIP_DIR_NAMES = {"__MACOSX", "$RECYCLE.BIN", "System Volume Information"}
SKIP_FILE_PREFIXES = ("~$", ".~lock", "._")

STATUS_OK = "Converted"
STATUS_CHECK = "Converted - please check"
STATUS_FAILED = "Failed"
STATUS_SKIPPED = "Skipped"
STATUS_KEY = {STATUS_OK: "ok", STATUS_CHECK: "check", STATUS_FAILED: "failed",
              STATUS_SKIPPED: "skipped"}

# Price guide used by --survey to print a rough quote (US dollars per file).
PRICE_PER_FILE_LOW = 0.50
PRICE_PER_FILE_HIGH = 2.00


@dataclass
class FileRecord:
    source: Path
    rel_path: str  # path inside the client's folder, with forward slashes
    ext: str
    size_bytes: int
    modified: datetime
    sha256: str = ""
    status: str = ""
    pages: int | None = None
    pdf_rel: str = ""  # relative to the output folder
    preview_rel: str = ""
    has_text: bool | None = None
    snippet: str = ""
    fonts: list[str] = field(default_factory=list)
    issues: list[tuple[str, str]] = field(default_factory=list)  # (short, full) wording
    duplicate_of: str = ""
    reused: bool = False  # PDF kept from an earlier run
    seconds: float = 0.0
    log: str = ""

    def add_issue(self, short: str, full: str) -> None:
        self.issues.append((short, full))

    @property
    def problem(self) -> str:
        """Full explanation, for index.xlsx and the log."""
        return " ".join(full for _, full in self.issues)

    @property
    def note(self) -> str:
        """Short version, for the contact sheet cards."""
        return "; ".join(short for short, _ in self.issues)

    @property
    def folder(self) -> str:
        parent = Path(self.rel_path).parent.as_posix()
        return "" if parent == "." else parent

    @property
    def name(self) -> str:
        return Path(self.rel_path).name


# ---------------------------------------------------------------------------
# Finding files
# ---------------------------------------------------------------------------

def parse_extensions(text: str) -> list[str]:
    wanted: list[str] = []
    for token in (t.strip().lower().lstrip(".") for t in text.split(",")):
        if not token:
            continue
        if token == "all":
            wanted.extend(MODULE_BY_EXT)
        elif token in EXTENSION_GROUPS:
            wanted.extend(EXTENSION_GROUPS[token])
        else:
            wanted.append(token)
    return sorted(set(wanted))


def is_inside(path: Path, folder: Path) -> bool:
    try:
        path.relative_to(folder)
        return True
    except ValueError:
        return False


def find_files(root: Path, extensions: list[str], exclude: Path | None) -> list[Path]:
    found: list[Path] = []
    for dirpath, dirnames, filenames in os.walk(root):
        current = Path(dirpath)
        dirnames[:] = sorted(
            d for d in dirnames
            if not d.startswith(".") and d not in SKIP_DIR_NAMES
            and not (exclude and is_inside((current / d).resolve(), exclude))
        )
        for name in sorted(filenames, key=str.lower):
            if name.startswith(SKIP_FILE_PREFIXES) or name.startswith("."):
                continue
            if Path(name).suffix.lower().lstrip(".") in extensions:
                found.append(current / name)
    return found


def sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def make_record(path: Path, root: Path) -> FileRecord:
    stat = path.stat()
    return FileRecord(
        source=path,
        rel_path=path.relative_to(root).as_posix(),
        ext=path.suffix.lower().lstrip("."),
        size_bytes=stat.st_size,
        modified=datetime.fromtimestamp(stat.st_mtime).replace(microsecond=0),
    )


# ---------------------------------------------------------------------------
# Survey (count before quoting)
# ---------------------------------------------------------------------------

def survey(files: list[Path], root: Path) -> None:
    records = [make_record(f, root) for f in files]
    print(f"Survey of {root} (nothing was converted)\n")
    if not records:
        print("No matching files found. Check the folder, or use --ext all.")
        return
    by_ext = Counter(r.ext for r in records)
    size_by_ext = Counter()
    for r in records:
        size_by_ext[r.ext] += r.size_bytes
    for ext, count in by_ext.most_common():
        print(f"  .{ext:<5} {count:>6} file{'s' if count != 1 else ' '}  "
              f"{size_by_ext[ext] / 1_048_576:8.1f} MB")
    seen: dict[str, str] = {}
    duplicates = empty = not_publisher = 0
    for r in records:
        if r.size_bytes == 0:
            empty += 1
            continue
        r.sha256 = sha256_of(r.source)
        if r.sha256 in seen:
            duplicates += 1
        else:
            seen[r.sha256] = r.rel_path
        if r.ext == "pub" and not has_ole_signature(r.source):
            not_publisher += 1
    folders = {r.folder for r in records}
    dates = sorted(r.modified for r in records)
    unique = len(records) - duplicates - empty
    print(f"\n  Folders with files:      {len(folders)}")
    print(f"  Exact duplicates:        {duplicates} (same content as another file)")
    print(f"  Empty files:             {empty}")
    print(f"  .pub files that do not look like Publisher 98+ files: {not_publisher}")
    print(f"  Dates (last modified):   {dates[0]:%Y-%m-%d} to {dates[-1]:%Y-%m-%d}")
    print(f"\n  Unique files to convert: {unique}")
    print(f"  At ${PRICE_PER_FILE_LOW:.2f}-${PRICE_PER_FILE_HIGH:.2f} a file that is about "
          f"${unique * PRICE_PER_FILE_LOW:,.0f}-${unique * PRICE_PER_FILE_HIGH:,.0f} "
          "for the archive part (templates are priced separately).")


# ---------------------------------------------------------------------------
# Converting
# ---------------------------------------------------------------------------

def looks_garbled(text: str) -> bool:
    """True when most letters are accented Latin-1 letters (À to ÿ). That is
    what Cyrillic, Greek or similar text looks like when it has been read with
    the wrong character set. Normal French or German text scores far lower."""
    letters = [c for c in text if c.isalpha()]
    if len(letters) < 6:
        return False
    odd = sum(1 for c in letters if "À" <= c <= "ÿ")
    return odd / len(letters) > 0.4


def unique_output(rel_path: str, ext: str, used: set[str]) -> str:
    """Output path (without extension) inside pdf/ or previews/. If two
    originals would get the same name (Newsletter.pub and Newsletter.docx),
    the second keeps its file type in the name."""
    base = Path(rel_path).with_suffix("").as_posix()
    candidate = base
    if candidate.lower() in used:
        candidate = f"{base}_{ext}"
        number = 2
        while candidate.lower() in used:
            candidate = f"{base}_{ext}_{number}"
            number += 1
    used.add(candidate.lower())
    return candidate


def convert_one(record: FileRecord, out_root: Path, lo: Soffice, args, used: set[str]) -> None:
    if record.size_bytes == 0:
        record.status = STATUS_SKIPPED
        record.add_issue("Empty file", "Empty file (0 bytes). Ask the client for another copy.")
        return
    if record.ext == "pub" and not has_ole_signature(record.source):
        record.status = STATUS_FAILED
        record.add_issue(
            "Not a readable Publisher file",
            "This does not look like a Publisher 98 or later file (its first bytes are "
            "wrong). It may be damaged, a very old Publisher version, or another kind of "
            "file with a .pub name. Ask the client for another copy.")
        return

    stem = unique_output(record.rel_path, record.ext, used)
    pdf_path = out_root / "pdf" / (stem + ".pdf")
    png_path = out_root / "previews" / (stem + ".png")
    record.pdf_rel = pdf_path.relative_to(out_root).as_posix()
    needs_check = False

    record.reused = (pdf_path.exists() and not args.overwrite
                     and pdf_path.stat().st_mtime >= record.source.stat().st_mtime)
    if not record.reused:
        result = lo.convert(
            record.source, pdf_path,
            target=pdf_target(record.ext, pdfa=args.pdfa),
            infilter=PUBLISHER_FILTER if record.ext == "pub" else None,
        )
        record.seconds = result.seconds
        record.log = result.log
        if not result.ok:
            record.status = STATUS_FAILED
            record.add_issue(result.short, result.message)
            record.pdf_rel = ""
            return
        expected = MODULE_LABEL.get(MODULE_BY_EXT.get(record.ext, ""), "")
        if expected and result.opened_as and result.opened_as != expected:
            needs_check = True
            record.add_issue(f"Opened as {result.opened_as} - check",
                             f"LibreOffice opened this as a {result.opened_as} document, not "
                             f"{expected}. Compare the PDF with the original.")

    try:
        info = read_pdf(pdf_path)
    except Exception as error:  # a broken PDF
        record.status = STATUS_FAILED
        record.add_issue("PDF unreadable", f"The PDF could not be read back: {error}")
        return
    if info.pages == 0:
        record.status = STATUS_FAILED
        record.add_issue("PDF has no pages", "The PDF has no pages.")
        return
    record.pages = info.pages
    record.has_text = info.has_text
    record.snippet = info.snippet
    record.fonts = info.fonts
    if not info.has_text:
        needs_check = True
        record.add_issue("No searchable text",
                         "No selectable text found, so searching will not find words in "
                         "this file (it may be all pictures).")
    elif looks_garbled(info.snippet):
        needs_check = True
        record.add_issue("Text may be garbled",
                         "The text looks garbled (letters such as \u00d0, \u00f1, \u00ea instead of "
                         "real words). This happens with some older Publisher files written "
                         "in non-Western languages. Compare with the original.")

    try:
        render_page(pdf_path, png_path, width_px=args.preview_width)
        record.preview_rel = png_path.relative_to(out_root).as_posix()
    except Exception as error:
        needs_check = True
        record.add_issue("No preview picture", f"Preview picture could not be made: {error}")

    if record.duplicate_of:
        record.add_issue(f"Same as {Path(record.duplicate_of).name}",
                         f"Exact copy of {record.duplicate_of}.")
    record.status = STATUS_CHECK if needs_check else STATUS_OK


# ---------------------------------------------------------------------------
# Writing the index spreadsheet
# ---------------------------------------------------------------------------

def write_index(records: list[FileRecord], out_root: Path, meta: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def clean(value):
        if isinstance(value, str):
            return ILLEGAL_CHARACTERS_RE.sub("", value)
        return value

    wb = Workbook()
    ws = wb.active
    ws.title = "Index"
    headers = ["#", "Original file", "Folder", "Date modified", "Status", "Pages",
               "Output PDF", "Preview picture", "Searchable text", "Opening words",
               "Fonts in PDF", "Problem or note", "Size (KB)"]
    widths = [5, 44, 26, 17, 24, 7, 44, 30, 11, 50, 30, 60, 10]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    fills = {
        STATUS_OK: ("E3F4E8", "1E6B35"),
        STATUS_CHECK: ("FFF2CC", "7A5200"),
        STATUS_FAILED: ("FDE2E1", "9C1C1C"),
        STATUS_SKIPPED: ("ECECEC", "444444"),
    }
    for number, r in enumerate(records, start=1):
        searchable = "" if r.has_text is None else ("Yes" if r.has_text else "No")
        row = [number, r.rel_path, r.folder or "(top folder)", r.modified, r.status,
               r.pages, r.pdf_rel, r.preview_rel, searchable, r.snippet,
               ", ".join(r.fonts), r.problem, round(r.size_bytes / 1024, 1)]
        ws.append([clean(v) for v in row])
        excel_row = ws.max_row
        ws.cell(excel_row, 4).number_format = "yyyy-mm-dd hh:mm"
        bg, fg = fills.get(r.status, ("FFFFFF", "000000"))
        status_cell = ws.cell(excel_row, 5)
        status_cell.fill = PatternFill("solid", fgColor=bg)
        status_cell.font = Font(bold=True, color=fg)
        for column, rel in ((7, r.pdf_rel), (8, r.preview_rel)):
            if rel:
                cell = ws.cell(excel_row, column)
                cell.hyperlink = rel
                cell.style = "Hyperlink"
        for column in (2, 10, 11, 12):
            ws.cell(excel_row, column).alignment = Alignment(wrap_text=True, vertical="top")
        for column in (1, 3, 4, 5, 6, 7, 8, 9, 13):
            ws.cell(excel_row, column).alignment = Alignment(vertical="top", wrap_text=column in (3, 7, 8))
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"

    summary = wb.create_sheet("Summary")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 70
    summary.append([meta["title"]])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    rows = [
        ("Folder converted", meta["source_name"]),
        ("Date of conversion", meta["generated"]),
        ("Files found", meta["counts"]["files"]),
        ("Converted", meta["counts"]["converted"]),
        ("Converted - please check", meta["counts"]["check"]),
        ("Failed", meta["counts"]["failed"]),
        ("Skipped (empty)", meta["counts"]["skipped"]),
        ("Exact duplicates (still converted)", meta["counts"]["duplicates"]),
        ("Total PDF pages", meta["counts"]["pages"]),
        ("PDF type", "PDF/A-2b (archive standard)" if meta["pdfa"] else "Standard PDF"),
        ("Converted with", meta["libreoffice"]),
    ]
    for key, value in rows:
        summary.append([key, value])
        summary.cell(summary.max_row, 1).font = Font(bold=True)
        summary.cell(summary.max_row, 2).alignment = Alignment(horizontal="left")
    summary.append([])
    notes = [
        "How to use this list",
        "Click a link in the 'Output PDF' column to open that PDF (keep this file in "
        "the same folder as the 'pdf' and 'previews' folders).",
        "Use the filter arrows in the header row to show, for example, only 2019 files "
        "or only rows that need checking.",
        "'Opening words' holds the first words of each file, so Ctrl+F finds a "
        "newsletter by its headline.",
        "Rows marked 'Failed' or 'please check' explain why in 'Problem or note'.",
    ]
    for line in notes:
        summary.append([line])
        summary.merge_cells(start_row=summary.max_row, start_column=1,
                            end_row=summary.max_row, end_column=2)
    summary.cell(summary.max_row - len(notes) + 1, 1).font = Font(bold=True)
    wb.properties.title = meta["title"]
    wb.properties.creator = "Publisher Rescue kit"
    path = out_root / "index.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Writing the contact sheet
# ---------------------------------------------------------------------------

CONTACT_SHEET_CSS = """
:root{--ink:#1d2433;--muted:#566070;--line:#d8dde6;--bg:#f4f5f7;--card:#fff;
--accent:#1f3a5f;--ok:#1e6b35;--ok-bg:#e3f4e8;--check:#7a5200;--check-bg:#fff2cc;
--failed:#9c1c1c;--failed-bg:#fde2e1;--skipped:#444;--skipped-bg:#ececec;color-scheme:light}
*{box-sizing:border-box}
body{margin:0;font:15px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;color:var(--ink);background:var(--bg)}
.wrap{max-width:1240px;margin:0 auto;padding:0 16px}
header{background:var(--card);border-bottom:1px solid var(--line);padding:28px 0 18px}
h1{font-size:1.65rem;line-height:1.2;margin:0 0 6px}
.sub{color:var(--muted);margin:0}
.stats{display:flex;flex-wrap:wrap;gap:8px;margin:16px 0 12px;padding:0;list-style:none}
.stats li{background:var(--bg);border:1px solid var(--line);border-radius:8px;padding:6px 12px}
.stats b{font-size:1.1rem}
.howto{margin:0;max-width:75ch;color:var(--muted)}
.toolbar{position:sticky;top:0;z-index:5;background:rgba(255,255,255,.97);border-bottom:1px solid var(--line);padding:10px 0}
.toolbar .wrap{display:flex;flex-wrap:wrap;gap:8px;align-items:center}
.toolbar input[type=search],.toolbar select{font:inherit;padding:7px 10px;border:1px solid #b9c0cc;border-radius:8px;background:#fff;color:var(--ink);min-width:0}
.toolbar input[type=search]{flex:1 1 220px}
.toolbar button{font:inherit;font-weight:600;padding:7px 14px;border-radius:8px;border:1px solid var(--accent);background:var(--accent);color:#fff;cursor:pointer}
.toolbar button:focus-visible,.toolbar input:focus-visible,.toolbar select:focus-visible,a:focus-visible,.pick input:focus-visible{outline:3px solid #f2b632;outline-offset:2px}
#shown,#pickcount{color:var(--muted);font-size:.9rem}
#copyout{width:100%;min-height:90px;margin-top:8px;font:13px/1.4 ui-monospace,Menlo,Consolas,monospace}
.folder h2{font-size:1rem;font-weight:600;margin:28px 0 10px;color:var(--muted);word-break:break-word}
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(180px,1fr));gap:14px}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;overflow:hidden;display:flex;flex-direction:column;margin:0}
.thumb{display:flex;align-items:center;justify-content:center;aspect-ratio:4/5;background:#e9ecf1;padding:10px;text-decoration:none}
.thumb img{max-width:100%;max-height:100%;object-fit:contain;background:#fff;box-shadow:0 1px 5px rgba(0,0,0,.2)}
.thumb.none{color:var(--muted);font-size:.9rem;text-align:center}
figcaption{padding:10px 12px 12px;display:flex;flex-direction:column;gap:4px;flex:1}
.name{font-weight:600;word-break:break-word;line-height:1.3}
.meta{color:var(--muted);font-size:.84rem}
.badge{align-self:flex-start;font-size:.76rem;font-weight:700;padding:2px 9px;border-radius:999px}
.b-ok{background:var(--ok-bg);color:var(--ok)}.b-check{background:var(--check-bg);color:var(--check)}
.b-failed{background:var(--failed-bg);color:var(--failed)}.b-skipped{background:var(--skipped-bg);color:var(--skipped)}
.note{font-size:.8rem;color:var(--muted);margin:0}
.pick{font-size:.86rem;display:flex;gap:6px;align-items:center;margin-top:auto;padding-top:6px;cursor:pointer}
.pick input{width:18px;height:18px;margin:0}
.empty{padding:40px 0;color:var(--muted)}
footer{color:var(--muted);font-size:.85rem;padding:28px 0 40px}
@media (max-width:480px){.grid{grid-template-columns:repeat(2,minmax(0,1fr));gap:10px}h1{font-size:1.35rem}}
@media print{.toolbar,.pick{display:none!important}body{background:#fff}header{border:0}
.grid{grid-template-columns:repeat(4,1fr)}.card{break-inside:avoid}.folder h2{break-after:avoid}}
"""

CONTACT_SHEET_JS = """
(function(){
  var q=document.getElementById('q'),st=document.getElementById('status');
  var cards=[].slice.call(document.querySelectorAll('.card'));
  var folders=[].slice.call(document.querySelectorAll('.folder'));
  var shown=document.getElementById('shown');
  function apply(){
    var term=q.value.trim().toLowerCase(),want=st.value,n=0;
    cards.forEach(function(c){
      var s=c.getAttribute('data-status');
      var ok=(!term||c.getAttribute('data-search').indexOf(term)>-1)&&
        (want==='all'||s===want||(want==='attention'&&(s==='check'||s==='failed'||s==='skipped')));
      c.hidden=!ok; if(ok){n++;}
    });
    folders.forEach(function(f){f.hidden=!f.querySelector('.card:not([hidden])');});
    shown.textContent=n+' of '+cards.length+' shown';
  }
  q.addEventListener('input',apply);st.addEventListener('change',apply);
  var key='publisher-rescue-ticks:'+document.title,saved={};
  try{saved=JSON.parse(localStorage.getItem(key)||'{}');}catch(e){saved={};}
  var boxes=[].slice.call(document.querySelectorAll('.pick input'));
  var count=document.getElementById('pickcount');
  function ticked(){return boxes.filter(function(b){return b.checked;}).map(function(b){return b.getAttribute('data-file');});}
  function update(){count.textContent=ticked().length+' ticked';}
  boxes.forEach(function(b){
    if(saved[b.getAttribute('data-file')]){b.checked=true;}
    b.addEventListener('change',function(){
      saved[b.getAttribute('data-file')]=b.checked;
      try{localStorage.setItem(key,JSON.stringify(saved));}catch(e){}
      update();
    });
  });
  document.getElementById('copy').addEventListener('click',function(){
    var list=ticked(),out=document.getElementById('copyout');
    var text=list.length?'Layouts we reuse. Please rebuild these as templates:\\n'+
      list.map(function(f){return '- '+f;}).join('\\n'):'No layouts ticked yet.';
    function show(){out.hidden=false;out.value=text;out.focus();out.select();}
    if(navigator.clipboard&&window.isSecureContext){
      navigator.clipboard.writeText(text).then(function(){
        count.textContent=list.length+' ticked - list copied, paste it into an email';
      },show);
    }else{show();}
  });
  apply();update();
})();
"""


def _url(rel: str) -> str:
    return quote(rel, safe="/")


def write_contact_sheet(records: list[FileRecord], out_root: Path, meta: dict) -> Path:
    esc = html.escape
    counts = meta["counts"]
    groups: dict[str, list[FileRecord]] = {}
    for r in records:
        groups.setdefault(r.folder, []).append(r)

    sections = []
    for folder in sorted(groups, key=lambda f: (f != "", f.lower())):
        cards = []
        for r in groups[folder]:
            key = STATUS_KEY.get(r.status, "failed")
            if r.preview_rel:
                thumb = (f'<a class="thumb" href="{_url(r.pdf_rel)}" target="_blank" '
                         f'rel="noopener" title="Open the PDF"><img src="{_url(r.preview_rel)}" '
                         f'alt="First page of {esc(r.name)}" loading="lazy"></a>')
            else:
                thumb = '<div class="thumb none">No preview</div>'
            bits = [r.modified.strftime("%d %b %Y")]
            if r.pages:
                bits.append(f"{r.pages} page{'s' if r.pages != 1 else ''}")
            bits.append(r.ext.upper())
            note = (f'<p class="note" title="{esc(r.problem)}">{esc(r.note)}</p>'
                    if r.issues else "")
            search = esc(" ".join([r.rel_path, r.snippet, r.status]).lower())
            pick = (f'<label class="pick"><input type="checkbox" data-file="{esc(r.rel_path)}"> '
                    'Rebuild as template</label>') if r.pdf_rel else ""
            cards.append(
                f'<figure class="card" data-status="{key}" data-search="{search}">{thumb}'
                f'<figcaption><span class="name">{esc(r.name)}</span>'
                f'<span class="meta">{esc(" · ".join(bits))}</span>'
                f'<span class="badge b-{key}">{esc(r.status)}</span>{note}{pick}'
                '</figcaption></figure>'
            )
        heading = esc(folder) if folder else "Top folder"
        sections.append(f'<section class="folder"><h2>{heading}</h2>'
                        f'<div class="grid">{"".join(cards)}</div></section>')

    body = "".join(sections) or '<p class="empty">No files were found.</p>'
    attention = counts["check"] + counts["failed"] + counts["skipped"]
    page = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(meta["title"])} - contact sheet</title>
<style>{CONTACT_SHEET_CSS}</style>
</head>
<body>
<header><div class="wrap">
<h1>{esc(meta["title"])}</h1>
<p class="sub">Contact sheet of the PDF archive, made {esc(meta["generated"])}</p>
<ul class="stats">
<li><b>{counts["files"]}</b> files</li>
<li><b>{counts["converted"] + counts["check"]}</b> converted</li>
<li><b>{counts["pages"]}</b> pages</li>
<li><b>{attention}</b> need attention</li>
</ul>
<p class="howto">Click a picture to open that PDF. Tick <b>Rebuild as template</b> on the
layouts you reuse most (your newsletter, weekly sheet, certificates), then press
<b>Copy my ticks</b> and paste the list into an email.</p>
</div></header>
<div class="toolbar"><div class="wrap">
<input id="q" type="search" placeholder="Search names, folders and opening words" aria-label="Search">
<select id="status" aria-label="Show">
<option value="all">All files</option>
<option value="ok">Converted</option>
<option value="attention">Need attention</option>
</select>
<button id="copy" type="button">Copy my ticks</button>
<span id="pickcount" aria-live="polite"></span>
<span id="shown" aria-live="polite"></span>
<textarea id="copyout" hidden aria-label="Your ticked layouts"></textarea>
</div></div>
<main class="wrap">{body}</main>
<footer class="wrap">Keep this page in the same folder as the <code>pdf</code> and
<code>previews</code> folders so the links keep working. The full list with details is in
<code>index.xlsx</code>.</footer>
<script>{CONTACT_SHEET_JS}</script>
</body>
</html>
"""
    path = out_root / "contact-sheet.html"
    path.write_text(page, encoding="utf-8")
    return path


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Convert a folder of Publisher (.pub) files to searchable PDFs, "
                    "with preview pictures, index.xlsx and contact-sheet.html.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Example: python scripts/convert_archive.py incoming archive --title \"St Mary's newsletters\"",
    )
    parser.add_argument("input", type=Path, help="folder with the client's files (subfolders included)")
    parser.add_argument("output", type=Path, nargs="?",
                        help="folder for the PDFs, previews, index.xlsx and contact sheet")
    parser.add_argument("--survey", action="store_true",
                        help="only count files, duplicates and dates (for quoting); convert nothing")
    parser.add_argument("--ext", default="pub",
                        help="file types to convert, comma separated: pub (default), word, "
                             "slides, drawing, all, or extensions like pub,docx,odg")
    parser.add_argument("--title", help="title for the index and contact sheet "
                                        "(default: the input folder's name)")
    parser.add_argument("--timeout", type=int, default=180,
                        help="seconds to wait for one file before giving up (default 180)")
    parser.add_argument("--no-pdfa", dest="pdfa", action="store_false",
                        help="make ordinary PDFs instead of PDF/A-2b archive PDFs")
    parser.add_argument("--preview-width", type=int, default=600,
                        help="width of preview pictures in pixels (default 600)")
    parser.add_argument("--overwrite", action="store_true",
                        help="convert again even if a PDF from an earlier run exists")
    parser.add_argument("--limit", type=int, default=0,
                        help="only do the first N files (for a quick trial run)")
    parser.add_argument("--soffice", help="path to LibreOffice's soffice program, if not found")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    in_root = args.input.expanduser().resolve()
    if not in_root.is_dir():
        print(f"Input folder not found: {in_root}", file=sys.stderr)
        return 2
    extensions = parse_extensions(args.ext)
    out_root = args.output.expanduser().resolve() if args.output else None
    files = find_files(in_root, extensions, exclude=out_root)

    if args.survey:
        survey(files, in_root)
        return 0
    if out_root is None:
        print("Please give an output folder (or use --survey).", file=sys.stderr)
        return 2
    if out_root == in_root:
        print("The output folder must be different from the input folder.", file=sys.stderr)
        return 2
    if args.limit:
        files = files[: args.limit]
    if not files:
        print(f"No {', '.join('.' + e for e in extensions)} files found in {in_root}.")
        return 0

    try:
        soffice = find_soffice(args.soffice)
    except OfficeToolsError as error:
        print(error, file=sys.stderr)
        return 2
    version = soffice_version(soffice)
    if args.pdfa and version and version < (7, 4):
        print("LibreOffice is older than 7.4, so ordinary PDFs will be made instead of PDF/A.")
        args.pdfa = False
    lo_text = soffice_version_text(soffice)

    out_root.mkdir(parents=True, exist_ok=True)
    title = args.title or in_root.name
    print(f"Converting {len(files)} file(s) from {in_root}\n  to {out_root}\n  with {lo_text}\n")

    records: list[FileRecord] = []
    used_names: set[str] = set()
    seen_hashes: dict[str, str] = {}
    interrupted = False
    with Soffice(soffice, timeout=args.timeout) as lo:
        try:
            for number, path in enumerate(files, start=1):
                record = make_record(path, in_root)
                if record.size_bytes:
                    record.sha256 = sha256_of(path)
                    if record.sha256 in seen_hashes:
                        record.duplicate_of = seen_hashes[record.sha256]
                    else:
                        seen_hashes[record.sha256] = record.rel_path
                print(f"[{number}/{len(files)}] {record.rel_path} ... ", end="", flush=True)
                convert_one(record, out_root, lo, args, used_names)
                records.append(record)
                detail = f"{record.pages} page(s)" if record.pages else record.problem
                if record.reused:
                    detail += ", kept from an earlier run"
                print(f"{record.status}" + (f" - {detail}" if detail else ""), flush=True)
        except KeyboardInterrupt:
            interrupted = True
            print("\nStopped. Writing the index for the files done so far; "
                  "run the same command again to carry on.")

    counts = {
        "files": len(records),
        "converted": sum(r.status == STATUS_OK for r in records),
        "check": sum(r.status == STATUS_CHECK for r in records),
        "failed": sum(r.status == STATUS_FAILED for r in records),
        "skipped": sum(r.status == STATUS_SKIPPED for r in records),
        "duplicates": sum(bool(r.duplicate_of) for r in records),
        "pages": sum(r.pages or 0 for r in records),
    }
    meta = {
        "title": title,
        "source_name": in_root.name,
        "generated": datetime.now().strftime("%d %B %Y, %H:%M"),
        "counts": counts,
        "pdfa": args.pdfa,
        "libreoffice": lo_text,
        "interrupted": interrupted,
    }
    index_path = write_index(records, out_root, meta)
    sheet_path = write_contact_sheet(records, out_root, meta)

    with open(out_root / "conversion-log.txt", "w", encoding="utf-8") as log:
        log.write(f"{title}\nConverted {meta['generated']} with {lo_text}\n\n")
        for r in records:
            log.write(f"{r.status:<26} {r.rel_path}\n")
            if r.problem:
                log.write(f"    {r.problem}\n")
            if r.status == STATUS_FAILED and r.log:
                log.write("    LibreOffice said: " + r.log.replace("\n", "\n    ") + "\n")
    summary = dict(meta)
    summary["failed_files"] = [{"file": r.rel_path, "problem": r.problem}
                               for r in records if r.status in (STATUS_FAILED, STATUS_SKIPPED)]
    summary["check_files"] = [{"file": r.rel_path, "problem": r.problem}
                              for r in records if r.status == STATUS_CHECK]
    (out_root / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print(f"\nDone: {counts['converted'] + counts['check']} of {counts['files']} converted "
          f"({counts['pages']} pages); {counts['check']} to check, {counts['failed']} failed, "
          f"{counts['skipped']} skipped.")
    print(f"  Index:         {index_path}\n  Contact sheet: {sheet_path}")
    return 130 if interrupted else 0


if __name__ == "__main__":
    sys.exit(main())
