#!/usr/bin/env python3
"""Create the FICTIONAL sample archive for St Aidan's Church, Wrenford.

Real .pub files cannot be generated without Microsoft Publisher, so this
makes stand-ins in other formats LibreOffice converts the same way:
Word newsletters and service sheets (.docx, .odt), posters (.odg) and a
certificate (.odp), styled like typical old Publisher files. It also adds
one deliberately damaged .pub file to show how failures are reported.

File dates are set to when each file would have been saved, so the index
shows realistic dates.

Example
  python scripts/make_sample_archive.py samples/st-aidans-wrenford/originals
"""

from __future__ import annotations

import argparse
import base64
import io
import os
import random
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Mm, Pt, RGBColor  # noqa: E402

import artwork  # noqa: E402
import docx_tools as dt  # noqa: E402
from office_tools import OLE_SIGNATURE, Soffice  # noqa: E402

CHURCH = "St Aidan's Church, Wrenford"

NEWSLETTERS = [
    {
        "file": "Newsletters/2019/Parish News Spring 2019.docx",
        "saved": datetime(2019, 3, 14, 19, 42),
        "issue": "Spring 2019  -  Issue 84",
        "colour": "0047AB",
        "art": "flowers",
        "stories": [
            ("Lent and Easter at St Aidan's",
             ["Lent begins on Ash Wednesday, 6 March, with a said service at 7.30 pm. "
              "The Lent lunches return to the church hall every Friday: soup, bread and "
              "cheese for a donation to the food bank.",
              "Easter Day is 21 April. Please bring a flower for the Easter garden."]),
            ("Bells ring again",
             ["After two years of silence the bells are ringing again, thanks to the "
              "repairs funded by last summer's fair. New ringers are always welcome on "
              "Wednesday evenings."]),
            ("Thank you, Margaret",
             ["After 25 years Margaret Bell is handing over the flower rota. Thank you, "
              "Margaret, for every arrangement."]),
        ],
        "diary": ["Wed 6 Mar  Ash Wednesday, 7.30 pm", "Sun 31 Mar  Mothering Sunday",
                  "Fri 19 Apr  Good Friday walk, 11 am", "Sun 21 Apr  Easter Day"],
    },
    {
        "file": "Newsletters/2019/Parish News Harvest 2019.docx",
        "saved": datetime(2019, 9, 24, 21, 10),
        "issue": "Harvest 2019  -  Issue 86",
        "colour": "B8860B",
        "art": "harvest",
        "stories": [
            ("Harvest Festival, 29 September",
             ["Our harvest service is at 10 am on Sunday 29 September, followed by a "
              "bring-and-share lunch. Tins and packets will go to the Wrenford food bank.",
              "The harvest supper is on Saturday 5 October in the hall. Tickets cost "
              "8 pounds from the parish office."]),
            ("Summer fair raises a record",
             ["The summer fair raised 3,412 pounds for the church roof. The dog show was "
              "won, for the third year running, by Biscuit."]),
            ("New toddler group",
             ["Little Wrens, a group for babies, toddlers and their grown-ups, starts on "
              "Wednesday mornings in October."]),
        ],
        "diary": ["Sun 29 Sep  Harvest Festival, 10 am", "Sat 5 Oct  Harvest supper, 7 pm",
                  "Wed 9 Oct  Little Wrens starts, 9.15 am"],
    },
    {
        "file": "Newsletters/2022/Parish News Advent 2022.docx",
        "saved": datetime(2022, 11, 20, 16, 5),
        "issue": "Advent 2022  -  Issue 99",
        "colour": "1B3A6B",
        "art": "candles",
        "stories": [
            ("Christmas at St Aidan's",
             ["Advent begins on Sunday 27 November. Our Christmas services are listed "
              "below; the crib service on Christmas Eve is especially for children.",
              "The Christmas tree festival returns this year with 30 trees decorated by "
              "local groups and businesses."]),
            ("Warm Space on Mondays",
             ["The church hall is open as a Warm Space every Monday from 10 am to 3 pm, "
              "with soup, newspapers and company. Everyone is welcome."]),
        ],
        "diary": ["Sun 27 Nov  Advent Sunday", "Sun 18 Dec  Carols by candlelight, 6.30 pm",
                  "Sat 24 Dec  Crib service, 4 pm", "Sun 25 Dec  Christmas Day, 10 am"],
    },
]

SERVICE_SHEETS = [
    {
        "file": "Service sheets/2024-03-10 Mothering Sunday.docx",
        "saved": datetime(2024, 3, 8, 11, 20),
        "title": "Mothering Sunday",
        "date": "Sunday 10 March 2024, 10.00 am",
        "art": "flowers",
        "hymns": [("Hymn", "192"), ("First reading", "1 Samuel 1.20-28"),
                  ("Psalm", "34"), ("Gospel", "Luke 2.33-35"), ("Hymn", "365"),
                  ("Final hymn", "16")],
        "convert_to": None,
    },
    {
        "file": "Service sheets/2024-12-24 Crib service.odt",
        "saved": datetime(2024, 12, 20, 15, 45),
        "title": "Crib Service",
        "date": "Christmas Eve, Tuesday 24 December 2024, 4.00 pm",
        "art": "candles",
        "hymns": [("Carol", "Once in royal David's city"), ("Story", "The shepherds"),
                  ("Carol", "Away in a manger"), ("Story", "The wise men"),
                  ("Carol", "O come, all ye faithful")],
        "convert_to": "odt",
    },
]


# ---------------------------------------------------------------------------
# Word stand-ins
# ---------------------------------------------------------------------------

def _font(run, name, size, bold=False, colour=None, italic=False):
    run.font.name = name
    rfonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
    rfonts.set(qn("w:cs"), name)
    rfonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if colour:
        run.font.color.rgb = RGBColor.from_string(colour)
    return run


def old_newsletter(spec: dict, path: Path) -> None:
    """A two-column newsletter in the style of an early-2000s Publisher file."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(12.7))
    dt.set_properties(doc, f"Parish News {spec['issue']}")
    doc.core_properties.author = "Parish Office"

    masthead = doc.add_paragraph()
    masthead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = masthead._p.get_or_add_pPr()
    dt.paragraph_shading(pPr, spec["colour"])
    masthead.paragraph_format.space_after = Pt(0)
    _font(masthead.add_run("PARISH NEWS"), "Arial", 40, True, "FFFFFF")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dt.paragraph_shading(sub._p.get_or_add_pPr(), spec["colour"])
    _font(sub.add_run(f"{CHURCH}   *   {spec['issue']}"), "Arial", 12, True, "FFFF66")
    sub.paragraph_format.space_after = Pt(8)

    body = doc.add_section(WD_SECTION.CONTINUOUS)
    dt.shrink_paragraph(doc.paragraphs[-1])
    dt.set_columns(body, 2, 14)
    cols = body._sectPr.find(qn("w:cols"))
    cols.set(qn("w:sep"), "1")  # a line between the columns, very Publisher

    for number, (headline, paragraphs) in enumerate(spec["stories"]):
        head = doc.add_paragraph()
        head.paragraph_format.space_before = Pt(6 if number else 0)
        head.paragraph_format.space_after = Pt(3)
        head.paragraph_format.keep_with_next = True
        _font(head.add_run(headline), "Arial", 15, True, "C00000")
        for text in paragraphs:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            para.paragraph_format.space_after = Pt(6)
            _font(para.add_run(text), "Times New Roman", 11.5)
        if number == 0:
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.add_run().add_picture(io.BytesIO(artwork.SCENES[spec["art"]]()), width=Mm(84))

    box_title = doc.add_paragraph()
    box_title.paragraph_format.space_before = Pt(10)
    box_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dt.paragraph_border(box_title._p.get_or_add_pPr(), top=(18, "000000", 4),
                        left=(18, "000000", 4), right=(18, "000000", 4))
    _font(box_title.add_run("DIARY DATES"), "Arial", 13, True, "000000")
    for index, line in enumerate(spec["diary"]):
        entry = doc.add_paragraph()
        entry.paragraph_format.space_after = Pt(1 if index < len(spec["diary"]) - 1 else 6)
        edges = {"left": (18, "000000", 4), "right": (18, "000000", 4)}
        if index == len(spec["diary"]) - 1:
            edges["bottom"] = (18, "000000", 4)
        dt.paragraph_border(entry._p.get_or_add_pPr(), **edges)
        _font(entry.add_run(line), "Times New Roman", 11)
    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(closing.add_run("Vicar: Revd Helen Marsh  -  Tel 01632 960412"), "Arial", 9,
          italic=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def old_service_sheet(spec: dict, path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(148), Mm(210)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(12))
    dt.set_properties(doc, spec["title"])
    doc.core_properties.author = "Parish Office"

    def centred(text, size, bold=False, colour=None, after=4, italic=False):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(after)
        _font(para.add_run(text), "Times New Roman", size, bold, colour, italic)
        return para

    centred("ST AIDAN'S CHURCH, WRENFORD", 13, True, "4B0082", 2)
    centred(spec["title"], 24, True, "4B0082", 2)
    centred(spec["date"], 11, italic=True, after=8)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(io.BytesIO(artwork.SCENES[spec["art"]]()), width=Mm(95))
    centred("Welcome! Please take this sheet home with you.", 11, after=10)
    for label, detail in spec["hymns"]:
        line = doc.add_paragraph()
        line.paragraph_format.tab_stops.add_tab_stop(Mm(124), WD_TAB_ALIGNMENT.RIGHT)
        line.paragraph_format.space_after = Pt(5)
        _font(line.add_run(label), "Times New Roman", 12, True)
        _font(line.add_run("\t" + detail), "Times New Roman", 12)
    centred("Tea and coffee are served in the hall after the service.", 10, italic=True,
            after=0).paragraph_format.space_before = Pt(12)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


# ---------------------------------------------------------------------------
# Posters (.odg) and certificate (.odp)
# ---------------------------------------------------------------------------

def _svg_image(png: bytes, x, y, w, h) -> str:
    data = base64.b64encode(png).decode("ascii")
    return (f'<image x="{x}" y="{y}" width="{w}" height="{h}" '
            f'xlink:href="data:image/png;base64,{data}" preserveAspectRatio="xMidYMid slice"/>')


def _svg(body: str, background: str) -> str:
    return ('<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
            'width="210mm" height="297mm" viewBox="0 0 210 297">'
            f'<rect x="0" y="0" width="210" height="297" fill="#{background}"/>{body}</svg>')


def _text(x, y, text, size, colour, family="Liberation Sans", weight="normal",
          style="normal") -> str:
    from html import escape
    return (f'<text x="{x}" y="{y}" font-family="{family}" font-size="{size}" '
            f'font-weight="{weight}" font-style="{style}" fill="#{colour}" '
            f'text-anchor="middle">{escape(text)}</text>')


POSTERS = [
    {
        "file": "Posters/Summer Fair 2023.odg",
        "saved": datetime(2023, 6, 2, 20, 30),
        "svg": lambda: _svg(
            '<rect x="10" y="10" width="190" height="62" rx="5" fill="#1F7A3F"/>'
            + _text(105, 30, "ST AIDAN'S CHURCH, WRENFORD", 7, "FFFFFF", weight="bold")
            + _text(105, 58, "SUMMER FAIR", 25, "FFE14D", weight="bold")
            + _svg_image(artwork.bunting_scene(900, 560), 15, 80, 180, 112)
            + _text(105, 210, "Saturday 8 July 2023", 13, "1F7A3F", weight="bold")
            + _text(105, 224, "12 noon to 4 pm, Vicarage Garden", 9, "333333")
            + _text(105, 245, "Cake stall  *  Tombola  *  Plants  *  Dog show  *  Teas", 7,
                    "C0392B", weight="bold")
            + _text(105, 270, "Free entry. All money raised goes to the church roof.", 7,
                    "333333", style="italic"),
            "FFF8E1"),
    },
    {
        "file": "Posters/Coffee morning poster.odg",
        "saved": datetime(2018, 2, 11, 14, 5),
        "svg": lambda: _svg(
            _text(105, 40, "COFFEE", 30, "6B3E26", "Liberation Serif", "bold")
            + _text(105, 66, "MORNING", 30, "6B3E26", "Liberation Serif", "bold")
            + _svg_image(artwork.coffee_scene(900, 560), 20, 78, 170, 106)
            + _text(105, 205, "Every Tuesday, 10 am to 12 noon", 11, "333333", weight="bold")
            + _text(105, 220, "St Aidan's Church Hall, Church Lane", 9, "333333")
            + _text(105, 245, "Homemade cakes  -  Fair-trade coffee  -  Good company", 7,
                    "6B3E26")
            + _text(105, 262, "All for 2 pounds!", 11, "C0392B", weight="bold"),
            "F6EBDD"),
    },
    {
        "file": "Posters/Carol Service 2024.odg",
        "saved": datetime(2024, 12, 2, 18, 15),
        "svg": lambda: _svg(
            _text(105, 32, "St Aidan's Church, Wrenford", 9, "F4E7B2", "Liberation Serif")
            + _text(105, 58, "Carols by", 22, "FFFFFF", "Liberation Serif", "bold", "italic")
            + _text(105, 82, "Candlelight", 26, "FFFFFF", "Liberation Serif", "bold", "italic")
            + _svg_image(artwork.candles_scene(900, 560), 15, 95, 180, 112)
            + _text(105, 228, "Sunday 22 December 2024 at 6.30 pm", 11, "F4E7B2", weight="bold")
            + _text(105, 245, "Mince pies and mulled wine afterwards", 8, "FFFFFF")
            + _text(105, 270, "Everyone welcome", 10, "F7B733", weight="bold"),
            "1B2A4A"),
    },
]


def make_certificate_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor as Rgb
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Mm as PMm, Pt as PPt

    prs = Presentation()
    prs.slide_width, prs.slide_height = PMm(297), PMm(210)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PMm(8), PMm(8), PMm(281), PMm(194))
    border.fill.background()
    border.line.color.rgb = Rgb.from_string("FF9900")
    border.line.width = PPt(12)
    border.shadow.inherit = False

    def text(top, height, words, size, colour, bold=False, font="Comic Sans MS"):
        box = slide.shapes.add_textbox(PMm(20), PMm(top), PMm(257), PMm(height))
        para = box.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = words
        run.font.size = PPt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = Rgb.from_string(colour)

    text(18, 25, "Well done!", 54, "CC0000", True)
    text(48, 15, "Wrens Club Summer Reading Challenge 2021", 24, "003399", True, "Arial")
    text(72, 12, "This is to certify that", 18, "333333", font="Times New Roman")
    text(88, 15, "....................................................", 24, "333333")
    text(106, 12, "read ......... books this summer!", 20, "333333", font="Times New Roman")
    slide.shapes.add_picture(io.BytesIO(artwork.books_scene(900, 560)), PMm(108), PMm(124),
                             PMm(80), PMm(50))
    text(178, 10, "Signed: Revd Helen Marsh          Date: September 2021", 14, "333333",
         font="Times New Roman")
    prs.core_properties.author = "Parish Office"
    prs.core_properties.title = "Reading certificate 2021"
    prs.core_properties.comments = ""
    prs.save(path)


def damaged_pub(path: Path) -> None:
    """Starts like a real Publisher file but the rest is noise, as when a file
    was half-copied from an old floppy disk or a failing hard drive."""
    rng = random.Random(2007)
    body = bytes(rng.getrandbits(8) for _ in range(61_440))
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(OLE_SIGNATURE + body)


def set_saved(path: Path, when: datetime) -> None:
    stamp = when.timestamp()
    os.utime(path, (stamp, stamp))


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Make the fictional stand-in 'client files' "
                                     "for the St Aidan's sample. Deletes the folder first.")
    parser.add_argument("out", nargs="?", default="samples/st-aidans-wrenford/originals",
                        help="folder to (re)create (default: samples/st-aidans-wrenford/originals)")
    out = Path(parser.parse_args(argv).out).expanduser().resolve()
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    made = []
    for spec in NEWSLETTERS:
        target = out / spec["file"]
        old_newsletter(spec, target)
        made.append((target, spec["saved"]))

    with tempfile.TemporaryDirectory() as tmp, Soffice() as lo:
        tmp_dir = Path(tmp)
        for spec in SERVICE_SHEETS:
            target = out / spec["file"]
            if spec["convert_to"]:
                draft = tmp_dir / (target.stem + ".docx")
                old_service_sheet(spec, draft)
                result = lo.convert(draft, target, target=spec["convert_to"])
                if not result.ok:
                    raise SystemExit(f"Could not make {target.name}: {result.message}")
            else:
                old_service_sheet(spec, target)
            made.append((target, spec["saved"]))
        for spec in POSTERS:
            target = out / spec["file"]
            svg = tmp_dir / (target.stem + ".svg")
            svg.write_text(spec["svg"](), encoding="utf-8")
            result = lo.convert(svg, target, target="odg")
            if not result.ok:
                raise SystemExit(f"Could not make {target.name}: {result.message}")
            made.append((target, spec["saved"]))
        cert_pptx = tmp_dir / "certificate.pptx"
        make_certificate_pptx(cert_pptx)
        target = out / "Certificates/Wrens Club reading certificate 2021.odp"
        result = lo.convert(cert_pptx, target, target="odp")
        if not result.ok:
            raise SystemExit(f"Could not make {target.name}: {result.message}")
        made.append((target, datetime(2021, 9, 8, 10, 30)))

    target = out / "Old PC backup/Parish magazine March 2007.pub"
    damaged_pub(target)
    made.append((target, datetime(2007, 3, 1, 9, 15)))

    for path, when in made:
        set_saved(path, when)
        print(f"Made {path.relative_to(out)}")
    print(f"\n{len(made)} sample files in {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
