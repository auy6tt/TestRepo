#!/usr/bin/env python3
"""Build the Publisher Rescue starter templates.

Makes, in A4 and US Letter sizes:
  - a two-column newsletter (Word)
  - a folded service bulletin: four A5 / half-Letter pages (Word)
  - a certificate of achievement (PowerPoint)
plus two business documents for you: a client intake checklist and a
delivery note (Word).

Every template uses named styles (Heading 1, Kicker, Caption, Box Text ...)
so clients restyle a whole document by changing one style.

Examples (run from kits/publisher-rescue/; jobs/ is git-ignored)
  python scripts/build_templates.py                       # blank templates into templates/
  python scripts/build_templates.py --paper letter --out jobs/stmarys/templates \
      --org "St Mary's School PTA" --primary 0B6E4F --accent E0A100 \
      --logo jobs/stmarys/logo.png
  python scripts/build_templates.py --content scripts/sample-data/st-aidans-wrenford.json \
      --paper a4 --out jobs/try-out/templates --pdf

With --content the templates are filled with the text in a JSON file (see
scripts/sample-data/st-aidans-wrenford.json for the format: sections "brand",
"newsletter", "bulletin", "certificate" and "delivery_note").
"""

from __future__ import annotations

import argparse
import copy
import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Emu, Inches, Mm, Pt, RGBColor  # noqa: E402

import artwork  # noqa: E402
import docx_tools as dt  # noqa: E402

KIT_DIR = Path(__file__).resolve().parent.parent

DEFAULT_BRAND = {
    "org": "[Organisation name]",
    "primary": "1F3A5F",      # deep blue: mastheads, headings
    "accent": "C8912E",       # warm gold: rules, bullets, seals
    "tint": "EEF2F7",         # pale background for boxes
    "ink": "1E2430",          # body text
    "muted": "5B6472",        # captions, bylines
    "heading_font": "Cambria",
    "body_font": "Calibri",
    "language": "en-GB",
    "logo": None,             # path to a logo picture, or "art:emblem:AB"
}

PAPER = {
    "a4": {
        "name": "A4", "page": (Mm(210), Mm(297)), "margin": Mm(15),
        "half_name": "A5", "half": (Mm(148), Mm(210)), "half_margin": Mm(12),
        "landscape": (Mm(297), Mm(210)),
    },
    "letter": {
        "name": "Letter", "page": (Inches(8.5), Inches(11)), "margin": Inches(0.6),
        "half_name": "HalfLetter", "half": (Inches(5.5), Inches(8.5)),
        "half_margin": Inches(0.5), "landscape": (Inches(11), Inches(8.5)),
    },
}

# ---------------------------------------------------------------------------
# Placeholder content for the blank templates. Text in [square brackets] is
# for the client to replace; the other text explains how to use the template.
# ---------------------------------------------------------------------------

NEWSLETTER_TEMPLATE = {
    "title": "[Newsletter name]",
    "tagline": "[Your organisation name]  ·  [Short strapline]",
    "issue": "[Issue 1]",
    "date": "[Spring 2026]",
    "issue_right": "[www.example.org]",
    "lead": {
        "kicker": "[Section label]",
        "headline": "[Main headline: keep it under ten words]",
        "byline": "[By name]",
        "paragraphs": [
            "Replace this text with your main story. Click in a paragraph and start "
            "typing. When you paste from elsewhere, use Paste > Keep Text Only so the "
            "newsletter keeps its own look.",
            "Text flows from the left column into the right column and on to the next "
            "page by itself. You never need to link text boxes as you did in Publisher.",
        ],
        "photo": None,
        "caption": "[Caption: who is in the photo and what is happening. Right-click the "
                   "picture and choose Change Picture. Only use photos of children with "
                   "permission.]",
        "paragraphs_after": [
            "Headlines use the Heading 1 style, the small label above a headline uses "
            "Kicker, and the line under a headline uses Byline. You will find them all in "
            "the Styles gallery on the Home tab.",
            "To change the look of every headline at once, right-click Heading 1 in the "
            "Styles gallery and choose Modify. The same works for every style.",
        ],
    },
    "quote": "[A short quote from the story, in the Quote style, to catch the eye]",
    "articles": [
        {
            "kicker": "[Section label]",
            "headline": "[Second story headline]",
            "byline": "",
            "paragraphs": [
                "[Your second story. Keep paragraphs short: three or four sentences "
                "read well in narrow columns.]",
                "[Add as many stories as you need. Copy a headline and its paragraphs, "
                "paste them below, then type over them.]",
            ],
        },
        {
            "kicker": "",
            "headline": "[A shorter item]",
            "byline": "",
            "paragraphs": ["[Use this space for a short item, a thank-you or a "
                           "volunteer request.]"],
        },
    ],
    "diary_title": "Dates for your diary",
    "diary": [
        ["[Sat 7 Mar]", "[Event, time and place]"],
        ["[Sun 15 Mar]", "[Event, time and place]"],
        ["[Wed 25 Mar]", "[Event, time and place]"],
        ["[Sat 4 Apr]", "[Event, time and place]"],
    ],
    "notices_title": "Notices",
    "notices": ["[A short notice]", "[Another notice]", "[A third notice]"],
    "contact_title": "Contact us",
    "contact": ["[Name], [role]: [phone]", "[Email address]", "[Postal address]",
                "[Website and social media]"],
    "footer": "[Organisation name]  ·  [Address]  ·  [Registered charity number]",
}

BULLETIN_TEMPLATE = {
    "org": "[Church name]",
    "place": "[Town or village]",
    "title": "[Name of the Sunday or service]",
    "date": "[Sunday 1 March 2026]",
    "service": "[10.00 am  ·  Service name]",
    "cover_picture": None,
    "welcome": "[A short welcome, for example: Welcome! If you are visiting, please make "
               "yourself known after the service. Large-print sheets and a hearing loop "
               "are available.]",
    "people": "[Leading: name  ·  Preaching: name]",
    "order_title": "Order of Service",
    "order_intro": "Words in bold are said by everyone. Please stand or sit as you are able.",
    "sections": [
        {"heading": "We gather", "items": [
            {"item": "Welcome", "detail": "[Name]"},
            {"item": "Hymn", "detail": "[123]", "note": "[Title or first line of the hymn]"},
            {"rubric": "[Please stand]"},
            {"leader": "[Words said by the leader]"},
            {"all": "[Words said by everyone]"},
            {"item": "Prayer", "detail": "[Name]"},
        ]},
        {"heading": "We listen", "items": [
            {"item": "First reading", "detail": "[Book 1.1-10]", "note": "Read by [name]"},
            {"item": "Psalm", "detail": "[23]"},
            {"item": "Second reading", "detail": "[Book 1.1-10]", "note": "Read by [name]"},
            {"item": "Sermon", "detail": "[Name]"},
        ]},
        {"heading": "We respond", "new_page": True, "items": [
            {"item": "Prayers", "detail": "[Name]"},
            {"item": "Hymn", "detail": "[456]", "note": "[Title or first line of the hymn]"},
            {"rubric": "[The collection is taken during this hymn]"},
            {"leader": "[Words said by the leader]"},
            {"all": "[Words said by everyone]"},
        ]},
        {"heading": "We go out", "items": [
            {"item": "Blessing", "detail": ""},
            {"item": "Final hymn", "detail": "[789]", "note": "[Title or first line of the hymn]"},
            {"leader": "[Words of sending out]"},
            {"all": "[Response]"},
            {"rubric": "[Please stay for refreshments after the service]"},
        ]},
    ],
    "week_title": "This week",
    "diary": [
        ["[Mon]", "[7.30 pm]", "[Event and place]"],
        ["[Tue]", "[10.00 am]", "[Event and place]"],
        ["[Wed]", "[7.00 pm]", "[Event and place]"],
        ["[Thu]", "[2.00 pm]", "[Event and place]"],
        ["[Sat]", "[10.00 am]", "[Event and place]"],
        ["[Sun]", "[10.00 am]", "[Next Sunday's service]"],
    ],
    "notices_title": "Notices",
    "notices": ["[A short notice]", "[Another notice]", "[A third notice]"],
    "prayer_title": "Please pray for",
    "prayer": "[Names or groups to pray for this week. Check people are happy to be named.]",
    "contact_title": "Contact",
    "contacts": ["[Minister]: [name]  ·  [phone]", "Office: [email]  ·  [opening hours]",
                 "[Website]  ·  [Registered charity number]"],
    "giving": "[How to give: card reader at the back, or scan the code]",
}

CERTIFICATE_TEMPLATE = {
    "org": "[Organisation name]",
    "title": "Certificate",
    "subtitle": "of Achievement",
    "presented": "This certificate is proudly presented to",
    "recipients": [{"name": "[Recipient name]",
                    "reason": "[for outstanding effort and achievement in ...]"}],
    "date": "Awarded on [date]",
    "seal": "[Year]",
    "signatures": [{"name": "[Name]", "role": "[Role, e.g. Head Teacher]"},
                   {"name": "[Name]", "role": "[Role, e.g. Club Leader]"}],
    "logo": None,
}

DELIVERY_TEMPLATE = {
    "client": "[Client organisation]",
    "contact": "[Contact name]",
    "from": "[Your name or business name]",
    "from_email": "[your email]",
    "date": "[Date]",
    "reference": "[Job reference]",
    "counts": {"files": "[n]", "converted": "[n]", "check": "[n]", "failed": "[n]",
               "pages": "[n]"},
    "attention": [["[file name]", "[what is wrong]", "[what we suggest]"]],
    "templates": [["[Two-column newsletter]", "[Word, A4]", "[newsletter.docx]"],
                  ["[Service bulletin, folded booklet]", "[Word, A5 pages]", "[bulletin.docx]"],
                  ["[Certificate of achievement]", "[PowerPoint, A4]", "[certificate.pptx]"]],
    "fonts": "[Fonts we replaced and why, for example: 'Gill Sans MT' was replaced by "
             "'Calibri', which every Office computer has.]",
    "deleted_on": "",          # fill in once done; the box is then ticked
    "delete_by": "[date]",     # used while deletion is still to come
    "support_days": "[14]",
}


# ---------------------------------------------------------------------------
# Shared Word set-up
# ---------------------------------------------------------------------------

def readable(color: str, background: str = "FFFFFF", target: float = 4.5) -> str:
    """Darken `color` until text in it is easy to read on `background`."""
    def lum(hex_color):
        values = []
        for i in (0, 2, 4):
            v = int(hex_color[i:i + 2], 16) / 255
            values.append(v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4)
        return 0.2126 * values[0] + 0.7152 * values[1] + 0.0722 * values[2]

    current = color
    for step in range(20):
        a, b = lum(current), lum(background)
        if (max(a, b) + 0.05) / (min(a, b) + 0.05) >= target:
            return current
        r, g, bl = (int(current[i:i + 2], 16) for i in (0, 2, 4))
        current = "".join(f"{int(v * 0.9):02X}" for v in (r, g, bl))
    return current


def start_document(brand: dict, page, margin, title: str, subject: str):
    doc = Document()
    dt.set_theme(doc, brand["heading_font"], brand["body_font"], {
        "dk2": brand["primary"], "lt2": brand["tint"], "accent1": brand["primary"],
        "accent2": brand["accent"],
    })
    section = doc.sections[0]
    section.page_width, section.page_height = page
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, margin)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(7)
    dt.set_properties(doc, title, subject, "template, Publisher Rescue")
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:  # the schema requires a percentage
        zoom.set(qn("w:percent"), "100")

    normal = doc.styles["Normal"]
    dt.style_text(normal, brand["body_font"], "minor", 10.5, color=brand["ink"])
    rpr = normal.element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), brand.get("language", "en-GB"))
    rpr.append(lang)
    fmt = normal.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(5)
    fmt.line_spacing = 1.1
    fmt.widow_control = True
    return doc


def text_width(page, margin):
    return Emu(int(page[0]) - 2 * int(margin))


def para_style(doc, name, base="Normal", font=None, theme=None, size=None, bold=None,
               italic=None, color=None, color_theme=None, caps=None, spacing=None,
               before=None, after=None, line=None, align=None, keep_next=None):
    style = dt.get_or_add_style(doc, name, base=base)
    dt.style_text(style, font, theme, size, bold, italic, color, color_theme, caps, spacing)
    fmt = style.paragraph_format
    if before is not None:
        fmt.space_before = Pt(before)
    if after is not None:
        fmt.space_after = Pt(after)
    if line is not None:
        fmt.line_spacing = line
    if align is not None:
        fmt.alignment = align
    if keep_next is not None:
        fmt.keep_with_next = keep_next
    return style


def image_stream(spec, base_dir, brand, fallback: bytes) -> io.BytesIO:
    data = artwork.picture(spec, base_dir, brand) if spec else fallback
    return io.BytesIO(data)


def add_text(paragraph, text: str, bold=False):
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    return run


def footer_with_page(section, text: str, width, style="Footer"):
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.style = style
    paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run(text)
    paragraph.add_run("\tPage ")
    dt.add_field(paragraph, "PAGE")


def boxed_table(doc, rows: int, cols: int, widths, fill: str, top_rule: str | None = None):
    table = doc.add_table(rows=rows, cols=cols)
    dt.table_borders(table)
    dt.table_cell_padding(table, 2, 6, 2, 6)
    dt.table_widths(table, widths)
    for row in table.rows:
        dt.row_cant_split(row)
        for cell in row.cells:
            dt.cell_shading(cell, fill)
    if top_rule:
        for cell in table.rows[0].cells:
            dt.cell_borders(cell, top=(18, top_rule))
    return table


def set_cell_text(cell, text: str, style: str, bold=False, align=None):
    paragraph = cell.paragraphs[0]
    paragraph.style = style
    add_text(paragraph, text, bold)
    if align is not None:
        paragraph.alignment = align
    return paragraph


# ---------------------------------------------------------------------------
# Newsletter
# ---------------------------------------------------------------------------

def newsletter_styles(doc, b, width):
    H, B = b["heading_font"], b["body_font"]
    accent_text = readable(b["accent"])
    para_style(doc, "Masthead Title", font=H, theme="major", size=38, bold=True,
               color="FFFFFF", before=0, after=0, line=1.0)
    para_style(doc, "Masthead Tagline", font=B, theme="minor", size=11, color="E6ECF4",
               spacing=0.3, before=2, after=0, line=1.0)
    para_style(doc, "Masthead Logo", before=0, after=0, line=1.0,
               align=WD_ALIGN_PARAGRAPH.CENTER)
    issue = para_style(doc, "Issue Line", font=B, theme="minor", size=9.5, bold=True,
                       color=b["primary"], before=5, after=12)
    dt.paragraph_border(issue.element.get_or_add_pPr(), bottom=(12, b["primary"], 4))
    issue.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    para_style(doc, "Kicker", font=B, theme="minor", size=8.5, bold=True, color=accent_text,
               caps=True, spacing=1.0, before=8, after=0, keep_next=True)
    para_style(doc, "Heading 1", font=H, theme="major", size=17, bold=True,
               color=b["primary"], color_theme="accent1", before=6, after=3, line=1.0,
               keep_next=True)
    para_style(doc, "Heading 2", font=B, theme="minor", size=10, bold=True, color=accent_text,
               caps=True, spacing=0.8, before=10, after=3, line=1.0, keep_next=True)
    para_style(doc, "Byline", font=B, theme="minor", size=9, italic=True, color=b["muted"],
               before=0, after=6)
    para_style(doc, "Picture", before=4, after=2, line=1.0, keep_next=True)
    para_style(doc, "Caption", font=B, theme="minor", size=8.5, bold=False, italic=True,
               color=b["muted"], before=0, after=8, line=1.0)
    quote = para_style(doc, "Quote", font=H, theme="major", size=13, italic=True,
                       color=b["primary"], before=8, after=10, line=1.1,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
    dt.paragraph_border(quote.element.get_or_add_pPr(), top=(12, b["accent"], 6),
                        bottom=(12, b["accent"], 6))
    para_style(doc, "Box Heading", font=B, theme="minor", size=9.5, bold=True,
               color=b["primary"], caps=True, spacing=0.8, before=2, after=3, line=1.0)
    para_style(doc, "Box Text", font=B, theme="minor", size=9.5, color=b["ink"], before=0,
               after=2, line=1.05)
    bullet = para_style(doc, "List Bullet", after=2)
    bullet.paragraph_format.left_indent = None
    dt.bullet_color(doc, b["accent"])
    footer = para_style(doc, "Footer", font=B, theme="minor", size=8, color=b["muted"],
                        before=0, after=0)
    footer.paragraph_format.tab_stops.clear_all()


def build_newsletter(content: dict, brand: dict, paper: str, path: Path, base_dir=None):
    P = PAPER[paper]
    page, margin = P["page"], P["margin"]
    width = text_width(page, margin)
    gap = Mm(7)
    column = Emu((int(width) - int(gap)) // 2)
    doc = start_document(brand, page, margin, f"{content['title']} - newsletter",
                         "Two-column newsletter")
    newsletter_styles(doc, brand, width)

    # Masthead: logo and title on a coloured band
    logo_width = Mm(26)
    band = doc.add_table(rows=1, cols=2)
    dt.table_borders(band)
    dt.table_cell_padding(band, 0, 0, 0, 0)
    dt.table_widths(band, [Emu(int(logo_width) + int(Mm(10))),
                           Emu(int(width) - int(logo_width) - int(Mm(10)))])
    logo_cell, title_cell = band.rows[0].cells
    for cell in (logo_cell, title_cell):
        dt.cell_shading(cell, brand["primary"])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    dt.cell_margins(logo_cell, 9, 10, 9, 4)
    dt.cell_margins(title_cell, 9, 8, 9, 10)
    logo_par = logo_cell.paragraphs[0]
    logo_par.style = "Masthead Logo"
    logo_spec = brand.get("logo")
    logo_par.add_run().add_picture(
        image_stream(logo_spec, base_dir, brand,
                     artwork.placeholder(400, 400, "Your logo", kind="logo")),
        width=logo_width)
    set_cell_text(title_cell, content["title"], "Masthead Title")
    title_cell.add_paragraph(content["tagline"], style="Masthead Tagline")

    issue = doc.add_paragraph(style="Issue Line")
    issue.add_run(f"{content['issue']}  ·  {content['date']}")
    issue.add_run("\t" + content["issue_right"])

    body = doc.add_section(WD_SECTION.CONTINUOUS)
    dt.shrink_paragraph(doc.paragraphs[-1])  # the paragraph holding the section break
    dt.set_columns(body, 2, gap.pt)

    lead = content["lead"]
    if lead.get("kicker"):
        doc.add_paragraph(lead["kicker"], style="Kicker")
    doc.add_paragraph(lead["headline"], style="Heading 1")
    if lead.get("byline"):
        doc.add_paragraph(lead["byline"], style="Byline")
    for text in lead["paragraphs"]:
        doc.add_paragraph(text)
    picture = doc.add_paragraph(style="Picture")
    picture.add_run().add_picture(
        image_stream(lead.get("photo"), base_dir, brand,
                     artwork.placeholder(800, 500, "Your photo here",
                                         "Right-click > Change Picture",
                                         tint=brand["tint"], ink=brand["primary"])),
        width=column)
    doc.add_paragraph(lead["caption"], style="Caption")
    for text in lead.get("paragraphs_after", []):
        doc.add_paragraph(text)
    if content.get("quote"):
        doc.add_paragraph(content["quote"], style="Quote")

    for article in content["articles"]:
        if article.get("kicker"):
            doc.add_paragraph(article["kicker"], style="Kicker")
        doc.add_paragraph(article["headline"], style="Heading 1")
        if article.get("byline"):
            doc.add_paragraph(article["byline"], style="Byline")
        for text in article["paragraphs"]:
            doc.add_paragraph(text)
        if article.get("photo"):
            pic = doc.add_paragraph(style="Picture")
            pic.add_run().add_picture(image_stream(article["photo"], base_dir, brand, b""),
                                      width=column)
            if article.get("caption"):
                doc.add_paragraph(article["caption"], style="Caption")

    # Dates for your diary
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    date_width = Mm(25)
    diary = boxed_table(doc, len(content["diary"]) + 1, 2,
                        [date_width, Emu(int(column) - int(date_width))],
                        brand["tint"], top_rule=brand["accent"])
    head = diary.rows[0].cells[0].merge(diary.rows[0].cells[1])
    set_cell_text(head, content["diary_title"], "Box Heading")
    dt.cell_margins(head, 6, 6, 2, 6)
    for row, (when, what) in zip(diary.rows[1:], content["diary"]):
        set_cell_text(row.cells[0], when, "Box Text", bold=True)
        set_cell_text(row.cells[1], what, "Box Text")
    last = diary.rows[-1].cells
    for cell in last:
        dt.cell_margins(cell, 2, 6, 6, 6)

    doc.add_paragraph(content["notices_title"], style="Heading 2")
    for notice in content["notices"]:
        doc.add_paragraph(notice, style="List Bullet")

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    box = boxed_table(doc, 1, 1, [column], brand["tint"], top_rule=brand["accent"])
    cell = box.rows[0].cells[0]
    dt.cell_margins(cell, 6, 8, 6, 8)
    set_cell_text(cell, content["contact_title"], "Box Heading")
    for line in content["contact"]:
        cell.add_paragraph(line, style="Box Text")

    # A continuous break after the columns makes Word even up the two columns
    # on the last page instead of leaving the right-hand one short.
    end = doc.add_section(WD_SECTION.CONTINUOUS)
    dt.shrink_paragraph(doc.paragraphs[-1])
    dt.set_columns(end, 1, 0)
    footer_with_page(doc.sections[0], content["footer"], width)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# Service bulletin (folded booklet: four half-size pages)
# ---------------------------------------------------------------------------

def bulletin_styles(doc, b, width):
    H, B = b["heading_font"], b["body_font"]
    accent_text = readable(b["accent"])
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(3)
    center = WD_ALIGN_PARAGRAPH.CENTER
    para_style(doc, "Cover Organisation", font=B, theme="minor", size=11, bold=True,
               color=accent_text, caps=True, spacing=2.0, before=4, after=0, align=center)
    para_style(doc, "Cover Place", font=B, theme="minor", size=10, color=b["muted"],
               spacing=0.5, before=0, after=10, align=center)
    para_style(doc, "Cover Picture", before=6, after=14, line=1.0, align=center)
    para_style(doc, "Cover Title", font=H, theme="major", size=21, bold=True,
               color=b["primary"], before=2, after=4, line=1.0, align=center)
    para_style(doc, "Cover Date", font=B, theme="minor", size=12, color=b["ink"],
               before=0, after=2, align=center)
    para_style(doc, "Cover Service", font=B, theme="minor", size=11, bold=True,
               color=accent_text, before=0, after=10, align=center)
    welcome = para_style(doc, "Cover Welcome", font=B, theme="minor", size=9.5, italic=True,
                         color=b["ink"], before=4, after=4, line=1.15, align=center)
    dt.paragraph_border(welcome.element.get_or_add_pPr(), top=(8, b["accent"], 6))
    para_style(doc, "Cover People", font=B, theme="minor", size=9, color=b["muted"],
               before=2, after=0, align=center)
    h1 = para_style(doc, "Heading 1", font=H, theme="major", size=15, bold=True,
                    color=b["primary"], color_theme="accent1", before=0, after=4, line=1.0,
                    align=center, keep_next=True)
    dt.paragraph_border(h1.element.get_or_add_pPr(), bottom=(8, b["accent"], 3))
    para_style(doc, "Heading 2", font=B, theme="minor", size=9, bold=True, color=accent_text,
               caps=True, spacing=1.2, before=8, after=2, line=1.0, keep_next=True)
    para_style(doc, "Intro", font=B, theme="minor", size=8.5, italic=True, color=b["muted"],
               before=4, after=2, align=center)
    item = para_style(doc, "Service Item", font=B, theme="minor", size=10, color=b["ink"],
                      before=3, after=0, keep_next=False)
    item.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT,
                                                 WD_TAB_LEADER.DOTS)
    detail = para_style(doc, "Service Detail", font=B, theme="minor", size=9.5, italic=True,
                        color=b["muted"], before=0, after=1)
    detail.paragraph_format.left_indent = Mm(4)
    rubric = para_style(doc, "Rubric", font=B, theme="minor", size=8.5, italic=True,
                        color=accent_text, before=2, after=1)
    rubric.paragraph_format.left_indent = Mm(4)
    leader = para_style(doc, "Liturgy", font=B, theme="minor", size=10, color=b["ink"],
                        before=1, after=1)
    leader.paragraph_format.left_indent = Mm(4)
    everyone = para_style(doc, "Liturgy All", base="Liturgy", bold=True, before=0, after=3)
    everyone.paragraph_format.left_indent = Mm(4)
    para_style(doc, "Box Heading", font=B, theme="minor", size=9, bold=True,
               color=b["primary"], caps=True, spacing=0.8, before=2, after=2, line=1.0)
    para_style(doc, "Box Text", font=B, theme="minor", size=9, color=b["ink"], before=0,
               after=1, line=1.05)
    bullet = para_style(doc, "List Bullet", size=9.5, after=1)
    bullet.paragraph_format.left_indent = None
    dt.bullet_color(doc, b["accent"])
    footer = para_style(doc, "Footer", font=B, theme="minor", size=8, color=b["muted"],
                        before=0, after=0, align=center)
    footer.paragraph_format.tab_stops.clear_all()


def build_bulletin(content: dict, brand: dict, paper: str, path: Path, base_dir=None):
    P = PAPER[paper]
    page, margin = P["half"], P["half_margin"]
    width = text_width(page, margin)
    doc = start_document(brand, page, margin, f"{content['org']} - service bulletin",
                         "Folded service bulletin")
    bulletin_styles(doc, brand, width)
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    footer = section.footer.paragraphs[0]
    footer.style = "Footer"
    footer.add_run(f"{content['org']}  ·  ")
    dt.add_field(footer, "PAGE")

    # Page 1: cover
    doc.add_paragraph(content["org"], style="Cover Organisation")
    doc.add_paragraph(content["place"], style="Cover Place")
    cover = doc.add_paragraph(style="Cover Picture")
    cover.add_run().add_picture(
        image_stream(content.get("cover_picture"), base_dir, brand,
                     artwork.placeholder(700, 560, "Cover picture",
                                         "Right-click > Change Picture",
                                         tint=brand["tint"], ink=brand["primary"])),
        width=Emu(int(int(width) * 0.74)))
    doc.add_paragraph(content["title"], style="Cover Title")
    doc.add_paragraph(content["date"], style="Cover Date")
    doc.add_paragraph(content["service"], style="Cover Service")
    doc.add_paragraph(content["welcome"], style="Cover Welcome")
    doc.add_paragraph(content["people"], style="Cover People")

    # Pages 2 and 3: order of service
    heading = doc.add_paragraph(content["order_title"], style="Heading 1")
    heading.paragraph_format.page_break_before = True
    doc.add_paragraph(content["order_intro"], style="Intro")
    for part in content["sections"]:
        heading = doc.add_paragraph(part["heading"], style="Heading 2")
        if part.get("new_page"):  # start page 3 here, so the booklet has four pages
            heading.paragraph_format.page_break_before = True
        for entry in part["items"]:
            if "item" in entry:
                line = doc.add_paragraph(style="Service Item")
                add_text(line, entry["item"])
                if entry.get("detail"):
                    line.add_run("\t")
                    add_text(line, entry["detail"], bold=True)
                if entry.get("note"):
                    line.paragraph_format.keep_with_next = True
                    doc.add_paragraph(entry["note"], style="Service Detail")
            elif "rubric" in entry:
                doc.add_paragraph(entry["rubric"], style="Rubric")
            elif "leader" in entry:
                doc.add_paragraph(entry["leader"], style="Liturgy")
            elif "all" in entry:
                doc.add_paragraph(entry["all"], style="Liturgy All")

    # Page 4: the week ahead
    heading = doc.add_paragraph(content["week_title"], style="Heading 1")
    heading.paragraph_format.page_break_before = True
    day_w, time_w = Mm(13), Mm(21)
    diary = boxed_table(doc, len(content["diary"]), 3,
                        [day_w, time_w, Emu(int(width) - int(day_w) - int(time_w))],
                        brand["tint"])
    for row, (day, time, what) in zip(diary.rows, content["diary"]):
        set_cell_text(row.cells[0], day, "Box Text", bold=True)
        set_cell_text(row.cells[1], time, "Box Text")
        set_cell_text(row.cells[2], what, "Box Text")
    for cell in diary.rows[0].cells:
        dt.cell_margins(cell, 5, 6, 1, 6)
    for cell in diary.rows[-1].cells:
        dt.cell_margins(cell, 1, 6, 5, 6)

    doc.add_paragraph(content["notices_title"], style="Heading 2")
    for notice in content["notices"]:
        doc.add_paragraph(notice, style="List Bullet")
    doc.add_paragraph(content["prayer_title"], style="Heading 2")
    doc.add_paragraph(content["prayer"])

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    qr_w = Mm(19)
    box = boxed_table(doc, 1, 2, [Emu(int(width) - int(qr_w) - int(Mm(4))),
                                  Emu(int(qr_w) + int(Mm(4)))],
                      brand["tint"], top_rule=brand["accent"])
    text_cell, code_cell = box.rows[0].cells
    dt.cell_margins(text_cell, 5, 6, 5, 4)
    dt.cell_margins(code_cell, 5, 2, 5, 6)
    code_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_text(text_cell, content["contact_title"], "Box Heading")
    for line in content["contacts"]:
        text_cell.add_paragraph(line, style="Box Text")
    text_cell.add_paragraph(content["giving"], style="Box Text").runs[0].italic = True
    code = code_cell.paragraphs[0]
    code.alignment = WD_ALIGN_PARAGRAPH.CENTER
    code.add_run().add_picture(
        image_stream(content.get("giving_code"), base_dir, brand,
                     artwork.placeholder(300, 300, "QR code", kind="qr",
                                         tint="FFFFFF", ink=brand["primary"])),
        width=qr_w)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# Certificate (PowerPoint)
# ---------------------------------------------------------------------------

def build_certificate(content: dict, brand: dict, paper: str, path: Path, base_dir=None):
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor as PptxRGB
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn as pqn
    from pptx.util import Emu as PEmu, Inches as PInches, Pt as PPt

    W, H = PAPER[paper]["landscape"]
    W, H = int(W), int(H)
    prs = Presentation()
    prs.slide_width, prs.slide_height = PEmu(W), PEmu(H)

    # Theme: fonts and colours, so Design > Variants in PowerPoint follows the brand
    theme_part = prs.slide_master.part.part_related_by(RT.THEME)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    root = etree.fromstring(theme_part.blob)
    scheme = root.find(".//a:clrScheme", ns)
    for slot, value in (("dk2", brand["primary"]), ("lt2", brand["tint"]),
                        ("accent1", brand["primary"]), ("accent2", brand["accent"])):
        el = scheme.find(f"a:{slot}", ns)
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{{{ns['a']}}}srgbClr").set("val", value)
    root.find(".//a:fontScheme/a:majorFont/a:latin", ns).set("typeface", brand["heading_font"])
    root.find(".//a:fontScheme/a:minorFont/a:latin", ns).set("typeface", brand["body_font"])
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                      standalone=True)

    def rgb(hex_color):
        return PptxRGB.from_string(hex_color)

    def y(inches):  # vertical positions are designed on an 8.5 inch page
        return PEmu(int(PInches(inches) * H / PInches(8.5)))

    def no_shadow(shape):
        """Remove the theme's shape style (shadows, outlines); we set colours ourselves."""
        style = shape._element.find(pqn("p:style"))
        if style is not None:
            shape._element.remove(style)

    def text_box(slide, left, top, width, height, text, size, font, color, bold=False,
                 italic=False, spacing=None, align=PP_ALIGN.CENTER, theme_col=None,
                 caps=False):
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(frame, side, 0)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.size = PPt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        if theme_col is not None:
            run.font.color.theme_color = theme_col
        else:
            run.font.color.rgb = rgb(color)
        if spacing:
            run.font._rPr.set("spc", str(int(spacing * 100)))
        if caps:
            run.font._rPr.set("cap", "all")
        return box

    # Decoration goes on the slide layout, so it cannot be moved by accident
    layout = prs.slide_layouts[6]  # "Blank"
    layout.name = "Certificate"
    scratch = prs.slides.add_slide(layout)
    deco = []

    def frame_rect(inset, color, weight):
        shape = scratch.shapes.add_shape(MSO_SHAPE.RECTANGLE, PEmu(inset), PEmu(inset),
                                         PEmu(W - 2 * inset), PEmu(H - 2 * inset))
        shape.fill.background()
        shape.line.color.rgb = rgb(color)
        shape.line.width = PPt(weight)
        no_shadow(shape)
        deco.append(shape)

    outer, inner = int(PInches(0.32)), int(PInches(0.5))
    frame_rect(outer, brand["primary"], 7)
    frame_rect(inner, brand["accent"], 1.5)
    size = int(PInches(0.34))
    for cx in (inner, W - inner):
        for cy in (inner, H - inner):
            gem = scratch.shapes.add_shape(MSO_SHAPE.DIAMOND, PEmu(cx - size // 2),
                                           PEmu(cy - size // 2), PEmu(size), PEmu(size))
            gem.fill.solid()
            gem.fill.fore_color.rgb = rgb(brand["accent"])
            gem.line.color.rgb = rgb("FFFFFF")
            gem.line.width = PPt(1.5)
            no_shadow(gem)
            deco.append(gem)
    tree = layout.shapes._spTree
    for shape in deco:
        tree.append(shape._element)  # moves it from the scratch slide to the layout
    next_id = 100
    for el in tree.iter(pqn("p:cNvPr")):
        el.set("id", str(next_id))
        next_id += 1
    # remove the scratch slide again
    slide_ids = prs.slides._sldIdLst
    scratch_id = slide_ids[-1]
    prs.part.drop_rel(scratch_id.rId)
    slide_ids.remove(scratch_id)

    accent_text = readable(brand["accent"])
    center_w = int(W * 0.72)
    left = PEmu((W - center_w) // 2)
    for person in content["recipients"]:
        slide = prs.slides.add_slide(layout)
        logo_size = int(PInches(0.95))
        logo_spec = content.get("logo") or brand.get("logo")
        logo_png = artwork.picture(logo_spec, base_dir, brand) if logo_spec else \
            artwork.placeholder(400, 400, "Your logo", kind="logo", tint="FFFFFF",
                                ink=brand["primary"])
        slide.shapes.add_picture(io.BytesIO(logo_png), PEmu((W - logo_size) // 2), y(0.72),
                                 PEmu(logo_size), PEmu(logo_size))
        text_box(slide, left, y(1.78), PEmu(center_w), y(0.36), content["org"], 13,
                 brand["body_font"], brand["primary"], bold=True, spacing=3, caps=True)
        text_box(slide, left, y(2.12), PEmu(center_w), y(0.85), content["title"].upper(), 50,
                 brand["heading_font"], brand["primary"], bold=True, spacing=8,
                 theme_col=MSO_THEME_COLOR.ACCENT_1)
        text_box(slide, left, y(2.92), PEmu(center_w), y(0.5), content["subtitle"], 26,
                 brand["heading_font"], accent_text, italic=True)
        text_box(slide, left, y(3.55), PEmu(center_w), y(0.4), content["presented"], 15,
                 brand["body_font"], brand["muted"])
        text_box(slide, left, y(3.98), PEmu(center_w), y(0.82), person["name"], 40,
                 brand["heading_font"], brand["ink"], italic=True)
        rule = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PEmu(int(W * 0.27)), y(4.86),
                                          PEmu(int(W * 0.73)), y(4.86))
        rule.line.color.rgb = rgb(brand["accent"])
        rule.line.width = PPt(1.25)
        no_shadow(rule)
        text_box(slide, PEmu(int(W * 0.17)), y(4.98), PEmu(int(W * 0.66)), y(0.72),
                 person["reason"], 16, brand["body_font"], brand["ink"])

        # Seal
        seal = int(PInches(1.28))
        sx, sy = (W - seal) // 2, int(y(5.86))
        star = slide.shapes.add_shape(MSO_SHAPE.STAR_24_POINT, PEmu(sx), PEmu(sy),
                                      PEmu(seal), PEmu(seal))
        star.fill.solid()
        star.fill.fore_color.rgb = rgb(brand["accent"])
        star.line.fill.background()
        no_shadow(star)
        core = int(seal * 0.66)
        disc = slide.shapes.add_shape(MSO_SHAPE.OVAL, PEmu(sx + (seal - core) // 2),
                                      PEmu(sy + (seal - core) // 2), PEmu(core), PEmu(core))
        disc.fill.solid()
        disc.fill.fore_color.rgb = rgb(brand["primary"])
        disc.line.color.rgb = rgb("FFFFFF")
        disc.line.width = PPt(1.5)
        no_shadow(disc)
        frame = disc.text_frame
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(frame, side, 0)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.word_wrap = False
        para = frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = content["seal"]
        run.font.size = PPt(15)
        run.font.bold = True
        run.font.name = brand["heading_font"]
        run.font.color.rgb = rgb("FFFFFF")
        text_box(slide, PEmu((W - int(PInches(3))) // 2), y(7.22), PEmu(int(PInches(3))),
                 y(0.3), content["date"], 11, brand["body_font"], brand["muted"])

        # Signatures
        sig_w = int(PInches(2.7))
        for index, signer in enumerate(content["signatures"][:2]):
            x0 = int(PInches(1.35)) if index == 0 else W - int(PInches(1.35)) - sig_w
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PEmu(x0), y(6.68),
                                              PEmu(x0 + sig_w), y(6.68))
            line.line.color.rgb = rgb(brand["ink"])
            line.line.width = PPt(0.75)
            no_shadow(line)
            text_box(slide, PEmu(x0), y(6.74), PEmu(sig_w), y(0.3), signer["name"], 12,
                     brand["body_font"], brand["ink"], bold=True)
            text_box(slide, PEmu(x0), y(7.02), PEmu(sig_w), y(0.28), signer["role"], 10.5,
                     brand["body_font"], brand["muted"])

        notes = slide.notes_slide.notes_text_frame
        notes.text = ("How to use this certificate: click any text to change it. For one "
                      "certificate per person, right-click this slide in the list on the "
                      "left and choose Duplicate Slide, then change the name. The border is "
                      "part of the layout (View > Slide Master) so it cannot move by "
                      "accident. Print at 100% (Fit to page off) on A4 or Letter card.")

    props = prs.core_properties
    props.title = f"{content['org']} - certificate"
    props.subject = "Certificate of achievement"
    props.author = ""
    props.last_modified_by = ""
    props.keywords = "template, Publisher Rescue"
    props.comments = ""
    props.revision = 1
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


# ---------------------------------------------------------------------------
# Business documents: intake checklist and delivery note
# ---------------------------------------------------------------------------

def business_styles(doc, b, width):
    H, B = b["heading_font"], b["body_font"]
    accent_text = readable(b["accent"])
    para_style(doc, "Title", font=H, theme="major", size=24, bold=True, color=b["primary"],
               color_theme="accent1", before=0, after=2, line=1.0)
    title = doc.styles["Title"]
    dt.paragraph_border(title.element.get_or_add_pPr(), bottom=(12, b["accent"], 4))
    para_style(doc, "Subtitle", font=B, theme="minor", size=11, italic=False,
               color=b["muted"], spacing=0, before=6, after=10)
    para_style(doc, "Heading 1", font=H, theme="major", size=14, bold=True,
               color=b["primary"], color_theme="accent1", before=14, after=4, line=1.0,
               keep_next=True)
    para_style(doc, "Heading 2", font=B, theme="minor", size=10.5, bold=True,
               color=accent_text, before=8, after=2, keep_next=True)
    para_style(doc, "Table Text", font=B, theme="minor", size=10, color=b["ink"], before=1,
               after=1, line=1.05)
    para_style(doc, "Table Label", base="Table Text", bold=True, color=b["primary"])
    para_style(doc, "Small Print", font=B, theme="minor", size=8.5, italic=True,
               color=b["muted"], before=6, after=0)
    check = para_style(doc, "Check Item", font=B, theme="minor", size=10.5, color=b["ink"],
                       before=1, after=3)
    check.paragraph_format.left_indent = Mm(7)
    check.paragraph_format.first_line_indent = Mm(-7)
    check.paragraph_format.tab_stops.add_tab_stop(Mm(7))
    bullet = para_style(doc, "List Bullet", after=2)
    bullet.paragraph_format.left_indent = None
    dt.bullet_color(doc, b["accent"])
    footer = para_style(doc, "Footer", font=B, theme="minor", size=8, color=b["muted"],
                        before=0, after=0)
    footer.paragraph_format.tab_stops.clear_all()


def grid_table(doc, rows, widths, b, header: list[str] | None = None, label_col=True):
    """A clean table with thin rules: optional header row, first column bold."""
    count = len(rows) + (1 if header else 0)
    table = doc.add_table(rows=count, cols=len(widths))
    dt.table_borders(table, color="C9D1DC", size_eighths=4, inside_h=True)
    dt.table_cell_padding(table, 3, 6, 3, 6)
    dt.table_widths(table, widths)
    start = 0
    if header:
        for cell, text in zip(table.rows[0].cells, header):
            dt.cell_shading(cell, b["primary"])
            par = set_cell_text(cell, text, "Table Text", bold=True)
            par.runs[0].font.color.rgb = RGBColor.from_string("FFFFFF")
        dt.repeat_header_row(table.rows[0])
        start = 1
    for row, values in zip(table.rows[start:], rows):
        dt.row_cant_split(row)
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            style = "Table Label" if (label_col and index == 0) else "Table Text"
            set_cell_text(cell, value, style)
    if len(table.rows) <= 10:  # keep short tables on one page
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def checklist(doc, items):
    for text in items:
        paragraph = doc.add_paragraph(style="Check Item")
        dt.add_checkbox(paragraph, size_pt=11)
        paragraph.add_run("\t" + text)


def build_intake_checklist(brand: dict, paper: str, path: Path):
    P = PAPER[paper]
    page, margin = P["page"], Mm(20) if paper == "a4" else Inches(0.8)
    width = text_width(page, margin)
    doc = start_document(brand, page, margin, "Client intake checklist",
                         "Publisher file rescue: client intake")
    business_styles(doc, brand, width)
    footer_with_page(doc.sections[0], "Publisher file rescue  ·  Client intake checklist", width)
    doc.add_paragraph("Client intake checklist", style="Title")
    doc.add_paragraph("Publisher file rescue. Fill this in on a call with the client, or "
                      "send it to them before you quote. Click a box to tick it.",
                      style="Subtitle")
    label = Mm(62)
    rest = Emu(int(width) - int(label))

    doc.add_paragraph("1. About the client", style="Heading 1")
    grid_table(doc, [[q, ""] for q in (
        "Organisation", "Contact name and role", "Email and phone",
        "Who approves and pays the invoice", "Number of sites (schools, parishes, offices)",
        "How they found you")], [label, rest], brand)

    doc.add_paragraph("2. The Publisher files", style="Heading 1")
    grid_table(doc, [[q, ""] for q in (
        "Where the files are now (PC, shared drive, USB stick, cloud folder)",
        "Rough number of .pub files", "Years covered",
        "Publisher versions, if known", "Files that already will not open",
        "Other files to include (Word, pictures)? Priced separately")],
        [label, rest], brand)
    checklist(doc, [
        "Client can still open some files in Publisher (useful for comparing results)",
        "Keep the same folder structure in the PDF archive",
        "Client has a list or examples of their most-used layouts",
    ])

    doc.add_paragraph("3. What they want", style="Heading 1")
    checklist(doc, [
        "PDF archive: searchable PDFs, preview pictures, index spreadsheet, contact sheet",
        "Rebuilt templates (list them below, most important first)",
        "Editing program: Word / PowerPoint / Google Docs / LibreOffice (circle)",
        "Paper size: A4 / Letter (circle)",
        "Printing: office printer / folded booklet / print shop (circle)",
    ])
    grid_table(doc, [[f"Template {n}", ""] for n in range(1, 6)], [label, rest], brand,
               header=["Template", "Example file name, and how often it is used"])

    doc.add_paragraph("4. Branding and fonts", style="Heading 1")
    checklist(doc, [
        "Logo file received (best: PNG with a clear background, or SVG)",
        "Colours agreed (hex codes, or 'match the old newsletter')",
        "Special fonts listed. If the client has no licence for a font, we use a free "
        "or standard substitute and say so in the delivery note",
        "Photos they have permission to reuse (especially photos of children)",
    ])

    doc.add_paragraph("5. Privacy and file handling", style="Heading 1")
    checklist(doc, [
        "Client confirms they may share these files with you for this work",
        "Files may contain personal details (names, addresses, photos); handled "
        "privately and never shared",
        "Transfer method agreed (shared folder link, not email attachments for big archives)",
        "Client agrees that you delete every copy within 7 days of delivery, and will "
        "confirm in writing",
        "Files or folders that must not be opened are listed: ________________",
    ])

    doc.add_paragraph("6. Price and timing", style="Heading 1")
    grid_table(doc, [
        ["Archive", "______ files  x  $______ per file  =  $______"],
        ["Templates", "______ templates  x  $______ each  =  $______"],
        ["Total quote", "$______   (deposit $______ before work starts)"],
        ["Deadline", ""],
        ["Delivery by", "Shared folder / USB stick / other: ________"],
    ], [label, rest], brand)
    doc.add_paragraph("Price guide: $0.50-2 per archived file and $25-75 per rebuilt template "
                      "(about $150-600 for a typical organisation). Do not charge for exact "
                      "duplicates.", style="Small Print")
    doc.add_paragraph("Agreed by (client): ______________________   Date: ____________",
                      style="Table Text").paragraph_format.space_before = Pt(14)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def build_delivery_note(content: dict, brand: dict, paper: str, path: Path):
    P = PAPER[paper]
    page, margin = P["page"], Mm(20) if paper == "a4" else Inches(0.8)
    width = text_width(page, margin)
    doc = start_document(brand, page, margin, f"Delivery note - {content['client']}",
                         "Publisher file rescue: delivery note")
    business_styles(doc, brand, width)
    footer_with_page(doc.sections[0], f"Delivery note  ·  {content['client']}", width)
    doc.add_paragraph("Delivery note", style="Title")
    doc.add_paragraph(f"Publisher file rescue for {content['client']}", style="Subtitle")
    label = Mm(40)
    rest = Emu(int(width) - int(label))
    grid_table(doc, [["To", f"{content['contact']}, {content['client']}"],
                     ["From", f"{content['from']}  ·  {content['from_email']}"],
                     ["Date", content["date"]], ["Reference", content["reference"]]],
               [label, rest], brand)

    doc.add_paragraph("What you are receiving", style="Heading 1")
    c = content["counts"]
    third = Emu(int(width) // 3)
    grid_table(doc, [
        ["PDF archive", f"{c['converted']} searchable PDFs, one per Publisher file, in the "
                        "same folders as your originals", "archive/pdf/"],
        ["Preview pictures", "A picture of the first page of every file", "archive/previews/"],
        ["Index spreadsheet", "Every file with its date, page count and status. Click a "
                              "link to open the PDF.", "archive/index.xlsx"],
        ["Contact sheet", "Thumbnails of every file. Opens in any web browser, with a "
                          "search box.", "archive/contact-sheet.html"],
        ["Templates", "Your most-used layouts, rebuilt so you can edit them",
         "templates/"],
    ], [Mm(36), Emu(int(width) - int(Mm(36)) - int(Mm(48))), Mm(48)], brand,
        header=["Item", "What it is", "Where to find it"])

    doc.add_paragraph("In numbers", style="Heading 1")
    grid_table(doc, [["Files received", str(c["files"])],
                     ["Converted to PDF", str(c["converted"])],
                     ["Converted, worth a quick check", str(c["check"])],
                     ["Could not be converted", str(c["failed"])],
                     ["Pages in the archive", str(c["pages"])]],
               [Emu(int(third * 2)), third], brand)

    doc.add_paragraph("Files that need your attention", style="Heading 1")
    if content["attention"]:
        grid_table(doc, content["attention"],
                   [Mm(52), Emu(int(width) - int(Mm(52)) - int(Mm(55))), Mm(55)], brand,
                   header=["File", "What we found", "What we suggest"])
    else:
        doc.add_paragraph("None. Every file converted cleanly.")

    doc.add_paragraph("Your templates", style="Heading 1")
    grid_table(doc, content["templates"],
               [Mm(60), Mm(40), Emu(int(width) - int(Mm(100)))], brand,
               header=["Template", "Format", "File"])

    doc.add_paragraph("Using your templates", style="Heading 1")
    for tip in (
        "Before you change a template, use File > Save As and give the copy a new name, "
        "so the original stays clean for next time.",
        "Type over any text in [square brackets]. The look comes from styles on the Home "
        "tab (Heading 1, Caption and so on), so the layout stays tidy.",
        "To swap a picture, right-click it and choose Change Picture.",
        "Newsletter text flows between the columns and on to the next page by itself.",
        "Service bulletin: each page is a half sheet. The PDF whose name ends in "
        "'print-booklet' is ready to print: double-sided, flip on the short edge, then "
        "fold. For a new week, save your edited copy as PDF and print it with the "
        "'Booklet' option in Adobe Acrobat Reader or your printer's settings.",
        "Certificates: right-click the slide in the list on the left, choose Duplicate "
        "Slide, and change the name. Repeat for each person.",
    ):
        doc.add_paragraph(tip, style="List Bullet")

    doc.add_paragraph("What to expect from the PDFs", style="Heading 1")
    doc.add_paragraph(
        "The PDFs are faithful copies for reading, searching, printing and sharing. They "
        "are not pixel-perfect copies: some fonts were replaced with close matches, and "
        "special effects such as WordArt, shadows and decorative borders can look a little "
        "different. Files marked 'worth a quick check' are listed in the index with the "
        "reason.")
    doc.add_paragraph("Fonts", style="Heading 2")
    doc.add_paragraph(content["fonts"])

    doc.add_paragraph("Your data", style="Heading 1")
    data = doc.add_paragraph(style="Check Item")
    if content.get("deleted_on"):
        dt.add_checkbox(data, checked=True, size_pt=11)
        data.add_run(f"\tAs agreed, we deleted every copy of your files from our computers "
                     f"and cloud storage on {content['deleted_on']}. We kept only this "
                     "delivery note and the invoice.")
    else:
        dt.add_checkbox(data, checked=False, size_pt=11)
        data.add_run(f"\tAs agreed, we will delete every copy of your files from our "
                     f"computers and cloud storage by {content['delete_by']}, once you have "
                     "checked everything, and confirm by email. We keep only this delivery "
                     "note and the invoice.")

    doc.add_paragraph("Help after delivery", style="Heading 1")
    doc.add_paragraph(f"If you spot a problem in the next {content['support_days']} days, reply "
                      "to this email and we will fix it free of charge. Thank you for "
                      f"choosing {content['from']}.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# Command line
# ---------------------------------------------------------------------------

# What to suggest to the client for each kind of problem convert_archive.py reports.
SUGGESTIONS = {
    "Not a readable Publisher file": "Send another copy if you have one, for example "
                                     "from a backup or an old email.",
    "LibreOffice could not open it": "Send another copy if you have one. If it still "
                                     "opens in Publisher, save it as PDF there.",
    "Took too long, stopped": "We can try again with more time. If that fails, a PDF "
                              "saved from Publisher is the fallback.",
    "Empty file": "The file is empty (0 bytes). Check backups for a good copy.",
    "Text may be garbled": "Compare the PDF with a printed copy. We can retype the "
                           "text if needed.",
    "No searchable text": "The page is all pictures, so search will not find words in "
                          "it. Text recognition (OCR) can be added on request.",
    "Font replaced": "A font was replaced, so some text can look different. See "
                     "'Fonts' below.",
}


def apply_summary(delivery: dict, summary_path: Path) -> None:
    """Fill the delivery note's numbers and problem list from summary.json."""
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    counts = summary["counts"]
    delivery["counts"] = {
        "files": counts["files"],
        "converted": counts["converted"] + counts["check"],
        "check": counts["check"],
        "failed": counts["failed"] + counts["skipped"],
        "pages": counts["pages"],
    }
    rows = []
    for item in summary.get("failed_files", []) + summary.get("check_files", []):
        notes = [n.strip() for n in item.get("note", "").split(";") if n.strip()]
        found = "; ".join(n for n in notes if not n.startswith("Same as")) or item["problem"]
        advice = " ".join(SUGGESTIONS[n] for n in notes if n in SUGGESTIONS) or \
            "Please compare this PDF with the original."
        rows.append([item["file"], found, advice])
    delivery["attention"] = rows


def merged(defaults: dict, override: dict | None) -> dict:
    result = copy.deepcopy(defaults)
    for key, value in (override or {}).items():
        result[key] = value
    return result


def render_pdfs(paths: list[Path]) -> None:
    from office_tools import Soffice, pdf_target

    with Soffice() as lo:
        for path in paths:
            result = lo.convert(path, path.with_suffix(".pdf"),
                                target=pdf_target(path.suffix, pdfa=False))
            status = "PDF saved" if result.ok else f"PDF failed: {result.message}"
            print(f"  {path.name}: {status}")


CONTENT_SECTIONS = ("brand", "newsletter", "bulletin", "certificate", "delivery_note")


def check_content(data, path: Path) -> dict | None:
    """The --content file must hold the sections of the sample file. A file that
    holds only the inside of "delivery_note" is used as that section."""
    if isinstance(data, dict) and any(key in data for key in CONTENT_SECTIONS):
        return data
    if isinstance(data, dict) and data and all(key in DELIVERY_TEMPLATE for key in data):
        print(f"Note: {path} has no \"delivery_note\": {{...}} around it; using it as the "
              "delivery note.")
        return {"delivery_note": data}
    print(f"{path} has none of the sections {', '.join(CONTENT_SECTIONS)}, so nothing "
          "would be filled in. Copy the layout of "
          "scripts/sample-data/st-aidans-wrenford.json, for example "
          "{\"delivery_note\": {\"client\": \"...\", ...}}.", file=sys.stderr)
    return None


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Build the newsletter, service bulletin "
                                     "and certificate templates, plus the intake "
                                     "checklist and delivery note.")
    parser.add_argument("--out", type=Path, default=KIT_DIR / "templates",
                        help="output folder (default: the kit's templates folder)")
    parser.add_argument("--paper", choices=["a4", "letter", "both"], default="both")
    parser.add_argument("--content", type=Path,
                        help="JSON file with real content (otherwise placeholders are used)")
    parser.add_argument("--summary", type=Path,
                        help="summary.json from convert_archive.py, to fill in the "
                             "delivery note's numbers and problem list")
    parser.add_argument("--only", default="newsletter,bulletin,certificate,intake,delivery",
                        help="comma-separated list of what to build")
    parser.add_argument("--org", help="organisation name")
    parser.add_argument("--primary", help="main colour as hex, e.g. 1F3A5F")
    parser.add_argument("--accent", help="accent colour as hex, e.g. C8912E")
    parser.add_argument("--heading-font", help="font for headings (default Cambria)")
    parser.add_argument("--body-font", help="font for body text (default Calibri)")
    parser.add_argument("--logo", help="logo picture file (PNG or JPEG)")
    parser.add_argument("--pdf", action="store_true",
                        help="also save a PDF of each file (needs LibreOffice) to check the layout")
    args = parser.parse_args(argv)

    data = json.loads(args.content.read_text(encoding="utf-8")) if args.content else {}
    if args.content:
        data = check_content(data, args.content)
        if data is None:
            return 2
    base_dir = args.content.parent if args.content else Path.cwd()
    brand = merged(DEFAULT_BRAND, data.get("brand"))
    for key, value in (("org", args.org), ("primary", args.primary), ("accent", args.accent),
                       ("heading_font", args.heading_font), ("body_font", args.body_font),
                       ("logo", args.logo)):
        if value:
            brand[key] = value.lstrip("#") if key in ("primary", "accent") else value
    if args.logo:
        brand["logo"] = str(Path(args.logo).expanduser().resolve())

    org = brand["org"]
    newsletter = merged(NEWSLETTER_TEMPLATE, data.get("newsletter"))
    bulletin = merged(BULLETIN_TEMPLATE, data.get("bulletin"))
    certificate = merged(CERTIFICATE_TEMPLATE, data.get("certificate"))
    delivery = merged(DELIVERY_TEMPLATE, data.get("delivery_note"))
    if args.summary:
        apply_summary(delivery, args.summary)
    if not data:
        if args.org:
            bulletin["org"] = certificate["org"] = org
            newsletter["footer"] = newsletter["footer"].replace("[Organisation name]", org)
            newsletter["tagline"] = newsletter["tagline"].replace("[Your organisation name]", org)

    wanted = {w.strip() for w in args.only.split(",")}
    papers = ["a4", "letter"] if args.paper == "both" else [args.paper]
    out = args.out.expanduser().resolve()
    made: list[Path] = []
    for paper in papers:
        P = PAPER[paper]
        if "newsletter" in wanted:
            made.append(build_newsletter(newsletter, brand, paper,
                                         out / f"newsletter-two-column-{P['name']}.docx", base_dir))
        if "bulletin" in wanted:
            made.append(build_bulletin(bulletin, brand, paper,
                                       out / f"service-bulletin-{P['half_name']}-booklet.docx",
                                       base_dir))
        if "certificate" in wanted:
            made.append(build_certificate(certificate, brand, paper,
                                          out / f"certificate-of-achievement-{P['name']}.pptx",
                                          base_dir))
    business_paper = papers[0]
    if "intake" in wanted and not data:
        made.append(build_intake_checklist(brand, business_paper,
                                           out / "client-intake-checklist.docx"))
    if "delivery" in wanted:
        made.append(build_delivery_note(delivery, brand, business_paper,
                                        out / "delivery-note.docx"))
    for path in made:
        print(f"Saved {path}")
    if args.pdf:
        print("Making PDFs to check the layout ...")
        render_pdfs(made)
    return 0


if __name__ == "__main__":
    sys.exit(main())
