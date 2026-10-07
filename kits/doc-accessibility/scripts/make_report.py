#!/usr/bin/env python3
"""Fill the "Document accessibility snapshot" report from the inventory.

Reads the Inventory sheet of inventory.xlsx, so any changes you made to the
'Suggested action' or 'Priority' columns are used, plus crawl_summary.json from
the crawl. Fills templates/snapshot-report-template.docx (or your own copy of
it) and, with --pdf, exports a tagged PDF with LibreOffice.

Read the result before sending it. Change anything that does not fit the client.

Example:
  python make_report.py --inventory inventory/inventory.xlsx --crawl-summary crawl/crawl_summary.json \\
      --client "Town of Example" --prepared-by "Your Name" --contact "you@example.com" \\
      --sector us-public --out report/snapshot.docx --pdf
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx import Document  # noqa: E402
from docx.table import _Row  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from build_templates import SECTOR_CLOSING, WHY_IT_MATTERS  # noqa: E402
from triage_pdfs import ACTION_MEANING, ACTIONS, COLUMNS, PRIORITY_ORDER, summarize  # noqa: E402

KIT = Path(__file__).resolve().parent.parent
PLACEHOLDER = re.compile(r"\{\{([A-Za-z_.]+)\}\}")


def read_inventory(path: Path) -> list[dict]:
    book = load_workbook(path, read_only=True, data_only=True)
    if "Inventory" not in book.sheetnames:
        sys.exit(f"{path} has no 'Inventory' sheet. Make it with triage_pdfs.py.")
    rows = book["Inventory"].iter_rows(values_only=True)
    headers = [str(h or "") for h in next(rows)]
    by_header = {header: key for key, header in COLUMNS}
    records = []
    for values in rows:
        if not any(values):
            continue
        rec = {by_header.get(h, h): ("" if v is None else v) for h, v in zip(headers, values)}
        rec["issues"] = [i for i in str(rec.get("issues", "")).split("; ") if i]
        for key in ("file", "type", "action", "priority", "tagged", "needs_ocr", "has_title", "has_lang", "duplicate_of"):
            rec[key] = str(rec.get(key, "") or "")
        records.append(rec)
    book.close()
    return records


def nice_date(iso: str | None) -> str:
    if not iso:
        return "[crawl date]"
    try:
        return dt.date.fromisoformat(iso).strftime("%d %B %Y").lstrip("0")
    except ValueError:
        return iso


def money(value: float) -> str:
    return f"${value:,.0f}"


def delay_text(crawl: dict) -> str:
    """The wait the crawler left between requests, for 'leaving at least ... between requests'."""
    seconds = crawl.get("effective_delay_seconds", crawl.get("delay_seconds"))
    try:
        seconds = float(seconds)
    except (TypeError, ValueError):
        return "[delay]"
    return "one second" if seconds == 1 else f"{seconds:g} seconds"


def summary_text(records: list[dict], numbers: dict, crawl: dict, site: str) -> str:
    pages_read = crawl.get("pages_crawled")
    parts = [f"We found {numbers['total_files']} documents ({numbers['total_pages']} pages) linked from "
             + (f"{pages_read} pages of {site}." if pages_read else f"{site}.")]
    if numbers["pdf_count"]:
        part = (f"{numbers['untagged']} of the {numbers['pdf_count']} PDFs have no tags, so people using screen "
                "readers cannot move through their headings, lists or tables")
        if numbers["scanned"]:
            part += f", and {numbers['scanned']} {'is a scan' if numbers['scanned'] == 1 else 'are scans'} with no real text"
        parts.append(part + ".")
    high = numbers["high_priority"]
    if high:
        parts.append(f"{high} {'file needs' if high == 1 else 'files need'} fixing first because people use "
                     f"{'it' if high == 1 else 'them'} to apply for a service or {'it is' if high == 1 else 'they are'} "
                     "linked from your home page.")
    archive_delete = sum(numbers["actions"].get(a, {}).get("files", 0) for a in ("Archive", "Delete"))
    sentence = f"We suggest fixing or converting {numbers['work_files']} files"
    if archive_delete:
        sentence += f" and archiving or deleting {archive_delete} more, which you decide with your ADA coordinator"
    parts.append(sentence + ".")
    return " ".join(parts)


def replace_in_paragraph(paragraph, values: dict) -> None:
    text = paragraph.text
    if "{{" not in text:
        return
    def fill(t):
        return PLACEHOLDER.sub(lambda m: str(values.get(m.group(1), m.group(0))), t)
    for run in paragraph.runs:
        if "{{" in run.text:
            run.text = fill(run.text)
    wanted = fill(text)
    if paragraph.text != wanted and paragraph.runs:  # a placeholder was split across runs
        paragraph.runs[0].text = wanted
        for run in paragraph.runs[1:]:
            run.text = ""


def all_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            yield from part.paragraphs


def expand_rows(doc, marker: str, items: list[dict]) -> None:
    for table in doc.tables:
        for row in table.rows:
            if marker in "".join(cell.text for cell in row.cells):
                template = row._tr
                for item in items:
                    new_tr = copy.deepcopy(template)
                    template.addprevious(new_tr)
                    for cell in _Row(new_tr, table).cells:
                        for paragraph in cell.paragraphs:
                            replace_in_paragraph(paragraph, item)
                template.getparent().remove(template)
                return


def replace_with_paragraphs(doc, marker: str, texts: list[str]) -> None:
    for paragraph in doc.paragraphs:
        if marker in paragraph.text:
            for text in texts:
                new_p = copy.deepcopy(paragraph._p)
                paragraph._p.addprevious(new_p)
            created = [p for p in doc.paragraphs if marker in p.text][: len(texts)]
            for new_paragraph, text in zip(created, texts):
                new_paragraph.runs[0].text = text
                for run in new_paragraph.runs[1:]:
                    run.text = ""
            paragraph._p.getparent().remove(paragraph._p)
            return


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Fill the snapshot report from inventory.xlsx.")
    parser.add_argument("--inventory", required=True, help="inventory.xlsx from triage_pdfs.py")
    parser.add_argument("--crawl-summary", help="crawl_summary.json from crawl_documents.py")
    parser.add_argument("--template", default=str(KIT / "templates" / "snapshot-report-template.docx"))
    parser.add_argument("--out", required=True, help="the .docx to write")
    parser.add_argument("--client", required=True, help="client name, as it should appear in the report")
    parser.add_argument("--site", help="website address to show in the report, also where the crawl started "
                        "(default: the crawl's start URL)")
    parser.add_argument("--prepared-by", default="[Your name]")
    parser.add_argument("--contact", default="[your email and phone]")
    parser.add_argument("--sector", choices=sorted(WHY_IT_MATTERS), default="us-public",
                        help="which 'Why this matters' text to use (default: us-public)")
    parser.add_argument("--rate-low", type=float, default=5)
    parser.add_argument("--rate-high", type=float, default=25)
    parser.add_argument("--max-files", type=int, default=10, help="rows in 'Files to fix first' (default: 10)")
    parser.add_argument("--pdf", action="store_true", help="also export a tagged PDF with LibreOffice")
    args = parser.parse_args(argv)

    records = read_inventory(Path(args.inventory))
    crawl = json.loads(Path(args.crawl_summary).read_text(encoding="utf-8")) if args.crawl_summary else {}
    numbers = summarize(records)
    site = args.site or crawl.get("start_url", "[website]")
    work_pages = numbers["work_pages"]

    values = {
        "CLIENT_NAME": args.client, "SITE_URL": site, "PREPARED_BY": args.prepared_by, "CONTACT": args.contact,
        "REPORT_DATE": dt.date.today().strftime("%d %B %Y").lstrip("0"),
        "CRAWL_DATE": nice_date(crawl.get("date")), "START_URL": site, "CRAWL_DELAY": delay_text(crawl),
        "DEPTH": crawl.get("depth", "[depth]"), "PAGES_CRAWLED": crawl.get("pages_crawled", "[number of]"),
        "TOTAL_FILES": numbers["total_files"], "TOTAL_PAGES": numbers["total_pages"],
        "PDF_COUNT": numbers["by_type"].get("PDF", 0), "WORD_COUNT": numbers["by_type"].get("Word", 0),
        "EXCEL_COUNT": numbers["by_type"].get("Excel", 0), "PPT_COUNT": numbers["by_type"].get("PowerPoint", 0),
        "UNTAGGED_COUNT": numbers["untagged"], "SCANNED_COUNT": numbers["scanned"],
        "NO_TITLE_COUNT": numbers["no_title"], "NO_LANG_COUNT": numbers["no_lang"], "FORMS_COUNT": numbers["forms"],
        "DUPLICATE_COUNT": numbers["duplicates"], "BROKEN_LINKS": crawl.get("broken_links", "not checked"),
        "VAGUE_LINKS": crawl.get("unclear_link_text", "not checked"),
        "WORK_FILES": numbers["work_files"], "WORK_PAGES": work_pages,
        "RATE_RANGE": f"{money(args.rate_low)} to {money(args.rate_high)}",
        "ESTIMATE_RANGE": f"{money(work_pages * args.rate_low)} to {money(work_pages * args.rate_high)}",
        "SUMMARY": summary_text(records, numbers, crawl, site),
    }

    doc = Document(args.template)
    for paragraph in list(doc.paragraphs):
        if "{{TEMPLATE_NOTE}}" in paragraph.text:
            paragraph._p.getparent().remove(paragraph._p)
    action_rows = [{"A.action": action, "A.files": data["files"], "A.pages": data["pages"],
                    "A.meaning": ACTION_MEANING.get(action, "")} for action, data in numbers["actions"].items()]
    expand_rows(doc, "{{A.", action_rows)
    first = [r for r in records if r["priority"] == "High" and r["action"] in ("Fix", "Convert to web page", "Review by hand")]
    first.sort(key=lambda r: (PRIORITY_ORDER.get(r["priority"], 3), ACTIONS.index(r["action"]) if r["action"] in ACTIONS else 9))
    file_rows = [{"F.file": r["file"], "F.pages": r["pages"], "F.issues": "; ".join(r["issues"]) or "-",
                  "F.action": r["action"]} for r in first[: args.max_files]]
    if not file_rows:
        file_rows = [{"F.file": "None found", "F.pages": "", "F.issues": "", "F.action": ""}]
    expand_rows(doc, "{{F.", file_rows)
    replace_with_paragraphs(doc, "{{WHY_IT_MATTERS}}", WHY_IT_MATTERS[args.sector] + [SECTOR_CLOSING])
    for paragraph in all_paragraphs(doc):
        replace_in_paragraph(paragraph, values)

    props = doc.core_properties
    props.title = f"Document accessibility snapshot: {args.client}"
    props.author = args.prepared_by
    props.last_modified_by = args.prepared_by
    props.created = props.modified = dt.datetime.now().replace(microsecond=0)
    props.language = "en-US"
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    leftover = sorted({m for p in all_paragraphs(Document(out)) for m in PLACEHOLDER.findall(p.text)})
    print(f"Wrote {out}")
    if leftover:
        print("Still to fill by hand: " + ", ".join("{{" + m + "}}" for m in leftover))
    if args.pdf:
        from docx_to_tagged_pdf import convert
        (pdf,) = convert([out], out.parent)
        print(f"Wrote {pdf} (tagged PDF)")
    print("Read the report before you send it. Never describe files as 'ADA compliant'.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
