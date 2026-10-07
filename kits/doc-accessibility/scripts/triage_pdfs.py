#!/usr/bin/env python3
"""Triage documents for accessibility work (step 1 of the service).

For every PDF it checks:
  - number of pages
  - tags (StructTreeRoot and MarkInfo), plus headings, images without alt
    text and tables without header cells inside the tags
  - real text, or scanned image (with or without an OCR text layer)
  - title (and whether the title bar shows it) and language (/Lang)
  - fillable form fields, bookmarks and security settings
and then suggests an action: Delete, Archive, Convert to web page, Fix, or
Keep (check only).

Word (.docx), PowerPoint (.pptx) and Excel (.xlsx) files get basic checks.
Older formats (.doc, .xls, .ppt, .rtf, .odt and similar) are listed so
nothing is forgotten.

Outputs in --out: inventory.xlsx, inventory.csv and summary.md.

These are automated checks only. They cannot show that a file is
accessible, and they are not a legal opinion.

Examples:
  python triage_pdfs.py downloads/ --out inventory/
  python triage_pdfs.py downloads/ --crawl crawl/documents.csv \\
      --out inventory/ --client "Town of Example"
  python triage_pdfs.py some-file.pdf      (prints the checks, writes nothing)
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

try:
    import pikepdf
    from pypdf import PdfReader
except ImportError as exc:  # pragma: no cover - friendly message only
    sys.exit(f"Missing Python package ({exc.name}). Run: pip install -r requirements.txt")

# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------

TYPE_BY_EXT = {
    ".pdf": "PDF",
    ".docx": "Word", ".docm": "Word", ".dotx": "Word", ".doc": "Word", ".dot": "Word",
    ".rtf": "Word", ".odt": "Word", ".wpd": "Word",
    ".pptx": "PowerPoint", ".pptm": "PowerPoint", ".ppsx": "PowerPoint", ".ppt": "PowerPoint",
    ".pps": "PowerPoint", ".odp": "PowerPoint",
    ".xlsx": "Excel", ".xlsm": "Excel", ".xls": "Excel", ".xlsb": "Excel", ".ods": "Excel",
}
MODERN_OFFICE = {".docx", ".docm", ".dotx", ".pptx", ".pptm", ".ppsx", ".xlsx", ".xlsm"}

MIN_TEXT_CHARS = 25          # fewer characters than this on a page = "no real text"
BIG_IMAGE_SHARE = 0.5        # an image covering half the page or more looks like a scan

# Words in a file name, link text or title that suggest people use the file
# to apply for, or take part in, a service. These files should be fixed first.
SERVICE_WORDS = re.compile(
    r"\b(applications?|forms?|permits?|registration|register|request|licen[cs]es?|"
    r"enrol(l)?ment|sign[\s_-]?up|apply|claims?|complaints?|appeals?|petitions?|waivers?)\b",
    re.I,
)
YEAR_RE = re.compile(r"(?<!\d)(19[89]\d|20[0-4]\d)(?!\d)")

ACTIONS = ["Fix", "Convert to web page", "Archive", "Delete", "Keep (check only)", "Review by hand"]
ACTION_MEANING = {
    "Fix": "Repair or rebuild the file so it works with screen readers, then a person checks it.",
    "Convert to web page": "Put the content on a normal web page instead of a document.",
    "Archive": "Old file. The client's ADA coordinator or lawyer decides if it can move to an archive section.",
    "Delete": "Duplicate or no longer needed. Remove it and fix the links.",
    "Keep (check only)": "Basic checks passed. A person still checks it before it counts as done.",
    "Review by hand": "Could not be checked automatically (password or damaged file).",
}
PRIORITY_ORDER = {"High": 0, "Medium": 1, "Low": 2}

# (key, column header) for the inventory. make_report.py reads these headers.
COLUMNS = [
    ("id", "ID"), ("file", "File"), ("type", "Type"), ("pages", "Pages"),
    ("action", "Suggested action"), ("priority", "Priority"), ("why", "Why"),
    ("issues", "Issues found"), ("tagged", "Tagged"), ("text", "Text"),
    ("needs_ocr", "Needs OCR"), ("has_title", "Has title"), ("title", "Title"),
    ("has_lang", "Has language"), ("lang", "Language"), ("form_fields", "Form fields"),
    ("headings", "Headings"), ("images_no_alt", "Images without alt text"),
    ("bookmarks", "Bookmarks"), ("year", "Year"), ("made_with", "Made with"),
    ("duplicate_of", "Duplicate of"), ("url", "URL"), ("found_on", "Found on page"),
    ("link_text", "Link text"), ("size_kb", "Size (KB)"),
    ("decision", "Client decision"), ("notes", "Notes"),
]
COLUMN_HELP = {
    "ID": "Row number. Other rows refer to it (for example 'Duplicate of').",
    "File": "File name as downloaded.",
    "Type": "PDF, Word, Excel or PowerPoint.",
    "Pages": "Pages (PDF and Word), slides (PowerPoint) or sheets (Excel). Word counts come from the file's saved statistics.",
    "Suggested action": "Fix, Convert to web page, Archive, Delete, Keep (check only) or Review by hand. A suggestion only: the client decides.",
    "Priority": "High = used to apply for or use a service, or linked from the home page. Medium = other files to fix. Low = archive, delete or check only.",
    "Why": "The reason for the suggested action.",
    "Issues found": "Problems the automated checks found. Problems they cannot find (reading order, alt text quality, colour contrast) need a person.",
    "Tagged": "PDF only. Yes = has a tag structure that screen readers use. No = untagged. Partly = tags exist but the file is not marked as tagged.",
    "Text": "PDF only. Real text, Scanned (image only, no text), Scanned with OCR text, Mixed, or No text found.",
    "Needs OCR": "Yes if some or all pages are images with no real text.",
    "Has title": "Yes if the file has a real title (not empty, not 'untitled', not just the file name).",
    "Title": "The title saved in the file.",
    "Has language": "Yes if the main language is set (screen readers need it to pronounce words).",
    "Language": "The language code saved in the file, for example en-US.",
    "Form fields": "Number of fillable form fields.",
    "Headings": "Headings found in the tags (PDF) or heading styles used (Word).",
    "Images without alt text": "Images with no text alternative. Decorative images should be marked as decorative instead.",
    "Bookmarks": "PDF only. Long documents should have bookmarks.",
    "Year": "Newest year found in the file name, link text or title. '(file date)' means it came from the file's own dates.",
    "Made with": "The program that made the file. If it was Word, ask the client for the Word file: fixing from the source is faster.",
    "Duplicate of": "This file is an exact copy of the file with that ID.",
    "URL": "Where the file is on the website (from the crawl).",
    "Found on page": "The first web page that links to it.",
    "Link text": "The text of that link.",
    "Size (KB)": "File size.",
    "Client decision": "For the client: keep, fix, convert, archive or delete.",
    "Notes": "Your notes.",
}


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def yes_no(value: bool) -> str:
    return "Yes" if value else "No"


def short_error(exc: Exception) -> str:
    text = str(exc).strip().splitlines()[0] if str(exc).strip() else exc.__class__.__name__
    # Drop a leading file path ("/path/file.pdf: unable to find trailer").
    text = re.sub(r"^.*[\\/][^:]*:\s*", "", text) or text
    return text[:160]


def sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def years_in(*texts: str) -> list[int]:
    this_year = dt.date.today().year
    found = []
    for text in texts:
        for match in YEAR_RE.findall(text or ""):
            year = int(match)
            if 1980 <= year <= this_year + 1:
                found.append(year)
    return found


def title_problem(title: str, file_name: str) -> str | None:
    """Return a short description if the title is missing or useless."""
    clean = (title or "").strip()
    if not clean:
        return "No title"
    low = clean.lower()
    placeholders = {"untitled", "untitled document", "document", "document1", "title", "pdf",
                    "microsoft word", "slide 1", "presentation1", "book1", "anonymous"}
    if low in placeholders:
        return f"Title is a placeholder ('{clean}')"
    stem = Path(file_name).stem.lower()
    looks_like_file = (
        re.search(r"\.(docx?|pdf|xlsx?|pptx?|rtf|odt|indd|tmp)$", low)
        or low.startswith(("microsoft word - ", "microsoft powerpoint - ", "microsoft excel - "))
        or (" " not in clean and (low == stem or re.search(r"[_.-]", clean)))
    )
    if looks_like_file:
        return f"Title looks like a file name ('{clean}')"
    return None


def placeholder_alt(text: str) -> bool:
    """True for alt text that says nothing: a file name, 'Picture 1', 'image' and so on."""
    clean = (text or "").strip().lower()
    return bool(
        re.fullmatch(r"(picture|image|graphic|figure|img|photo|chart|logo|untitled|object|shape)[\s_-]*\d*", clean)
        or re.search(r"\.(png|jpe?g|gif|bmp|tiff?|svg|emf|wmf|webp)$", clean)
        or re.fullmatch(r"(img|dsc|dscn|pxl|scan)[_-]?\d+", clean)
    )


def plural(count: int, word: str, many: str | None = None) -> str:
    return f"{count} {word if count == 1 else (many or word + 's')}"


def pick_sample_pages(total: int, wanted: int) -> list[int]:
    if total <= wanted:
        return list(range(total))
    step = (total - 1) / (wanted - 1)
    return sorted({round(i * step) for i in range(wanted)})


def new_record(path: Path) -> dict:
    ext = path.suffix.lower()
    return {
        "id": 0, "file": path.name, "path": str(path), "ext": ext,
        "type": TYPE_BY_EXT.get(ext, "Other"), "pages": "", "action": "", "priority": "",
        "why": "", "issues": [], "tagged": "n/a", "text": "", "needs_ocr": "No",
        "has_title": "No", "title": "", "has_lang": "No", "lang": "", "form_fields": 0,
        "headings": "", "images_no_alt": "", "bookmarks": "", "year": "", "year_value": None,
        "made_with": "", "duplicate_of": "", "url": "", "found_on": "", "found_on_depth": "",
        "link_text": "", "times_linked": 0, "size_kb": round(path.stat().st_size / 1024, 1), "decision": "",
        "notes": "", "sha256": "", "readable": True, "meta_years": [],
        # PDF details used by the rules
        "struct": {}, "display_title": False, "xfa": False, "fields_no_name": 0,
        "a11y_blocked": False, "scan_pages": 0, "checked_pages": 0, "pdfua_claim": False,
        "title_issue": None,
    }


# ---------------------------------------------------------------------------
# PDF checks
# ---------------------------------------------------------------------------

STANDARD_STRUCT = {
    "Document", "DocumentFragment", "Part", "Art", "Sect", "Div", "BlockQuote", "Caption",
    "TOC", "TOCI", "Index", "NonStruct", "Private", "P", "H", "H1", "H2", "H3", "H4", "H5",
    "H6", "L", "LI", "Lbl", "LBody", "Table", "TR", "TH", "TD", "THead", "TBody", "TFoot",
    "Span", "Quote", "Note", "Reference", "BibEntry", "Code", "Link", "Annot", "Ruby", "RB",
    "RT", "RP", "Warichu", "WT", "WP", "Figure", "Formula", "Form", "Title", "FENote", "Sub",
    "Em", "Strong", "Aside", "Artifact",
}
HEADING_ROLES = {"H", "H1", "H2", "H3", "H4", "H5", "H6"}


def analyze_structure(struct_root, limit: int = 50000) -> dict:
    """Count tag types in the structure tree (resolving custom names via RoleMap)."""
    out = {"elements": 0, "headings": 0, "figures": 0, "figures_no_alt": 0, "figures_bad_alt": 0,
           "tables": 0, "th": 0, "lists": 0, "paragraphs": 0, "truncated": False}
    rolemap = {}
    rm = struct_root.get("/RoleMap")
    if isinstance(rm, pikepdf.Dictionary):
        for key in rm.keys():
            try:
                rolemap[str(key)[1:]] = str(rm[key])[1:]
            except Exception:
                pass

    def resolve(name: str) -> str:
        seen = set()
        while name not in STANDARD_STRUCT and name in rolemap and name not in seen:
            seen.add(name)
            name = rolemap[name]
        return name

    stack = [struct_root.get("/K")]
    visited = set()
    while stack:
        node = stack.pop()
        if node is None:
            continue
        if isinstance(node, pikepdf.Array):
            stack.extend(node[i] for i in range(len(node) - 1, -1, -1))
            continue
        if not isinstance(node, pikepdf.Dictionary):
            continue  # marked-content id (a number)
        if node.is_indirect:
            if node.objgen in visited:
                continue
            visited.add(node.objgen)
        tag = node.get("/S")
        if tag is None:
            continue  # marked-content or object reference, not an element
        out["elements"] += 1
        if out["elements"] > limit:
            out["truncated"] = True
            break
        role = resolve(str(tag)[1:])
        if role in HEADING_ROLES:
            out["headings"] += 1
        elif role == "Figure":
            out["figures"] += 1
            alt = str(node.get("/Alt") or "").strip()
            actual = str(node.get("/ActualText") or "").strip()
            if not (alt or actual):
                out["figures_no_alt"] += 1
            elif placeholder_alt(alt or actual):
                out["figures_bad_alt"] += 1
        elif role == "P":
            out["paragraphs"] += 1
        elif role == "Table":
            out["tables"] += 1
        elif role == "TH":
            out["th"] += 1
        elif role == "L":
            out["lists"] += 1
        kids = node.get("/K")
        if kids is not None:
            stack.append(kids)
    return out


def _resources_of(page_dict):
    node, hops = page_dict, 0
    while node is not None and hops < 32:
        res = node.get("/Resources")
        if res is not None:
            return res
        node = node.get("/Parent")
        hops += 1
    return None


def _det(values) -> float:
    a, b, c, d = (float(v) for v in list(values)[:4])
    return a * d - b * c


def content_stats(owner, resources, det: float = 1.0, depth: int = 0, stats: dict | None = None,
                  chain: frozenset = frozenset()) -> dict:
    """Walk a content stream. Adds up the area covered by images and counts
    visible and invisible (OCR layer, render mode 3) text drawing operations."""
    if stats is None:
        stats = {"image_area": 0.0, "visible_text": 0, "invisible_text": 0, "ops": 0}
    if depth > 4:
        return stats
    try:
        instructions = pikepdf.parse_content_stream(owner)
    except Exception:
        return stats
    xobjects = resources.get("/XObject") if isinstance(resources, pikepdf.Dictionary) else None
    saved = []
    cur_det, render_mode = det, 0
    for inst in instructions:
        stats["ops"] += 1
        if stats["ops"] > 300000:
            break
        op = str(inst.operator)
        try:
            if op == "q":
                saved.append((cur_det, render_mode))
            elif op == "Q":
                if saved:
                    cur_det, render_mode = saved.pop()
            elif op == "cm":
                cur_det *= _det(inst.operands)
            elif op == "Tr":
                render_mode = int(inst.operands[0])
            elif op in ("Tj", "TJ", "'", '"'):
                if render_mode == 3:
                    stats["invisible_text"] += 1
                else:
                    stats["visible_text"] += 1
            elif op == "INLINE IMAGE":
                stats["image_area"] += abs(cur_det)
            elif op == "Do" and xobjects is not None:
                xobj = xobjects.get(str(inst.operands[0]))
                if xobj is None:
                    continue
                subtype = str(xobj.get("/Subtype", ""))
                if subtype == "/Image":
                    stats["image_area"] += abs(cur_det)
                elif subtype == "/Form":
                    key = xobj.objgen if xobj.is_indirect else None
                    if key is not None and key in chain:
                        continue
                    matrix = xobj.get("/Matrix")
                    form_det = cur_det * (_det(matrix) if matrix is not None else 1.0)
                    form_res = xobj.get("/Resources")
                    content_stats(xobj, form_res if form_res is not None else resources,
                                  form_det, depth + 1, stats, chain | {key})
        except Exception:
            continue
    return stats


def classify_page(chars: int, image_share: float, stats: dict) -> str:
    big_image = image_share >= BIG_IMAGE_SHARE
    if chars >= MIN_TEXT_CHARS:
        if big_image and stats["invisible_text"] > 0 and stats["visible_text"] == 0:
            return "ocr"
        return "text"
    return "scan" if big_image else "empty"


def analyze_pdf(path: Path, rec: dict, sample_pages: int) -> None:
    try:
        pdf = pikepdf.open(path)
    except pikepdf.PasswordError:
        rec["readable"] = False
        rec["issues"].append("Needs a password to open")
        return
    except Exception as exc:
        rec["readable"] = False
        rec["issues"].append(f"Could not open the file ({short_error(exc)})")
        return

    with pdf:
        root = pdf.Root
        rec["pages"] = len(pdf.pages)

        # Title, language, metadata
        info = pdf.docinfo if "/Info" in pdf.trailer else {}
        title = str(info.get("/Title", "") or "").strip() if info else ""
        try:
            meta = pdf.open_metadata()
            if not title:
                title = str(meta.get("dc:title", "") or "").strip()
            rec["pdfua_claim"] = any("pdfua/ns/id" in str(key) for key in meta.keys())
        except Exception:
            pass
        rec["title"] = title
        rec["title_issue"] = title_problem(title, path.name)
        rec["has_title"] = yes_no(rec["title_issue"] is None)
        lang = str(root.get("/Lang", "") or "").strip()
        rec["lang"] = lang
        rec["has_lang"] = yes_no(bool(lang))
        prefs = root.get("/ViewerPreferences")
        rec["display_title"] = bool(prefs.get("/DisplayDocTitle", False)) if isinstance(prefs, pikepdf.Dictionary) else False
        creator = str(info.get("/Creator", "") or "").strip() if info else ""
        producer = str(info.get("/Producer", "") or "").strip() if info else ""
        rec["made_with"] = " / ".join(x for x in (creator, producer) if x)[:120]
        for key in ("/CreationDate", "/ModDate"):
            if info and key in info:
                rec["meta_years"] += years_in(str(info.get(key)))

        # Tags
        mark_info = root.get("/MarkInfo")
        marked = bool(mark_info.get("/Marked", False)) if isinstance(mark_info, pikepdf.Dictionary) else False
        struct_root = root.get("/StructTreeRoot")
        struct = analyze_structure(struct_root) if isinstance(struct_root, pikepdf.Dictionary) else {}
        rec["struct"] = struct
        elements = struct.get("elements", 0)
        if elements and marked:
            rec["tagged"] = "Yes"
        elif elements:
            rec["tagged"] = "Partly"
        else:
            rec["tagged"] = "No"
        rec["headings"] = struct.get("headings", 0) if elements else ""
        rec["images_no_alt"] = (struct.get("figures_no_alt", 0) + struct.get("figures_bad_alt", 0)) if elements else ""
        if marked and not elements:
            rec["issues"].append("Marked as tagged but has no tags")

        # Forms
        acro = root.get("/AcroForm")
        if isinstance(acro, pikepdf.Dictionary):
            rec["xfa"] = "/XFA" in acro
            count, no_name = 0, 0
            stack = [(f, 0) for f in (acro.get("/Fields") or [])]
            while stack:
                field, depth = stack.pop()
                if not isinstance(field, pikepdf.Dictionary) or depth > 20:
                    continue
                kids = field.get("/Kids")
                child_fields = [k for k in kids if isinstance(k, pikepdf.Dictionary) and "/T" in k] if kids is not None else []
                if child_fields:
                    stack.extend((k, depth + 1) for k in child_fields)
                    continue
                count += 1
                tooltip = field.get("/TU")
                if tooltip is None or not str(tooltip).strip():
                    no_name += 1
            rec["form_fields"] = count
            rec["fields_no_name"] = no_name

        # Bookmarks and security
        outlines = root.get("/Outlines")
        rec["bookmarks"] = yes_no(isinstance(outlines, pikepdf.Dictionary) and outlines.get("/First") is not None)
        if pdf.is_encrypted:
            try:
                # Bit 10 of /P is "extract text for accessibility". Newer files may have
                # it cleared even though modern readers ignore it, so report it either way.
                p_value = int(pdf.trailer.Encrypt.get("/P", -1))
                rec["a11y_blocked"] = not (p_value & 512) or not pdf.allow.accessibility
            except Exception:
                pass

        # Real text or scanned image? Check a sample of pages.
        try:
            reader = PdfReader(str(path), strict=False)
            if reader.is_encrypted:
                reader.decrypt("")
        except Exception:
            reader = None
        kinds = []
        for index in pick_sample_pages(len(pdf.pages), sample_pages):
            page = pdf.pages[index]
            chars = 0
            if reader is not None:
                try:
                    chars = len(re.sub(r"\s+", "", reader.pages[index].extract_text() or ""))
                except Exception:
                    chars = 0
            try:
                box = [float(v) for v in page.mediabox]
                area = abs((box[2] - box[0]) * (box[3] - box[1])) or 1.0
            except Exception:
                area = 612.0 * 792.0
            stats = content_stats(page, _resources_of(page.obj))
            if reader is None and (stats["visible_text"] or stats["invisible_text"]):
                chars = MIN_TEXT_CHARS  # text extraction failed; fall back to "the page draws text"
            kinds.append(classify_page(chars, min(stats["image_area"] / area, 1.0), stats))

    counts = Counter(kinds)
    checked = len(kinds)
    rec["checked_pages"] = checked
    rec["scan_pages"] = counts["scan"]
    if checked == 0:
        rec["text"] = "No pages"
    elif counts["scan"] == checked:
        rec["text"] = "Scanned, no text"
    elif counts["ocr"] and counts["ocr"] + counts["empty"] == checked:
        rec["text"] = "Scanned with OCR text"
    elif counts["scan"]:
        rec["text"] = "Mixed"
    elif counts["ocr"]:
        rec["text"] = "Mixed"
    elif counts["empty"] == checked:
        rec["text"] = "No text found"
    else:
        rec["text"] = "Real text"
    rec["needs_ocr"] = yes_no(counts["scan"] > 0)

    # Plain-English issues
    issues = rec["issues"]
    if rec["tagged"] == "No":
        issues.append("Not tagged (screen readers get no structure)")
    elif rec["tagged"] == "Partly":
        issues.append("Has tags but is not marked as tagged")
    if rec["text"] == "Scanned, no text":
        issues.append("Scanned image with no real text (needs OCR)")
    elif counts["scan"]:
        issues.append(f"{counts['scan']} of {checked} checked pages are images with no real text")
    elif rec["text"] == "Scanned with OCR text":
        issues.append("Scanned pages with an OCR text layer (check the text is right)")
    elif rec["text"] == "No text found":
        issues.append("No real text found (blank, or text drawn as shapes)")
    if rec["title_issue"]:
        issues.append(rec["title_issue"])
    if not rec["lang"]:
        issues.append("No language set")
    if rec["form_fields"]:
        text = plural(rec["form_fields"], "fillable form field")
        if rec["fields_no_name"]:
            text += f" ({rec['fields_no_name']} with no name for screen readers)"
        issues.append(text)
    if rec["xfa"]:
        issues.append("XFA form (must be rebuilt as a standard form)")
    if rec["tagged"] in ("Yes", "Partly"):
        struct = rec["struct"]
        if not rec["display_title"]:
            issues.append("Title bar shows the file name, not the title")
        long_enough = (isinstance(rec["pages"], int) and rec["pages"] >= 2) or struct.get("paragraphs", 0) >= 6
        if struct.get("headings", 0) == 0 and long_enough:
            issues.append("No headings in the tags")
        if struct.get("figures_no_alt"):
            issues.append(plural(struct["figures_no_alt"], "image") + " without alt text")
        if struct.get("figures_bad_alt"):
            issues.append(plural(struct["figures_bad_alt"], "image") + " with a file name or placeholder as alt text")
        if struct.get("tables") and not struct.get("th"):
            issues.append("Tables have no header cells")
        if isinstance(rec["pages"], int) and struct.get("elements", 0) < rec["pages"]:
            issues.append(f"Tags look incomplete ({struct.get('elements', 0)} tags for {rec['pages']} pages)")
    if isinstance(rec["pages"], int) and rec["pages"] >= 10 and rec["bookmarks"] == "No":
        issues.append("No bookmarks (long document)")
    if rec["a11y_blocked"]:
        issues.append("Security settings may block screen readers")


# ---------------------------------------------------------------------------
# Office checks (basic)
# ---------------------------------------------------------------------------

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _read_xml(archive: zipfile.ZipFile, name: str, max_bytes: int = 60_000_000):
    try:
        info = archive.getinfo(name)
    except KeyError:
        return None
    if info.file_size > max_bytes:
        return None
    return ET.fromstring(archive.read(name))


def _office_props(archive: zipfile.ZipFile) -> dict:
    props = {}
    for part in ("docProps/core.xml", "docProps/app.xml"):
        root = _read_xml(archive, part)
        if root is not None:
            for element in root:
                props[element.tag.split("}")[-1]] = (element.text or "").strip()
    return props


def _decorative(element) -> bool:
    return any(child.tag.endswith("}decorative") and child.get("val") in ("1", "true")
               for child in element.iter())


def analyze_docx(path: Path, rec: dict) -> None:
    with zipfile.ZipFile(path) as archive:
        props = _office_props(archive)
        styles = _read_xml(archive, "word/styles.xml")
        document = _read_xml(archive, "word/document.xml")
    heading_styles, default_lang = set(), ""
    if styles is not None:
        for style in styles.iter(W + "style"):
            if style.get(W + "type") != "paragraph":
                continue
            name_el = style.find(W + "name")
            name = (name_el.get(W + "val") if name_el is not None else "") or ""
            ppr = style.find(W + "pPr")
            if re.match(r"(?i)^(heading\s*\d|title)$", name) or (ppr is not None and ppr.find(W + "outlineLvl") is not None):
                heading_styles.add(style.get(W + "styleId"))
        lang_el = styles.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr/{W}lang")
        if lang_el is not None:
            default_lang = lang_el.get(W + "val") or ""
    headings = images = no_alt = tables = tables_no_header = 0
    doc_lang = ""
    if document is not None:
        for para in document.iter(W + "p"):
            ppr = para.find(W + "pPr")
            if ppr is None:
                continue
            style = ppr.find(W + "pStyle")
            if (style is not None and style.get(W + "val") in heading_styles) or ppr.find(W + "outlineLvl") is not None:
                headings += 1
        for drawing in document.iter(WP + "docPr"):
            images += 1
            descr = (drawing.get("descr") or "").strip()
            if (not descr or placeholder_alt(descr)) and not _decorative(drawing):
                no_alt += 1
        for table in document.iter(W + "tbl"):
            tables += 1
            first_row = table.find(W + "tr")
            if first_row is None or first_row.find(f"{W}trPr/{W}tblHeader") is None:
                tables_no_header += 1
        lang_el = next(document.iter(W + "lang"), None)
        if lang_el is not None:
            doc_lang = lang_el.get(W + "val") or ""
    rec["title"] = props.get("title", "")
    rec["title_issue"] = title_problem(rec["title"], path.name)
    rec["has_title"] = yes_no(rec["title_issue"] is None)
    rec["lang"] = default_lang or doc_lang or props.get("language", "")
    rec["has_lang"] = yes_no(bool(rec["lang"]))
    rec["pages"] = int(props["Pages"]) if props.get("Pages", "").isdigit() else ""
    rec["headings"] = headings
    rec["images_no_alt"] = no_alt
    rec["made_with"] = props.get("Application", "")
    for key in ("created", "modified"):
        rec["meta_years"] += years_in(props.get(key, ""))
    issues = rec["issues"]
    if rec["title_issue"]:
        issues.append(rec["title_issue"])
    if not rec["lang"]:
        issues.append("No language set")
    if headings == 0:
        issues.append("No heading styles used")
    if no_alt:
        issues.append(plural(no_alt, "image") + " without alt text (or with a file name as alt text)")
    if tables_no_header:
        issues.append(plural(tables_no_header, "table") + " without a header row")


def analyze_pptx(path: Path, rec: dict) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(str(path))

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(shape.shapes)
            else:
                yield shape

    no_title = pictures = no_alt = 0
    for slide in deck.slides:
        title_shape = slide.shapes.title
        if title_shape is None or not title_shape.text_frame.text.strip():
            no_title += 1
        for shape in walk(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
                props = shape._element.find(f".//{P_NS}cNvPr")
                descr = (props.get("descr") or "").strip() if props is not None else ""
                if (not descr or placeholder_alt(descr)) and not (props is not None and _decorative(props)):
                    no_alt += 1
    core = deck.core_properties
    rec["pages"] = len(deck.slides)
    rec["title"] = core.title or ""
    rec["title_issue"] = title_problem(rec["title"], path.name)
    rec["has_title"] = yes_no(rec["title_issue"] is None)
    rec["lang"] = core.language or ""
    rec["has_lang"] = "Yes" if rec["lang"] else "n/a"  # PowerPoint stores language per text run
    rec["headings"] = "n/a"
    rec["images_no_alt"] = no_alt
    with zipfile.ZipFile(path) as archive:
        rec["made_with"] = _office_props(archive).get("Application", "")
    for value in (core.created, core.modified):
        if value:
            rec["meta_years"].append(value.year)
    if rec["title_issue"]:
        rec["issues"].append(rec["title_issue"])
    if no_title:
        rec["issues"].append(plural(no_title, "slide") + " without a title")
    if no_alt:
        rec["issues"].append(plural(no_alt, "picture") + " without alt text (or with a file name as alt text)")


def analyze_xlsx(path: Path, rec: dict) -> None:
    from openpyxl import load_workbook

    book = load_workbook(str(path), read_only=True)
    names = list(book.sheetnames)
    props = book.properties
    book.close()
    rec["pages"] = len(names)
    rec["title"] = props.title or ""
    rec["title_issue"] = title_problem(rec["title"], path.name)
    rec["has_title"] = yes_no(rec["title_issue"] is None)
    rec["lang"] = props.language or ""
    rec["has_lang"] = "Yes" if rec["lang"] else "n/a"  # rarely set in Excel files
    rec["headings"] = "n/a"
    with zipfile.ZipFile(path) as archive:
        rec["made_with"] = _office_props(archive).get("Application", "")
    for value in (props.created, props.modified):
        if value:
            rec["meta_years"].append(value.year)
    if rec["title_issue"]:
        rec["issues"].append(rec["title_issue"])
    default_names = [n for n in names if re.fullmatch(r"(?i)sheet\d+", n)]
    if default_names:
        rec["issues"].append(f"Sheet names not changed from the default ({', '.join(default_names)})")


# ---------------------------------------------------------------------------
# Suggested action
# ---------------------------------------------------------------------------

def suggest_action(rec: dict, archive_before: int, short_pages: int) -> None:
    """Fill in action, why and priority. The first rule that matches wins."""
    pages = rec["pages"] if isinstance(rec["pages"], int) else None
    words = " ".join([rec["file"], rec.get("link_text", ""), rec.get("title", ""), rec.get("url", "")])
    service_match = SERVICE_WORDS.search(words.replace("_", " ").replace("-", " "))
    is_service = bool(rec.get("form_fields")) or bool(rec.get("xfa")) or bool(service_match)
    text_years = years_in(rec["file"], rec.get("url", ""), rec.get("link_text", ""), rec.get("title", ""))
    if text_years:
        rec["year_value"] = max(text_years)
        rec["year"] = str(rec["year_value"])
    elif rec["meta_years"]:
        rec["year_value"] = max(rec["meta_years"])
        rec["year"] = f"{rec['year_value']} (file date)"
    year = rec["year_value"]
    app = {"Word": "Word", "PowerPoint": "PowerPoint", "Excel": "Excel"}.get(rec["type"], "the program that made it")

    def done(action, why, priority):
        rec["action"], rec["why"], rec["priority"] = action, why, priority

    if not rec["readable"]:
        return done("Review by hand", f"Could not be checked automatically: {'; '.join(rec['issues'])}. Open it by hand.", "Medium")
    if rec["duplicate_of"]:
        return done("Delete", f"Exact copy of ID {rec['duplicate_of']}. Keep one copy and point every link to it.", "Low")
    if year is not None and year < archive_before and not is_service:
        return done("Archive", f"Dated {year}. If it is kept only for reference and not changed, the client's ADA "
                    "coordinator or lawyer may decide it can go in an archive section. Otherwise fix or delete it.", "Low")
    if is_service:
        if rec.get("form_fields"):
            what = plural(rec["form_fields"], "fillable field")
        elif rec.get("xfa"):
            what = "an XFA form"
        else:
            what = f"'{service_match.group(0)}' in its name or link"
        return done("Fix", f"People use it to apply for or use a service ({what}). Fix it first, "
                    "or replace it with an accessible web form.", "High")
    if rec["ext"] not in MODERN_OFFICE and rec["type"] != "PDF":
        return done("Fix", "Old file format. Open it, save it in a current format, then check it.", "Medium")
    if rec["type"] != "PDF":
        if rec["issues"]:
            article = "an" if app[0].lower() in "aeiou" else "a"
            return done("Fix", f"Fix it in {app} with the built-in Accessibility Checker. Then keep it as {article} "
                        f"{app} file, export a tagged PDF, or make it a web page.", "Medium")
        return done("Keep (check only)", "The basic checks found no problems. A person still checks it.", "Low")
    if rec["needs_ocr"] == "Yes":
        return done("Fix", "Scanned image with no real text. Run OCR, check the text, then tag it (or retype it).", "Medium")
    if rec["tagged"] != "Yes" and pages is not None and pages <= short_pages:
        return done("Convert to web page", f"Short ({pages} page{'s' if pages != 1 else ''}) and untagged. A web page is "
                    "often cheaper to make accessible and easier to read on a phone.", "Medium")
    if rec["tagged"] != "Yes":
        return done("Fix", "No tags. Best route: rebuild it from the source file (ask the client for the Word "
                    "file) and export a tagged PDF.", "Medium")
    if rec["issues"]:
        return done("Fix", "Tagged, but needs repair: " + "; ".join(rec["issues"]).lower() + ".", "Medium")
    return done("Keep (check only)", "Tags, title and language found. A person still checks reading order, "
                "alt text and tables.", "Low")


# ---------------------------------------------------------------------------
# Running the checks
# ---------------------------------------------------------------------------

def analyze_file(path: Path, sample_pages: int = 15) -> dict:
    rec = new_record(path)
    rec["sha256"] = sha256_of(path)
    try:
        if rec["ext"] == ".pdf":
            analyze_pdf(path, rec, sample_pages)
        elif rec["ext"] in (".docx", ".docm", ".dotx"):
            analyze_docx(path, rec)
        elif rec["ext"] in (".pptx", ".pptm", ".ppsx"):
            analyze_pptx(path, rec)
        elif rec["ext"] in (".xlsx", ".xlsm"):
            analyze_xlsx(path, rec)
        else:
            rec["issues"].append("Old file format: cannot be checked automatically")
    except zipfile.BadZipFile:
        rec["readable"] = False
        rec["issues"].append("Damaged or not really an Office file")
    except Exception as exc:
        rec["readable"] = False
        rec["issues"].append(f"Could not be checked ({short_error(exc)})")
    return rec


def triage_one(path: Path, archive_before: int | None = None, short_pages: int = 1,
               sample_pages: int = 15) -> dict:
    """Check one file and suggest an action (used by the other scripts)."""
    if archive_before is None:
        archive_before = dt.date.today().year - 3
    rec = analyze_file(Path(path), sample_pages)
    suggest_action(rec, archive_before, short_pages)
    return rec


def find_files(inputs: list[str]) -> list[Path]:
    files = []
    for item in inputs:
        path = Path(item)
        if path.is_dir():
            files += [p for p in path.rglob("*") if p.is_file() and p.suffix.lower() in TYPE_BY_EXT]
        elif path.is_file():
            files.append(path)
        else:
            print(f"Not found: {item}", file=sys.stderr)
    return sorted(set(files), key=lambda p: str(p).lower())


def mark_duplicates(records: list[dict]) -> None:
    """Number the rows and mark exact copies. The copy to keep is the one with
    the most links to it, then the shortest file name."""
    groups: dict[str, list[dict]] = {}
    for index, rec in enumerate(records, start=1):
        rec["id"] = index
        groups.setdefault(rec["sha256"], []).append(rec)
    for group in groups.values():
        if len(group) < 2:
            continue
        keeper = min(group, key=lambda r: (-r.get("times_linked", 0), len(r["file"]), r["file"].lower()))
        for rec in group:
            if rec is not keeper:
                rec["duplicate_of"] = keeper["id"]


def load_crawl(csv_path: Path | None) -> tuple[dict, list]:
    """Read the crawler's documents.csv. Returns {saved file name: row} and the
    rows for links that were not downloaded (broken, blocked, off-site)."""
    if not csv_path:
        return {}, []
    by_file, not_checked = {}, []
    with open(csv_path, newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            if row.get("saved_as"):
                by_file[row["saved_as"]] = row
            else:
                not_checked.append(row)
    return by_file, not_checked


def summarize(records: list[dict], not_checked: list | None = None) -> dict:
    """Numbers for the summary sheet, summary.md and the snapshot report."""
    def num(value):
        try:
            return int(value)
        except (TypeError, ValueError):
            return 0

    pdfs = [r for r in records if r["type"] == "PDF"]
    summary = {
        "total_files": len(records),
        "total_pages": sum(num(r["pages"]) for r in records),
        "unknown_pages": sum(1 for r in records if num(r["pages"]) == 0),
        "by_type": Counter(r["type"] for r in records),
        "pdf_count": len(pdfs),
        "untagged": sum(1 for r in pdfs if r["tagged"] in ("No", "Partly")),
        "scanned": sum(1 for r in pdfs if r["needs_ocr"] == "Yes"),
        "no_title": sum(1 for r in records if r["has_title"] == "No"),
        "no_lang": sum(1 for r in records if r["has_lang"] == "No" and r["type"] in ("PDF", "Word")),
        "forms": sum(1 for r in records if num(r["form_fields"]) > 0),
        "duplicates": sum(1 for r in records if r["duplicate_of"]),
        "high_priority": sum(1 for r in records if r["priority"] == "High"),
        "actions": {},
        "not_checked": len(not_checked or []),
        "broken_links": sum(1 for r in (not_checked or []) if str(r.get("http_status", "")).startswith(("4", "5"))),
    }
    for action in ACTIONS:
        rows = [r for r in records if r["action"] == action]
        if rows:
            summary["actions"][action] = {"files": len(rows), "pages": sum(num(r["pages"]) for r in rows)}
    work = [r for r in records if r["action"] in ("Fix", "Convert to web page", "Review by hand")]
    summary["work_files"] = len(work)
    summary["work_pages"] = sum(max(num(r["pages"]), 1) for r in work)
    return summary


def sort_key(rec: dict):
    return (PRIORITY_ORDER.get(rec["priority"], 3), ACTIONS.index(rec["action"]) if rec["action"] in ACTIONS else 9,
            rec["file"].lower())


def cell_value(rec: dict, key: str):
    value = rec.get(key, "")
    if key == "issues":
        return "; ".join(value)
    return value


# ---------------------------------------------------------------------------
# Writing the outputs
# ---------------------------------------------------------------------------

def write_csv(records: list[dict], path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow([header for _, header in COLUMNS])
        for rec in records:
            writer.writerow([cell_value(rec, key) for key, _ in COLUMNS])


def write_xlsx(records: list[dict], summary: dict, not_checked: list, path: Path, client: str,
               rate_low: float, rate_high: float, source: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    book = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DCE6F1")
    wrap = Alignment(wrap_text=True, vertical="top")

    # Summary sheet
    sheet = book.active
    sheet.title = "Summary"
    sheet["A1"] = "Document accessibility inventory"
    sheet["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Client", client or ""),
        ("Checked on", dt.date.today().isoformat()),
        ("Files checked from", source),
        ("", ""),
        ("Key numbers", "Count"),
        ("Documents checked", summary["total_files"]),
        ("Total pages (PDF and Word pages, slides, sheets)", summary["total_pages"]),
        ("PDF files", summary["by_type"].get("PDF", 0)),
        ("Word files", summary["by_type"].get("Word", 0)),
        ("Excel files", summary["by_type"].get("Excel", 0)),
        ("PowerPoint files", summary["by_type"].get("PowerPoint", 0)),
        ("PDFs without tags", summary["untagged"]),
        ("PDFs with scanned pages and no real text", summary["scanned"]),
        ("Files without a proper title", summary["no_title"]),
        ("PDF and Word files without a language", summary["no_lang"]),
        ("Files with fillable form fields", summary["forms"]),
        ("Exact duplicates", summary["duplicates"]),
        ("High-priority files", summary["high_priority"]),
        ("Document links not checked (see 'Not checked' sheet)", summary["not_checked"]),
        ("", ""),
        ("Suggested action", "Files"),
    ]
    action_start = None
    for row_index, (label, value) in enumerate(rows, start=3):
        sheet.cell(row=row_index, column=1, value=label)
        sheet.cell(row=row_index, column=2, value=value)
        if label in ("Key numbers", "Suggested action"):
            for col in (1, 2, 3):
                sheet.cell(row=row_index, column=col).font = bold
                sheet.cell(row=row_index, column=col).fill = head_fill
            if label == "Suggested action":
                sheet.cell(row=row_index, column=3, value="Pages")
                action_start = row_index + 1
    row_index = action_start
    for action in ACTIONS:
        if action in summary["actions"]:
            sheet.cell(row=row_index, column=1, value=action)
            sheet.cell(row=row_index, column=2, value=summary["actions"][action]["files"])
            sheet.cell(row=row_index, column=3, value=summary["actions"][action]["pages"])
            row_index += 1
    row_index += 1
    estimate_low = summary["work_pages"] * rate_low
    estimate_high = summary["work_pages"] * rate_high
    notes = [
        f"Pages to fix or convert: {summary['work_pages']} (files marked Fix, Convert to web page or Review by hand; "
        "files with an unknown page count are counted as 1 page).",
        f"Rough estimate at ${rate_low:,.0f} to ${rate_high:,.0f} per page: ${estimate_low:,.0f} to ${estimate_high:,.0f}. "
        "An estimate only, before the client decides what to delete or archive.",
        "These counts come from the automated checks when this file was made. If you change the "
        "'Suggested action' column, run make_report.py: it recounts from the Inventory sheet.",
        "Automated checks find only some problems. They cannot judge reading order, alt text quality, "
        "colour contrast or whether tables make sense. A person checks every fixed file.",
        "This is not a legal opinion or a certification. Exceptions (for example archived content) are "
        "decided by the client's ADA coordinator or lawyer.",
    ]
    sheet.cell(row=row_index, column=1, value="Notes").font = bold
    for offset, note in enumerate(notes, start=1):
        cell = sheet.cell(row=row_index + offset, column=1, value=note)
        cell.alignment = wrap
    sheet.column_dimensions["A"].width = 95
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 10

    # Inventory sheet
    inv = book.create_sheet("Inventory")
    headers = [header for _, header in COLUMNS]
    inv.append(headers)
    for rec in records:
        inv.append([cell_value(rec, key) for key, _ in COLUMNS])
    widths = {"ID": 6, "File": 34, "Type": 11, "Pages": 8, "Suggested action": 20, "Priority": 10,
              "Why": 50, "Issues found": 50, "Title": 30, "URL": 45, "Found on page": 40,
              "Link text": 28, "Made with": 28, "Client decision": 18, "Notes": 30}
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        inv.column_dimensions[letter].width = widths.get(header, 13)
    for row in inv.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    url_col = headers.index("URL") + 1
    for row_number in range(2, inv.max_row + 1):
        cell = inv.cell(row=row_number, column=url_col)
        if isinstance(cell.value, str) and cell.value.startswith("http"):
            cell.hyperlink = cell.value
    if records:
        table = Table(displayName="Inventory", ref=f"A1:{get_column_letter(len(headers))}{len(records) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        inv.add_table(table)
    inv.freeze_panes = "C2"

    # Not checked sheet
    nc = book.create_sheet("Not checked")
    nc.append(["URL", "Type", "Found on page", "Link text", "Why it was not checked"])
    for row in not_checked:
        nc.append([row.get("url", ""), row.get("file_type", ""), row.get("first_found_on", ""),
                   row.get("first_link_text", ""), row.get("note", "") or f"HTTP {row.get('http_status', '')}"])
    for letter, width in zip("ABCDE", (55, 10, 45, 30, 50)):
        nc.column_dimensions[letter].width = width
    for cell in nc[1]:
        cell.font = bold
        cell.fill = head_fill

    # How to read sheet
    help_sheet = book.create_sheet("How to read")
    help_sheet.append(["Column", "What it means"])
    for header in headers:
        help_sheet.append([header, COLUMN_HELP.get(header, "")])
    for cell in help_sheet[1]:
        cell.font = bold
        cell.fill = head_fill
    help_sheet.column_dimensions["A"].width = 24
    help_sheet.column_dimensions["B"].width = 110
    for row in help_sheet.iter_rows(min_row=2):
        row[1].alignment = wrap

    book.properties.title = f"Document accessibility inventory{' - ' + client if client else ''}"
    book.properties.language = "en-US"
    book.save(path)


def summary_markdown(records: list[dict], summary: dict, client: str, source: str,
                     rate_low: float, rate_high: float) -> str:
    lines = [f"# Document inventory summary{': ' + client if client else ''}", "",
             f"Checked on {dt.date.today().isoformat()} from `{source}`. Automated checks only.", "",
             "## Key numbers", "", "| Measure | Count |", "|---|---|"]
    for label, value in [
        ("Documents checked", summary["total_files"]),
        ("Total pages (slides and sheets included)", summary["total_pages"]),
        ("PDF / Word / Excel / PowerPoint", " / ".join(str(summary["by_type"].get(t, 0)) for t in ("PDF", "Word", "Excel", "PowerPoint"))),
        ("PDFs without tags", summary["untagged"]),
        ("PDFs with scanned pages and no real text", summary["scanned"]),
        ("Files without a proper title", summary["no_title"]),
        ("PDF and Word files without a language", summary["no_lang"]),
        ("Files with fillable form fields", summary["forms"]),
        ("Exact duplicates", summary["duplicates"]),
        ("Document links not checked (broken, blocked or off-site)", summary["not_checked"]),
    ]:
        lines.append(f"| {label} | {value} |")
    lines += ["", "## Suggested actions", "", "| Action | Files | Pages |", "|---|---|---|"]
    for action, data in summary["actions"].items():
        lines.append(f"| {action} | {data['files']} | {data['pages']} |")
    lines += ["", f"Pages to fix or convert: **{summary['work_pages']}**. At ${rate_low:,.0f} to ${rate_high:,.0f} per page "
              f"that is roughly **${summary['work_pages'] * rate_low:,.0f} to ${summary['work_pages'] * rate_high:,.0f}** "
              "(estimate only).", "", "## Files by priority", "",
              "| ID | File | Pages | Action | Priority | Issues |", "|---|---|---|---|---|---|"]
    for rec in sorted(records, key=sort_key):
        issues = "; ".join(rec["issues"]) or "none found"
        lines.append(f"| {rec['id']} | {rec['file']} | {rec['pages']} | {rec['action']} | {rec['priority']} | {issues} |")
    lines += ["", "## Notes", "",
              "- Automated checks find only some problems. A person checks reading order, alt text, tables and "
              "colour contrast in every fixed file (NVDA screen reader, plus PAC or Acrobat on Windows).",
              "- Suggested actions are suggestions. The client's ADA coordinator or lawyer decides on exceptions "
              "such as archived content.",
              "- This is not a legal opinion or a certification of compliance.", ""]
    return "\n".join(lines)


def print_record(rec: dict) -> None:
    unit = {"PowerPoint": "slide", "Excel": "sheet"}.get(rec["type"], "page")
    count = plural(rec["pages"], unit) if isinstance(rec["pages"], int) else f"? {unit}s"
    print(f"\n{rec['file']}  ({rec['type']}, {count}, {rec['size_kb']} KB)")
    if rec["type"] == "PDF":
        struct = rec.get("struct") or {}
        print(f"  Tagged:        {rec['tagged']}" + (f"  ({struct.get('elements', 0)} tags, {struct.get('headings', 0)} headings, "
                                                   f"{struct.get('figures', 0)} images, {struct.get('tables', 0)} tables)" if struct.get("elements") else ""))
        print(f"  Text:          {rec['text']}")
        print(f"  Form fields:   {rec['form_fields']}")
        print(f"  Bookmarks:     {rec['bookmarks']}")
    print(f"  Title:         {rec['title'] or '(none)'}" + ("" if rec["has_title"] == "Yes" else "  <- problem"))
    print(f"  Language:      {rec['lang'] or '(none)'}")
    print(f"  Issues:        {'; '.join(rec['issues']) or 'none found by the automated checks'}")
    print(f"  Suggestion:    {rec['action']} ({rec['priority']} priority) - {rec['why']}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Check PDFs (and basic Office files) for accessibility problems and suggest an action for each.",
        epilog="Automated checks only: a person must still check every file you fix.")
    parser.add_argument("inputs", nargs="+", help="files or folders to check")
    parser.add_argument("--out", help="folder for inventory.xlsx, inventory.csv and summary.md "
                        "(without it, results are only printed)")
    parser.add_argument("--crawl", help="documents.csv from crawl_documents.py, to add URLs and link text")
    parser.add_argument("--client", default="", help="client name for the inventory")
    parser.add_argument("--archive-before", type=int, default=dt.date.today().year - 3,
                        help="files dated before this year are suggested for Archive (default: %(default)s)")
    parser.add_argument("--short-pages", type=int, default=1,
                        help="untagged PDFs with this many pages or fewer are suggested for Convert to web page "
                             "(default: %(default)s)")
    parser.add_argument("--sample-pages", type=int, default=15,
                        help="pages checked per PDF for real text vs scan (default: %(default)s)")
    parser.add_argument("--rate-low", type=float, default=5, help="low price per page for the estimate (default: 5)")
    parser.add_argument("--rate-high", type=float, default=25, help="high price per page for the estimate (default: 25)")
    args = parser.parse_args(argv)

    files = find_files(args.inputs)
    if not files:
        print("No PDF, Word, Excel or PowerPoint files found.", file=sys.stderr)
        return 1
    crawl_rows, not_checked = load_crawl(Path(args.crawl) if args.crawl else None)

    records = []
    for number, path in enumerate(files, start=1):
        print(f"[{number}/{len(files)}] {path.name}", file=sys.stderr)
        rec = analyze_file(path, args.sample_pages)
        row = crawl_rows.get(path.name)
        if row:
            rec["url"] = row.get("url", "")
            rec["found_on"] = row.get("first_found_on", "")
            rec["found_on_depth"] = row.get("found_on_depth", "")
            rec["link_text"] = row.get("first_link_text", "")
            rec["times_linked"] = int(row.get("times_linked") or 0)
        records.append(rec)
    mark_duplicates(records)
    for rec in records:
        suggest_action(rec, args.archive_before, args.short_pages)
        if rec["action"] in ("Fix", "Convert to web page") and str(rec.get("found_on_depth")) == "0":
            rec["priority"] = "High"
            rec["why"] += " Linked from the home page."

    summary = summarize(records, not_checked)
    source = ", ".join(args.inputs)
    if not args.out:
        for rec in records:
            print_record(rec)
        return 0

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    write_csv(records, out / "inventory.csv")
    write_xlsx(records, summary, not_checked, out / "inventory.xlsx", args.client, args.rate_low, args.rate_high, source)
    (out / "summary.md").write_text(summary_markdown(records, summary, args.client, source, args.rate_low, args.rate_high),
                                    encoding="utf-8")
    print(f"\nChecked {summary['total_files']} files ({summary['total_pages']} pages).")
    for action, data in summary["actions"].items():
        print(f"  {action:<20} {data['files']:>4} files {data['pages']:>6} pages")
    print(f"PDFs without tags: {summary['untagged']}  |  scanned: {summary['scanned']}  |  "
          f"no title: {summary['no_title']}  |  no language: {summary['no_lang']}  |  forms: {summary['forms']}")
    print(f"Wrote {out / 'inventory.xlsx'}, {out / 'inventory.csv'} and {out / 'summary.md'}")
    print("Reminder: automated checks only. Do not call any file 'compliant'.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
