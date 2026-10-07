#!/usr/bin/env python3
"""Build the two Office templates in templates/:

  snapshot-report-template.docx   the client-facing "Document accessibility snapshot"
  remediation-log-template.xlsx   the per-file log you fill in while fixing

You can edit both files directly in Word or Excel. Keep the {{PLACEHOLDERS}} in
the report if you want make_report.py to fill it. Run this script again only if
you want to start over from the default text (it overwrites both files).

The report template is itself accessible: real headings, table header rows,
a title and a language in its properties.
"""
from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from office_helpers import (add_page_number_footer, add_table, patch_app_properties,  # noqa: E402
                            set_base_font, set_properties)

KIT = Path(__file__).resolve().parent.parent
TEMPLATES = KIT / "templates"
REPORT = TEMPLATES / "snapshot-report-template.docx"
LOG = TEMPLATES / "remediation-log-template.xlsx"
APP_NAME = "Document accessibility kit"

# Sector paragraphs. make_report.py puts the chosen one where {{WHY_IT_MATTERS}} is.
WHY_IT_MATTERS = {
    "us-public": [
        "The US Justice Department's rule under Title II of the Americans with Disabilities Act (ADA) requires "
        "state and local governments, including school districts, public colleges and special districts, to make "
        "their web content meet WCAG 2.1 Level AA. Documents posted online, such as PDFs, Word files and "
        "spreadsheets, are part of that web content.",
        "Deadlines as reported when this snapshot was written (please verify the current dates, as they have "
        "changed before): 26 April 2027 for public entities serving 50,000 or more people, and 26 April 2028 "
        "for smaller ones. Special district governments were grouped with the smaller entities in the 2024 rule.",
    ],
    "us-health": [
        "Health care providers that receive funding from the US Department of Health and Human Services are "
        "covered by Section 504 of the Rehabilitation Act. The department's Section 504 rule also uses WCAG 2.1 "
        "Level AA for websites, including the documents posted on them.",
        "Deadlines depend on the size of the organization. Please verify the current dates with your compliance "
        "lead or lawyer.",
    ],
    "eu": [
        "The European Accessibility Act has applied since 28 June 2025 to many products and services sold in "
        "the EU, such as online shops, banking, transport and e-books. Its harmonised standard (EN 301 549) points "
        "to WCAG 2.1 Level AA, which covers documents you offer online.",
        "Very small service businesses (microenterprises) are exempt. Your legal adviser can confirm which parts "
        "apply to you.",
    ],
}
SECTOR_CLOSING = ("Accessible documents also help everyone else: people with low vision, people using a phone, "
                  "and anyone who needs to search, copy or translate the text.")


def build_report_template(path: Path) -> None:
    doc = Document()
    set_properties(doc, title="Document accessibility snapshot: {{CLIENT_NAME}}", author=APP_NAME,
                   language="en-US", subject="Document accessibility snapshot",
                   when=dt.datetime(2026, 10, 1, 9, 0))
    set_base_font(doc, "Calibri", 11)
    for level, size in ((1, 20), (2, 14)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    doc.add_heading("Document accessibility snapshot", level=1)
    doc.add_paragraph("{{CLIENT_NAME}}").runs[0].bold = True
    doc.add_paragraph("Website: {{SITE_URL}}")
    doc.add_paragraph("Prepared by {{PREPARED_BY}} on {{REPORT_DATE}}")
    note = doc.add_paragraph("{{TEMPLATE_NOTE}} How to use this template: run scripts/make_report.py to fill it "
                             "from the inventory, or replace every {{...}} by hand. Delete this paragraph before "
                             "sending.")
    note.runs[0].italic = True

    doc.add_heading("Summary", level=2)
    doc.add_paragraph("{{SUMMARY}}")

    doc.add_heading("Key numbers", level=2)
    add_table(doc, ["What we counted", "Number"], [
        ["Documents found", "{{TOTAL_FILES}}"],
        ["Total pages (slides and spreadsheet tabs included)", "{{TOTAL_PAGES}}"],
        ["PDF files", "{{PDF_COUNT}}"],
        ["Word files", "{{WORD_COUNT}}"],
        ["Excel files", "{{EXCEL_COUNT}}"],
        ["PowerPoint files", "{{PPT_COUNT}}"],
        ["PDFs without tags (screen readers cannot follow their structure)", "{{UNTAGGED_COUNT}}"],
        ["PDFs with scanned pages and no real text", "{{SCANNED_COUNT}}"],
        ["Files without a proper title", "{{NO_TITLE_COUNT}}"],
        ["PDF and Word files without a language setting", "{{NO_LANG_COUNT}}"],
        ["Files with fillable form fields", "{{FORMS_COUNT}}"],
        ["Exact duplicates", "{{DUPLICATE_COUNT}}"],
        ["Broken links to documents", "{{BROKEN_LINKS}}"],
        ["Links with unclear text, such as \"click here\"", "{{VAGUE_LINKS}}"],
    ], widths=[4.9, 1.4])

    doc.add_heading("Suggested action for each file", level=2)
    doc.add_paragraph("Every file is listed in the attached spreadsheet with a suggested action. These are "
                      "suggestions only: you decide what happens to each file, and your ADA coordinator or lawyer "
                      "decides whether any exception applies.")
    add_table(doc, ["Suggested action", "Files", "Pages", "What it means"],
              [["{{A.action}}", "{{A.files}}", "{{A.pages}}", "{{A.meaning}}"]], widths=[1.6, 0.7, 0.7, 3.3])

    doc.add_heading("Files to fix first", level=2)
    doc.add_paragraph("These files are used to apply for or use a service, or are linked from your home page.")
    add_table(doc, ["File", "Pages", "What we found", "Suggested action"],
              [["{{F.file}}", "{{F.pages}}", "{{F.issues}}", "{{F.action}}"]], widths=[1.9, 0.6, 2.6, 1.2])

    doc.add_heading("Estimate for fixing", level=2)
    doc.add_paragraph("{{WORK_FILES}} files ({{WORK_PAGES}} pages) are marked Fix or Convert to web page. At "
                      "{{RATE_RANGE}} per page, fixing them would cost about {{ESTIMATE_RANGE}}. This is an "
                      "estimate. We send a fixed quote after you decide which files to keep, archive or delete. "
                      "Files we can fix from the original Word files cost less.")

    doc.add_heading("Why this matters", level=2)
    doc.add_paragraph("{{WHY_IT_MATTERS}}")

    doc.add_heading("What this snapshot is and is not", level=2)
    for line in [
        "It lists the documents we could reach by following links from {{START_URL}}, up to {{DEPTH}} clicks "
        "deep, on {{CRAWL_DATE}}. Files behind logins, inside search tools or loaded by scripts may be missing.",
        "The checks are automated. They find common problems such as missing tags, scanned pages, and missing "
        "titles or language settings. They cannot tell whether the reading order, alt text or tables make sense. "
        "A person checks those when the files are fixed.",
        "It is not a legal opinion, a full audit of your website, or a certification. We do not certify files "
        "as \"ADA compliant\".",
        "Your ADA coordinator or lawyer decides whether any exception, such as the one for archived content, "
        "applies to a file.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Next steps", level=2)
    for line in [
        "Review the spreadsheet and mark each file: keep, fix, convert, archive or delete.",
        "Send us the original Word files where you have them. Fixing from the source is faster and costs less.",
        "We send a fixed quote. After you approve it, we fix the files, test each one with automated checks and "
        "a screen reader, and give you a log of what we did and the test results.",
    ]:
        doc.add_paragraph(line, style="List Number")

    doc.add_heading("How we checked", level=2)
    doc.add_paragraph("We read {{PAGES_CRAWLED}} web pages on {{CRAWL_DATE}}, following your site's robots.txt "
                      "rules and leaving at least one second between requests. For each file we checked tags, "
                      "real text versus scanned images, title, language, form fields, bookmarks and security "
                      "settings, using open-source tools.")
    doc.add_paragraph("Questions: {{PREPARED_BY}}, {{CONTACT}}")

    add_page_number_footer(doc, prefix="Document accessibility snapshot, {{CLIENT_NAME}}. Page ")
    doc.sections[0].footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.save(path)
    patch_app_properties(path, APP_NAME)


LOG_COLUMNS = [
    ("File ID", 8), ("File name", 32), ("URL", 36), ("Pages", 7),
    ("Before: tagged", 10), ("Before: text", 14), ("Before: title", 24), ("Before: language", 10),
    ("Before: veraPDF PDF/UA-1", 22),
    ("Work done", 40), ("Tools used", 26),
    ("After: tagged", 10), ("After: title", 24), ("After: language", 10), ("After: veraPDF PDF/UA-1", 22),
    ("PAC result", 14), ("Reading order (NVDA)", 14), ("Alt text", 12), ("Tables", 12), ("Links", 12),
    ("Forms", 12), ("Colour and contrast", 14), ("Checked by", 14), ("Check date", 12), ("Status", 24),
    ("Hours", 7), ("Notes", 40),
]
CHECK_VALUES = '"Pass,Fail,Not applicable,To do"'
PAC_VALUES = '"Pass,Pass with notes,Fail,To do"'
STATUS_VALUES = ('"Not started,In progress,Basic fixes only,Ready for human check,'
                 'Checked - ready to deliver,Delivered,Client to decide"')
LOG_HELP = [
    ("What this is", "One row per file you fix. It is your record of what you did and what you tested, and the "
     "client gets a copy with the fixed files."),
    ("Before / After", "Copy these from triage_pdfs.py and validate_pdfs.py, or let make_remediation_log.py fill "
     "them. 'veraPDF PDF/UA-1' is the machine check result (Pass, or the number of rules failed)."),
    ("Work done", "Plain words: 'Rebuilt in Word with headings and lists, exported tagged PDF', 'OCR, then tagged in "
     "Acrobat', 'Added alt text to 3 images (client approved the chart text)'."),
    ("Human checks", "PAC result, Reading order, Alt text, Tables, Links, Forms, Colour and contrast: use the "
     "drop-down (Pass, Fail, Not applicable, To do). Follow templates/human-check-checklist.md."),
    ("Checked by / Check date", "The person who did the screen-reader check, and when. Never fill these in for a "
     "check that did not happen."),
    ("Status", "Basic fixes only = title or language set, still untagged. In progress = tagged, but the automated "
     "checks still find problems. Ready for human check = tagged and the automated checks pass. "
     "Checked - ready to deliver = the human checks passed too."),
    ("Wording", "Write 'tested against WCAG 2.1 AA checks on [date]' or 'passed PAC and a screen-reader check'. "
     "Never write 'ADA compliant' or 'certified'."),
]


def build_log_template(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Log"
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for col, (name, width) in enumerate(LOG_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=name)
        cell.font, cell.fill = bold, fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "C2"
    last = sheet.cell(row=1, column=len(LOG_COLUMNS)).column_letter
    sheet.auto_filter.ref = f"A1:{last}1"
    headers = [name for name, _ in LOG_COLUMNS]

    def column(name):
        return sheet.cell(row=1, column=headers.index(name) + 1).column_letter

    checks = DataValidation(type="list", formula1=CHECK_VALUES, allow_blank=True)
    pac = DataValidation(type="list", formula1=PAC_VALUES, allow_blank=True)
    status = DataValidation(type="list", formula1=STATUS_VALUES, allow_blank=True)
    for validation in (checks, pac, status):
        sheet.add_data_validation(validation)
    for name in ("Reading order (NVDA)", "Alt text", "Tables", "Links", "Forms", "Colour and contrast"):
        checks.add(f"{column(name)}2:{column(name)}2000")
    pac.add(f"{column('PAC result')}2:{column('PAC result')}2000")
    status.add(f"{column('Status')}2:{column('Status')}2000")

    help_sheet = book.create_sheet("How to use")
    help_sheet.append(["Topic", "What to do"])
    for cell in help_sheet[1]:
        cell.font, cell.fill = bold, fill
    for topic, text in LOG_HELP:
        help_sheet.append([topic, text])
    help_sheet.column_dimensions["A"].width = 24
    help_sheet.column_dimensions["B"].width = 110
    for row in help_sheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    book.properties.title = "Document remediation log"
    book.properties.creator = APP_NAME
    book.properties.language = "en-US"
    book.save(path)


def main() -> int:
    TEMPLATES.mkdir(exist_ok=True)
    build_report_template(REPORT)
    build_log_template(LOG)
    print(f"Wrote {REPORT.relative_to(KIT)} and {LOG.relative_to(KIT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
