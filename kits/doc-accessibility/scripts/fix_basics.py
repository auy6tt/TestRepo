#!/usr/bin/env python3
"""Safe, automatic fixes for PDFs: title, language, title bar and (optionally) OCR.

What it changes, and nothing else:
  - Title: sets a real title (document info and XMP metadata) when the file has
    none, a placeholder like "untitled", or a file name as its title
  - Language: sets the document language (/Lang) when it is missing
  - Title bar: tells PDF readers to show the title instead of the file name
    (ViewerPreferences /DisplayDocTitle)
  - OCR, with --ocr: adds a text layer to scanned pages using ocrmypdf, if it and
    the tesseract program are installed

It never adds or changes tags, reading order, alt text, tables or form
fields. Those need a person. A file is NOT accessible just because this script
ran on it, and the log says what is still to do.

Originals are never changed. Fixed copies go to --out, with fix-log.csv.

Where titles come from (the first that exists):
  1. --titles CSV with columns file,title (make one with --suggest-titles)
  2. --title (only when fixing one file)
  3. the file's own title, if it is a real title
  4. the link text from the crawl (--crawl documents.csv), if it says enough
  5. the file name, tidied up ("council-agenda-2026-03-10" becomes
     "Council agenda, March 10, 2026")
Titles from 4 and 5 are marked "check" in the log.

Examples:
  python fix_basics.py downloads/ --crawl crawl/documents.csv --suggest-titles titles.csv
  python fix_basics.py downloads/ --out fixed-basics/ --titles titles.csv --lang en-US --ocr
  python fix_basics.py agenda.pdf --out fixed/ --title "Town Council Agenda, March 10, 2026"
"""
from __future__ import annotations

import argparse
import calendar
import csv
import datetime as dt
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pikepdf  # noqa: E402
from pypdf import PdfReader  # noqa: E402

from triage_pdfs import title_problem, triage_one  # noqa: E402

TESSERACT_LANGS = {"en": "eng", "es": "spa", "fr": "fra", "de": "deu", "it": "ita", "pt": "por", "nl": "nld",
                   "pl": "pol", "sv": "swe", "da": "dan", "fi": "fin", "nb": "nor", "no": "nor", "cs": "ces",
                   "el": "ell", "hu": "hun", "ro": "ron", "ru": "rus", "uk": "ukr", "zh": "chi_sim", "ja": "jpn",
                   "ko": "kor", "vi": "vie", "ar": "ara", "tl": "tgl", "ht": "hat"}
VAGUE = {"click here", "here", "download", "pdf", "link", "view", "open", "file", "document", "agenda", "minutes",
         "more", "read more", "details"}


def pretty_dates(text: str) -> str:
    def full(match):
        year, month, day = int(match.group(1)), int(match.group(2)), int(match.group(3))
        return f", {calendar.month_name[month]} {day}, {year}"

    def month_only(match):
        return f", {calendar.month_name[int(match.group(2))]} {match.group(1)}"

    text = re.sub(r"[\s_-]*\b((?:19|20)\d{2})[ _-](0[1-9]|1[0-2])[ _-](0[1-9]|[12]\d|3[01])\b", full, text)
    text = re.sub(r"[\s_-]*\b((?:19|20)\d{2})[ _-](0[1-9]|1[0-2])\b", month_only, text)
    text = re.sub(r"\bfy\s?(\d{2,4})\b", lambda m: "FY" + m.group(1), text, flags=re.I)
    return text


def title_from_filename(name: str) -> str:
    stem = Path(name).stem
    stem = re.sub(r"[_\s-]+(final|copy|print|v\d+|rev\d*|draft)$", "", stem, flags=re.I)
    words = pretty_dates(re.sub(r"[_]+", " ", stem))
    words = re.sub(r"(?<=[A-Za-z])-(?=[A-Za-z])", " ", words)
    words = re.sub(r"\s+", " ", words).strip(" ,-")
    return words[:1].upper() + words[1:]


def title_from_link_text(text: str) -> str:
    clean = re.sub(r"\([^)]*\b(pdf|word|excel|powerpoint|docx?|xlsx?|pptx?|kb|mb|pages?)\b[^)]*\)", "", text or "", flags=re.I)
    clean = re.sub(r"\[[^\]]*\]", "", clean)
    clean = re.sub(r"\s+", " ", clean).strip(" -:,.")
    if len(clean.split()) < 3 or clean.lower() in VAGUE:
        return ""
    return clean[:1].upper() + clean[1:]


def first_text(path: Path, limit: int = 160) -> str:
    try:
        text = PdfReader(str(path), strict=False).pages[0].extract_text() or ""
    except Exception:
        return ""
    return re.sub(r"\s+", " ", text).strip()[:limit]


def load_titles(csv_path: str | None) -> dict:
    if not csv_path:
        return {}
    with open(csv_path, newline="", encoding="utf-8") as handle:
        return {row["file"].strip(): row["title"].strip() for row in csv.DictReader(handle)
                if row.get("file") and row.get("title", "").strip()}


def load_link_texts(csv_path: str | None) -> dict:
    if not csv_path:
        return {}
    texts = {}
    with open(csv_path, newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            if row.get("saved_as"):
                texts[row["saved_as"]] = row.get("first_link_text", "")
    return texts


def load_skips(path: str | None) -> dict:
    """Files the inventory marks Delete or Archive (the client's decision wins)."""
    if not path:
        return {}
    if path.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        rows = book["Inventory"].iter_rows(values_only=True)
        headers = [str(h or "") for h in next(rows)]
        records = [dict(zip(headers, ["" if v is None else str(v) for v in values])) for values in rows]
        book.close()
    else:
        with open(path, newline="", encoding="utf-8") as handle:
            records = list(csv.DictReader(handle))
    skips = {}
    for row in records:
        decision = (row.get("Client decision") or "").strip().lower()
        action = (row.get("Suggested action") or "").strip()
        if any(word in decision for word in ("delete", "archive", "remove")):
            skips[row.get("File", "")] = f"client decision: {row.get('Client decision')}"
        elif not decision and action in ("Delete", "Archive"):
            skips[row.get("File", "")] = f"suggested action: {action}"
    return skips


def choose_title(path: Path, existing: str, args, titles: dict, link_texts: dict, single: bool) -> tuple[str, str]:
    if path.name in titles:
        return titles[path.name], "titles CSV"
    if single and args.title:
        return args.title, "--title"
    if existing and not title_problem(existing, path.name):
        return existing, "kept (already set)"
    from_link = title_from_link_text(link_texts.get(path.name, ""))
    if from_link:
        return from_link, "link text (check)"
    return title_from_filename(path.name), "file name (check)"


def ocr_ready() -> tuple[bool, str]:
    try:
        import ocrmypdf  # noqa: F401
    except ImportError:
        return False, "ocrmypdf is not installed (pip install ocrmypdf)"
    if not shutil.which("tesseract"):
        return False, "the tesseract program is not installed (Linux: sudo apt-get install -y tesseract-ocr)"
    return True, ""


def run_ocr(source: Path, target: Path, lang: str, title: str) -> tuple[bool, str]:
    tess = "+".join(TESSERACT_LANGS.get(part.split("-")[0].lower(), "eng") for part in lang.split("+"))
    command = [sys.executable, "-m", "ocrmypdf", "--output-type", "pdf", "--skip-text", "--deskew",
               "-l", tess, "--title", title, str(source), str(target)]
    result = subprocess.run(command, capture_output=True, text=True, timeout=3600)
    if result.returncode != 0:
        lines = (result.stderr or result.stdout or "").strip().splitlines()
        return False, lines[-1][:200] if lines else f"ocrmypdf exit code {result.returncode}"
    return True, f"OCR done with ocrmypdf (tesseract language: {tess})"


def set_metadata(source: Path, target: Path, title: str, lang: str, force_lang: bool) -> dict:
    with pikepdf.open(source) as pdf:
        root = pdf.Root
        before = {"lang": str(root.get("/Lang", "") or ""),
                  "bar": bool(root.get("/ViewerPreferences", {}).get("/DisplayDocTitle", False))
                  if isinstance(root.get("/ViewerPreferences"), pikepdf.Dictionary) else False}
        with pdf.open_metadata(set_pikepdf_as_editor=False) as meta:
            meta["dc:title"] = title
        pdf.docinfo["/Title"] = title
        lang_after = before["lang"]
        if lang and (force_lang or not before["lang"]):
            root.Lang = pikepdf.String(lang)
            lang_after = lang
        if not isinstance(root.get("/ViewerPreferences"), pikepdf.Dictionary):
            root.ViewerPreferences = pikepdf.Dictionary()
        root.ViewerPreferences.DisplayDocTitle = True
        pdf.save(target)
    return {"lang_before": before["lang"], "lang_after": lang_after, "bar_before": before["bar"]}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Safe automatic PDF fixes (title, language, title bar, OCR) with a log.")
    parser.add_argument("inputs", nargs="+", help="PDF files or folders")
    parser.add_argument("--out", help="folder for the fixed copies and fix-log.csv")
    parser.add_argument("--lang", default="en-US", help="language to set when missing (default: en-US)")
    parser.add_argument("--force-lang", action="store_true", help="replace an existing language too")
    parser.add_argument("--title", help="title to use (one input file only)")
    parser.add_argument("--titles", help="CSV with columns file,title")
    parser.add_argument("--crawl", help="documents.csv from crawl_documents.py (link text helps make titles)")
    parser.add_argument("--suggest-titles", metavar="CSV", help="write suggested titles to this CSV for review, then stop")
    parser.add_argument("--inventory", help="inventory.xlsx or inventory.csv: skip files marked Delete or Archive")
    parser.add_argument("--ocr", action="store_true", help="run OCR on scanned files (needs ocrmypdf and tesseract)")
    parser.add_argument("--dry-run", action="store_true", help="show what would change, write nothing")
    args = parser.parse_args(argv)

    files = []
    for item in args.inputs:
        path = Path(item)
        files += sorted(path.rglob("*.pdf")) + sorted(path.rglob("*.PDF")) if path.is_dir() else [path]
    files = [f for f in dict.fromkeys(files) if f.is_file() and f.suffix.lower() == ".pdf"]
    if not files:
        print("No PDF files found. (Fix Word, Excel and PowerPoint files in Office with its Accessibility Checker.)")
        return 1
    single = len(files) == 1
    if args.title and not single:
        print("--title works with one file only. Use --titles for many files.", file=sys.stderr)
        return 1
    titles, link_texts = load_titles(args.titles), load_link_texts(args.crawl)
    skips = load_skips(args.inventory)

    if args.suggest_titles:
        with open(args.suggest_titles, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(["file", "title", "current_title", "source", "first_text_on_page_1"])
            for path in files:
                rec = triage_one(path)
                title, source = choose_title(path, rec["title"], args, titles, link_texts, single)
                writer.writerow([path.name, title, rec["title"], source, first_text(path)])
        print(f"Wrote {args.suggest_titles}. Edit the 'title' column, then run again with --titles {args.suggest_titles}")
        return 0
    if not args.out and not args.dry_run:
        print("Give --out (a folder for the fixed copies) or --dry-run.", file=sys.stderr)
        return 1

    can_ocr, ocr_problem = ocr_ready() if args.ocr else (False, "")
    out = Path(args.out) if args.out else None
    if out:
        out.mkdir(parents=True, exist_ok=True)
    rows = []
    for path in files:
        rec = triage_one(path)
        row = {"file": path.name, "output": "", "changes": "", "title_before": rec["title"], "title_after": "",
               "title_source": "", "lang_before": rec["lang"], "lang_after": "", "title_bar_before": "",
               "title_bar_after": "", "ocr": "", "tagged": rec["tagged"], "still_to_do": "",
               "checked_by_person": "No", "date": dt.date.today().isoformat()}
        if path.name in skips:
            row["changes"] = f"Skipped ({skips[path.name]})"
            row["still_to_do"] = "Nothing, unless the client decides to keep it"
            rows.append(row)
            continue
        if not rec["readable"]:
            row["changes"] = "Not changed: " + "; ".join(rec["issues"])
            row["still_to_do"] = "Open by hand"
            rows.append(row)
            continue
        with pikepdf.open(path) as probe:
            encrypted = probe.is_encrypted
        if encrypted:
            row["changes"] = "Not changed: the file has security settings. Ask the client for an unprotected copy."
            row["still_to_do"] = "Get an unprotected copy, then fix"
            rows.append(row)
            continue
        title, source = choose_title(path, rec["title"], args, titles, link_texts, single)
        row["title_after"], row["title_source"] = title, source
        changes, todo = [], []
        if args.dry_run:
            row["changes"] = f"Would set title ({source})" + ("" if rec["lang"] else f", language {args.lang}")
            rows.append(row)
            continue
        target = out / path.name
        with tempfile.TemporaryDirectory() as tmp:
            source_file = path
            if rec["needs_ocr"] == "Yes":
                if args.ocr and can_ocr:
                    ocr_file = Path(tmp) / path.name
                    ok, message = run_ocr(path, ocr_file, args.lang, title)
                    row["ocr"] = message if ok else f"OCR failed: {message}"
                    if ok:
                        source_file = ocr_file
                        changes.append("added OCR text layer")
                        todo.append("Check the OCR text against the scan (names, numbers, dates)")
                    else:
                        todo.append("Needs OCR")
                else:
                    row["ocr"] = "Needed, not run" + (f" ({ocr_problem})" if ocr_problem else " (use --ocr)")
                    todo.append("Needs OCR")
            ocr_used = source_file != path
            result = set_metadata(source_file, target, title, args.lang,
                                  args.force_lang or (ocr_used and not rec["lang"]))
        if title != rec["title"]:
            changes.append("set title")
        if result["lang_after"] != rec["lang"]:
            changes.append(f"set language to {result['lang_after']}")
        if not result["bar_before"]:
            changes.append("title bar shows the title")
        row.update(output=str(target), lang_after=result["lang_after"], title_bar_before="Yes" if result["bar_before"] else "No",
                   title_bar_after="Yes", changes="; ".join(changes) or "no change needed")
        if rec["tagged"] != "Yes":
            todo.append("Still untagged: needs tags, reading order and alt text (rebuild from the Word file, or tag "
                        "in Acrobat), then a person checks it")
        else:
            todo.append("A person checks reading order, alt text and tables (NVDA, plus PAC or Acrobat)")
        if "check" in source:
            todo.insert(0, "Check the title")
        row["still_to_do"] = "; ".join(todo)
        rows.append(row)

    for row in rows:
        print(f"{row['file']}: {row['changes']}")
        if row["title_after"] and row["title_after"] != row["title_before"]:
            print(f"    title: '{row['title_before'] or '(none)'}' -> '{row['title_after']}'  [{row['title_source']}]")
        if row["ocr"]:
            print(f"    OCR: {row['ocr']}")
        print(f"    still to do: {row['still_to_do']}")
    if out and not args.dry_run:
        with open(out / "fix-log.csv", "w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        print(f"\nWrote {len(rows)} rows to {out / 'fix-log.csv'}")
    print("These are basic fixes only. No file is accessible until it is tagged and checked by a person.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
