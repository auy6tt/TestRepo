#!/usr/bin/env python3
"""
Self-test for the weekly-digest kit. Runs offline in about a minute.

  python tests/selftest.py

It copies the fictional sample sites to a temporary folder, serves them on
this computer and checks that:

  - robots.txt rules are read correctly (wildcards, longest match, groups)
  - the first run saves snapshots, and the "week 41" run finds the changes:
    a new agenda PDF, a replaced PDF, a new RSS notice, a new open-data
    record, a changed legislation page, a page that disappeared (404) and a
    page that robots.txt blocks
  - unchanged pages are not downloaded again (HTTP 304) and a repeat run
    reports no changes
  - API keys written as ${NAME} are sent to the site but never saved
  - private individuals' records and unlisted fields never reach the snapshots
  - --dry-run, --only and broken settings behave
  - the build refuses unchecked items and makes valid HTML, text, PDF, XLSX
    and ICS files

Nothing in the kit folder is changed. Exit code 0 means every check passed.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

KIT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(KIT / "scripts"))

import run_demo  # noqa: E402
from watch_sources import RobotsRules, find_keywords, html_to_extracted  # noqa: E402

SAMPLE = KIT / "samples" / "battery-storage-ohio-week41"
UA = "WeeklyDigestWatcher/1.0 (+mailto:test@example.org)"
results: list[tuple[bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((bool(ok), name))
    print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f"  ({detail})" if detail and not ok else ""))


def run(args: list[str]) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable] + args, capture_output=True, text=True)


def by_id(report: dict, key: str) -> dict:
    return {entry["id"]: entry for entry in report[key]}


def unit_tests() -> None:
    print("Rules and text extraction")
    rules = RobotsRules("User-agent: *\nDisallow: /search/\nAllow: /search/help\nDisallow: /*?print=\n", UA)
    check("robots: Disallow blocks a folder", not rules.allows("https://a.example/search/q"))
    check("robots: the longer Allow rule wins", rules.allows("https://a.example/search/help"))
    check("robots: * wildcard in a rule", not rules.allows("https://a.example/page?print=1"))
    check("robots: other pages allowed", rules.allows("https://a.example/page"))
    anchored = RobotsRules("User-agent: *\nDisallow: /*.pdf$\n", UA)
    check("robots: $ anchors the end", not anchored.allows("https://a.example/x.pdf")
          and anchored.allows("https://a.example/x.pdf?v=2"))
    groups = RobotsRules("User-agent: WeeklyDigestWatcher\nDisallow: /private/\n\n"
                         "User-agent: *\nDisallow: /\n", UA)
    check("robots: our own group overrides *", groups.allows("https://a.example/public")
          and not groups.allows("https://a.example/private/x"))
    check("robots: empty Disallow allows all", RobotsRules("User-agent: *\nDisallow:\n", UA).allows("https://a.example/x"))
    check("robots: crawl-delay is read", RobotsRules("User-agent: *\nCrawl-delay: 4\n", UA).crawl_delay == 4.0)

    html = (b"<html><head><title>T</title></head><body><nav>Menu</nav><main><p>Hello\n  world</p>"
            b"<table><tr><td>A</td><td><a href='/docs/x.pdf#p2'>B<br>C</a></td></tr></table>"
            b"<script>bad()</script></main><footer>Foot</footer></body></html>")
    ext = html_to_extracted(html, "https://site.example/dir/page")
    check("html: source line breaks become spaces, cells joined", ext.lines == ["Hello world", "A | B; C"], str(ext.lines))
    check("html: links made absolute, fragment dropped",
          [l["url"] for l in ext.links] == ["https://site.example/docs/x.pdf"], str(ext.links))
    check("keywords: whole words, plurals, wildcards",
          find_keywords(["Battery storage", "Bessemer Road"], ["battery", "BESS"]) == ["battery"]
          and find_keywords(["new batteries"], ["batter*"]) == ["batter*"])


def watcher_tests(tmp: Path) -> None:
    digest = tmp / "digest"
    digest.mkdir()
    shutil.copy(SAMPLE / "sources.yaml", digest / "sources.yaml")
    live = tmp / "live"
    server = run_demo.DemoServer(live)
    watch = str(KIT / "scripts" / "watch_sources.py")
    base = [watch, str(digest / "sources.yaml"), "--map-url", server.map_url, "--delay", "0.05"]

    def watch_run(name: str, *extra: str) -> tuple[int, dict]:
        out = tmp / f"{name}.json"
        proc = run(base + ["--out", str(out)] + list(extra))
        report = json.loads(out.read_text()) if out.exists() else {}
        if proc.returncode not in (0, 1):
            print(proc.stdout, proc.stderr)
        return proc.returncode, report

    try:
        print("Watcher: week 40 (first run)")
        run_demo.sync_site(SAMPLE / "site" / "week40", live, run_demo.WEEK40_TIME)
        code, rep = watch_run("week40")
        s = rep.get("summary", {})
        check("first run exits 0", code == 0, str(code))
        check("first run saves 7 snapshots", s.get("first_run") == 7, str(s))
        check("robots.txt blocks the search page", [e["id"] for e in rep.get("skipped", [])]
              == ["tarrow-county-site-search-battery-storage"], str(rep.get("skipped")))
        check("the blocked page was never requested", not any("/search/" in r["path"] for r in server.log))
        check("robots.txt is read before pages", server.log and server.log[0]["path"].endswith("/robots.txt"))
        check("every request sent our user agent", all(r["user_agent"].startswith("WeeklyDigestWatcher/") for r in server.log))
        snaps = (digest / "data" / "snapshots").rglob("*")
        text = " ".join(p.read_text(errors="ignore") for p in snaps if p.is_file())
        records = json.loads((digest / "data" / "snapshots" / "halden-county-open-data-zoning-applications-json"
                              / "records.json").read_text())
        check("private individuals' records are never stored",
              "Resident (sample)" not in text and "V-2026-109" not in records
              and not any("applicant_type: individual" in r["text"] for r in records.values()))
        check("fields not listed (phone numbers) are never stored", "555-01" not in text)

        print("Watcher: week 41 (changes)")
        run_demo.sync_site(SAMPLE / "site" / "week41", live, run_demo.WEEK41_TIME)
        server.log.clear()
        code, rep = watch_run("week41")
        s = rep.get("summary", {})
        changed = by_id(rep, "changed") if rep else {}
        check("week 41 counts: 5 changed, 1 unchanged, 1 skipped, 1 error",
              (s.get("changed"), s.get("unchanged"), s.get("skipped"), s.get("errors")) == (5, 1, 1, 1), str(s))
        agendas = changed.get("tarrow-county-commissioners-agendas", {})
        check("new agenda PDF link found", any(l["url"].endswith("2026-10-13-agenda.pdf") for l in agendas.get("new_links", [])))
        docs = agendas.get("documents", [])
        check("new agenda PDF fetched and its text saved",
              docs and any("Resolution 2026-118" in line for line in docs[0]["excerpt"] + [""])
              or (docs and "2026-118" in (digest / docs[0]["text_file"]).read_text()))
        check("ignore_patterns drop the 'Page last updated' line",
              not any("Page last updated" in line for line in agendas.get("added", []) + agendas.get("removed", [])))
        check("keywords found in the agenda change", "moratorium" in agendas.get("keyword_hits", []))
        pdf = changed.get("kestrel-regional-planning-commission-current-agenda-pdf", {})
        check("replaced PDF: new case found with its page number",
              any(line.startswith("[p. 1]") and "RZ-26-014" in line for line in pdf.get("added", [])))
        rss = changed.get("wrenfield-township-public-notices-rss", {})
        check("RSS: one matching notice, one filtered out by keywords",
              len(rss.get("new_records", [])) == 1 and rss.get("other_records_not_matching_keywords") == 1)
        js = changed.get("halden-county-open-data-zoning-applications-json", {})
        check("JSON: new BESS record found, others filtered",
              [r["id"] for r in js.get("new_records", [])] == ["CU-2026-031"]
              and js.get("other_records_not_matching_keywords") == 2)
        city = changed.get("city-of-marlowe-falls-pending-legislation", {})
        check("selector limits the HTML page to one section",
              not any("Resolution 2026-22" in line for line in city.get("added", []) + city.get("removed", [])))
        check("amended ordinance PDF followed", any("600 kWh" in " ".join(d["excerpt"]) for d in city.get("documents", [])))
        errors = {e["id"]: e for e in rep.get("errors", [])}
        moved = errors.get("halden-county-zoning-commission-hearings-page", {})
        check("moved page reported as 404 with a hint", "404" in moved.get("error", "") and moved.get("hint"))
        bza = [r for r in server.log if "/bza/decisions/" in r["path"]]
        check("unchanged page answered 304 (not downloaded again)", bza and bza[-1]["status"] == 304, str(bza))
        check("changes.md written next to changes.json", (tmp / "week41.md").exists())

        print("Watcher: repeat run, dry run, --only, broken settings, site down")
        code, rep = watch_run("again")
        check("repeat run finds no changes", rep["summary"]["changed"] == 0 and rep["summary"]["unchanged"] == 6,
              str(rep["summary"]))
        page = live / "www.tarrowcounty.example" / "commissioners" / "agendas" / "index.html"
        page.write_text(page.read_text().replace("</tbody>", "<tr><td>October 20, 2026</td><td>Special session</td>"
                                                 "<td>October 8, 2026</td><td>-</td><td>-</td></tr></tbody>"))
        snapshot = digest / "data" / "snapshots" / "tarrow-county-commissioners-agendas" / "current.txt"
        before = hashlib.sha256(snapshot.read_bytes()).hexdigest()
        code, rep = watch_run("dry", "--dry-run", "--only", "commissioners")
        check("--only checks just the matching source", rep["summary"]["sources"] == 1, str(rep["summary"]))
        check("--dry-run reports the change", rep["summary"]["changed"] == 1)
        check("--dry-run leaves snapshots alone", hashlib.sha256(snapshot.read_bytes()).hexdigest() == before)
        code, rep = watch_run("after-dry", "--only", "commissioners")
        check("the change is still reported after a dry run", rep["summary"]["changed"] == 1)

        bad = tmp / "bad.yaml"
        bad.write_text("sources:\n  - name: Broken\n    type: xml\n  - name: Broken\n    url: https://a.example/\n    type: html\n")
        proc = run([watch, str(bad)])
        check("broken sources file stops with exit code 2", proc.returncode == 2
              and "'type' must be one of" in proc.stderr and "'url' must start" in proc.stderr, proc.stderr[-300:])
    finally:
        server.stop()
    code, rep = watch_run("down")
    check("site down: every source is an error and exit code is 1",
          code == 1 and rep["summary"]["errors"] == rep["summary"]["sources"], f"{code} {rep.get('summary')}")


def feature_tests(tmp: Path) -> None:
    print("Watcher: API keys and keyword filters")
    site = tmp / "feature-site"
    api = site / "api.example" / "v1"
    api.mkdir(parents=True)
    (api / "items.json").write_text('{"results": [{"id": "A1", "title": "Battery storage bid"}]}')
    digest = tmp / "feature-digest"
    digest.mkdir()
    (digest / "sources.yaml").write_text(
        "settings:\n  user_agent: 'WeeklyDigestWatcher/1.0 (+mailto:test@example.org)'\n"
        "  keywords: [battery]\nsources:\n  - name: Keyed API\n"
        "    url: https://api.example/v1/items.json?api_key=${SELFTEST_KEY}\n"
        "    type: json\n    items_path: results\n    id_field: id\n    terms_checked: 2026-10-07\n")
    server = run_demo.DemoServer(site)
    watch = [str(KIT / "scripts" / "watch_sources.py"), str(digest / "sources.yaml"),
             "--map-url", server.map_url, "--delay", "0"]
    try:
        import os
        env = {**os.environ, "SELFTEST_KEY": "not-a-real-key-42"}
        missing = {k: v for k, v in os.environ.items() if k != "SELFTEST_KEY"}
        proc = subprocess.run([sys.executable] + watch + ["--out", str(tmp / "k0.json")],
                              capture_output=True, text=True, env=missing)
        check("a missing ${KEY} is reported clearly", "SELFTEST_KEY is not set" in proc.stdout, proc.stdout[-300:])
        subprocess.run([sys.executable] + watch + ["--out", str(tmp / "k1.json")], capture_output=True, env=env)
        (api / "items.json").write_text('{"results": [{"id": "A1", "title": "Battery storage bid"}, '
                                        '{"id": "A2", "title": "Paving bid"}]}')
        later = (api / "items.json").stat().st_mtime + 60  # a newer "last modified" date, like a real update
        os.utime(api / "items.json", (later, later))
        subprocess.run([sys.executable] + watch + ["--out", str(tmp / "k2.json")], capture_output=True, env=env)
        check("the key is sent to the site", any("not-a-real-key-42" in r["path"] for r in server.log))
        written = " ".join(p.read_text(errors="ignore") for p in list(tmp.glob("k*.*")) +
                           list((digest / "data").rglob("*")) if p.is_file())
        check("the key never appears in reports or snapshots", "not-a-real-key-42" not in written)
        report = json.loads((tmp / "k2.json").read_text())
        check("changes without keywords are listed apart from 'changed'",
              report["summary"]["changed"] == 0 and report["summary"]["no_keyword_match"] == 1, str(report["summary"]))
    finally:
        server.stop()


def build_tests(tmp: Path) -> None:
    print("Build")
    build = str(KIT / "scripts" / "build_digest.py")
    items = SAMPLE / "items.reviewed.yaml"
    out = tmp / "out"
    proc = run([build, str(items), "--out", str(out)])
    check("build exits 0", proc.returncode == 0, proc.stderr[-400:])
    name = "battery-storage-zoning-watch-ohio-2026-10-08"
    files = {ext: out / f"{name}.{ext}" for ext in ("html", "txt", "pdf", "xlsx", "ics")}
    check("all five files written", all(f.exists() and f.stat().st_size > 500 for f in files.values()))
    if not all(f.exists() for f in files.values()):
        return

    html = files["html"].read_text()
    check("email has no <style>, <script> or external files",
          "<style" not in html and "<script" not in html and not re.search(r'\s(src|href)="(?!https?:|mailto:)', html))
    check("email is styled inline", html.count('style="') > 100)
    check("email is mobile-friendly (viewport, 640px max width)",
          'name="viewport"' in html and "max-width:640px" in html)
    check("approved items in, rejected item out", html.count("Checked against the source") == 5
          and "overlay district back on" not in html)
    check("every item links to its official source",
          all(u in html for u in ("2026-10-13-agenda.pdf", "ZTA-2026-03-hearing-notice.pdf",
                                  "current-agenda.pdf", "CU-2026-031", "ord-2026-41-amended.pdf")))
    check("internal review notes are not published", "Rejected." not in html and "review_notes" not in html)
    check("plain-text version has the dates", "Tue, Oct 13, 9:30 AM ET" in files["txt"].read_text())

    from pypdf import PdfReader
    reader = PdfReader(str(files["pdf"]))
    pdf_text = " ".join(page.extract_text() or "" for page in reader.pages)
    check("PDF has the issue, page numbers and every item",
          "Battery Storage Zoning Watch" in pdf_text and "Page 1 of" in pdf_text
          and all(t in pdf_text for t in ("Resolution 2026-118", "ZTA-2026-03", "RZ-26-014", "CU-2026-031", "2026-41")),
          f"{len(reader.pages)} pages")

    from openpyxl import load_workbook
    wb = load_workbook(files["xlsx"])
    items_ws = wb["Items"]
    check("tracker has About, Items and Dates sheets", wb.sheetnames == ["About", "Items", "Dates"], str(wb.sheetnames))
    check("tracker has one row per approved item", items_ws.max_row == 6, str(items_ws.max_row))
    check("tracker 'Days left' is a live formula", str(items_ws["G2"].value).startswith("=IF(E2"))
    check("tracker links to the sources", items_ws["L2"].hyperlink is not None
          and items_ws["L2"].hyperlink.target.startswith("https://"))

    raw = files["ics"].read_bytes()
    lines = raw.split(b"\r\n")
    check("calendar uses CRLF line endings only", raw.count(b"\n") == raw.count(b"\r\n"))
    check("calendar lines are at most 75 bytes", max(len(l) for l in lines) <= 75)
    text = raw.decode().replace("\r\n ", "")
    check("calendar has 5 events with matching BEGIN/END",
          text.count("BEGIN:VEVENT") == 5 == text.count("END:VEVENT") and text.count("BEGIN:VCALENDAR") == 1)
    check("calendar times are converted to UTC (9:30 AM EDT = 13:30Z)", "DTSTART:20261013T133000Z" in text)

    unchecked = tmp / "unchecked.yaml"
    unchecked.write_text(items.read_text().replace("  - id: W41-03\n    status: approved",
                                                   "  - id: W41-03\n    status: needs_review", 1))
    proc = run([build, str(unchecked), "--out", str(tmp / "gate")])
    check("unchecked items stop the build (exit code 3)", proc.returncode == 3, str(proc.returncode))
    proc = run([build, str(unchecked), "--out", str(tmp / "gate"), "--preview", "--formats", "html"])
    preview = tmp / "gate" / f"{name}-PREVIEW.html"
    check("--preview builds a clearly marked draft", proc.returncode == 0 and preview.exists()
          and "PREVIEW, NOT CHECKED" in preview.read_text())


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="weekly-digest-selftest-"))
    try:
        unit_tests()
        watcher_tests(tmp)
        feature_tests(tmp)
        build_tests(tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    failed = [name for ok, name in results if not ok]
    print(f"\n{len(results) - len(failed)} of {len(results)} checks passed.")
    if failed:
        print("Failed: " + "; ".join(failed))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
