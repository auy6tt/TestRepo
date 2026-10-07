#!/usr/bin/env python3
"""
Build the files for one issue from a reviewed items file.

In the output folder it makes:

  <name>.html   the email, with every style inline, ready to paste or import
                into your email tool
  <name>.txt    the same email as plain text
  <name>.pdf    the same issue as a print-ready PDF
  <name>.xlsx   a tracker spreadsheet: one row per item, plus a dates sheet
  <name>.ics    a calendar file with every hearing, vote and deadline

Only items with status "approved" go into the files. If any item is still
"draft" or "needs_review" the build stops, so nothing unchecked goes out by
mistake. Add --preview to build a clearly marked draft for your own review.

Run it from the kits/weekly-digest folder:

  python scripts/build_digest.py digests/ohio-bess/issues/2026-W41/items.yaml
  python scripts/build_digest.py path/to/items.yaml --preview
  python scripts/build_digest.py path/to/items.yaml --formats html,pdf

The items format is explained in templates/items.example.yaml and checked
against templates/items.schema.json.
"""

from __future__ import annotations

import argparse
import datetime as dt
import glob
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path
from urllib.parse import quote, urlsplit
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

try:
    import jinja2
    import jsonschema
    import yaml
    from markupsafe import Markup
except ImportError as exc:  # pragma: no cover - message for the user
    sys.exit(
        f"Missing Python package: {exc.name}\n"
        "Install the kit's requirements first, for example:\n"
        "  python3 -m venv .venv && .venv/bin/pip install -r requirements.txt"
    )

KIT = Path(__file__).resolve().parent.parent
SCHEMA_PATH = KIT / "templates" / "items.schema.json"
EMAIL_TEMPLATE = KIT / "templates" / "digest-email.html.j2"
FORMATS = ("html", "txt", "pdf", "xlsx", "ics")
UNCHECKED = ("draft", "needs_review")

DEFAULT_DISCLAIMER = (
    "This brief summarises public documents for information only. It is not legal advice and "
    "does not tell you what you must do. Dates and details can change after it is sent, so "
    "confirm with the official source before you act, and ask a qualified professional about "
    "your own situation."
)

TZ_LABELS = {
    "America/New_York": "ET", "America/Detroit": "ET", "America/Indiana/Indianapolis": "ET",
    "America/Chicago": "CT", "America/Denver": "MT", "America/Phoenix": "MST",
    "America/Los_Angeles": "PT", "America/Anchorage": "AKT", "Pacific/Honolulu": "HT",
    "Europe/London": "UK time", "Europe/Paris": "CET", "Europe/Berlin": "CET",
    "Europe/Brussels": "CET", "UTC": "UTC",
}
TZ_NAMES = {"ET": "Eastern Time", "CT": "Central Time", "MT": "Mountain Time",
            "PT": "Pacific Time", "UTC": "UTC"}

SANS = "-apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif"
SERIF = "Georgia, 'Times New Roman', Times, serif"

COLORS = {
    "page_bg": "#eef1f5", "ink": "#1c2733", "muted": "#5c6b7a", "rule": "#dfe5ec",
    "rule_light": "#edf1f5", "brand": "#16365c", "accent": "#0f766e", "tile_bg": "#f3f6f9",
    "amber_bg": "#fff7e8", "amber_rule": "#f1d39b", "amber_ink": "#8a4b00", "link": "#0b5cad",
    "ok": "#1f7a4d", "danger": "#b42318", "danger_bg": "#fef3f2", "foot_bg": "#f7f9fb",
    "high": "#c2410c",
}

PERSONAL_DATA_PATTERNS = [
    (re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+"), "an email address"),
    (re.compile(r"(?<!\d)(?:\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]\d{3}[\s.-]\d{4}(?!\d)"), "a phone number"),
]
ADVICE_PATTERN = re.compile(
    r"\b(you|your company|your business) (must|need to|should|are required to|have to|owe)\b"
    r"|\bwe (recommend|advise)\b", re.IGNORECASE)


class BuildError(Exception):
    pass


# --------------------------------------------------------------------------
# Loading and checking
# --------------------------------------------------------------------------

class _Loader(yaml.SafeLoader):
    """yaml.safe_load, but dates stay as text so they can be checked exactly."""


_Loader.yaml_implicit_resolvers = {
    key: [(tag, rx) for tag, rx in resolvers if tag != "tag:yaml.org,2002:timestamp"]
    for key, resolvers in yaml.SafeLoader.yaml_implicit_resolvers.items()
}


def load_items(path: Path) -> dict:
    try:
        text = path.read_text(encoding="utf-8")
    except FileNotFoundError:
        raise BuildError(f"File not found: {path}") from None
    try:
        data = json.loads(text) if path.suffix.lower() == ".json" else yaml.load(text, Loader=_Loader)
    except (yaml.YAMLError, ValueError) as exc:
        raise BuildError(f"{path} could not be read: {exc}") from None
    if not isinstance(data, dict):
        raise BuildError(f"{path} must start with 'issue:' and 'items:'. See templates/items.example.yaml.")
    return data


def _where(path, doc) -> str:
    parts = list(path)
    if len(parts) >= 2 and parts[0] == "items" and isinstance(parts[1], int):
        item = doc["items"][parts[1]] if parts[1] < len(doc.get("items") or []) else {}
        label = f"item {parts[1] + 1}"
        if isinstance(item, dict) and item.get("id"):
            label += f" ({item['id']})"
        rest = ".".join(str(p) for p in parts[2:])
        return f"{label}{', ' + rest if rest else ''}"
    return ".".join(str(p) for p in parts) or "file"


def _friendly(err) -> str:
    field = str(err.absolute_path[-1]) if err.absolute_path else ""
    if err.validator == "required":
        match = re.search(r"'(.+?)' is a required property", err.message)
        name = match.group(1) if match else err.message
        if name in ("checked_by", "checked_on"):
            return (f"missing '{name}': approved items need checked_by (your initials) and "
                    "checked_on (the date you checked the source)")
        return f"missing '{name}'"
    if err.validator == "additionalProperties":
        match = re.search(r"\((.+?) (?:was|were) unexpected\)", err.message)
        return f"unknown field {match.group(1) if match else ''} (check the spelling)"
    if err.validator == "pattern":
        if "more_dates" in [str(p) for p in err.absolute_path]:
            return f"{err.instance!r} must look like 2026-10-13 or 2026-10-13T09:30"
        if field in ("date", "checked_on"):
            return f"{err.instance!r} must be a date like 2026-10-08"
        if field == "deadline" or (field == "date" and "more_dates" in str(err.absolute_path)):
            return f"{err.instance!r} must look like 2026-10-13 or 2026-10-13T09:30"
        if field in ("source_url", "website"):
            return f"{err.instance!r} must start with https://"
        if field == "id":
            return f"{err.instance!r}: use only letters, numbers, dots, dashes and underscores"
        if field == "contact_email":
            return f"{err.instance!r} is not an email address"
        if field == "accent_color":
            return f"{err.instance!r} must be a color like #16365c"
        return f"{err.instance!r} is not in the expected format"
    if err.validator == "enum":
        return f"{err.instance!r} must be one of: {', '.join(map(str, err.validator_value))}"
    if err.validator == "minLength":
        return f"too short (at least {err.validator_value} characters)"
    if err.validator == "maxLength":
        return f"too long ({len(err.instance)} characters, the limit is {err.validator_value})"
    if err.validator == "type":
        wanted = err.validator_value
        wanted = " or ".join(wanted) if isinstance(wanted, list) else wanted
        return f"must be {'text' if wanted == 'string' else wanted}"
    return err.message


def parse_when(value: str) -> tuple[dt.date, dt.time | None]:
    value = value.strip().replace(" ", "T", 1)
    if "T" in value:
        moment = dt.datetime.fromisoformat(value)
        return moment.date(), moment.time().replace(second=0, microsecond=0)
    return dt.date.fromisoformat(value), None


def check(doc: dict) -> tuple[list[str], list[str]]:
    """Schema check plus a few common-sense checks. Returns (errors, warnings)."""
    schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    validator = jsonschema.Draft202012Validator(schema)
    errors = [f"{_where(e.absolute_path, doc)}: {_friendly(e)}"
              for e in sorted(validator.iter_errors(doc), key=lambda e: [str(p) for p in e.absolute_path])]
    if errors:
        return errors, []

    warnings: list[str] = []
    issue = doc["issue"]
    try:
        ZoneInfo(issue.get("timezone") or "America/New_York")
    except (ZoneInfoNotFoundError, ValueError):
        errors.append(f"issue.timezone: {issue.get('timezone')!r} is not a time zone name like America/New_York")
    try:
        issue_date = dt.date.fromisoformat(issue["date"])
    except ValueError:
        errors.append(f"issue.date: {issue['date']!r} is not a real date")
        issue_date = None

    seen = set()
    contact = (issue.get("contact_email") or "").lower()
    for number, item in enumerate(doc["items"], start=1):
        label = f"item {number} ({item['id']})"
        if item["id"] in seen:
            errors.append(f"{label}: the id {item['id']!r} is used twice")
        seen.add(item["id"])
        try:
            item_date = dt.date.fromisoformat(item["date"])
            if item.get("checked_on"):
                dt.date.fromisoformat(item["checked_on"])
        except ValueError:
            errors.append(f"{label}: a date is not a real date")
            continue
        whens = ([("deadline", item["deadline"])] if item.get("deadline") else []) + \
                [(f"more_dates {d['label']!r}", d["date"]) for d in item.get("more_dates") or []]
        for name, value in whens:
            try:
                day, _ = parse_when(value)
            except ValueError:
                errors.append(f"{label}: {name} {value!r} is not a real date or time")
                continue
            if day < item_date:
                warnings.append(f"{label}: {name} {value} is before the item's date {item['date']}")
            if issue_date and day < issue_date and item["status"] != "rejected":
                warnings.append(f"{label}: {name} {value} has already passed on the send date; "
                                "it is left out of 'Coming up' and the calendar")
        if item["status"] == "rejected":
            continue
        text = " ".join(str(item.get(k) or "") for k in ("title", "summary", "why_it_matters")) + " " + \
            " ".join(str(v) for v in (item.get("details") or {}).values())
        for pattern, what in PERSONAL_DATA_PATTERNS:
            for found in pattern.findall(text):
                if found.lower() != contact:
                    warnings.append(f"{label}: contains {what} ({found}). Remove it unless it belongs "
                                    "to a public office or a business, not a private person.")
        advice = ADVICE_PATTERN.search(text)
        if advice:
            warnings.append(f"{label}: '{advice.group(0)}' sounds like advice. Describe what the "
                            "document says instead of telling readers what to do.")
    return errors, warnings


# --------------------------------------------------------------------------
# Display helpers
# --------------------------------------------------------------------------

def fmt_date(day: dt.date) -> str:
    return f"{day:%b} {day.day}, {day.year}"


def fmt_date_long(day: dt.date) -> str:
    return f"{day:%A}, {day:%B} {day.day}, {day.year}"


def fmt_day(day: dt.date) -> str:
    return f"{day:%a}, {day:%b} {day.day}"


def fmt_time(moment: dt.time) -> str:
    hour = moment.hour % 12 or 12
    return f"{hour}:{moment.minute:02d} {'AM' if moment.hour < 12 else 'PM'}"


def paragraphs(text: str | None) -> list[str]:
    if not text:
        return []
    return [" ".join(p.split()) for p in re.split(r"\n\s*\n", str(text).strip()) if p.strip()]


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def mix(color: str, other: str, amount: float) -> str:
    a = [int(color[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(other[i:i + 2], 16) for i in (1, 3, 5)]
    return "#" + "".join(f"{round(x + (y - x) * amount):02x}" for x, y in zip(a, b))


def css_string(text: str) -> str:
    return re.sub(r"[<>\n\r]", " ", text).replace("\\", "\\\\").replace('"', '\\"')


def make_styles(c: dict, print_mode: bool) -> dict:
    """Inline styles for the template. The print version uses a smaller type scale."""
    pad = "20px" if print_mode else "24px"
    scale = 0.86 if print_mode else 1.0

    def text(size, weight=400, lh=1.5, color=None, family=SANS, extra=""):
        size = round(size * scale, 1)
        return (f"margin:0; font-family:{family}; font-size:{size:g}px; line-height:{lh}; "
                f"font-weight:{weight}; color:{color or c['ink']};" + (f" {extra}" if extra else ""))

    page_bg = "#ffffff" if print_mode else c["page_bg"]
    st = {
        "body": f"margin:0; padding:0; background-color:{page_bg}; -webkit-text-size-adjust:100%; -ms-text-size-adjust:100%;",
        "outer": f"background-color:{page_bg}; border-collapse:collapse;",
        "outer_cell": "padding:0;" if print_mode else "padding:24px 10px;",
        "card": ("background-color:#ffffff; border-collapse:separate;" if print_mode else
                 f"max-width:640px; background-color:#ffffff; border:1px solid {c['rule']}; "
                 "border-radius:10px; border-collapse:separate; overflow:hidden;"),
        "preview_bar": text(13, 700, 1.4, "#ffffff", extra=f"background-color:{c['danger']}; padding:10px {pad};"),
        "sample_bar": text(12, 600, 1.45, c["muted"], extra=f"background-color:#f1f4f7; padding:9px {pad}; border-bottom:1px solid {c['rule']};"),
        "masthead": f"background-color:{c['brand']}; padding:26px {pad} 24px;",
        "kicker": text(11, 700, 1.2, c["brand_soft"], extra="letter-spacing:0.16em; text-transform:uppercase; margin-bottom:10px;"),
        "h1": text(28, 700, 1.18, "#ffffff", SERIF),
        "subtitle": text(15, 400, 1.45, c["brand_text"], extra="margin-top:8px;"),
        "block_intro": f"padding:24px {pad} 4px;",
        "label": text(12, 700, 1.3, c["accent"], extra="letter-spacing:0.12em; text-transform:uppercase; margin-bottom:8px;"),
        "intro": text(16, 400, 1.6, extra="margin-bottom:12px;"),
        "block_stats": f"padding:6px {pad} 20px;",
        "tile": f"background-color:{c['tile_bg']}; border-radius:8px; padding:12px 6px 11px;",
        "tile_gap": "font-size:1px; line-height:1px;",
        "stat_value": text(26, 700, 1.1, c["brand"]),
        "stat_label": text(11, 600, 1.3, c["muted"], extra="letter-spacing:0.05em; text-transform:uppercase; margin-top:4px;"),
        "block_dates": f"padding:0 {pad} 22px;",
        "dates_box": f"background-color:{c['amber_bg']}; border:1px solid {c['amber_rule']}; border-radius:8px; border-collapse:separate;",
        "dates_title": text(12, 700, 1.3, c["amber_ink"], extra="letter-spacing:0.12em; text-transform:uppercase; padding:14px 18px 6px;"),
        "dates_row": f"padding:8px 18px 9px; border-bottom:1px solid {c['amber_rule']};",
        "dates_row_last": "padding:8px 18px 9px;",
        "date_when": text(12, 700, 1.4, c["amber_ink"], extra="letter-spacing:0.04em; text-transform:uppercase; margin-bottom:2px;"),
        "date_what": text(14, 400, 1.45),
        "date_where": text(14, 400, 1.45, c["muted"]),
        "dates_note": text(12, 400, 1.5, c["muted"], extra=f"padding:8px 18px 14px; border-top:1px solid {c['amber_rule']};"),
        "block_section": f"padding:10px {pad} 12px;",
        "section": text(13, 700, 1.3, c["brand"], extra=f"letter-spacing:0.1em; text-transform:uppercase; border-bottom:2px solid {c['brand']}; padding-bottom:6px;"),
        "section_count": text(13, 400, 1.3, c["muted"]),
        "block_item": f"padding:0 {pad} {12 if print_mode else 16}px;",
        "item": f"background-color:#ffffff; border:1px solid {c['rule']}; border-left:4px solid {c['accent']}; border-radius:6px; border-collapse:separate;",
        "item_high": f"background-color:#ffffff; border:1px solid {c['rule']}; border-left:4px solid {c['high']}; border-radius:6px; border-collapse:separate;",
        "item_cell": "padding:13px 16px 11px;" if print_mode else "padding:16px 18px 14px;",
        "item_meta": text(12, 600, 1.4, c["muted"], extra="letter-spacing:0.04em; text-transform:uppercase; margin-bottom:6px;"),
        "badge": text(11, 700, 1.2, "#ffffff", extra=f"background-color:{c['high']}; border-radius:3px; padding:2px 6px; letter-spacing:0.06em;"),
        "item_title": text(18, 700, 1.35, extra="margin-bottom:8px;"),
        "item_title_link": f"color:{c['ink']}; text-decoration:none;",
        "item_text": text(15, 400, 1.6, extra="margin-bottom:10px;"),
        "why_label": f"color:{c['brand']};",
        "details": "margin:2px 0 12px; border-collapse:collapse;",
        "details_key": text(13, 600, 1.4, c["muted"], extra=f"padding:6px 10px 6px 0; border-top:1px solid {c['rule_light']};"),
        "details_value": text(13, 400, 1.4, extra=f"padding:6px 0; border-top:1px solid {c['rule_light']};"),
        "pill_table": "margin:0 0 8px; border-collapse:separate;",
        "pill": text(13, 600, 1.4, c["amber_ink"], extra=f"background-color:{c['amber_bg']}; border:1px solid {c['amber_rule']}; border-radius:4px; padding:6px 10px;"),
        "pill_where": text(13, 400, 1.4, c["amber_ink"]),
        "pill_past": text(13, 600, 1.4, c["muted"], extra=f"background-color:{c['tile_bg']}; border:1px solid {c['rule']}; border-radius:4px; padding:6px 10px;"),
        "source": text(13, 400, 1.5, c["muted"], extra="margin-top:6px;"),
        "link": f"color:{c['link']}; text-decoration:underline;",
        "checked": text(12, 600, 1.5, c["ok"], extra="margin-top:4px;"),
        "unchecked": text(12, 700, 1.5, c["danger"], extra=f"margin-top:6px; background-color:{c['danger_bg']}; padding:4px 8px;"),
        "quiet": text(15, 400, 1.6, extra=f"padding:4px {pad} 22px;"),
        "footer": f"padding:22px {pad} 24px; background-color:{c['foot_bg']}; border-top:1px solid {c['rule']};",
        "footer_head": text(13, 700, 1.4, extra="margin-bottom:4px;"),
        "footer_text": text(13, 400, 1.6, c["muted"], extra="margin-bottom:14px;"),
        "footer_last": text(13, 400, 1.6, c["muted"]),
        "legal": text(12, 400, 1.6, c["muted"], extra="margin-top:16px; text-align:center;"),
        "legal_link": f"color:{c['muted']}; text-decoration:underline;",
    }
    return {k: Markup(v) for k, v in st.items()}


# --------------------------------------------------------------------------
# The issue, ready to display
# --------------------------------------------------------------------------

def build_view(doc: dict, preview: bool, sources_checked: int | None, paper: str) -> dict:
    issue = dict(doc["issue"])
    schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    for key in schema["properties"]["issue"]["properties"]:
        issue.setdefault(key, None)  # optional fields the template may ask for
    tz_name = issue.get("timezone") or "America/New_York"
    tz = ZoneInfo(tz_name)
    tz_label = issue.get("timezone_label") or TZ_LABELS.get(tz_name, "")
    issue_date = dt.date.fromisoformat(issue["date"])
    if issue.get("sources_checked") is None and sources_checked is not None:
        issue["sources_checked"] = sources_checked

    colors = dict(COLORS)
    if issue.get("accent_color"):
        colors["brand"] = issue["accent_color"]
    colors["brand_soft"] = mix(colors["brand"], "#ffffff", 0.62)
    colors["brand_text"] = mix(colors["brand"], "#ffffff", 0.84)

    wanted = ("approved",) + (UNCHECKED if preview else ())
    chosen = [i for i in doc["items"] if i["status"] in wanted]

    def when(value: str) -> dict:
        day, moment = parse_when(value)
        text = f"{day:%a}, {fmt_date(day)}"
        local = None
        if moment:
            text += f", {fmt_time(moment)}" + (f" {tz_label}" if tz_label else "")
            local = dt.datetime.combine(day, moment, tzinfo=tz)
        return {"day": day, "time": moment, "local": local, "text": text, "past": day < issue_date}

    views = []
    for item in chosen:
        raw_dates = []
        if item.get("deadline"):
            raw_dates.append({"label": item.get("deadline_label") or "Date", "value": item["deadline"],
                              "location": item.get("deadline_location")})
        for extra in item.get("more_dates") or []:
            raw_dates.append({"label": extra["label"], "value": extra["date"], "location": extra.get("location")})
        dates = []
        for d in raw_dates:
            w = when(d["value"])
            dates.append({**d, **w, "when": w["text"]})
        dates.sort(key=lambda d: (d["day"], d["time"] or dt.time(0)))
        item_date = dt.date.fromisoformat(item["date"])
        meta_parts = [item.get("jurisdiction"), fmt_date(item_date)]
        views.append({
            "id": item["id"],
            "title": item["title"].strip(),
            "short_title": (item.get("short_title") or item["title"]).strip(),
            "summary_paras": paragraphs(item["summary"]),
            "summary_text": " ".join(paragraphs(item["summary"])),
            "why": " ".join(paragraphs(item.get("why_it_matters"))),
            "category": (item.get("category") or "Updates").strip(),
            "jurisdiction": item.get("jurisdiction") or "",
            "meta": " \u00b7 ".join(p for p in meta_parts if p),
            "date": item_date,
            "source_url": item["source_url"],
            "source_name": item["source_name"],
            "source_ref": item.get("source_ref") or "",
            "details": [(str(k), str(v)) for k, v in (item.get("details") or {}).items()],
            "dates": dates,
            "tags": item.get("tags") or [],
            "high": item.get("priority") == "high",
            "status": item["status"],
            "unchecked": item["status"] in UNCHECKED,
            "checked_on": item.get("checked_on"),
            "checked_display": fmt_date(dt.date.fromisoformat(item["checked_on"])) if item.get("checked_on") else "",
        })

    order = list(doc.get("sections") or [])
    names = order + sorted({v["category"] for v in views} - set(order))
    sections = [{"name": n, "entries": [v for v in views if v["category"] == n]} for n in names]
    sections = [s for s in sections if s["entries"]]

    upcoming = []
    for v in views:
        for d in v["dates"]:
            if not d["past"]:
                place = re.sub(r"^(city|town|village|county|borough|township) of ", "",
                               v["jurisdiction"].split(",")[0].strip(), flags=re.IGNORECASE)
                upcoming.append({
                    "day": fmt_day(d["day"]),
                    "time": (fmt_time(d["time"]) + (f" {tz_label}" if tz_label else "")) if d["time"] else "",
                    "label": d["label"], "what": v["short_title"],
                    "where": "" if place.lower() in v["short_title"].lower() else v["jurisdiction"],
                    "sort": (d["day"], d["time"] or dt.time(0)), "item": v, "date": d,
                })
    upcoming.sort(key=lambda d: d["sort"])

    stats = [{"value": len(views), "label": "item this week" if len(views) == 1 else "items this week"},
             {"value": len(upcoming), "label": "date coming up" if len(upcoming) == 1 else "dates coming up"}]
    if issue.get("sources_checked"):
        stats.append({"value": issue["sources_checked"], "label": "sources checked"})

    subtitle_parts = [issue.get("region"), issue.get("edition"), fmt_date_long(issue_date)]
    if issue.get("number"):
        subtitle_parts.append(f"Issue {issue['number']}")
    subtitle = " \u00b7 ".join(str(p) for p in subtitle_parts if p)

    intro = paragraphs(issue.get("intro"))
    issue["intro_paras"] = intro
    if intro:
        first = re.split(r"(?<=[.!?])\s", intro[0])[0]
        preheader = first if len(first) <= 150 else first[:147] + "..."
    else:
        preheader = "; ".join(v["title"] for v in views[:3])

    n = issue.get("sources_checked")
    method = f"Each week we check {n} official sources" if n else "Each week we check official sources"
    if issue.get("sources_note"):
        method += f" ({issue['sources_note']})"
    method += (" for changes. A person reads every item against its source document before the brief "
               "is sent. We link to the official documents rather than copying them.")

    base_name = issue.get("file_name") or slugify(f"{issue['title']} {issue.get('region') or ''}") + "-" + issue["date"]
    if preview:
        base_name += "-PREVIEW"
    website = issue.get("website") or ""
    host = urlsplit(website).netloc if website else ""
    domain = host or "weekly-digest.invalid"
    manage = issue.get("manage_url") or f"mailto:{issue['contact_email']}?subject={quote('Manage my subscription')}"
    tz_long = TZ_NAMES.get(tz_label, tz_label)
    footer_bits = [issue["title"], issue.get("edition") or fmt_date(issue_date), issue["brand"], "Not legal advice"]

    return {
        "issue": issue, "preview": preview, "subtitle": subtitle, "preheader": preheader,
        "stats": stats, "dates": upcoming, "sections": sections, "items": views,
        "unchecked_count": sum(1 for v in views if v["unchecked"]),
        "method_text": method, "disclaimer": issue.get("disclaimer") or DEFAULT_DISCLAIMER,
        "base_name": base_name, "ics_name": base_name + ".ics",
        "tz": tz, "tz_name": tz_name, "tz_label": tz_label,
        "tz_note": f"Times are {tz_long}." if tz_long else "",
        "issue_date": issue_date, "colors": colors, "website_host": host, "domain": domain,
        "manage_url": manage, "paper": paper,
        "print_footer": Markup(css_string(" \u00b7 ".join(str(b) for b in footer_bits if b))),
        "fonts": {"sans": Markup(SANS), "serif": Markup(SERIF)},
    }


# --------------------------------------------------------------------------
# Outputs
# --------------------------------------------------------------------------

def render_html(view: dict, print_mode: bool, template_path: Path) -> str:
    env = jinja2.Environment(loader=jinja2.FileSystemLoader(str(template_path.parent)),
                             autoescape=True, trim_blocks=False, lstrip_blocks=False,
                             undefined=jinja2.StrictUndefined)
    template = env.get_template(template_path.name)
    html = template.render(**view, print_mode=print_mode, st=make_styles(view["colors"], print_mode))
    return "\n".join(line.rstrip() for line in html.splitlines()) + "\n"


def render_text(view: dict) -> str:
    issue, width = view["issue"], 72
    out: list[str] = []

    def para(text: str, indent: str = "", hang: str = "") -> None:
        out.extend(textwrap.wrap(text, width, initial_indent=indent,
                                 subsequent_indent=indent + hang) or [indent])

    if view["preview"]:
        out += ["*** PREVIEW, NOT CHECKED. DO NOT SEND. ***", ""]
    if issue.get("sample_notice"):
        para(issue["sample_notice"])
        out.append("")
    out += [issue["title"].upper(), view["subtitle"], issue["brand"], "=" * width, ""]
    if issue["intro_paras"]:
        out += ["THIS WEEK", ""]
        for p in issue["intro_paras"]:
            para(p)
            out.append("")
    if view["dates"]:
        out += ["COMING UP", ""]
        for d in view["dates"]:
            when = d["day"] + (f", {d['time']}" if d["time"] else "")
            para(f"- {when}: {d['label']}: {d['what']}" + (f" ({d['where']})" if d["where"] else ""), "", "  ")
        out += ["  (All dates are in the attached calendar file.)", ""]
    if not view["sections"]:
        para("Nothing new this week. We checked every source and found no relevant changes.")
        out.append("")
    for section in view["sections"]:
        out += [section["name"].upper(), "-" * len(section["name"]), ""]
        for it in section["entries"]:
            para(("[TOP ITEM] " if it["high"] else "") + it["title"], "")
            para(it["meta"], "  ")
            for p in it["summary_paras"]:
                para(p, "  ")
            if it["why"]:
                para(f"Why it matters: {it['why']}", "  ")
            for key, value in it["details"]:
                para(f"{key}: {value}", "  ")
            for d in it["dates"]:
                para(f"{d['label']}: {d['when']}" + (f", {d['location']}" if d.get("location") else ""), "  ")
            para(f"Source: {it['source_name']}" + (f", {it['source_ref']}" if it["source_ref"] else ""), "  ")
            out.append(f"  {it['source_url']}")
            if it["unchecked"]:
                out.append(f"  NOT CHECKED YET ({it['status']})")
            else:
                out.append(f"  Checked against the source on {it['checked_display']}.")
            out.append("")
    out += ["=" * width, "HOW THIS BRIEF IS MADE"]
    para(view["method_text"])
    out += ["", "NOT LEGAL ADVICE"]
    para(view["disclaimer"])
    if issue.get("attribution"):
        out.append("")
        para(issue["attribution"])
    out.append("")
    para(f"Questions or corrections: reply to this email or write to {issue['contact_email']}.")
    out.append("")
    para(issue["brand"] + (f", {issue['mailing_address']}" if issue.get("mailing_address") else ""))
    out.append(f"Manage or cancel your subscription: {view['manage_url']}")
    return "\n".join(out).rstrip() + "\n"


def find_chromium() -> str | None:
    """A Chromium or Chrome program on this computer, if there is one."""
    if os.environ.get("CHROMIUM_PATH") and Path(os.environ["CHROMIUM_PATH"]).exists():
        return os.environ["CHROMIUM_PATH"]
    for name in ("chromium", "chromium-browser", "google-chrome", "google-chrome-stable", "chrome"):
        found = shutil.which(name)
        if found:
            return found
    roots = [os.environ.get("PLAYWRIGHT_BROWSERS_PATH", ""), str(Path.home() / ".cache" / "ms-playwright"),
             "/opt/pw-browsers"]
    for root in roots:
        if not root:
            continue
        for pattern in ("chromium-*/chrome-linux*/chrome", "chromium_headless_shell-*/chrome-linux*/headless_shell",
                        "chromium-*/chrome-mac*/Chromium.app/Contents/MacOS/Chromium",
                        "chromium-*/chrome-win*/chrome.exe"):
            matches = sorted(glob.glob(os.path.join(root, pattern)), reverse=True)
            if matches:
                return matches[0]
    mac = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
    return mac if Path(mac).exists() else None


def html_to_pdf(html_text: str, out_path: Path, paper: str) -> str:
    """Print the HTML to PDF with headless Chromium. Returns how it was done."""
    problems = []
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        sync_playwright = None
        problems.append("the Python package 'playwright' is not installed")
    chrome = find_chromium()
    if sync_playwright:
        # First Playwright's own browser; if that is missing, any Chromium found on this computer.
        for executable in [None] + ([chrome] if chrome else []):
            try:
                with sync_playwright() as p:
                    browser = p.chromium.launch(executable_path=executable) if executable else p.chromium.launch()
                    page = browser.new_page()
                    page.set_content(html_text, wait_until="load")
                    page.pdf(path=str(out_path), format=paper, print_background=True,
                             prefer_css_page_size=True)
                    browser.close()
                return "Playwright and Chromium" + (f" ({executable})" if executable else "")
            except Exception as exc:  # try the next way
                problems.append(f"Playwright: {(str(exc).strip().splitlines() or ['error'])[0][:200]}")
    if chrome:
        with tempfile.TemporaryDirectory() as tmp:
            page_file = Path(tmp) / "issue.html"
            page_file.write_text(html_text, encoding="utf-8")
            cmd = [chrome, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                   f"--print-to-pdf={out_path}", "--virtual-time-budget=5000"]
            if hasattr(os, "geteuid") and os.geteuid() == 0:
                cmd.append("--no-sandbox")
            result = subprocess.run(cmd + [page_file.as_uri()], capture_output=True, text=True, timeout=120)
            if result.returncode == 0 and out_path.exists() and out_path.stat().st_size > 0:
                return f"Chromium command line ({chrome})"
            problems.append(f"Chromium command line: exit code {result.returncode}")
    else:
        problems.append("no Chromium or Chrome browser was found")
    raise BuildError("Could not make the PDF: " + "; ".join(problems) +
                     ". Install a browser with: python -m playwright install chromium")


def add_pdf_metadata(path: Path, view: dict) -> None:
    try:
        from pypdf import PdfReader, PdfWriter
        writer = PdfWriter(clone_from=PdfReader(str(path)))
        issue = view["issue"]
        writer.add_metadata({
            "/Title": f"{issue['title']}: {issue.get('edition') or fmt_date(view['issue_date'])}",
            "/Author": issue["brand"],
            "/Subject": view["subtitle"],
            "/Creator": "weekly-digest build_digest.py",
        })
        with open(path, "wb") as handle:
            writer.write(handle)
    except Exception:
        pass  # metadata is a nice extra; the PDF is fine without it


def build_xlsx(view: dict, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties
    from openpyxl.worksheet.datavalidation import DataValidation

    issue = view["issue"]
    brand = view["colors"]["brand"].lstrip("#").upper()
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=brand)
    input_fill = PatternFill("solid", fgColor="FFF8D6")
    link_font = Font(name="Arial", size=10, color="0B5CAD", underline="single")
    thin = Side(style="thin", color="D5DCE4")
    border = Border(bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    tz_suffix = f" ({view['tz_label']})" if view["tz_label"] else ""
    status_choices = "New,Watching,Acting on it,Done,Not relevant"

    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)

    def header(ws, columns):
        for col, (title, width) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
        ws.print_title_rows = "1:1"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    def set_when(cell, d):
        if d is None:
            cell.value = None
            return
        if d["time"]:
            cell.value = dt.datetime.combine(d["day"], d["time"])
            cell.number_format = "ddd mmm d, yyyy h:mm AM/PM"
        else:
            cell.value = d["day"]
            cell.number_format = "ddd mmm d, yyyy"

    def days_left_rules(ws, column, last_row):
        rng = f"{column}2:{column}{max(last_row, 2)}"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"AND(ISNUMBER({column}2),{column}2<0)"], font=Font(name="Arial", color="8A949E")))
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"AND(ISNUMBER({column}2),{column}2>=0,{column}2<=7)"],
            fill=PatternFill("solid", fgColor="FDE2E1"), font=Font(name="Arial", bold=True, color="9B1C1C")))
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"AND(ISNUMBER({column}2),{column}2>7,{column}2<=14)"],
            fill=PatternFill("solid", fgColor="FFF1D6")))

    # Items sheet: one row per item.
    ws = wb.active
    ws.title = "Items"
    columns = [("ID", 9), ("Title", 44), ("Jurisdiction", 18), ("Category", 18),
               (f"Next date{tz_suffix}", 22), ("What happens", 24), ("Days left", 9),
               ("Summary", 70), ("Why it matters", 48), ("Record date", 13), ("Source", 36),
               ("Link", 12), ("Tags", 22), ("Checked on", 12), ("Your status", 14), ("Your notes", 32)]
    header(ws, columns)
    status_rule = DataValidation(type="list", formula1=f'"{status_choices}"', allow_blank=True,
                                 showDropDown=False)
    status_rule.prompt, status_rule.promptTitle = "Pick a status for your own tracking.", "Your status"
    ws.add_data_validation(status_rule)
    row = 1
    for it in view["items"]:
        row += 1
        nxt = next((d for d in it["dates"] if not d["past"]), it["dates"][-1] if it["dates"] else None)
        values = [it["id"], it["title"], it["jurisdiction"], it["category"], None,
                  nxt["label"] if nxt else "", None, it["summary_text"], it["why"],
                  it["date"], it["source_name"] + (f", {it['source_ref']}" if it["source_ref"] else ""),
                  "Open source", ", ".join(it["tags"]),
                  dt.date.fromisoformat(it["checked_on"]) if it["checked_on"] else ("NOT CHECKED" if it["unchecked"] else None),
                  "New", None]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font, cell.alignment, cell.border = font, wrap, border
        set_when(ws.cell(row=row, column=5), nxt)
        ws.cell(row=row, column=7, value=f'=IF(E{row}="","",INT(E{row})-TODAY())').number_format = "0"
        ws.cell(row=row, column=7).alignment = top
        ws.cell(row=row, column=10).number_format = "mmm d, yyyy"
        ws.cell(row=row, column=14).number_format = "mmm d, yyyy"
        ws.cell(row=row, column=2).font = bold
        link = ws.cell(row=row, column=12)
        link.hyperlink, link.font = it["source_url"], link_font
        for col in (15, 16):
            ws.cell(row=row, column=col).fill = input_fill
        status_rule.add(f"O{row}")
        longest = max(len(it["summary_text"]) / 75, len(it["why"]) / 52, len(it["title"]) / 46, 1)
        ws.row_dimensions[row].height = min(15 * (int(longest) + 1) + 4, 300)
    days_left_rules(ws, "G", row)

    # Dates sheet: every date, soonest first.
    dates_ws = wb.create_sheet("Dates")
    date_cols = [("Item", 9), ("Title", 46), (f"Date{tz_suffix}", 24), ("What happens", 34),
                 ("Days left", 9), ("Jurisdiction", 18), ("Where", 40), ("Link", 12)]
    header(dates_ws, date_cols)
    every = sorted(((it, d) for it in view["items"] for d in it["dates"]),
                   key=lambda pair: (pair[1]["day"], pair[1]["time"] or dt.time(0)))
    drow = 1
    for it, d in every:
        drow += 1
        values = [it["id"], it["title"], None, d["label"], None, it["jurisdiction"], d.get("location") or "",
                  "Open source"]
        for col, value in enumerate(values, start=1):
            cell = dates_ws.cell(row=drow, column=col, value=value)
            cell.font, cell.alignment, cell.border = font, wrap, border
        set_when(dates_ws.cell(row=drow, column=3), d)
        dates_ws.cell(row=drow, column=5, value=f'=IF(C{drow}="","",INT(C{drow})-TODAY())').number_format = "0"
        link = dates_ws.cell(row=drow, column=8)
        link.hyperlink, link.font = it["source_url"], link_font
        dates_ws.row_dimensions[drow].height = 30 if len(it["title"]) > 46 else 16
    days_left_rules(dates_ws, "E", drow)

    # About sheet: what this is and how to use it.
    about = wb.create_sheet("About")
    about.column_dimensions["A"].width = 24
    about.column_dimensions["B"].width = 100
    about.page_setup.orientation = "landscape"
    about.page_setup.fitToWidth, about.page_setup.fitToHeight = 1, 0
    about.sheet_properties.pageSetUpPr.fitToPage = True
    rows = [
        (issue["title"], ""),
        ("Issue", view["subtitle"]),
        ("From", issue["brand"]),
        ("Items in this issue", len(view["items"])),
        ("Official sources checked", issue.get("sources_checked") or ""),
        ("", ""),
        ("How to use this file", "Items has one row per item. Dates lists every hearing, vote and "
                                 "deadline, soonest first. Click 'Open source' to read the official document."),
        ("Your columns", "The yellow columns (Your status, Your notes) are for you. Pick a status from "
                         "the list and add notes; nothing else needs editing."),
        ("Days left", "Counted from today's date by a formula, so it updates each time you open the file. "
                      "Red: 7 days or less. Amber: 8 to 14 days. Grey: the date has passed."),
        ("Times", f"Times are {TZ_NAMES.get(view['tz_label'], view['tz_name'])}." if view["tz_label"] else view["tz_name"]),
        ("", ""),
        ("Not legal advice", view["disclaimer"]),
        ("How it is made", view["method_text"]),
        ("Contact", issue["contact_email"]),
    ]
    if issue.get("sample_notice"):
        rows.insert(1, ("Sample", issue["sample_notice"]))
    if issue.get("attribution"):
        rows.append(("Data credits", issue["attribution"]))
    for number, (key, value) in enumerate(rows, start=1):
        a, b = about.cell(row=number, column=1, value=key), about.cell(row=number, column=2, value=value)
        a.font, b.font = bold, font
        a.alignment = b.alignment = wrap
        if isinstance(value, str) and len(value) > 95:
            about.row_dimensions[number].height = 15 * (len(value) // 95 + 1) + 2
    about["A1"].font = Font(name="Arial", size=14, bold=True, color=brand)
    about.row_dimensions[1].height = 22
    if view["preview"]:
        about.insert_rows(1)
        about["A1"] = "PREVIEW, NOT CHECKED. Do not send."
        about["A1"].font = Font(name="Arial", size=12, bold=True, color="B42318")
    wb.move_sheet("About", offset=-2)
    wb.active = wb.sheetnames.index("Items")
    wb.save(path)


def ics_escape(text: str) -> str:
    return (str(text).replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,")
            .replace("\r\n", "\\n").replace("\n", "\\n"))


def ics_fold(line: str) -> str:
    """Split long lines at 75 bytes as the calendar standard (RFC 5545) requires."""
    data = line.encode("utf-8")
    parts, limit = [], 75
    while len(data) > limit:
        cut = limit
        while cut > 0 and (data[cut] & 0xC0) == 0x80:  # never split a multi-byte character
            cut -= 1
        parts.append(data[:cut])
        data = data[cut:]
        limit = 74  # continuation lines start with a space
    parts.append(data)
    return "\r\n ".join(p.decode("utf-8") for p in parts)


def build_ics(view: dict, path: Path) -> int:
    issue = view["issue"]
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//weekly-digest kit//build_digest.py//EN",
             "CALSCALE:GREGORIAN", "METHOD:PUBLISH",
             f"X-WR-CALNAME:{ics_escape(issue['title'] + ': ' + (issue.get('edition') or issue['date']))}",
             f"X-WR-TIMEZONE:{view['tz_name']}"]
    count = 0
    pairs = sorted(((it, d) for it in view["items"] for d in it["dates"] if not d["past"]),
                   key=lambda pair: (pair[1]["day"], pair[1]["time"] or dt.time(0)))
    for it, d in pairs:
        count += 1
        key = f"{it['id']}|{d['label']}|{d['value']}"
        uid = f"{slugify(issue['title'])}-{it['id'].lower()}-{hashlib.sha1(key.encode()).hexdigest()[:10]}@{view['domain']}"
        lines += ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{stamp}"]
        if d["local"]:
            start = d["local"].astimezone(dt.timezone.utc)
            end = start + dt.timedelta(hours=1)
            lines += [f"DTSTART:{start:%Y%m%dT%H%M%SZ}", f"DTEND:{end:%Y%m%dT%H%M%SZ}"]
            trigger = "-P1D"
        else:
            lines += [f"DTSTART;VALUE=DATE:{d['day']:%Y%m%d}",
                      f"DTEND;VALUE=DATE:{d['day'] + dt.timedelta(days=1):%Y%m%d}"]
            trigger = "-PT15H"  # 9:00 the day before
        prefix = "NOT CHECKED: " if it["unchecked"] else ""
        description = (f"{it['summary_text']}\n\nSource: {it['source_name']}"
                       + (f", {it['source_ref']}" if it["source_ref"] else "")
                       + f"\n{it['source_url']}\n\nFrom {issue['title']}"
                       + (f", {issue['edition']}" if issue.get("edition") else "")
                       + f" ({issue['brand']}). Check the official source before you rely on this "
                         "date. Not legal advice.")
        lines += [f"SUMMARY:{ics_escape(prefix + d['label'] + ': ' + it['short_title'])}",
                  f"DESCRIPTION:{ics_escape(description)}",
                  f"URL:{it['source_url']}",
                  f"CATEGORIES:{ics_escape(it['category'])}",
                  "TRANSP:TRANSPARENT"]
        if d.get("location"):
            lines.append(f"LOCATION:{ics_escape(d['location'])}")
        lines += ["BEGIN:VALARM", "ACTION:DISPLAY",
                  f"DESCRIPTION:{ics_escape('Reminder: ' + d['label'])}",
                  f"TRIGGER:{trigger}", "END:VALARM", "END:VEVENT"]
    lines.append("END:VCALENDAR")
    with open(path, "w", encoding="utf-8", newline="") as handle:
        handle.write("\r\n".join(ics_fold(line) for line in lines) + "\r\n")
    return count


# --------------------------------------------------------------------------
# Command line
# --------------------------------------------------------------------------

def sources_checked_from(changes_path: Path) -> int | None:
    try:
        summary = json.loads(changes_path.read_text(encoding="utf-8"))["summary"]
        return int(summary["changed"]) + int(summary["unchanged"]) + int(summary["first_run"])
    except (OSError, ValueError, KeyError, TypeError):
        print(f"Warning: could not read the summary in {changes_path}")
        return None


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass
    parser = argparse.ArgumentParser(description="Build the email, PDF, tracker and calendar for one issue.")
    parser.add_argument("items", help="the reviewed items file (.yaml or .json)")
    parser.add_argument("--out", help="output folder (default: 'output' next to the items file)")
    parser.add_argument("--formats", default=",".join(FORMATS),
                        help=f"comma-separated list from: {', '.join(FORMATS)} (default: all)")
    parser.add_argument("--preview", action="store_true",
                        help="include items that are not checked yet, clearly marked; for your own review only")
    parser.add_argument("--changes", help="changes.json from the watcher, used to count the sources checked")
    parser.add_argument("--paper", default="Letter", choices=["Letter", "A4"], help="PDF paper size")
    parser.add_argument("--template", default=str(EMAIL_TEMPLATE), help="email template to use")
    args = parser.parse_args(argv)

    items_path = Path(args.items)
    formats = [f.strip().lower() for f in args.formats.split(",") if f.strip()]
    unknown = [f for f in formats if f not in FORMATS]
    if unknown:
        print(f"Unknown format(s): {', '.join(unknown)}. Choose from: {', '.join(FORMATS)}", file=sys.stderr)
        return 2
    try:
        doc = load_items(items_path)
    except BuildError as exc:
        print(exc, file=sys.stderr)
        return 2

    print(f"Checking {items_path}")
    errors, warnings = check(doc)
    if errors:
        print("The items file has problems:", file=sys.stderr)
        for error in errors:
            print(f"  - {error}", file=sys.stderr)
        return 2
    counts = {}
    for item in doc["items"]:
        counts[item["status"]] = counts.get(item["status"], 0) + 1
    print("  " + ", ".join(f"{n} {status}" for status, n in sorted(counts.items())) if counts else "  no items")
    for warning in warnings:
        print(f"  Warning: {warning}")

    unchecked = [i for i in doc["items"] if i["status"] in UNCHECKED]
    if unchecked and not args.preview:
        print(f"\nStopped: {len(unchecked)} item(s) are not checked yet:", file=sys.stderr)
        for item in unchecked:
            print(f"  - {item['id']} ({item['status']}): {item['title']}", file=sys.stderr)
        print("Check each one against its source, then set status: approved with checked_by and "
              "checked_on, or status: rejected.\nTo look at a draft first, run again with --preview.",
              file=sys.stderr)
        return 3

    sources_checked = sources_checked_from(Path(args.changes)) if args.changes else None
    view = build_view(doc, args.preview, sources_checked, args.paper)
    if not view["items"]:
        print("  Note: no approved items, so this builds a 'nothing new this week' issue.")
    out_dir = Path(args.out) if args.out else items_path.resolve().parent / "output"
    out_dir.mkdir(parents=True, exist_ok=True)
    base = out_dir / view["base_name"]
    template = Path(args.template)
    written, failed = [], False

    if "html" in formats:
        base.with_suffix(".html").write_text(render_html(view, False, template), encoding="utf-8")
        written.append(base.with_suffix(".html"))
    if "txt" in formats:
        base.with_suffix(".txt").write_text(render_text(view), encoding="utf-8")
        written.append(base.with_suffix(".txt"))
    if "pdf" in formats:
        try:
            how = html_to_pdf(render_html(view, True, template), base.with_suffix(".pdf"), args.paper)
            add_pdf_metadata(base.with_suffix(".pdf"), view)
            written.append(base.with_suffix(".pdf"))
            print(f"  PDF printed with {how}")
        except BuildError as exc:
            print(f"  {exc}", file=sys.stderr)
            failed = True
    if "xlsx" in formats:
        build_xlsx(view, base.with_suffix(".xlsx"))
        written.append(base.with_suffix(".xlsx"))
    if "ics" in formats:
        events = build_ics(view, base.with_suffix(".ics"))
        written.append(base.with_suffix(".ics"))
        print(f"  Calendar: {events} event{'s' if events != 1 else ''}")

    print(f"\nWrote {len(written)} file(s) to {out_dir}:")
    for path in written:
        print(f"  {path.name}  ({max(1, round(path.stat().st_size / 1024))} KB)")
    if args.preview:
        print("\nThis is a PREVIEW with unchecked items. Do not send it.")
    else:
        print("\nNext: open the HTML and the PDF, click every link, import the .ics into a test "
              "calendar, then send the issue yourself.")
    return 4 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
