#!/usr/bin/env python3
"""Turn a simple Markdown report into a Word (.docx) file.

Usage:
    python md_to_docx.py report.md report.docx [--footer "Prepared for Acme Home Goods"]

Handles the Markdown used in the kit's templates:
    # Title, ## Heading, ### Subheading, #### Small heading
    paragraphs, **bold**, *italic*, `code`, [links](https://...)
    bullet lists (- or *), numbered lists (1.), one level of nesting, task boxes (- [ ] / - [x])
    tables (| a | b |), > quotes, ``` code blocks, and --- lines
HTML comments (<!-- ... -->) are left out, so you can keep instructions to yourself in the template.
Needs python-docx (pip install -r requirements.txt).
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is not installed. Run: pip install -r requirements.txt")

INLINE = re.compile(
    r"(\*\*(?P<b>.+?)\*\*)|(`(?P<c>[^`]+)`)|(\[(?P<lt>[^\]]+)\]\((?P<lu>[^)\s]+)\))|(\*(?P<i>[^*\s][^*]*?)\*)")
TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
LIST_ITEM = re.compile(r"^(?P<indent>\s*)(?P<marker>[-*+]|\d+[.)])\s+(?P<text>.*)$")
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
CHAR_WIDTH = 60000  # rough width of one character of 9.5 pt table text, in EMU (914400 per inch)


# --------------------------------------------------------------------------- inline text
def add_hyperlink(paragraph, text: str, url: str):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text: str, bold: bool = False, size=None):
    text = text.replace("\\|", "|")
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            _run(paragraph, text[pos:m.start()], bold=bold, size=size)
        if m.group("b") is not None:
            add_inline(paragraph, m.group("b"), bold=True, size=size)
        elif m.group("c") is not None:
            _run(paragraph, m.group("c"), bold=bold, size=size, mono=True)
        elif m.group("lt") is not None:
            add_hyperlink(paragraph, m.group("lt"), m.group("lu"))
        elif m.group("i") is not None:
            _run(paragraph, m.group("i"), bold=bold, size=size, italic=True)
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], bold=bold, size=size)


def _run(paragraph, text, bold=False, italic=False, mono=False, size=None):
    run = paragraph.add_run(text)
    run.bold, run.italic = bold or None, italic or None
    if mono:
        run.font.name = "Consolas"
        run.font.size = Pt(size.pt - 1) if size else Pt(9.5)
    elif size:
        run.font.size = size
    return run


# --------------------------------------------------------------------------- blocks
def shade(cell, hex_color: str):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    props.append(shd)


def split_row(line: str) -> list:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


def add_table(doc, lines: list, usable_width):
    header = split_row(lines[0])
    aligns = []
    body = lines[1:]
    if body and TABLE_SEP.match(body[0]):
        for spec in split_row(body[0]):
            aligns.append("right" if spec.endswith(":") and not spec.startswith(":")
                          else "center" if spec.startswith(":") and spec.endswith(":") else "left")
        body = body[1:]
    rows = [header] + [split_row(b) for b in body]
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    def measure(text):
        """Rough printed width of a cell and of its longest word (code text is wider)."""
        text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)
        code = re.findall(r"`([^`]*)`", text)
        plain = re.sub(r"`[^`]*`", " ", text).replace("*", "")
        def pieces(t):  # Word and LibreOffice can break lines after "/" and "-"
            return [w for w in re.split(r"(?<=[/-])|\s+", t) if w]
        words = [len(w) for w in pieces(plain)] + [1.05 * len(w) for part in code for w in pieces(part)]
        return len(plain) + 1.05 * sum(len(part) for part in code), max(words or [0])
    # Each column first gets room for its longest word; the rest is shared by amount of text.
    available = usable_width / CHAR_WIDTH
    mins, wants = [], []
    for c in range(ncols):
        sizes = [measure(r[c]) if c < len(r) else (0, 0) for r in rows]
        lens = [size for size, _ in sizes]
        mins.append(max(word for _, word in sizes) + 3)
        wants.append(min(max(0.4 * max(lens) + 0.6 * sum(lens) / len(lens), 6), 60))
    wants = [max(w, m) for w, m in zip(wants, mins)]
    if sum(mins) >= available:            # even the longest words do not fit: shrink evenly
        chars = [available * m / sum(mins) for m in mins]
    elif sum(wants) >= available:         # grow each column from its minimum towards what it wants
        grow = [w - m for w, m in zip(wants, mins)]
        chars = [m + (available - sum(mins)) * g / sum(grow) for m, g in zip(mins, grow)]
    else:                                 # everything fits: share the page in proportion
        chars = [available * w / sum(wants) for w in wants]
    table.autofit = False
    for c in range(ncols):
        width = int(chars[c] * CHAR_WIDTH)
        table.columns[c].width = width          # the grid (used by LibreOffice and Google Docs)
        for r in range(len(rows)):
            table.cell(r, c).width = width      # each cell (used by Word)
    for r, row in enumerate(rows):
        for c in range(ncols):
            cell = table.cell(r, c)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            if c < len(aligns) and aligns[c] != "left":
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if aligns[c] == "right" else WD_ALIGN_PARAGRAPH.CENTER
            add_inline(para, row[c] if c < len(row) else "", bold=(r == 0), size=Pt(9.5))
            if r == 0:
                shade(cell, "DDEBF7")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_rule(doc):
    para = doc.add_paragraph()
    props = para._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "A6A6A6")):
        bottom.set(qn(key), value)
    border.append(bottom)
    props.append(border)


def add_page_number_footer(doc, text: str):
    para = doc.sections[0].footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        _run(para, text + "  ·  ", size=Pt(8.5))
    _run(para, "Page ", size=Pt(8.5))
    run = para.add_run()
    run.font.size = Pt(8.5)
    for kind, instr in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        run._r.append(el)


# --------------------------------------------------------------------------- converter
def convert(md_text: str, out_path: Path, footer: str = ""):
    md_text = re.sub(r"<!--.*?-->", "", md_text, flags=re.S)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = ACCENT
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1)
    usable = section.page_width - section.left_margin - section.right_margin
    doc.core_properties.author = ""
    doc.core_properties.comments = ""
    doc.core_properties.last_modified_by = ""

    lines = md_text.split("\n")
    i, para_buf, title_done = 0, [], False
    number_counters = {}

    def flush():
        nonlocal para_buf
        if para_buf:
            add_inline(doc.add_paragraph(), " ".join(s.strip() for s in para_buf))
            para_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(12)
            for n, code_line in enumerate(code_lines):
                run = _run(para, code_line, mono=True)
                if n < len(code_lines) - 1:
                    run.add_break()
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            flush()
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and not title_done:
                doc.add_heading(text, level=0)
                doc.core_properties.title = re.sub(r"[*`]", "", text)
                title_done = True
            else:
                doc.add_heading(text.replace("**", ""), level=min(max(level - 1, 1), 4))
            number_counters.clear()
            i += 1
            continue
        if stripped.startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block, usable)
            continue
        item = LIST_ITEM.match(line)
        if item:
            flush()
            depth = 1 if len(item.group("indent").replace("\t", "    ")) >= 2 else 0
            text = item.group("text")
            text = re.sub(r"^\[ \]\s*", "☐ ", text)
            text = re.sub(r"^\[[xX]\]\s*", "☑ ", text)
            if item.group("marker")[0].isdigit():
                number_counters[depth] = number_counters.get(depth, 0) + 1
                for deeper in [d for d in number_counters if d > depth]:
                    del number_counters[deeper]
                para = doc.add_paragraph()
                fmt = para.paragraph_format
                fmt.left_indent = Pt(18 + 18 * depth)
                fmt.first_line_indent = Pt(-18)
                fmt.space_after = Pt(3)
                _run(para, f"{number_counters[depth]}.\t")
                fmt.tab_stops.add_tab_stop(Pt(18 + 18 * depth))
            else:
                para = doc.add_paragraph(style="List Bullet 2" if depth else "List Bullet")
                para.paragraph_format.space_after = Pt(3)
            # continuation lines of the same item
            i += 1
            while i < len(lines) and lines[i].strip() and lines[i].startswith((" ", "\t")) \
                    and not LIST_ITEM.match(lines[i]) and not lines[i].strip().startswith("|"):
                text += " " + lines[i].strip()
                i += 1
            add_inline(para, text)
            continue
        if not stripped:
            flush()
            if i + 1 < len(lines) and not LIST_ITEM.match(lines[i + 1]):
                number_counters.clear()
            i += 1
            continue
        if stripped.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            add_inline(doc.add_paragraph(style="Intense Quote"), " ".join(quote))
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            flush()
            add_rule(doc)
            i += 1
            continue
        para_buf.append(line)
        i += 1
    flush()
    add_page_number_footer(doc, footer)
    doc.save(out_path)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Convert a simple Markdown report to .docx")
    ap.add_argument("markdown", help="input .md file")
    ap.add_argument("docx", nargs="?", help="output .docx (default: same name as the input)")
    ap.add_argument("--footer", default="", help="text for the footer, next to the page number")
    args = ap.parse_args(argv)
    src = Path(args.markdown)
    if not src.exists():
        print(f"Not found: {src}", file=sys.stderr)
        return 1
    out = Path(args.docx) if args.docx else src.with_suffix(".docx")
    convert(src.read_text(encoding="utf-8"), out, args.footer)
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
